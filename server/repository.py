import re
import io
import json
from datetime import date, timedelta, datetime, timezone
import csv,io,html,base64,mimetypes
from pathlib import Path
from .database import connect, write_transaction
from .services import inventory_service
from .services import bank_pdf_parser, purchase_service, accounting_service, cash_service, returns_service
from .security import hash_password, verify_password, new_token, utc_now, expires_at, is_expired
from .config import TOKEN_HOURS, DATA_DIR, IS_POSTGRES
from . import licensing, multi_unit, subscription
def user_dict(r):
    return {"id":r["id"],"username":r["username"],"full_name":r["full_name"],
    "is_active":bool(r["is_active"]),"role":{"id":r["role_id"],"code":r["role_code"],
    "name":r["role_name"],"permissions":json.loads(r["permissions_json"] or "[]")},
    "created_at":r["created_at"],"updated_at":r["updated_at"]}
def audit(uid,action,etype=None,eid=None,details=None,ip=None,conn=None):
    """Best-effort audit logging; audit failure must never abort login or posting."""
    q="""INSERT INTO audit_logs(user_id,action,entity_type,entity_id,details_json,client_ip,created_at)
    VALUES(?,?,?,?,?,?,?)"""
    p=(uid,action,etype,str(eid) if eid is not None else None,json.dumps(details or {},ensure_ascii=False),ip,utc_now())
    try:
        if conn:
            conn.execute(q,p)
        else:
            with write_transaction() as tx:
                tx.execute(q,p)
        return True
    except Exception:
        return False
def login(username,password,ip,client_name="browser",device_id=None,user_agent=None):
    c=connect()
    try:
        r=c.execute("""SELECT u.*,r.code role_code,r.name role_name,r.permissions_json
        FROM users u JOIN roles r ON r.id=u.role_id WHERE u.username=? COLLATE NOCASE""",(username.strip(),)).fetchone()
    finally: c.close()
    if not r or not r["is_active"] or not verify_password(password,r["password_hash"]):
        audit(r["id"] if r else None,"LOGIN_FAILED","user",r["id"] if r else username,{"username":username},ip); return None
    device_id=str(device_id or '').strip()[:160]
    if not device_id:
        device_id=f"legacy:{ip}:{str(client_name or 'browser')[:60]}"
    token=new_token(); now=utc_now()
    with write_transaction() as tx:
        status=licensing.current_license_status(tx)
        # Bersihkan sesi kedaluwarsa dan sesi tanpa heartbeat lebih dari 10 detik.
        tx.execute("DELETE FROM sessions WHERE expires_at<=?",(now,))
        cutoff=(datetime.now(timezone.utc)-timedelta(seconds=licensing.DEVICE_IDLE_SECONDS)).isoformat()
        tx.execute("DELETE FROM sessions WHERE last_seen_at<?",(cutoff,))
        # Satu browser/device hanya memiliki satu sesi aktif agar refresh login tidak memakai slot ganda.
        tx.execute("DELETE FROM sessions WHERE device_id=?",(device_id,))
        active=int(tx.execute("SELECT COUNT(DISTINCT COALESCE(NULLIF(device_id,''),client_ip||':'||client_name)) FROM sessions WHERE expires_at>?",(now,)).fetchone()[0])
        limit=int(status.get("device_limit") or status.get("max_users") or 1)
        if active>=limit:
            raise ValueError(f"Batas akses maksimal {limit} device/browser aktif telah tercapai. Logout dari device lain; sesi yang tidak aktif akan dilepas otomatis dalam {licensing.DEVICE_IDLE_SECONDS} detik.")
        tx.execute("""INSERT INTO sessions(token,user_id,client_ip,client_name,device_id,user_agent,created_at,expires_at,last_seen_at)
        VALUES(?,?,?,?,?,?,?,?,?)""",(token,r["id"],ip,client_name,device_id,str(user_agent or '')[:300],now,expires_at(TOKEN_HOURS),now))
        audit(r["id"],"LOGIN_SUCCESS","session",token[:8],{"device_id":device_id,"client_name":client_name},ip,tx)
    return {"token":token,"user":user_dict(r)}

def logout(token):
    if not token:return False
    with write_transaction() as tx:
        row=tx.execute("SELECT user_id,device_id FROM sessions WHERE token=?",(token,)).fetchone()
        tx.execute("DELETE FROM sessions WHERE token=?",(token,))
        if row:audit(row["user_id"],"LOGOUT","session",str(token)[:8],{"device_id":row["device_id"]},None,tx)
    return True
def heartbeat(token):
    if not token:return False
    with write_transaction() as tx:
        row=tx.execute("SELECT token FROM sessions WHERE token=?",(token,)).fetchone()
        if not row:return False
        tx.execute("UPDATE sessions SET last_seen_at=? WHERE token=?",(utc_now(),token))
    return True

def list_active_sessions():
    cutoff=(datetime.now(timezone.utc)-timedelta(seconds=licensing.DEVICE_IDLE_SECONDS)).isoformat()
    now=utc_now()
    c=connect()
    try:
        rows=c.execute("""SELECT s.token,s.client_ip,s.client_name,s.device_id,s.user_agent,s.created_at,s.last_seen_at,u.username,u.full_name
        FROM sessions s JOIN users u ON u.id=s.user_id
        WHERE s.expires_at>? AND s.last_seen_at>=? ORDER BY s.last_seen_at DESC""",(now,cutoff)).fetchall()
        return [{"session_id":r["token"][:12],"client_ip":r["client_ip"],"client_name":r["client_name"],"device_id":r["device_id"],"user_agent":r["user_agent"],"created_at":r["created_at"],"last_seen_at":r["last_seen_at"],"username":r["username"],"full_name":r["full_name"]} for r in rows]
    finally:c.close()

def force_logout_session(session_id):
    session_id=str(session_id or '').strip()
    if not session_id:return False
    with write_transaction() as tx:
        row=tx.execute("SELECT token,user_id,device_id FROM sessions WHERE substr(token,1,12)=?",(session_id,)).fetchone()
        if not row:return False
        tx.execute("DELETE FROM sessions WHERE token=?",(row["token"],))
        audit(row["user_id"],"FORCE_LOGOUT","session",session_id,{"device_id":row["device_id"]},None,tx)
    return True

def authenticate(token):
    if not token: return None
    c=connect()
    try:
        r=c.execute("""SELECT s.expires_at,u.*,r.code role_code,r.name role_name,r.permissions_json
        FROM sessions s JOIN users u ON u.id=s.user_id JOIN roles r ON r.id=u.role_id WHERE s.token=?""",(token,)).fetchone()
        if not r or not r["is_active"] or is_expired(r["expires_at"]): return None
        c.execute("UPDATE sessions SET last_seen_at=? WHERE token=?",(utc_now(),token)); c.commit()
        return user_dict(r)
    finally: c.close()
def allowed(user,perm): return "*" in user["role"]["permissions"] or perm in user["role"]["permissions"]
def license_status():
    c=connect()
    try:
        return licensing.current_license_status(c)
    finally:
        c.close()

def subscription_admin_status(actor=None):
    c=connect()
    try:
        st=subscription.status(c); st['plans']=subscription.catalog(); return st
    finally:c.close()

def subscription_admin_activate(actor,d,ip):
    if not actor or actor.get('role',{}).get('code')!='ADMIN': raise ValueError('Hanya Administrator yang dapat mengatur langganan.')
    subscription.validate_admin_key(d.get('admin_key'))
    with write_transaction() as tx:
        st=subscription.activate(tx,d.get('plan_code'),d.get('addon_users',0),d.get('customer_name',''),d.get('notes',''),d.get('starts_at'))
        audit(actor['id'],'SUBSCRIPTION_ACTIVATED','subscription','1',{'plan_code':st.get('plan_code'),'addon_users':st.get('addon_users'),'expires_at':st.get('expires_at')},ip,tx)
        return st

def subscription_admin_reset_trial(actor,ip,admin_key=None):
    if not actor or actor.get('role',{}).get('code')!='ADMIN': raise ValueError('Hanya Administrator yang dapat mengatur trial.')
    subscription.validate_admin_key(admin_key)
    with write_transaction() as tx:
        st=subscription.reset_trial(tx); audit(actor['id'],'TRIAL_RESET','subscription','1',{},ip,tx); return st
def list_roles():
    c=connect()
    try:
        return [{"id":r["id"],"code":r["code"],"name":r["name"],"permissions":json.loads(r["permissions_json"])} for r in c.execute("SELECT * FROM roles ORDER BY id")]
    finally:c.close()
def list_users():
    c=connect()
    try:
        return [user_dict(r) for r in c.execute("""SELECT u.*,r.code role_code,r.name role_name,r.permissions_json
        FROM users u JOIN roles r ON r.id=u.role_id ORDER BY u.username""")]
    finally:c.close()
def create_user(actor,d,ip):
    username=str(d.get("username","")).strip(); name=str(d.get("full_name","")).strip()
    password=str(d.get("password","")); role=str(d.get("role_code","")).strip().upper()
    if not username or not name or len(password)<8 or not role: raise ValueError("Data user belum lengkap.")
    now=utc_now()
    with write_transaction() as tx:
        licensing.enforce_user_limit(tx, activating_new_user=True)
        rr=tx.execute("SELECT id FROM roles WHERE code=?",(role,)).fetchone()
        if not rr: raise ValueError("Role tidak ditemukan.")
        cur=tx.execute("""INSERT INTO users(username,full_name,password_hash,role_id,is_active,created_at,updated_at)
        VALUES(?,?,?,?,1,?,?)""",(username,name,hash_password(password),rr["id"],now,now))
        audit(actor["id"],"USER_CREATED","user",cur.lastrowid,{"username":username,"role_code":role},ip,tx)
        return cur.lastrowid
def list_audit(limit=100):
    limit=max(1,min(int(limit),500)); c=connect()
    try:
        return [{"id":r["id"],"username":r["username"],"action":r["action"],"client_ip":r["client_ip"],"created_at":r["created_at"]}
        for r in c.execute("""SELECT a.*,u.username FROM audit_logs a LEFT JOIN users u ON u.id=a.user_id
        ORDER BY a.id DESC LIMIT ?""",(limit,))]
    finally:c.close()



def _date_value(value,label):
    value=str(value or "").strip()[:10]
    try: return date.fromisoformat(value).isoformat()
    except Exception: raise ValueError(f"{label} tidak valid.")

def _decimal(value, field, allow_negative=False):
    try:
        number=float(value or 0)
    except (TypeError,ValueError):
        raise ValueError(f"{field} harus berupa angka.")
    if not allow_negative and number < 0:
        raise ValueError(f"{field} tidak boleh negatif.")
    return round(number,4)

def list_categories(active_only=False):
    c=connect()
    try:
        sql="SELECT * FROM item_categories" + (" WHERE is_active=1" if active_only else "") + " ORDER BY code COLLATE NOCASE, name COLLATE NOCASE"
        return [dict(r) | {"is_active":bool(r["is_active"])} for r in c.execute(sql)]
    finally:c.close()

def create_category(actor,d,ip):
    code=str(d.get("code","")).strip().upper(); name=str(d.get("name","")).strip()
    if not code or not name: raise ValueError("Kode dan nama kategori wajib diisi.")
    now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("INSERT INTO item_categories(code,name,is_active,created_at,updated_at) VALUES(?,?,1,?,?)",(code,name,now,now))
        audit(actor["id"],"CATEGORY_CREATED","category",cur.lastrowid,{"code":code,"name":name},ip,tx)
        return cur.lastrowid

def list_units(active_only=False):
    c=connect()
    try:
        sql="SELECT * FROM units" + (" WHERE is_active=1" if active_only else "") + " ORDER BY code COLLATE NOCASE, name COLLATE NOCASE"
        return [dict(r) | {"is_active":bool(r["is_active"])} for r in c.execute(sql)]
    finally:c.close()

def create_unit(actor,d,ip):
    code=str(d.get("code","")).strip().upper(); name=str(d.get("name","")).strip()
    decimals=int(d.get("decimals",0))
    if not code or not name: raise ValueError("Kode dan nama satuan wajib diisi.")
    if decimals<0 or decimals>4: raise ValueError("Desimal satuan harus 0 sampai 4.")
    now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("INSERT INTO units(code,name,decimals,is_active,created_at,updated_at) VALUES(?,?,?,1,?,?)",(code,name,decimals,now,now))
        audit(actor["id"],"UNIT_CREATED","unit",cur.lastrowid,{"code":code,"name":name},ip,tx)
        return cur.lastrowid

def _product_dict(r):
    keys=set(r.keys())
    return {"id":r["id"],"sku":r["sku"],"barcode":r["barcode"],"name":r["name"],
    "category_id":r["category_id"],"category_name":r["category_name"],
    "brand_id":r["brand_id"] if "brand_id" in keys else None,
    "brand_name":r["brand_name"] if "brand_name" in keys else None,
    "unit_id":r["unit_id"],"unit_code":r["unit_code"],"product_type":r["product_type"],
    "purchase_price":float(r["purchase_price"]),"selling_price":float(r["selling_price"]),
    "stock_qty":float(r["realtime_stock_qty"] if "realtime_stock_qty" in keys else r["stock_qty"]),"stock_value":float(r["realtime_stock_value"] if "realtime_stock_value" in keys else 0),"minimum_stock":float(r["minimum_stock"]),
    "opening_stock_qty":float(r["opening_stock_qty"] or 0) if "opening_stock_qty" in keys else 0.0,
    "opening_warehouse_id":r["opening_warehouse_id"] if "opening_warehouse_id" in keys else None,
    "opening_balance_date":r["opening_balance_date"] if "opening_balance_date" in keys else None,
    "inventory_account_id":r["inventory_account_id"] if "inventory_account_id" in keys else None,
    "sales_account_id":r["sales_account_id"] if "sales_account_id" in keys else None,
    "cogs_account_id":r["cogs_account_id"] if "cogs_account_id" in keys else None,"inventory_account_name":r["inventory_account_name"] if "inventory_account_name" in keys else None,"sales_account_name":r["sales_account_name"] if "sales_account_name" in keys else None,"cogs_account_name":r["cogs_account_name"] if "cogs_account_name" in keys else None,
    "service_id":r["service_id"] if "service_id" in keys else None,
    "notes":r["notes"],"is_active":bool(r["is_active"]),
    "created_at":r["created_at"],"updated_at":r["updated_at"]}

def list_service_categories(active_only=False):
    c=connect()
    try:
        sql="SELECT * FROM service_categories"
        if active_only:sql+=" WHERE is_active=1"
        sql+=" ORDER BY code,name"
        return [dict(x) for x in c.execute(sql).fetchall()]
    finally:c.close()

def create_service_category(actor,d,ip):
    code=str(d.get("code") or "").strip().upper()
    name=str(d.get("name") or "").strip()
    if not code or not name:raise ValueError("Kode dan nama kategori jasa wajib diisi.")
    now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("""INSERT INTO service_categories(
          code,name,notes,is_active,created_at,updated_at
        ) VALUES(?,?,?,?,?,?)""",(
          code,name,str(d.get("notes") or "").strip() or None,
          1 if d.get("is_active",True) else 0,now,now))
        audit(actor["id"],"SERVICE_CATEGORY_CREATED","service_category",cur.lastrowid,
          {"code":code,"name":name},ip,tx)
        return cur.lastrowid

def update_service_category(actor,category_id,d,ip):
    category_id=int(category_id)
    code=str(d.get("code") or "").strip().upper()
    name=str(d.get("name") or "").strip()
    if not code or not name:raise ValueError("Kode dan nama kategori jasa wajib diisi.")
    with write_transaction() as tx:
        if not tx.execute("SELECT 1 FROM service_categories WHERE id=?",(category_id,)).fetchone():
            raise ValueError("Kategori jasa tidak ditemukan.")
        tx.execute("""UPDATE service_categories SET code=?,name=?,notes=?,
          is_active=?,updated_at=? WHERE id=?""",(
          code,name,str(d.get("notes") or "").strip() or None,
          1 if d.get("is_active",True) else 0,utc_now(),category_id))
        audit(actor["id"],"SERVICE_CATEGORY_UPDATED","service_category",category_id,
          {"code":code,"name":name},ip,tx)
    return next(x for x in list_service_categories(False) if x["id"]==category_id)

def delete_service_category(actor,category_id,ip):
    category_id=int(category_id)
    with write_transaction() as tx:
        row=tx.execute("SELECT * FROM service_categories WHERE id=?",(category_id,)).fetchone()
        if not row:raise ValueError("Kategori jasa tidak ditemukan.")
        if tx.execute("SELECT 1 FROM services WHERE category_id=? LIMIT 1",(category_id,)).fetchone():
            raise ValueError("Kategori masih digunakan oleh Master Jasa dan tidak dapat dihapus.")
        tx.execute("DELETE FROM service_categories WHERE id=?",(category_id,))
        audit(actor["id"],"SERVICE_CATEGORY_DELETED","service_category",category_id,
          {"code":row["code"],"name":row["name"]},ip,tx)
    return True

def _service_account(tx,value,side):
    try:account_id=int(value)
    except Exception:raise ValueError(
      "Akun Pembelian wajib dipilih." if side=="PURCHASE" else "Akun Penjualan wajib dipilih.")
    row=tx.execute("""SELECT id,code,name,account_type,COALESCE(account_subtype,'') account_subtype
      FROM chart_of_accounts WHERE id=? AND is_active=1""",(account_id,)).fetchone()
    if not row:raise ValueError("Akun jasa tidak ditemukan atau nonaktif.")
    account_type=(row["account_type"] or "").upper()
    subtype=(row["account_subtype"] or "").upper()
    if side=="PURCHASE":
        allowed=account_type in ("EXPENSE","ASSET") or subtype in ("HPP","OPERATING_EXPENSE")
        if not allowed:raise ValueError("Akun Pembelian Jasa hanya boleh tipe Beban, HPP, atau Aset.")
    else:
        allowed=account_type in ("REVENUE","ASSET") or subtype=="REVENUE"
        if not allowed:raise ValueError("Akun Penjualan Jasa hanya boleh tipe Pendapatan atau Aset.")
    return account_id

def list_services(q="",active=None,category_id=None):
    where=[];params=[]
    if q:
        where.append("(s.service_code LIKE ? OR s.service_name LIKE ?)")
        term=f"%{q}%";params.extend([term,term])
    if active is not None:
        where.append("s.is_active=?");params.append(1 if active else 0)
    if category_id not in (None,""):
        where.append("s.category_id=?");params.append(int(category_id))
    clause=(" WHERE "+" AND ".join(where)) if where else ""
    c=connect()
    try:
        rows=c.execute("""SELECT s.*,sc.code category_code,sc.name category_name,
          u.code unit_code,u.name unit_name,
          pa.code purchase_account_code,pa.name purchase_account_name,
          sa.code sales_account_code,sa.name sales_account_name
          FROM services s
          LEFT JOIN service_categories sc ON sc.id=s.category_id
          LEFT JOIN units u ON u.id=s.unit_id
          LEFT JOIN chart_of_accounts pa ON pa.id=s.purchase_account_id
          LEFT JOIN chart_of_accounts sa ON sa.id=s.sales_account_id"""+
          clause+" ORDER BY s.service_code,s.service_name",params).fetchall()
        result=[]
        for row in rows:
            item=dict(row)
            try:item["purchase_price"]=float(item.get("purchase_price") or 0)
            except (TypeError,ValueError):item["purchase_price"]=0.0
            try:item["selling_price"]=float(item.get("selling_price") or 0)
            except (TypeError,ValueError):item["selling_price"]=0.0
            try:item["tax_percent"]=float(item.get("tax_percent") or 0)
            except (TypeError,ValueError):item["tax_percent"]=0.0
            item["is_active"]=bool(item["is_active"])
            result.append(item)
        return result
    finally:c.close()

def _service_shadow_create(tx,code,name,unit_id,purchase_price,selling_price,
                           purchase_account_id,sales_account_id,is_active,now):
    shadow_sku="@SVC:"+code
    cur=tx.execute("""INSERT INTO products(
      sku,barcode,name,category_id,brand_id,unit_id,product_type,
      purchase_price,selling_price,stock_qty,minimum_stock,notes,
      inventory_account_id,sales_account_id,cogs_account_id,
      is_active,created_at,updated_at
    ) VALUES(?,NULL,?,NULL,NULL,?,'SERVICE',?,?,0,0,?,
      NULL,?,?,?, ?,?)""",(
      shadow_sku,name,unit_id,str(purchase_price),str(selling_price),
      "Internal Master Jasa",sales_account_id,purchase_account_id,
      1 if is_active else 0,now,now))
    return cur.lastrowid

def _service_shadow_update(tx,product_id,code,name,unit_id,purchase_price,selling_price,
                           purchase_account_id,sales_account_id,is_active,now):
    tx.execute("""UPDATE products SET sku=?,name=?,unit_id=?,product_type='SERVICE',
      purchase_price=?,selling_price=?,stock_qty=0,minimum_stock=0,
      inventory_account_id=NULL,sales_account_id=?,cogs_account_id=?,
      is_active=?,updated_at=? WHERE id=?""",(
      "@SVC:"+code,name,unit_id,str(purchase_price),str(selling_price),
      sales_account_id,purchase_account_id,1 if is_active else 0,now,product_id))

def _service_unit_or_default(tx,value):
    try:
        unit_id=int(value) if value not in (None,'') else None
    except Exception:
        unit_id=None
    if unit_id:
        row=tx.execute("SELECT id FROM units WHERE id=? AND is_active=1",(unit_id,)).fetchone()
        if row:return int(row['id'])
        raise ValueError("Satuan jasa tidak ditemukan atau nonaktif.")
    row=tx.execute("SELECT id FROM units WHERE code='PCS' AND is_active=1 ORDER BY id LIMIT 1").fetchone()
    if row:return int(row['id'])
    row=tx.execute("SELECT id FROM units WHERE is_active=1 ORDER BY id LIMIT 1").fetchone()
    if row:return int(row['id'])
    raise ValueError("Belum ada master satuan aktif. Tambahkan minimal satu satuan untuk penyimpanan internal jasa.")

def create_service(actor,d,ip):
    code=str(d.get("service_code") or "").strip().upper()
    name=str(d.get("service_name") or "").strip()
    if not code or not name:raise ValueError("Kode dan nama jasa wajib diisi.")
    unit_id=None
    category_id=_optional_int(d.get("category_id"))
    purchase_price=_decimal(d.get("purchase_price",0),"Harga beli jasa")
    selling_price=_decimal(d.get("selling_price",0),"Harga jual jasa")
    tax_percent=_decimal(d.get("tax_percent",0),"PPN jasa")
    if tax_percent>100:raise ValueError("PPN jasa tidak boleh lebih dari 100%.")
    now=utc_now()
    with write_transaction() as tx:
        unit_id=_service_unit_or_default(tx,d.get('unit_id'))
        if category_id and not tx.execute(
          "SELECT 1 FROM service_categories WHERE id=? AND is_active=1",(category_id,)).fetchone():
            raise ValueError("Kategori jasa tidak ditemukan atau nonaktif.")
        purchase_account_id=_service_account(tx,d.get("purchase_account_id"),"PURCHASE")
        sales_account_id=_service_account(tx,d.get("sales_account_id"),"SALES")
        active=1 if d.get("is_active",True) else 0
        product_id=_service_shadow_create(tx,code,name,unit_id,purchase_price,selling_price,
          purchase_account_id,sales_account_id,active,now)
        cur=tx.execute("""INSERT INTO services(
          service_code,service_name,category_id,unit_id,purchase_price,selling_price,
          tax_percent,purchase_account_id,sales_account_id,product_id,notes,
          is_active,created_at,updated_at
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(
          code,name,category_id,unit_id,str(purchase_price),str(selling_price),
          str(tax_percent),purchase_account_id,sales_account_id,product_id,
          str(d.get("notes") or "").strip() or None,active,now,now))
        audit(actor["id"],"SERVICE_CREATED","service",cur.lastrowid,
          {"service_code":code,"service_name":name},ip,tx)
        return cur.lastrowid

def update_service(actor,service_id,d,ip):
    service_id=int(service_id)
    current=next((x for x in list_services() if x["id"]==service_id),None)
    if not current:raise ValueError("Jasa tidak ditemukan.")
    payload=dict(current);payload.update(d)
    code=str(payload.get("service_code") or "").strip().upper()
    name=str(payload.get("service_name") or "").strip()
    if not code or not name:raise ValueError("Kode dan nama jasa wajib diisi.")
    unit_id=None
    category_id=_optional_int(payload.get("category_id"))
    purchase_price=_decimal(payload.get("purchase_price",0),"Harga beli jasa")
    selling_price=_decimal(payload.get("selling_price",0),"Harga jual jasa")
    tax_percent=_decimal(payload.get("tax_percent",0),"PPN jasa")
    if tax_percent>100:raise ValueError("PPN jasa tidak boleh lebih dari 100%.")
    with write_transaction() as tx:
        unit_id=_service_unit_or_default(tx,payload.get('unit_id'))
        if category_id and not tx.execute(
          "SELECT 1 FROM service_categories WHERE id=? AND is_active=1",(category_id,)).fetchone():
            raise ValueError("Kategori jasa tidak ditemukan atau nonaktif.")
        purchase_account_id=_service_account(tx,payload.get("purchase_account_id"),"PURCHASE")
        sales_account_id=_service_account(tx,payload.get("sales_account_id"),"SALES")
        now=utc_now();active=1 if payload.get("is_active",True) else 0
        product_id=current.get("product_id")
        if not product_id:
            product_id=_service_shadow_create(tx,code,name,unit_id,purchase_price,selling_price,
              purchase_account_id,sales_account_id,active,now)
        else:
            _service_shadow_update(tx,int(product_id),code,name,unit_id,purchase_price,
              selling_price,purchase_account_id,sales_account_id,active,now)
        tx.execute("""UPDATE services SET service_code=?,service_name=?,category_id=?,
          unit_id=?,purchase_price=?,selling_price=?,tax_percent=?,purchase_account_id=?,
          sales_account_id=?,product_id=?,notes=?,is_active=?,updated_at=? WHERE id=?""",(
          code,name,category_id,unit_id,str(purchase_price),str(selling_price),
          str(tax_percent),purchase_account_id,sales_account_id,product_id,
          str(payload.get("notes") or "").strip() or None,active,now,service_id))
        audit(actor["id"],"SERVICE_UPDATED","service",service_id,
          {"service_code":code,"service_name":name},ip,tx)
    return next(x for x in list_services() if x["id"]==service_id)

def delete_service(actor,service_id,ip):
    service_id=int(service_id)
    with write_transaction() as tx:
        row=tx.execute("SELECT * FROM services WHERE id=?",(service_id,)).fetchone()
        if not row:raise ValueError("Jasa tidak ditemukan.")
        product_id=row["product_id"] if "product_id" in row.keys() else None
        used=bool(product_id and (
          tx.execute("SELECT 1 FROM sales_items WHERE product_id=? LIMIT 1",(product_id,)).fetchone()
          or tx.execute("SELECT 1 FROM purchase_items WHERE product_id=? LIMIT 1",(product_id,)).fetchone()
        ))
        if used:
            tx.execute("UPDATE services SET is_active=0,updated_at=? WHERE id=?",(utc_now(),service_id))
            tx.execute("UPDATE products SET is_active=0,updated_at=? WHERE id=?",(utc_now(),product_id))
        else:
            tx.execute("DELETE FROM services WHERE id=?",(service_id,))
            if product_id:tx.execute("DELETE FROM products WHERE id=?",(product_id,))
        audit(actor["id"],"SERVICE_DELETED","service",service_id,
          {"service_code":row["service_code"],"service_name":row["service_name"]},ip,tx)
    return True


def list_products(q="",active=None,low_stock=False):
    where=[]; params=[]
    if q:
        where.append("(p.sku LIKE ? OR p.name LIKE ? OR COALESCE(p.barcode,'') LIKE ?)")
        term=f"%{q}%"; params += [term,term,term]
    if active is not None: where.append("p.is_active=?"); params.append(1 if active else 0)
    if low_stock: where.append("p.product_type='STOCK' AND COALESCE((SELECT SUM(ibx.quantity) FROM inventory_balances ibx WHERE ibx.product_id=p.id),0)<=p.minimum_stock")
    clause=(" WHERE "+" AND ".join(where)) if where else ""
    c=connect()
    try:
        rows=c.execute("""SELECT p.*,c.name category_name,b.name brand_name,u.code unit_code,s.id service_id,
        COALESCE((SELECT SUM(ib.quantity) FROM inventory_balances ib WHERE ib.product_id=p.id),0) realtime_stock_qty,
        COALESCE((SELECT SUM(COALESCE(ib.book_value,ib.quantity*ib.average_cost)) FROM inventory_balances ib WHERE ib.product_id=p.id),0) realtime_stock_value,
        ia.code||' - '||ia.name inventory_account_name,sa.code||' - '||sa.name sales_account_name,ha.code||' - '||ha.name cogs_account_name
        FROM products p LEFT JOIN item_categories c ON c.id=p.category_id
        LEFT JOIN brands b ON b.id=p.brand_id JOIN units u ON u.id=p.unit_id
        LEFT JOIN services s ON s.product_id=p.id LEFT JOIN chart_of_accounts ia ON ia.id=p.inventory_account_id
        LEFT JOIN chart_of_accounts sa ON sa.id=p.sales_account_id LEFT JOIN chart_of_accounts ha ON ha.id=p.cogs_account_id"""+clause+" ORDER BY p.sku COLLATE NOCASE,p.name COLLATE NOCASE",params)
        out=[_product_dict(r) for r in rows]
        for item in out:
            item["price_levels"]=product_price_map(item["id"])
            item["units"]=multi_unit.product_units(c,item["id"],item["unit_id"],item["purchase_price"],item["selling_price"])
            base_stock=Decimal(str(item.get("stock_qty") or 0))
            ordered_units=sorted(item["units"],key=lambda unit:(0 if unit.get("is_base") else 1,int(unit.get("id") or unit.get("unit_id") or 0)))
            item["warehouse_stocks"]=[{"warehouse_id":x["warehouse_id"],"warehouse_code":x["warehouse_code"],"warehouse_name":x["warehouse_name"],"quantity":float(x["quantity"] or 0),"average_cost":float(x["average_cost"] or 0),"stock_value":float(x["book_value"] if "book_value" in x.keys() else (x["quantity"] or 0)*(x["average_cost"] or 0))} for x in c.execute("SELECT ib.warehouse_id,w.code warehouse_code,w.name warehouse_name,ib.quantity,ib.average_cost,COALESCE(ib.book_value,ib.quantity*ib.average_cost) book_value FROM inventory_balances ib JOIN warehouses w ON w.id=ib.warehouse_id WHERE ib.product_id=? ORDER BY w.code",(item["id"],)).fetchall()]
            item["stock_units"]=[]
            for unit in ordered_units[:3]:
                ratio=Decimal(str(unit.get("conversion_ratio") or 1))
                equivalent=(base_stock/ratio) if ratio else Decimal("0")
                item["stock_units"].append({
                    "unit_id":unit.get("unit_id"),
                    "unit_code":unit.get("unit_code") or unit.get("unit_name") or "-",
                    "unit_name":unit.get("unit_name") or unit.get("unit_code") or "-",
                    "conversion_ratio":float(ratio),
                    "quantity":float(equivalent),
                    "decimals":int(unit.get("decimals") or 0),
                    "is_base":bool(unit.get("is_base")),
                    "purchase_price":float(unit.get("purchase_price") or 0),
                    "selling_price":float(unit.get("selling_price") or 0),
                })
        return out
    finally:c.close()

def get_product(product_id):
    c=connect()
    try:
        r=c.execute("""SELECT p.*,c.name category_name,b.name brand_name,u.code unit_code,s.id service_id,
        COALESCE((SELECT SUM(ib.quantity) FROM inventory_balances ib WHERE ib.product_id=p.id),0) realtime_stock_qty,
        COALESCE((SELECT SUM(COALESCE(ib.book_value,ib.quantity*ib.average_cost)) FROM inventory_balances ib WHERE ib.product_id=p.id),0) realtime_stock_value,
        ia.code||' - '||ia.name inventory_account_name,sa.code||' - '||sa.name sales_account_name,ha.code||' - '||ha.name cogs_account_name
        FROM products p LEFT JOIN item_categories c ON c.id=p.category_id
        LEFT JOIN brands b ON b.id=p.brand_id JOIN units u ON u.id=p.unit_id
        LEFT JOIN services s ON s.product_id=p.id LEFT JOIN chart_of_accounts ia ON ia.id=p.inventory_account_id
        LEFT JOIN chart_of_accounts sa ON sa.id=p.sales_account_id LEFT JOIN chart_of_accounts ha ON ha.id=p.cogs_account_id WHERE p.id=?""",(product_id,)).fetchone()
        
        if not r:return None
        out=_product_dict(r)
        out["price_levels"]=product_price_map(product_id)
        out["units"]=multi_unit.product_units(c,product_id,out["unit_id"],out["purchase_price"],out["selling_price"])
        base_stock=Decimal(str(out.get("stock_qty") or 0))
        ordered_units=sorted(out["units"],key=lambda unit:(0 if unit.get("is_base") else 1,int(unit.get("id") or unit.get("unit_id") or 0)))
        out["warehouse_stocks"]=[{"warehouse_id":x["warehouse_id"],"warehouse_code":x["warehouse_code"],"warehouse_name":x["warehouse_name"],"quantity":float(x["quantity"] or 0),"average_cost":float(x["average_cost"] or 0),"stock_value":float(x["book_value"] if "book_value" in x.keys() else (x["quantity"] or 0)*(x["average_cost"] or 0))} for x in c.execute("SELECT ib.warehouse_id,w.code warehouse_code,w.name warehouse_name,ib.quantity,ib.average_cost,COALESCE(ib.book_value,ib.quantity*ib.average_cost) book_value FROM inventory_balances ib JOIN warehouses w ON w.id=ib.warehouse_id WHERE ib.product_id=? ORDER BY w.code",(product_id,)).fetchall()]
        out["stock_units"]=[]
        for unit in ordered_units[:3]:
            ratio=Decimal(str(unit.get("conversion_ratio") or 1))
            out["stock_units"].append({"unit_id":unit.get("unit_id"),"unit_code":unit.get("unit_code") or unit.get("unit_name") or "-","unit_name":unit.get("unit_name") or unit.get("unit_code") or "-","conversion_ratio":float(ratio),"quantity":float(base_stock/ratio) if ratio else 0,"decimals":int(unit.get("decimals") or 0),"is_base":bool(unit.get("is_base")),"purchase_price":float(unit.get("purchase_price") or 0),"selling_price":float(unit.get("selling_price") or 0)})
        return out
    finally:c.close()


def _product_account(tx,value,default_code,allowed_subtypes,label):
    if value in (None,""):
        row=tx.execute("SELECT id,account_subtype FROM chart_of_accounts WHERE code=? AND is_active=1",(default_code,)).fetchone()
    else:
        try: account_id=int(value)
        except Exception: raise ValueError(f"{label} tidak valid.")
        row=tx.execute("SELECT id,account_subtype FROM chart_of_accounts WHERE id=? AND is_active=1",(account_id,)).fetchone()
    if not row or row["account_subtype"] not in allowed_subtypes:
        raise ValueError(f"{label} harus memilih akun aktif dengan tipe yang sesuai.")
    return int(row["id"])

def create_product(actor,d,ip):
    sku=str(d.get("sku","")).strip().upper()
    name=str(d.get("name","")).strip()
    barcode=str(d.get("barcode") or "").strip() or None
    ptype=str(d.get("product_type","STOCK")).strip().upper()
    if not sku: raise ValueError("Kode Barang / SKU wajib diisi.")
    if not name: raise ValueError("Nama Barang / Jasa wajib diisi.")
    if ptype not in ("STOCK","SERVICE"): raise ValueError("Jenis item harus Barang Stok atau Jasa.")
    try: unit_id=int(d.get("unit_id"))
    except (TypeError,ValueError): raise ValueError("Satuan wajib dipilih.")
    cat=d.get("category_id")
    brand=d.get("brand_id")
    try: category_id=int(cat) if cat not in (None,"") else None
    except (TypeError,ValueError): raise ValueError("Kategori tidak valid.")
    try: brand_id=int(brand) if brand not in (None,"") else None
    except (TypeError,ValueError): raise ValueError("Merk tidak valid.")
    purchase=_decimal(d.get("purchase_price",0),"Harga beli")
    selling=_decimal(d.get("selling_price",0),"Harga jual")
    minimum=_decimal(d.get("minimum_stock",0),"Stok minimum")
    initial=_decimal(d.get("initial_stock",0),"Stok awal",allow_negative=False) if ptype=="STOCK" else 0
    opening_balance_date=_date_value(d.get("opening_balance_date") or utc_now()[:10],"Tanggal saldo awal")
    now=utc_now()
    with write_transaction() as tx:
        inventory_account_id=_product_account(tx,d.get("inventory_account_id"),"1200",("ASSET",),"Akun Persediaan")
        sales_account_id=_product_account(tx,d.get("sales_account_id"),"4000",("REVENUE",),"Akun Penjualan")
        cogs_account_id=_product_account(tx,d.get("cogs_account_id"),"5000",("HPP",),"Akun HPP")
        if not tx.execute("SELECT 1 FROM units WHERE id=? AND is_active=1",(unit_id,)).fetchone():
            raise ValueError("Satuan tidak ditemukan atau sudah nonaktif.")
        if category_id and not tx.execute("SELECT 1 FROM item_categories WHERE id=? AND is_active=1",(category_id,)).fetchone():
            raise ValueError("Kategori tidak ditemukan atau sudah nonaktif.")
        if brand_id and not tx.execute("SELECT 1 FROM brands WHERE id=? AND is_active=1",(brand_id,)).fetchone():
            raise ValueError("Merk tidak ditemukan atau sudah nonaktif.")
        cur=tx.execute("""INSERT INTO products(
            sku,barcode,name,category_id,brand_id,unit_id,product_type,
            purchase_price,selling_price,stock_qty,minimum_stock,notes,
            inventory_account_id,sales_account_id,cogs_account_id,opening_balance_date,opening_stock_qty,opening_warehouse_id,
            is_active,created_at,updated_at
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1,?,?)""",
        (sku,barcode,name,category_id,brand_id,unit_id,ptype,purchase,selling,0,minimum,
         str(d.get("notes","")).strip() or None,inventory_account_id,sales_account_id,
         cogs_account_id,opening_balance_date,str(initial),int(d.get("warehouse_id") or inventory_service.get_default_warehouse_id(tx)) if ptype=="STOCK" else None,now,now))
        pid=cur.lastrowid
        if initial:
            try: warehouse_id=int(d.get("warehouse_id") or inventory_service.get_default_warehouse_id(tx))
            except (TypeError,ValueError): raise ValueError("Gudang stok awal tidak valid.")
            inventory_service.stock_in(
                tx, product_id=pid, warehouse_id=warehouse_id, quantity=initial,
                unit_cost=purchase, reference_type="OPENING", reference_no="OPENING",
                reason="Stok awal", user_id=actor["id"], created_at=now
            )
            opening_value=(Decimal(str(initial))*Decimal(str(purchase))).quantize(Decimal("0.01"))
            if opening_value>0:
                accounting_service.post_opening_balance(tx,entity_type="PRODUCT",entity_id=pid,reference_no=sku,amount=opening_value,user_id=actor["id"],opening_date=opening_balance_date,inventory_account_id=inventory_account_id)
        _save_product_price_levels(tx,pid,d.get("price_levels"))
        multi_unit.save_product_units(tx,pid,unit_id,d.get("units"),purchase,selling)
        audit(actor["id"],"PRODUCT_CREATED","product",pid,
              {"sku":sku,"name":name,"brand_id":brand_id,"initial_stock":initial},ip,tx)
        return pid

def update_product(actor,product_id,d,ip):
    allowed_fields={"sku","barcode","name","category_id","brand_id","unit_id","product_type","purchase_price","selling_price","minimum_stock","notes","is_active","inventory_account_id","sales_account_id","cogs_account_id","opening_balance_date"}
    current=get_product(product_id)
    if not current: raise ValueError("Barang tidak ditemukan.")
    fields=[];params=[];details={}
    for key in allowed_fields:
        if key not in d:continue
        value=d[key]
        if key in ("sku","name"):
            value=str(value).strip().upper() if key=="sku" else str(value).strip()
            if not value:raise ValueError(f"{key} tidak boleh kosong.")
        elif key=="barcode":value=str(value or "").strip() or None
        elif key in ("purchase_price","selling_price","minimum_stock"):value=_decimal(value,key)
        elif key in ("category_id","brand_id","unit_id"):value=int(value) if value not in (None,"") else None
        elif key in ("inventory_account_id","sales_account_id","cogs_account_id"):value=int(value)
        elif key=="opening_balance_date":value=_date_value(value or current.get("opening_balance_date") or utc_now()[:10],"Tanggal saldo awal")
        elif key=="product_type":
            value=str(value).upper()
            if value not in ("STOCK","SERVICE"):raise ValueError("Jenis barang tidak valid.")
            if value=="SERVICE" and current["stock_qty"]!=0:raise ValueError("Barang dengan stok tidak dapat diubah menjadi jasa.")
        elif key=="is_active":value=1 if bool(value) else 0
        elif key=="notes":value=str(value).strip() or None
        fields.append(f"{key}=?");params.append(value);details[key]=value
    opening_requested=d.get("initial_stock",None)
    with write_transaction() as tx:
        final=dict(details);category_id=final.get("category_id",current.get("category_id"));brand_id=final.get("brand_id",current.get("brand_id"));unit_id=final.get("unit_id",current.get("unit_id"))
        if unit_id is None or not tx.execute("SELECT 1 FROM units WHERE id=? AND is_active=1",(unit_id,)).fetchone():raise ValueError("Satuan tidak ditemukan atau sudah nonaktif.")
        if category_id and not tx.execute("SELECT 1 FROM item_categories WHERE id=? AND is_active=1",(category_id,)).fetchone():raise ValueError("Kategori tidak ditemukan atau sudah nonaktif.")
        if brand_id and not tx.execute("SELECT 1 FROM brands WHERE id=? AND is_active=1",(brand_id,)).fetchone():raise ValueError("Merk tidak ditemukan atau sudah nonaktif.")
        for key,subtype,label in (("inventory_account_id","ASSET","Akun Persediaan"),("sales_account_id","REVENUE","Akun Penjualan"),("cogs_account_id","HPP","Akun HPP")):
            aid=final.get(key,current.get(key))
            if not tx.execute("SELECT 1 FROM chart_of_accounts WHERE id=? AND account_subtype=? AND is_active=1",(aid,subtype)).fetchone():raise ValueError(f"{label} tidak valid.")
        if fields:
            fields.append("updated_at=?");params.append(utc_now());params.append(product_id)
            tx.execute(f"UPDATE products SET {', '.join(fields)} WHERE id=?",params)
        # Saldo awal diedit IN-PLACE: selalu satu movement OPENING, tanpa OPENING_EDIT/koreksi baru.
        if opening_requested is not None and str(final.get("product_type",current.get("product_type")))=="STOCK":
            target=Decimal(str(_decimal(opening_requested,"Stok awal",allow_negative=False)))
            warehouse_id=int(d.get("warehouse_id") or current.get("opening_warehouse_id") or inventory_service.get_default_warehouse_id(tx))
            cost=Decimal(str(_decimal(final.get("purchase_price",current.get("purchase_price",0)),"Harga beli")))
            opening_date=str(final.get("opening_balance_date",current.get("opening_balance_date") or utc_now()[:10]))[:10]
            rows=tx.execute("SELECT * FROM inventory_transactions WHERE product_id=? AND reference_type IN ('OPENING','OPENING_EDIT') ORDER BY id",(product_id,)).fetchall()
            keep=rows[0] if rows else None
            for r in rows[1:]:tx.execute("DELETE FROM inventory_transactions WHERE id=?",(r['id'],))
            if target>0:
                created=(keep['created_at'] if keep else opening_date+'T00:00:00Z')
                if keep:
                    tx.execute("""UPDATE inventory_transactions SET product_id=?,warehouse_id=?,movement_type='IN',quantity_change=?,unit_cost=?,
                      reference_type='OPENING',reference_no='OPENING',reason='Stok awal',created_at=? WHERE id=?""",
                      (product_id,warehouse_id,str(target),str(cost),created,keep['id']))
                else:
                    tx.execute("""INSERT INTO inventory_transactions(product_id,warehouse_id,movement_type,quantity_change,quantity_before,quantity_after,unit_cost,average_cost_before,average_cost_after,reference_type,reference_no,reason,user_id,created_at)
                      VALUES(?,?,'IN',?,0,0,?,0,0,'OPENING','OPENING','Stok awal',?,?)""",(product_id,warehouse_id,str(target),str(cost),actor['id'],opening_date+'T00:00:00Z'))
            elif keep:
                tx.execute("DELETE FROM inventory_transactions WHERE id=?",(keep['id'],))
            tx.execute("UPDATE products SET opening_stock_qty=?,opening_warehouse_id=?,opening_balance_date=?,updated_at=? WHERE id=?",(str(target),warehouse_id,opening_date,utc_now(),product_id))
            replay=_replay_inventory_state_tx(tx)
            opening_value=(target*cost).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)
            _replace_opening_product_journal_tx(tx,product_id,current['sku'],opening_date,int(final.get('inventory_account_id',current.get('inventory_account_id'))),opening_value,actor['id'])
            # Perubahan opening cost dapat mengubah moving-average dan HPP Sales sesudahnya.
            synced=_sync_sale_cogs_journals_tx(tx)
            details['opening_stock_qty']=float(target);details['opening_value']=float(opening_value);details['sale_cogs_resynced']=synced
        if not fields and opening_requested is None:raise ValueError("Tidak ada perubahan.")
        target_warehouse=d.get("warehouse_id")
        old_warehouse=current.get("opening_warehouse_id")
        if target_warehouse not in (None,"") and str(final.get("product_type",current.get("product_type")))=="STOCK":
            target_warehouse=int(target_warehouse)
            if old_warehouse and int(old_warehouse)!=target_warehouse:
                inventory_service.reassign_warehouse_balance(tx,product_id=product_id,source_warehouse_id=int(old_warehouse),
                  target_warehouse_id=target_warehouse,reason="Perubahan gudang master barang",user_id=actor["id"],created_at=utc_now())
            tx.execute("UPDATE products SET opening_warehouse_id=?,updated_at=? WHERE id=?",(target_warehouse,utc_now(),product_id))
            details["opening_warehouse_id"]=target_warehouse
        _save_product_price_levels(tx,product_id,d.get("price_levels"))
        multi_unit.save_product_units(tx,product_id,unit_id,d.get("units",current.get("units")),final.get("purchase_price",current["purchase_price"]),final.get("selling_price",current["selling_price"]))
        audit(actor["id"],"PRODUCT_UPDATED","product",product_id,details,ip,tx)

def adjust_stock(actor,d,ip):
    try:
        product_id=int(d.get("product_id"))
        warehouse_id=int(d.get("warehouse_id"))
    except Exception:
        raise ValueError("Barang dan gudang wajib dipilih.")
    qty_change=_decimal(d.get("qty_change"),"Perubahan stok",allow_negative=True)
    if qty_change==0: raise ValueError("Perubahan stok tidak boleh nol.")
    reason=str(d.get("reason","")).strip()
    if not reason: raise ValueError("Alasan penyesuaian wajib diisi.")
    reference=str(d.get("reference_no","")).strip() or None
    try: adjustment_account_id=int(d.get("adjustment_account_id"))
    except Exception: raise ValueError("Akun penyesuaian wajib dipilih.")
    now=utc_now()
    with write_transaction() as tx:
        department_id,project_id=_transaction_dimensions(tx,d)
        adjustment_account=tx.execute("""SELECT id,account_subtype FROM chart_of_accounts
          WHERE id=? AND is_active=1""",(adjustment_account_id,)).fetchone()
        if not adjustment_account or adjustment_account["account_subtype"] in ("CASH_BANK","RECEIVABLE","PAYABLE"):
            raise ValueError("Akun penyesuaian tidak valid.")
        product=tx.execute("""SELECT id,sku,name,inventory_account_id FROM products
          WHERE id=? AND is_active=1""",(product_id,)).fetchone()
        if not product or not product["inventory_account_id"]:
            raise ValueError("Akun persediaan barang belum diatur.")
        result=inventory_service.adjust(
            tx, product_id=product_id, warehouse_id=warehouse_id,
            quantity_change=qty_change, reason=reason, reference_no=reference,
            user_id=actor["id"], created_at=now,department_id=department_id,project_id=project_id
        )
        unit_cost=Decimal(str(result.get("unit_cost") or result["average_cost_before"]))
        value=(Decimal(str(abs(qty_change)))*unit_cost).quantize(Decimal("0.01"))
        if value>0:
            if qty_change>0:
                lines=[{"account_id":product["inventory_account_id"],"debit":value},
                       {"account_id":adjustment_account_id,"credit":value}]
            else:
                lines=[{"account_id":adjustment_account_id,"debit":value},
                       {"account_id":product["inventory_account_id"],"credit":value}]
            accounting_service.post_journal(tx,journal_date=now[:10],
                description=f"Penyesuaian stok {product['sku']}",source_type="STOCK_ADJUSTMENT",
                source_id=result["transaction_id"],reference_no=reference,lines=lines,user_id=actor["id"],department_id=department_id,project_id=project_id)
        audit(actor["id"],"STOCK_ADJUSTED","product",product_id,{
            "warehouse_id":warehouse_id,"qty_change":qty_change,
            "qty_before":result["quantity_before"],"qty_after":result["quantity_after"],
            "reason":reason},ip,tx)
        return result

def stock_movements(product_id=None,limit=100):
    limit=max(1,min(int(limit),500)); params=[]; where=""
    if product_id is not None: where=" WHERE m.product_id=?"; params.append(int(product_id))
    params.append(limit); c=connect()
    try:
        return [dict(r) for r in c.execute("""SELECT m.*,p.sku,p.name product_name,u.username FROM stock_movements m
        JOIN products p ON p.id=m.product_id JOIN users u ON u.id=m.user_id"""+where+" ORDER BY m.id DESC LIMIT ?",params)]
    finally:c.close()

def inventory_summary():
    c=connect()
    try:
        x=c.execute("""SELECT (SELECT COUNT(*) FROM products WHERE product_type='STOCK') total_products,
        (SELECT COUNT(*) FROM products WHERE product_type='STOCK' AND is_active=1) active_products,
        (SELECT COUNT(*) FROM products p WHERE p.product_type='STOCK' AND p.is_active=1
          AND COALESCE((SELECT SUM(ib.quantity) FROM inventory_balances ib WHERE ib.product_id=p.id),0)<=p.minimum_stock) low_stock,
        COALESCE((SELECT SUM(COALESCE(ib.book_value,ib.quantity*ib.average_cost)) FROM inventory_balances ib JOIN products p ON p.id=ib.product_id
          WHERE p.product_type='STOCK' AND p.is_active=1),0) stock_value""").fetchone()
        return {"total_products":x["total_products"] or 0,"active_products":x["active_products"] or 0,"low_stock":x["low_stock"] or 0,"stock_value":float(x["stock_value"] or 0)}
    finally:c.close()


def _partner_dict(r):
    return {"id":r["id"],"partner_type":r["partner_type"],"code":r["code"],"name":r["name"],
    "tax_id":r["tax_id"],"phone":r["phone"],"email":r["email"],"address":r["address"],
    "city":r["city"],"credit_limit":float(r["credit_limit"] or 0),
    "payment_term_days":int(r["payment_term_days"] or 0),"price_level_id":r["price_level_id"] if "price_level_id" in r.keys() else None,"receivable_account_id":r["receivable_account_id"] if "receivable_account_id" in r.keys() else None,"payable_account_id":r["payable_account_id"] if "payable_account_id" in r.keys() else None,"opening_balance":float(r["opening_balance"] or 0),"opening_balance_date":r["opening_balance_date"],"partner_group":r["partner_group"],"contact_person":r["contact_person"],"notes":r["notes"],
    "is_active":bool(r["is_active"]),"created_at":r["created_at"],"updated_at":r["updated_at"]}

def list_partners(q="", partner_type=None, active=None):
    where=[]; params=[]
    if q:
        where.append("(code LIKE ? OR name LIKE ? OR COALESCE(phone,'') LIKE ? OR COALESCE(email,'') LIKE ? OR COALESCE(tax_id,'') LIKE ?)")
        term=f"%{q}%"; params += [term,term,term,term,term]
    if partner_type:
        pt=str(partner_type).upper()
        if pt not in ("CUSTOMER","SUPPLIER","BOTH"): raise ValueError("Jenis partner tidak valid.")
        if pt=="CUSTOMER": where.append("partner_type IN ('CUSTOMER','BOTH')")
        elif pt=="SUPPLIER": where.append("partner_type IN ('SUPPLIER','BOTH')")
        else: where.append("partner_type='BOTH'")
    if active is not None: where.append("is_active=?"); params.append(1 if active else 0)
    clause=(" WHERE "+" AND ".join(where)) if where else ""
    c=connect()
    try:
        return [_partner_dict(r) for r in c.execute("SELECT * FROM business_partners"+clause+" ORDER BY code COLLATE NOCASE, name COLLATE NOCASE",params)]
    finally:c.close()

def get_partner(partner_id):
    c=connect()
    try:
        r=c.execute("SELECT * FROM business_partners WHERE id=?",(partner_id,)).fetchone()
        return _partner_dict(r) if r else None
    finally:c.close()

def _partner_values(d, partial=False):
    out={}
    if not partial or "partner_type" in d:
        pt=str(d.get("partner_type","")).strip().upper()
        if pt not in ("CUSTOMER","SUPPLIER","BOTH"): raise ValueError("Jenis partner harus CUSTOMER, SUPPLIER, atau BOTH.")
        out["partner_type"]=pt
    for key,label in (("code","Kode"),("name","Nama")):
        if not partial or key in d:
            value=str(d.get(key,"")).strip()
            if key=="code": value=value.upper()
            if not value: raise ValueError(f"{label} wajib diisi.")
            out[key]=value
    for key in ("tax_id","phone","email","address","city","partner_group","contact_person","notes"):
        if not partial or key in d: out[key]=str(d.get(key,"")).strip() or None
    if not partial or "credit_limit" in d: out["credit_limit"]=_decimal(d.get("credit_limit",0),"Batas kredit")
    if not partial or "opening_balance" in d: out["opening_balance"]=_decimal(d.get("opening_balance",0),"Saldo awal")
    if not partial or "opening_balance_date" in d:
        out["opening_balance_date"]=_date_value(d.get("opening_balance_date") or utc_now()[:10],"Tanggal saldo awal")
    if not partial or "payment_term_days" in d:
        try: days=int(d.get("payment_term_days",0) or 0)
        except: raise ValueError("Termin pembayaran harus berupa bilangan bulat.")
        if days<0 or days>3650: raise ValueError("Termin pembayaran tidak valid.")
        out["payment_term_days"]=days
    if not partial or "price_level_id" in d:
        value=d.get("price_level_id");out["price_level_id"]=int(value) if value not in (None,"") else None
    for key,subtype,label in (("receivable_account_id","RECEIVABLE","Akun Piutang"),("payable_account_id","PAYABLE","Akun Hutang")):
        if not partial or key in d:
            value=d.get(key);out[key]=int(value) if value not in (None,"") else None
    if "is_active" in d: out["is_active"]=1 if bool(d["is_active"]) else 0
    return out

def create_partner(actor,d,ip):
    v=_partner_values(d);now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("""INSERT INTO business_partners(partner_type,code,name,tax_id,phone,email,address,city,
        credit_limit,payment_term_days,opening_balance,opening_balance_date,partner_group,contact_person,price_level_id,receivable_account_id,payable_account_id,notes,is_active,created_at,updated_at)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1,?,?)""",
        (v["partner_type"],v["code"],v["name"],v["tax_id"],v["phone"],v["email"],v["address"],v["city"],
         v["credit_limit"],v["payment_term_days"],v.get("opening_balance",0),v.get("opening_balance_date"),v.get("partner_group"),v.get("contact_person"),v.get("price_level_id"),v.get("receivable_account_id"),v.get("payable_account_id"),v["notes"],now,now))
        pid=cur.lastrowid;opening=_decimal(v.get("opening_balance",0),"Saldo awal")
        if v.get("receivable_account_id"):
            a=tx.execute("SELECT account_subtype FROM chart_of_accounts WHERE id=? AND is_active=1",(v["receivable_account_id"],)).fetchone()
            if not a or a["account_subtype"]!="RECEIVABLE": raise ValueError("Akun Piutang pelanggan harus bertipe Piutang.")
        if v.get("payable_account_id"):
            a=tx.execute("SELECT account_subtype FROM chart_of_accounts WHERE id=? AND is_active=1",(v["payable_account_id"],)).fetchone()
            if not a or a["account_subtype"]!="PAYABLE": raise ValueError("Akun Hutang pemasok harus bertipe Hutang.")
        if opening>0:
            if v["partner_type"]=="CUSTOMER":
                accounting_service.post_opening_balance(tx,entity_type="CUSTOMER",entity_id=pid,reference_no=v["code"],amount=opening,user_id=actor["id"],opening_date=v.get("opening_balance_date"),partner_account_id=v.get("receivable_account_id"))
            elif v["partner_type"]=="SUPPLIER":
                accounting_service.post_opening_balance(tx,entity_type="SUPPLIER",entity_id=pid,reference_no=v["code"],amount=opening,user_id=actor["id"],opening_date=v.get("opening_balance_date"),partner_account_id=v.get("payable_account_id"))
            else:raise ValueError("Saldo awal untuk Keduanya tidak didukung. Buat pelanggan dan pemasok terpisah.")
        audit(actor["id"],"PARTNER_CREATED","partner",pid,{"code":v["code"],"name":v["name"],"partner_type":v["partner_type"],"opening_balance":float(opening)},ip,tx)
        return pid

def update_partner(actor,partner_id,d,ip):
    if not get_partner(partner_id): raise ValueError("Pelanggan/supplier tidak ditemukan.")
    v=_partner_values(d,True)
    if not v: raise ValueError("Tidak ada perubahan.")
    fields=[]; params=[]
    for k,val in v.items(): fields.append(f"{k}=?"); params.append(val)
    fields.append("updated_at=?"); params.append(utc_now()); params.append(partner_id)
    with write_transaction() as tx:
        old=tx.execute("SELECT * FROM business_partners WHERE id=?",(partner_id,)).fetchone()
        for key,subtype,label in (("receivable_account_id","RECEIVABLE","Akun Piutang"),("payable_account_id","PAYABLE","Akun Hutang")):
            aid=v.get(key) if key in v else old[key] if key in old.keys() else None
            if aid:
                a=tx.execute("SELECT account_subtype FROM chart_of_accounts WHERE id=? AND is_active=1",(aid,)).fetchone()
                if not a or a["account_subtype"]!=subtype: raise ValueError(label+" partner tidak sesuai tipe akun.")
        tx.execute(f"UPDATE business_partners SET {', '.join(fields)} WHERE id=?",params)
        if "receivable_account_id" in v and v.get("receivable_account_id"):
            tx.execute("""UPDATE journal_lines SET account_id=? WHERE partner_id=? AND account_id IN
              (SELECT id FROM chart_of_accounts WHERE account_subtype='RECEIVABLE')""",(v["receivable_account_id"],partner_id))
        if "payable_account_id" in v and v.get("payable_account_id"):
            tx.execute("""UPDATE journal_lines SET account_id=? WHERE partner_id=? AND account_id IN
              (SELECT id FROM chart_of_accounts WHERE account_subtype='PAYABLE')""",(v["payable_account_id"],partner_id))
        if "opening_balance" in v or "opening_balance_date" in v or "partner_type" in v:
            ptype=v.get("partner_type",old["partner_type"]);opening=_decimal(v.get("opening_balance",old["opening_balance"]),"Saldo awal")
            tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE source_type IN ('OPENING_CUSTOMER','OPENING_SUPPLIER') AND CAST(source_id AS TEXT)=? AND status='POSTED'",(utc_now(),str(partner_id)))
            if opening>0 and ptype in ("CUSTOMER","SUPPLIER"):
                accounting_service.post_opening_balance(tx,entity_type=ptype,entity_id=partner_id,reference_no=v.get("code",old["code"]),amount=opening,user_id=actor["id"],opening_date=v.get("opening_balance_date",old["opening_balance_date"]),partner_account_id=(v.get("receivable_account_id",old["receivable_account_id"]) if ptype=="CUSTOMER" else v.get("payable_account_id",old["payable_account_id"])))
        audit(actor["id"],"PARTNER_UPDATED","partner",partner_id,v,ip,tx)

def partner_summary():
    c=connect()
    try:
        r=c.execute("""SELECT COUNT(*) total,
        SUM(CASE WHEN is_active=1 AND partner_type IN ('CUSTOMER','BOTH') THEN 1 ELSE 0 END) customers,
        SUM(CASE WHEN is_active=1 AND partner_type IN ('SUPPLIER','BOTH') THEN 1 ELSE 0 END) suppliers,
        SUM(CASE WHEN is_active=0 THEN 1 ELSE 0 END) inactive
        FROM business_partners""").fetchone()
        return {"total":r["total"] or 0,"customers":r["customers"] or 0,"suppliers":r["suppliers"] or 0,"inactive":r["inactive"] or 0}
    finally:c.close()


# ---------------- PENJUALAN ----------------
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import date

MONEY = Decimal("0.01")
QTY = Decimal("0.0001")

def _normalize_decimal_text(value):
    """Normalize legacy Indonesian/English formatted numeric text.

    Examples: 10.000.000 -> 10000000, 10.000.000,50 -> 10000000.50,
    10000000.00 remains 10000000.00.
    """
    if value in (None, ""):
        return "0"
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    text=str(value).strip().replace("Rp", "").replace("rp", "").replace(" ", "")
    if not text:
        return "0"
    neg=text.startswith("-")
    if neg: text=text[1:]
    if "," in text and "." in text:
        # Separator terakhir dianggap desimal; separator lain adalah ribuan.
        if text.rfind(",") > text.rfind("."):
            text=text.replace(".", "").replace(",", ".")
        else:
            text=text.replace(",", "")
    elif "," in text:
        parts=text.split(",")
        if len(parts)==2 and len(parts[1]) in (1,2):
            text=parts[0].replace(".", "")+"."+parts[1]
        else:
            text="".join(parts)
    elif "." in text:
        parts=text.split(".")
        if len(parts)>2 or (len(parts)==2 and len(parts[1])==3):
            text="".join(parts)
    return ("-" if neg else "")+text

def _d(value, label, minimum=None):
    try:
        result = Decimal(_normalize_decimal_text(value))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{label} tidak valid.")
    if not result.is_finite():
        raise ValueError(f"{label} tidak valid.")
    if minimum is not None and result < Decimal(str(minimum)):
        raise ValueError(f"{label} tidak boleh kurang dari {minimum}.")
    return result

def _money(value): return _d(value, "Nilai").quantize(MONEY, rounding=ROUND_HALF_UP)
def _qty(value): return _d(value, "Jumlah", "0.0001").quantize(QTY, rounding=ROUND_HALF_UP)

def _next_invoice(tx, sale_date):
    try:
        period = str(sale_date)[:7].replace("-", "")
        if len(period) != 6: raise ValueError
    except Exception:
        raise ValueError("Tanggal penjualan tidak valid.")
    key = "SALES-" + period
    now = utc_now()
    tx.execute("""INSERT INTO document_sequences(sequence_key,current_value,updated_at)
                  VALUES(?,0,?) ON CONFLICT(sequence_key) DO NOTHING""", (key, now))
    tx.execute("UPDATE document_sequences SET current_value=current_value+1,updated_at=? WHERE sequence_key=?", (now,key))
    number = tx.execute("SELECT current_value FROM document_sequences WHERE sequence_key=?", (key,)).fetchone()["current_value"]
    return f"PJ-{period}-{number:06d}"

def _suggest_invoice(tx,sale_date):
    period=str(sale_date)[:7].replace('-','');prefix=f"PJ-{period}-"
    row=tx.execute("SELECT invoice_no FROM sales WHERE invoice_no LIKE ? ORDER BY invoice_no DESC LIMIT 1",(prefix+'%',)).fetchone()
    n=int(row['invoice_no'].rsplit('-',1)[-1])+1 if row and row['invoice_no'] else 1
    return f"{prefix}{n:06d}"

def _sale_dict(r):
    keys=set(r.keys())
    return {"id":r["id"],"invoice_no":r["invoice_no"],"delivery_no":r["delivery_no"] if "delivery_no" in keys else None,"sale_date":r["sale_date"],
    "customer_id":r["customer_id"],"customer_code":r["customer_code"],"customer_name":r["customer_name"],
    "salesperson_id":r["salesperson_id"] if "salesperson_id" in keys else None,
    "salesperson_name":r["salesperson_name"] if "salesperson_name" in keys else None,
    "warehouse_id":r["warehouse_id"] if "warehouse_id" in keys else None,
    "warehouse_name":r["warehouse_name"] if "warehouse_name" in keys else None,
    "payment_type":r["payment_type"],"payment_method":r["payment_method"] if "payment_method" in keys else r["payment_type"],
    "cash_account_id":r["cash_account_id"] if "cash_account_id" in keys else None,
    "cash_account_name":r["cash_account_name"] if "cash_account_name" in keys else None,
    "due_date":r["due_date"] if "due_date" in keys else None,
    "tax_percent":float(r["tax_percent"] if "tax_percent" in keys else 0),
    "tax_amount":float(r["tax_amount"] if "tax_amount" in keys else 0),
    "subtotal":float(r["subtotal"]),"discount_amount":float(r["discount_amount"]),
    "total_amount":float(r["total_amount"]),"paid_amount":float(r["paid_amount"]),"balance_due":float(r["balance_due"]),
    "notes":r["notes"],"status":r["status"],"created_by":r["created_by"],"created_at":r["created_at"]}

def _discount_amount(raw, gross, index):
    mode=str(raw.get("discount_mode","AMOUNT")).upper()
    value=_money(raw.get("discount_value",raw.get("discount_amount",0)))
    if mode=="PERCENT":
        if value>100: raise ValueError(f"Diskon persen baris {index} tidak boleh lebih dari 100%.")
        amount=(gross*value/Decimal("100")).quantize(MONEY,rounding=ROUND_HALF_UP)
        return value,amount
    if value>gross: raise ValueError(f"Diskon item baris {index} melebihi nilai item.")
    pct=(value/gross*Decimal("100")).quantize(MONEY,rounding=ROUND_HALF_UP) if gross else Decimal("0")
    return pct,value

def _next_delivery_no(tx,doc_date):
    prefix=f"SJ-{str(doc_date)[:7].replace('-','')}"
    row=tx.execute("SELECT delivery_no FROM sales WHERE delivery_no LIKE ? ORDER BY id DESC LIMIT 1",(prefix+'-%',)).fetchone()
    seq=int(row["delivery_no"].rsplit('-',1)[-1])+1 if row and row["delivery_no"] else 1
    return f"{prefix}-{seq:06d}"

def document_number_suggestions(doc_date=None):
    doc_date=str(doc_date or date.today().isoformat())[:10]
    c=connect()
    try:
        return {
          "invoice_no":_suggest_invoice(c,doc_date),
          "delivery_no":_next_delivery_no(c,doc_date),
          "purchase_no":purchase_service.next_purchase_no(c,doc_date),
          "goods_receipt_no":purchase_service.next_goods_receipt_no(c,doc_date),
          "journal_no":accounting_service.next_journal_no(c,doc_date),
        }
    finally:c.close()

def _transaction_dimensions(tx,data):
    department_id=data.get("department_id") or None
    project_id=data.get("project_id") or None
    if department_id is not None:
        try:department_id=int(department_id)
        except Exception:raise ValueError("Departemen tidak valid.")
        if not tx.execute("SELECT 1 FROM departments WHERE id=? AND is_active=1",(department_id,)).fetchone():
            raise ValueError("Departemen tidak ditemukan atau nonaktif.")
    if project_id is not None:
        try:project_id=int(project_id)
        except Exception:raise ValueError("Proyek tidak valid.")
        if not tx.execute("SELECT 1 FROM projects WHERE id=? AND is_active=1",(project_id,)).fetchone():
            raise ValueError("Proyek tidak ditemukan atau nonaktif.")
    return department_id,project_id


def _save_sale_invoice_materials(tx,sale_id,materials):
    tx.execute("DELETE FROM sale_invoice_materials WHERE sale_id=?",(int(sale_id),))
    if not isinstance(materials,list):return
    for idx,m in enumerate(materials):
        if not isinstance(m,dict):continue
        name=str(m.get("material_name") or m.get("name") or "").strip()
        if not name:continue
        qty=float(m.get("qty") or 0);unit_cost=float(m.get("unit_cost") or 0)
        total=float(m.get("total_cost") or (qty*unit_cost))
        tx.execute("""INSERT INTO sale_invoice_materials(sale_id,project_material_issue_id,product_id,sku,material_name,qty,unit_code,unit_cost,total_cost,issue_date,issue_no,sort_order)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",(int(sale_id),m.get("project_material_issue_id") or None,m.get("product_id") or None,
          str(m.get("sku") or ""),name,qty,str(m.get("unit_code") or ""),unit_cost,total,
          str(m.get("issue_date") or "")[:10] or None,str(m.get("issue_no") or ""),idx))
def _sale_invoice_materials(c,sale_id):
    return [dict(x) for x in c.execute("SELECT * FROM sale_invoice_materials WHERE sale_id=? ORDER BY sort_order,id",(int(sale_id),)).fetchall()]

def _create_sale_tx(tx, actor, d, ip):
    items=d.get("items")
    if not isinstance(items,list) or not items: raise ValueError("Minimal satu item penjualan wajib diisi.")
    sale_date=str(d.get("sale_date") or date.today().isoformat())[:10]
    method=str(d.get("payment_method",d.get("payment_type","CASH"))).upper()
    if method not in ("CASH","TRANSFER","CREDIT"): raise ValueError("Metode pembayaran harus Tunai, Transfer, atau Kredit.")
    payment_type="CREDIT" if method=="CREDIT" else "CASH"
    customer_id=d.get("customer_id") or None; salesperson_id=d.get("salesperson_id") or None
    sp=None
    notes=str(d.get("notes","")).strip() or None
    header_mode=str(d.get("discount_mode","AMOUNT")).upper(); header_value=_money(d.get("discount_value",d.get("discount_amount",0)))
    tax_percent=_money(d.get("tax_percent",0)); paid_requested=_money(d.get("paid_amount",0))
    if tax_percent>100: raise ValueError("PPN tidak boleh lebih dari 100%.")
    now=utc_now()
    licensing.enforce_transaction_limit(tx)
    department_id,project_id=_transaction_dimensions(tx,d)
    warehouse_id=int(d.get("warehouse_id") or inventory_service.get_default_warehouse_id(tx))
    customer=None
    if customer_id:
        customer=tx.execute("""SELECT * FROM business_partners WHERE id=? AND is_active=1
            AND partner_type IN ('CUSTOMER','BOTH')""",(int(customer_id),)).fetchone()
        if not customer: raise ValueError("Pelanggan tidak ditemukan atau tidak aktif.")
    if salesperson_id:
        sp=tx.execute("SELECT id,commission_percent FROM salespersons WHERE id=? AND is_active=1",(int(salesperson_id),)).fetchone()
        if not sp: raise ValueError("Salesman tidak ditemukan atau tidak aktif.")
    allow_negative_stock=allowed(actor,"inventory.negative_stock")
    normalized=[]; subtotal=Decimal("0")
    for index,raw in enumerate(items,1):
        try: product_id=int(raw.get("product_id"))
        except Exception: raise ValueError(f"Produk pada baris {index} tidak valid.")
        product=tx.execute("SELECT * FROM products WHERE id=? AND is_active=1",(product_id,)).fetchone()
        if not product: raise ValueError(f"Produk pada baris {index} tidak ditemukan atau tidak aktif.")
        description=str(raw.get("description") or "").strip()
        if len(description)>255: raise ValueError(f"Deskripsi invoice pada baris {index} maksimal 255 karakter.")
        invoice_description=description or product["name"]
        unit=multi_unit.resolve(tx,product,raw.get("unit_id"),raw.get("qty"),raw.get("unit_price"),"SELL"); entered_qty=unit["entered_qty"]; qty=unit["base_qty"]; price=_money(unit["price"]); gross=(entered_qty*price).quantize(MONEY,rounding=ROUND_HALF_UP)
        discount_pct,line_discount=_discount_amount(raw,gross,index); line_total=gross-line_discount
        if product["product_type"]=="STOCK" and not allow_negative_stock:
            before,_avg=inventory_service.get_balance(tx,warehouse_id,product["id"])
            if before-qty<0: raise ValueError(f"Stok {product['sku']} - {product['name']} tidak cukup. Tersedia {before}.")
        normalized.append((product,invoice_description,qty,price,discount_pct,line_discount,line_total,unit)); subtotal+=line_total
    subtotal=subtotal.quantize(MONEY,rounding=ROUND_HALF_UP)
    if header_mode=="PERCENT":
        if header_value>100: raise ValueError("Diskon transaksi persen tidak boleh lebih dari 100%.")
        header_discount=(subtotal*header_value/Decimal("100")).quantize(MONEY,rounding=ROUND_HALF_UP)
    else: header_discount=header_value
    if header_discount>subtotal: raise ValueError("Diskon transaksi melebihi subtotal.")
    taxable=subtotal-header_discount; tax_amount=(taxable*tax_percent/Decimal("100")).quantize(MONEY,rounding=ROUND_HALF_UP)
    total=(taxable+tax_amount).quantize(MONEY,rounding=ROUND_HALF_UP)
    cash_account_id=d.get("cash_account_id") or None
    if method in ("CASH","TRANSFER"):
        if not cash_account_id: raise ValueError("Akun kas/bank wajib dipilih.")
        if paid_requested<total: raise ValueError("Pembayaran kurang dari total transaksi.")
        paid=total; balance=Decimal("0"); due_date=None
    else:
        if paid_requested>total: raise ValueError("Pembayaran awal kredit melebihi total transaksi.")
        paid=paid_requested; balance=(total-paid).quantize(MONEY,rounding=ROUND_HALF_UP)
        if not customer_id: raise ValueError("Penjualan kredit wajib memilih pelanggan.")
        term=int(d.get("payment_term_days") or customer["payment_term_days"] or 0)
        due_date=str(d.get("due_date") or (date.fromisoformat(sale_date)+timedelta(days=term)).isoformat())[:10]
        if paid>0 and not cash_account_id: raise ValueError("Akun kas/bank wajib dipilih untuk pembayaran awal.")
        credit_limit=_money(customer["credit_limit"]); outstanding=_money(tx.execute("SELECT COALESCE(SUM(balance_due),0) value FROM sales WHERE customer_id=? AND status='POSTED'",(int(customer_id),)).fetchone()["value"])
        if credit_limit>0 and outstanding+balance>credit_limit: raise ValueError("Transaksi melebihi batas kredit pelanggan.")
    invoice=str(d.get("invoice_no") or "").strip() or _next_invoice(tx,sale_date)
    sales_order_id=d.get("sales_order_id") or None
    from . import order_dp
    order_dp.validate_order_fulfillment(tx,"sales",sales_order_id,[(p[0]["id"],p[2]) for p in normalized])
    delivery_no=str(d.get("delivery_no") or "").strip() or _next_delivery_no(tx,sale_date)
    if tx.execute("SELECT 1 FROM sales WHERE invoice_no=? COLLATE NOCASE",(invoice,)).fetchone():
        raise ValueError("Nomor Invoice sudah digunakan.")
    if tx.execute("SELECT 1 FROM sales WHERE delivery_no=? COLLATE NOCASE",(delivery_no,)).fetchone():
        raise ValueError("Nomor Surat Jalan sudah digunakan.")
    cur=tx.execute("""INSERT INTO sales(invoice_no,delivery_no,sale_date,customer_id,salesperson_id,warehouse_id,department_id,project_id,payment_type,payment_method,cash_account_id,due_date,tax_percent,tax_amount,subtotal,discount_amount,total_amount,paid_amount,balance_due,notes,status,user_id,created_at,updated_at)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'POSTED',?,?,?)""",
        (invoice,delivery_no,sale_date,customer_id,salesperson_id,warehouse_id,department_id,project_id,payment_type,method,cash_account_id,due_date,str(tax_percent),str(tax_amount),str(subtotal),str(header_discount),str(total),str(paid),str(balance),notes,actor["id"],now,now))
    sale_id=cur.lastrowid
    if sales_order_id: tx.execute("UPDATE sales SET sales_order_id=? WHERE id=?",(int(sales_order_id),sale_id))
    cogs_total=Decimal("0.00")
    revenue_groups={};cogs_groups={}
    line_discount_total=sum((x[5] for x in normalized),Decimal("0.00")).quantize(MONEY,rounding=ROUND_HALF_UP)
    total_sales_discount=(line_discount_total+header_discount).quantize(MONEY,rounding=ROUND_HALF_UP)
    commission_percent=_money(sp["commission_percent"] if sp else 0)
    commission_amount=(taxable*commission_percent/Decimal("100")).quantize(MONEY,rounding=ROUND_HALF_UP) if commission_percent>0 else Decimal("0.00")
    for line_index,(product,invoice_description,qty,price,discount_pct,line_discount,line_total,unit) in enumerate(normalized):
        gross_revenue=(line_total+line_discount).quantize(MONEY,rounding=ROUND_HALF_UP)
        sales_aid=product["sales_account_id"] or accounting_service.account_id(tx,"4000")
        revenue_groups[sales_aid]=revenue_groups.get(sales_aid,Decimal("0.00"))+gross_revenue
        tx.execute("""INSERT INTO sales_items(sale_id,product_id,sku,product_name,product_type,qty,entered_qty,unit_id,unit_code,conversion_ratio,unit_price,discount_percent,discount_amount,line_total,purchase_price_snapshot) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (sale_id,product["id"],product["sku"],invoice_description,product["product_type"],str(qty),str(unit["entered_qty"]),unit["unit_id"],unit["unit_code"],str(unit["ratio"]),str(price),str(discount_pct),str(line_discount),str(line_total),product["purchase_price"]))
        if product["product_type"]=="STOCK":
            stock_result=inventory_service.stock_out(tx,product_id=product["id"],warehouse_id=warehouse_id,quantity=qty,reference_type="SALE",reference_no=invoice,reason="Penjualan",user_id=actor["id"],created_at=now,allow_negative=allow_negative_stock)
            item_cogs=Decimal(str(stock_result.get("cost_of_goods",0)))
            actual_unit_cogs=(item_cogs/qty).quantize(MONEY,rounding=ROUND_HALF_UP) if qty else Decimal("0.00")
            tx.execute("UPDATE sales_items SET purchase_price_snapshot=? WHERE sale_id=? AND product_id=? AND id=(SELECT MAX(id) FROM sales_items WHERE sale_id=? AND product_id=?)",(str(actual_unit_cogs),sale_id,product["id"],sale_id,product["id"]))
            cogs_total += item_cogs
            key=(product["cogs_account_id"] or accounting_service.account_id(tx,"5000"),
                 product["inventory_account_id"] or accounting_service.account_id(tx,"1200"))
            cogs_groups[key]=cogs_groups.get(key,Decimal("0.00"))+item_cogs
    if paid>0:
        cash_service.post(tx,account_id=int(cash_account_id),transaction_date=sale_date,transaction_type="IN",amount=paid,description=f"Penerimaan penjualan {invoice}",reference_no=invoice,user_id=actor["id"],department_id=department_id,project_id=project_id)
    accounting_service.post_sale(tx,sale_id=sale_id,invoice_no=invoice,sale_date=sale_date,
        taxable_amount=taxable,tax_amount=tax_amount,paid_amount=paid,balance_due=balance,
        cash_account_id=cash_account_id,cogs_amount=cogs_total,customer_id=customer_id,
        user_id=actor["id"],department_id=department_id,project_id=project_id,
        revenue_lines=[{"account_id":aid,"amount":amount} for aid,amount in revenue_groups.items() if amount>0],
        cogs_lines=[{"cogs_account_id":key[0],"inventory_account_id":key[1],"amount":amount}
                    for key,amount in cogs_groups.items() if amount>0],
        discount_amount=total_sales_discount,commission_amount=commission_amount,receivable_account_id=(customer["receivable_account_id"] if customer else None))
    order_dp.close_order_if_fulfilled(tx,"sales",sales_order_id)
    # Presentation-only material snapshot: no stock movement and no journal posting.
    _save_sale_invoice_materials(tx,sale_id,d.get("invoice_materials"))
    audit(actor["id"],"SALE_CREATED","sale",sale_id,{"invoice_no":invoice,"total":float(total),"method":method,"tax":float(tax_amount)},ip,tx)
    return {"id":sale_id,"invoice_no":invoice,"delivery_no":delivery_no,"subtotal":float(subtotal),"discount_amount":float(header_discount),"tax_amount":float(tax_amount),"total_amount":float(total),"paid_amount":float(paid),"change_amount":float(max(Decimal("0"),paid_requested-total)),"balance_due":float(balance),"due_date":due_date}

def create_sale(actor, d, ip):
    with write_transaction() as tx:
        return _create_sale_tx(tx,actor,d,ip)
def list_sales(q="", date_from=None, date_to=None, limit=200):
    limit=max(1,min(int(limit),500)); where=["s.status='POSTED'"]; params=[]
    if q:
        like=f"%{q.strip()}%"; where.append("(s.invoice_no LIKE ? OR COALESCE(bp.name,'') LIKE ?)"); params += [like,like]
    if date_from: where.append("s.sale_date>=?"); params.append(str(date_from)[:10])
    if date_to: where.append("s.sale_date<=?"); params.append(str(date_to)[:10])
    sql="""SELECT s.*,bp.code customer_code,bp.name customer_name,bp.address customer_address,bp.city customer_city,bp.phone customer_phone,bp.email customer_email,u.username created_by,
        sp.name salesperson_name,w.name warehouse_name,ca.name cash_account_name
        FROM sales s LEFT JOIN business_partners bp ON bp.id=s.customer_id
        LEFT JOIN salespersons sp ON sp.id=s.salesperson_id LEFT JOIN warehouses w ON w.id=s.warehouse_id
        LEFT JOIN cash_accounts ca ON ca.id=s.cash_account_id JOIN users u ON u.id=s.user_id"""
    if where: sql += " WHERE "+" AND ".join(where)
    sql += " ORDER BY s.sale_date DESC,s.id DESC LIMIT ?"; params.append(limit)
    c=connect()
    try:return [_sale_dict(r) for r in c.execute(sql,params).fetchall()]
    finally:c.close()

def get_sale(sale_id):
    c=connect()
    try:
        r=c.execute("""SELECT s.*,bp.code customer_code,bp.name customer_name,bp.address customer_address,bp.city customer_city,bp.phone customer_phone,bp.email customer_email,u.username created_by,
            sp.name salesperson_name,w.name warehouse_name,ca.name cash_account_name
            FROM sales s LEFT JOIN business_partners bp ON bp.id=s.customer_id
            LEFT JOIN salespersons sp ON sp.id=s.salesperson_id LEFT JOIN warehouses w ON w.id=s.warehouse_id
            LEFT JOIN cash_accounts ca ON ca.id=s.cash_account_id JOIN users u ON u.id=s.user_id
            WHERE s.id=?""",(sale_id,)).fetchone()
        if not r:return None
        result=_sale_dict(r)
        result["items"]=[{"id":x["id"],"product_id":x["product_id"],"sku":x["sku"],"product_name":x["product_name"],
            "product_type":x["product_type"],"qty":float(x["entered_qty"] if x["entered_qty"] is not None else x["qty"]),"base_qty":float(x["qty"]),"unit_id":x["unit_id"],"unit_code":x["unit_code"],"conversion_ratio":float(x["conversion_ratio"] or 1),"unit_price":float(x["unit_price"]),
            "discount_percent":float(x["discount_percent"]),"discount_amount":float(x["discount_amount"]),"line_total":float(x["line_total"])}
            for x in c.execute("SELECT * FROM sales_items WHERE sale_id=? ORDER BY id",(sale_id,)).fetchall()]
        result["invoice_materials"]=_sale_invoice_materials(c,sale_id)
        return result
    finally:c.close()

def sales_summary(date_from=None,date_to=None):
    where=["status='POSTED'"];params=[]
    if date_from:where.append("sale_date>=?");params.append(str(date_from)[:10])
    if date_to:where.append("sale_date<=?");params.append(str(date_to)[:10])
    c=connect()
    try:
        r=c.execute(f"""SELECT COUNT(*) transaction_count,COALESCE(SUM(total_amount),0) sales_total,
            COALESCE(SUM(paid_amount),0) paid_total,COALESCE(SUM(balance_due),0) receivable_total
            FROM sales WHERE {' AND '.join(where)}""",params).fetchone()
        return {"transaction_count":r["transaction_count"],"sales_total":float(r["sales_total"]),
                "paid_total":float(r["paid_total"]),"receivable_total":float(r["receivable_total"])}
    finally:c.close()


def _warehouse_dict(r):
    return {"id":r["id"],"code":r["code"],"name":r["name"],"address":r["address"],
            "is_default":bool(r["is_default"]),"is_active":bool(r["is_active"]),
            "created_at":r["created_at"],"updated_at":r["updated_at"]}

def list_warehouses(active_only=False):
    c=connect()
    try:
        sql="SELECT * FROM warehouses"
        if active_only: sql+=" WHERE is_active=1"
        sql+=" ORDER BY is_default DESC,name"
        return [_warehouse_dict(r) for r in c.execute(sql).fetchall()]
    finally:c.close()

def create_warehouse(actor,d,ip):
    code=str(d.get("code","")).strip().upper()
    name=str(d.get("name","")).strip()
    address=str(d.get("address","")).strip() or None
    is_default=1 if bool(d.get("is_default")) else 0
    if not code or not name: raise ValueError("Kode dan nama gudang wajib diisi.")
    now=utc_now()
    with write_transaction() as tx:
        if is_default: tx.execute("UPDATE warehouses SET is_default=0")
        cur=tx.execute("""INSERT INTO warehouses(code,name,address,is_default,is_active,created_at,updated_at)
                          VALUES(?,?,?,?,1,?,?)""",(code,name,address,is_default,now,now))
        audit(actor["id"],"WAREHOUSE_CREATED","warehouse",cur.lastrowid,
              {"code":code,"name":name,"is_default":bool(is_default)},ip,tx)
        return cur.lastrowid


def update_warehouse(actor,warehouse_id,d,ip):
    warehouse_id=int(warehouse_id)
    code=str(d.get("code","")).strip().upper();name=str(d.get("name","")).strip()
    address=str(d.get("address","")).strip() or None
    is_default=1 if bool(d.get("is_default")) else 0
    is_active=1 if bool(d.get("is_active",True)) else 0
    if not code or not name:raise ValueError("Kode dan nama gudang wajib diisi.")
    with write_transaction() as tx:
        old=tx.execute("SELECT * FROM warehouses WHERE id=?",(warehouse_id,)).fetchone()
        if not old:raise ValueError("Gudang tidak ditemukan.")
        if old["is_default"] and not is_default:
            other=tx.execute("SELECT 1 FROM warehouses WHERE id<>? AND is_default=1 AND is_active=1",(warehouse_id,)).fetchone()
            if not other:raise ValueError("Pilih gudang default lain sebelum melepas status default gudang ini.")
        if is_default:
            tx.execute("UPDATE warehouses SET is_default=0 WHERE id<>?",(warehouse_id,))
            is_active=1
        tx.execute("UPDATE warehouses SET code=?,name=?,address=?,is_default=?,is_active=?,updated_at=? WHERE id=?",
                   (code,name,address,is_default,is_active,utc_now(),warehouse_id))
        audit(actor["id"],"WAREHOUSE_UPDATED","warehouse",warehouse_id,{"code":code,"name":name,"is_default":bool(is_default),"is_active":bool(is_active)},ip,tx)
    return next(x for x in list_warehouses(False) if x["id"]==warehouse_id)


def delete_warehouse(actor,warehouse_id,ip):
    warehouse_id=int(warehouse_id)
    with write_transaction() as tx:
        row=tx.execute("SELECT * FROM warehouses WHERE id=?",(warehouse_id,)).fetchone()
        if not row:raise ValueError("Gudang tidak ditemukan.")
        if row["is_default"]:raise ValueError("Gudang default tidak dapat dihapus. Jadikan gudang lain sebagai default terlebih dahulu.")
        refs=0
        checks=[("inventory_balances","warehouse_id"),("inventory_transactions","warehouse_id"),("sales","warehouse_id"),("purchases","warehouse_id"),("purchase_returns","warehouse_id"),("sales_returns","warehouse_id")]
        for table,col in checks:
            try:
                if tx.execute(f"SELECT 1 FROM {table} WHERE {col}=? LIMIT 1",(warehouse_id,)).fetchone():refs+=1
            except Exception:pass
        if refs:
            tx.execute("UPDATE warehouses SET is_active=0,updated_at=? WHERE id=?",(utc_now(),warehouse_id))
            action="WAREHOUSE_DEACTIVATED"
        else:
            tx.execute("DELETE FROM warehouses WHERE id=?",(warehouse_id,))
            action="WAREHOUSE_DELETED"
        audit(actor["id"],action,"warehouse",warehouse_id,{"code":row["code"],"name":row["name"]},ip,tx)
    return {"id":warehouse_id,"deactivated":bool(refs)}

def inventory_balances(warehouse_id=None,q="",low_stock=False):
    where=["p.product_type='STOCK'"];params=[]
    if warehouse_id not in (None,""):
        where.append("w.id=?");params.append(int(warehouse_id))
    if q:
        where.append("(p.sku LIKE ? OR p.name LIKE ?)")
        term=f"%{q}%";params += [term,term]
    if low_stock:
        where.append("b.quantity<=p.minimum_stock")
    c=connect()
    try:
        rows=c.execute("""SELECT w.id warehouse_id,w.code warehouse_code,w.name warehouse_name,
          p.id product_id,p.sku,p.name product_name,u.code unit_code,
          COALESCE(b.quantity,0) quantity,COALESCE(b.average_cost,p.purchase_price,0) average_cost,
          p.minimum_stock,
          COALESCE(b.quantity,0)*COALESCE(b.average_cost,p.purchase_price,0) stock_value
          FROM warehouses w CROSS JOIN products p
          JOIN units u ON u.id=p.unit_id
          LEFT JOIN inventory_balances b ON b.warehouse_id=w.id AND b.product_id=p.id
          WHERE """+" AND ".join(where)+" ORDER BY w.name,p.name",params).fetchall()
        return [{**dict(r),"quantity":float(r["quantity"]),"average_cost":float(r["average_cost"]),
                 "minimum_stock":float(r["minimum_stock"]),"stock_value":float(r["stock_value"])}
                for r in rows]
    finally:c.close()

def inventory_card(product_id=None,warehouse_id=None,limit=200):
    limit=max(1,min(int(limit),1000));where=[];params=[]
    if product_id not in (None,""):where.append("t.product_id=?");params.append(int(product_id))
    if warehouse_id not in (None,""):where.append("t.warehouse_id=?");params.append(int(warehouse_id))
    sql="""SELECT t.*,p.sku,p.name product_name,w.code warehouse_code,w.name warehouse_name,u.username
           FROM inventory_transactions t
           JOIN products p ON p.id=t.product_id
           JOIN warehouses w ON w.id=t.warehouse_id
           JOIN users u ON u.id=t.user_id"""
    if where:sql+=" WHERE "+" AND ".join(where)
    sql+=" ORDER BY t.id DESC LIMIT ?";params.append(limit)
    c=connect()
    try:
        return [{**dict(r),"quantity_change":float(r["quantity_change"]),
                 "quantity_before":float(r["quantity_before"]),"quantity_after":float(r["quantity_after"]),
                 "unit_cost":float(r["unit_cost"]),"average_cost_before":float(r["average_cost_before"]),
                 "average_cost_after":float(r["average_cost_after"])}
                for r in c.execute(sql,params).fetchall()]
    finally:c.close()

def transfer_stock(actor,d,ip):
    try:
        product_id=int(d.get("product_id"))
        source_id=int(d.get("source_warehouse_id"))
        target_id=int(d.get("target_warehouse_id"))
    except Exception:
        raise ValueError("Barang, gudang asal, dan gudang tujuan wajib dipilih.")
    amount=_decimal(d.get("quantity"),"Jumlah transfer")
    reference=str(d.get("reference_no","")).strip() or None
    reason=str(d.get("reason","")).strip() or "Transfer gudang"
    now=utc_now()
    with write_transaction() as tx:
        licensing.enforce_transaction_capacity(tx,1)
        result=inventory_service.transfer(
            tx,product_id=product_id,source_warehouse_id=source_id,
            target_warehouse_id=target_id,quantity=amount,reference_no=reference,
            reason=reason,user_id=actor["id"],created_at=now
        )
        audit(actor["id"],"STOCK_TRANSFERRED","product",product_id,{
            "source_warehouse_id":source_id,"target_warehouse_id":target_id,
            "quantity":amount,"reference_no":reference},ip,tx)
        usage_id=result.get("out_transaction_id") or result.get("id") or f"{product_id}-{now}"
        licensing.record_transaction_usage(tx,event_key=f"STOCK-TRANSFER-{usage_id}",event_type="STOCK_TRANSFER",
          reference_no=reference,source_table="inventory_transactions",source_id=usage_id,units=1)
        return result

def inventory_engine_summary():
    c=connect()
    try:
        r=c.execute("""SELECT
          (SELECT COUNT(*) FROM warehouses WHERE is_active=1) warehouse_count,
          (SELECT COUNT(*) FROM inventory_balances b JOIN products p ON p.id=b.product_id
             WHERE p.is_active=1 AND b.quantity<=p.minimum_stock) low_stock_lines,
          (SELECT COALESCE(SUM(quantity*average_cost),0) FROM inventory_balances) stock_value,
          (SELECT COALESCE(SUM(quantity),0) FROM inventory_balances) total_quantity
        """).fetchone()
        return {"warehouse_count":r["warehouse_count"],"low_stock_lines":r["low_stock_lines"],
                "stock_value":float(r["stock_value"]),"total_quantity":float(r["total_quantity"])}
    finally:c.close()


def create_purchase(actor, data, ip):
    with write_transaction() as tx:
        return purchase_service.create_purchase(
            tx,
            actor=actor,
            data=data,
            client_ip=ip,
            audit_callback=audit,
        )

def _purchase_dict(row):
    return {
        "id": row["id"],
        "purchase_no": row["purchase_no"],
        "supplier_invoice_no": row["supplier_invoice_no"] if "supplier_invoice_no" in row.keys() else None,
        "goods_receipt_no": row["goods_receipt_no"] if "goods_receipt_no" in row.keys() else None,
        "purchase_date": row["purchase_date"],
        "supplier_id": row["supplier_id"],
        "supplier_code": row["supplier_code"],
        "supplier_name": row["supplier_name"],
        "warehouse_id": row["warehouse_id"],
        "warehouse_code": row["warehouse_code"],
        "warehouse_name": row["warehouse_name"],
        "department_id": row["department_id"] if "department_id" in row.keys() else None,
        "project_id": row["project_id"] if "project_id" in row.keys() else None,
        "payment_type": row["payment_type"],
        "payment_method": row["payment_method"] if "payment_method" in row.keys() else row["payment_type"],
        "cash_account_id": row["cash_account_id"] if "cash_account_id" in row.keys() else None,
        "cash_account_name": row["cash_account_name"] if "cash_account_name" in row.keys() else None,
        "due_date": row["due_date"] if "due_date" in row.keys() else None,
        "tax_percent": float(row["tax_percent"] if "tax_percent" in row.keys() else 0),
        "tax_amount": float(row["tax_amount"] if "tax_amount" in row.keys() else 0),
        "subtotal": float(row["subtotal"]),
        "discount_amount": float(row["discount_amount"]),
        "total_amount": float(row["total_amount"]),
        "paid_amount": float(row["paid_amount"]),
        "balance_due": float(row["balance_due"]),
        "notes": row["notes"],
        "status": row["status"],
        "created_by": row["created_by"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }

def list_purchases(q="", date_from=None, date_to=None, limit=200):
    limit = max(1, min(int(limit), 500))
    where = ["p.status='POSTED'"]
    params = []
    if q:
        like = f"%{q.strip()}%"
        where.append("(p.purchase_no LIKE ? OR bp.name LIKE ?)")
        params += [like, like]
    if date_from:
        where.append("p.purchase_date>=?")
        params.append(str(date_from)[:10])
    if date_to:
        where.append("p.purchase_date<=?")
        params.append(str(date_to)[:10])

    sql = """SELECT p.*,bp.code supplier_code,bp.name supplier_name,
                    w.code warehouse_code,w.name warehouse_name,ca.name cash_account_name,
                    u.username created_by
             FROM purchases p
             JOIN business_partners bp ON bp.id=p.supplier_id
             JOIN warehouses w ON w.id=p.warehouse_id
             LEFT JOIN cash_accounts ca ON ca.id=p.cash_account_id
             JOIN users u ON u.id=p.user_id"""
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY p.purchase_date DESC,p.id DESC LIMIT ?"
    params.append(limit)

    conn = connect()
    try:
        return [_purchase_dict(row) for row in conn.execute(sql, params).fetchall()]
    finally:
        conn.close()

def get_purchase(purchase_id):
    conn = connect()
    try:
        row = conn.execute(
            """SELECT p.*,bp.code supplier_code,bp.name supplier_name,bp.address supplier_address,bp.city supplier_city,bp.phone supplier_phone,bp.email supplier_email,
                      w.code warehouse_code,w.name warehouse_name,ca.name cash_account_name,
                      u.username created_by
               FROM purchases p
               JOIN business_partners bp ON bp.id=p.supplier_id
               JOIN warehouses w ON w.id=p.warehouse_id
               LEFT JOIN cash_accounts ca ON ca.id=p.cash_account_id
               JOIN users u ON u.id=p.user_id
               WHERE p.id=?""",
            (purchase_id,),
        ).fetchone()
        if not row:
            return None
        result = _purchase_dict(row)
        result["items"] = [
            {
                "id": item["id"],
                "product_id": item["product_id"],
                "sku": item["sku"],
                "product_name": item["product_name"],
                "product_type": item["product_type"],
                "qty": float(item["entered_qty"] if item["entered_qty"] is not None else item["qty"]),
                "base_qty": float(item["qty"]),
                "unit_id": item["unit_id"],"unit_code":item["unit_code"],"conversion_ratio":float(item["conversion_ratio"] or 1),
                "unit_cost": float(item["unit_cost"]),
                "discount_amount": float(item["discount_amount"]),
                "line_total": float(item["line_total"]),
            }
            for item in conn.execute(
                "SELECT * FROM purchase_items WHERE purchase_id=? ORDER BY id",
                (purchase_id,),
            ).fetchall()
        ]
        return result
    finally:
        conn.close()

def purchases_summary(date_from=None, date_to=None):
    where = ["status='POSTED'"]
    params = []
    if date_from:
        where.append("purchase_date>=?")
        params.append(str(date_from)[:10])
    if date_to:
        where.append("purchase_date<=?")
        params.append(str(date_to)[:10])

    conn = connect()
    try:
        row = conn.execute(
            f"""SELECT COUNT(*) transaction_count,
                       COALESCE(SUM(total_amount),0) purchase_total,
                       COALESCE(SUM(paid_amount),0) paid_total,
                       COALESCE(SUM(balance_due),0) payable_total
                FROM purchases
                WHERE {' AND '.join(where)}""",
            params,
        ).fetchone()
        return {
            "transaction_count": row["transaction_count"],
            "purchase_total": float(row["purchase_total"]),
            "paid_total": float(row["paid_total"]),
            "payable_total": float(row["payable_total"]),
        }
    finally:
        conn.close()


def cash_accounts(active_only=False):
    conn = connect()
    try:
        sql = "SELECT * FROM cash_accounts"
        if active_only:
            sql += " WHERE is_active=1"
        sql += " ORDER BY account_type,name"
        return [
            {
                **dict(row),
                "opening_balance": float(row["opening_balance"]),
                "current_balance": float(row["current_balance"]),
                "is_active": bool(row["is_active"]),
            }
            for row in conn.execute(sql).fetchall()
        ]
    finally:
        conn.close()


def _cash_default_coa(tx, account_type):
    code="1010" if str(account_type).upper()=="BANK" else "1000"
    row=tx.execute("SELECT id FROM chart_of_accounts WHERE code=? AND is_active=1",(code,)).fetchone()
    if not row: raise ValueError(f"COA kas/bank default {code} tidak ditemukan.")
    return int(row["id"])

def _post_cash_opening(tx, actor, cash_account_id, opening, transaction_date=None):
    opening=Decimal(str(opening or 0))
    if opening==0:return
    cash_row=tx.execute("SELECT code,name,coa_account_id,account_type FROM cash_accounts WHERE id=?",(cash_account_id,)).fetchone()
    cash_coa=int(cash_row["coa_account_id"] or _cash_default_coa(tx,cash_row["account_type"]))
    opening_coa=accounting_service.account_id(tx,"3200")
    tdate=(transaction_date or utc_now()[:10])[:10]
    if opening>0:
        result=cash_service.post(tx,account_id=cash_account_id,transaction_date=tdate,transaction_type="IN",amount=opening,description="Saldo awal",reference_no="OPENING",user_id=actor["id"],allow_negative=True)
        lines=[{"account_id":cash_coa,"debit":opening},{"account_id":opening_coa,"credit":opening}]
    else:
        amt=abs(opening)
        result=cash_service.post(tx,account_id=cash_account_id,transaction_date=tdate,transaction_type="OUT",amount=amt,description="Saldo awal negatif",reference_no="OPENING",user_id=actor["id"],allow_negative=True)
        lines=[{"account_id":opening_coa,"debit":amt},{"account_id":cash_coa,"credit":amt}]
    accounting_service.post_journal(tx,journal_date=tdate,description=f"Saldo awal {cash_row['name']}",source_type="CASH_OPENING",source_id=cash_account_id,reference_no=f"OPENING-{cash_row['code']}",lines=lines,user_id=actor["id"])

def create_cash_account(actor, data, ip):
    code=str(data.get("code","")).strip().upper();name=str(data.get("name","")).strip();account_type=str(data.get("account_type","CASH")).upper()
    if not code or not name:raise ValueError("Kode dan nama akun wajib diisi.")
    if account_type not in ("CASH","BANK"):raise ValueError("Jenis akun tidak valid.")
    opening=cash_service.money(data.get("opening_balance",0),"Saldo awal");now=utc_now()
    with write_transaction() as tx:
        coa_account_id=int(data.get('coa_account_id') or _cash_default_coa(tx,account_type))
        cursor=tx.execute("""INSERT INTO cash_accounts(code,name,account_type,bank_name,account_number,opening_balance,current_balance,is_active,created_at,updated_at,coa_account_id) VALUES(?,?,?,?,?,?,0,1,?,?,?)""",(code,name,account_type,str(data.get("bank_name","")).strip() or None,str(data.get("account_number","")).strip() or None,str(opening),now,now,coa_account_id))
        _post_cash_opening(tx,actor,cursor.lastrowid,opening,data.get('opening_balance_date'))
        audit(actor["id"],"CASH_ACCOUNT_CREATED","cash_account",cursor.lastrowid,{"code":code,"name":name,"account_type":account_type,"opening_balance":float(opening)},ip,tx)
        return cursor.lastrowid

def update_cash_account(actor, account_id, data, ip):
    account_id=int(account_id);now=utc_now()
    with write_transaction() as tx:
        old=tx.execute("SELECT * FROM cash_accounts WHERE id=?",(account_id,)).fetchone()
        if not old:raise ValueError("Akun kas/bank tidak ditemukan.")
        code=str(data.get('code') or old['code']).strip().upper();name=str(data.get('name') or old['name']).strip();atype=str(data.get('account_type') or old['account_type']).upper()
        opening=cash_service.money(data.get('opening_balance',old['opening_balance']),"Saldo awal")
        coa_account_id=int(data.get('coa_account_id') or old['coa_account_id'] or _cash_default_coa(tx,atype))
        tx.execute("UPDATE cash_accounts SET code=?,name=?,account_type=?,bank_name=?,account_number=?,opening_balance=?,coa_account_id=?,updated_at=? WHERE id=?",(code,name,atype,str(data.get('bank_name') or '').strip() or None,str(data.get('account_number') or '').strip() or None,str(opening),coa_account_id,now,account_id))
        # Repair khusus impor lama: bila saldo awal ada tetapi jurnal opening belum pernah dibuat, buat sekarang tanpa menggandakan mutasi yang sudah ada.
        j=tx.execute("SELECT id FROM journal_entries WHERE source_type='CASH_OPENING' AND source_id=? AND status='POSTED'",(str(account_id),)).fetchone()
        if opening!=0 and not j:
            existing=tx.execute("SELECT id FROM cash_transactions WHERE account_id=? AND reference_no='OPENING' ORDER BY id LIMIT 1",(account_id,)).fetchone()
            if existing:
                opening_coa=accounting_service.account_id(tx,'3200');amt=abs(opening)
                lines=([{'account_id':coa_account_id,'debit':amt},{'account_id':opening_coa,'credit':amt}] if opening>0 else [{'account_id':opening_coa,'debit':amt},{'account_id':coa_account_id,'credit':amt}])
                accounting_service.post_journal(tx,journal_date=now[:10],description=f"Saldo awal {name}",source_type='CASH_OPENING',source_id=account_id,reference_no=f"OPENING-{code}",lines=lines,user_id=actor['id'])
            else:_post_cash_opening(tx,actor,account_id,opening,data.get('opening_balance_date'))
        audit(actor['id'],'CASH_ACCOUNT_UPDATED','cash_account',account_id,{'code':code,'opening_balance':float(opening)},ip,tx)
    return account_id

def create_cash_transaction(actor, data, ip):
    transaction_date = str(data.get("transaction_date") or "")[:10]
    if len(transaction_date) != 10:
        raise ValueError("Tanggal wajib diisi.")
    try:
        account_id = int(data.get("account_id"))
    except Exception:
        raise ValueError("Akun kas/bank wajib dipilih.")
    transaction_type = str(data.get("transaction_type", "IN")).upper()
    if transaction_type not in ("IN","OUT"): raise ValueError("Jenis transaksi kas tidak valid.")
    try: counter_account_id=int(data.get("counter_account_id"))
    except Exception: raise ValueError("Akun lawan transaksi wajib dipilih.")
    description = str(data.get("description", "")).strip()
    if not description:
        raise ValueError("Keterangan wajib diisi.")
    with write_transaction() as tx:
        department_id,project_id=_transaction_dimensions(tx,data)
        counter=tx.execute("""SELECT id,account_subtype FROM chart_of_accounts
          WHERE id=? AND is_active=1""",(counter_account_id,)).fetchone()
        if not counter or counter["account_subtype"] in ("CASH_BANK","RECEIVABLE","PAYABLE"):
            raise ValueError("Akun lawan kas/bank tidak valid.")
        result = cash_service.post(
            tx,
            account_id=account_id,
            transaction_date=transaction_date,
            transaction_type=transaction_type,
            amount=data.get("amount"),
            description=description,
            reference_no=str(data.get("reference_no", "")).strip() or None,
            user_id=actor["id"],
            allow_negative=True,department_id=department_id,project_id=project_id,
        )
        cash_row=tx.execute("SELECT coa_account_id FROM cash_accounts WHERE id=?",(account_id,)).fetchone()
        cash_coa=cash_row["coa_account_id"] if cash_row and cash_row["coa_account_id"] else accounting_service.account_id(tx,"1000")
        amount=Decimal(str(result["amount"]))
        lines=([{"account_id":cash_coa,"debit":amount},{"account_id":counter_account_id,"credit":amount}]
               if transaction_type=="IN" else
               [{"account_id":counter_account_id,"debit":amount},{"account_id":cash_coa,"credit":amount}])
        accounting_service.post_journal(tx,journal_date=transaction_date,
          description=description,source_type="CASH_TRANSACTION",source_id=result["id"],
          reference_no=str(data.get("reference_no","")).strip() or result["transaction_no"],
          lines=lines,user_id=actor["id"],department_id=department_id,project_id=project_id)
        audit(
            actor["id"], "CASH_TRANSACTION_CREATED", "cash_transaction",
            result["id"],
            {
                "transaction_no": result["transaction_no"],
                "account_id": account_id,
                "transaction_type": transaction_type,
                "amount": result["amount"],
            },
            ip, tx,
        )
        return result


def create_cash_transfer(actor, data, ip):
    transaction_date = str(data.get("transaction_date") or "")[:10]
    if len(transaction_date) != 10:
        raise ValueError("Tanggal wajib diisi.")
    try:
        source_id = int(data.get("source_account_id"))
        target_id = int(data.get("target_account_id"))
    except Exception:
        raise ValueError("Akun asal dan tujuan wajib dipilih.")
    description = str(data.get("description", "")).strip() or "Transfer antar akun"
    with write_transaction() as tx:
        result = cash_service.transfer(
            tx,
            source_account_id=source_id,
            target_account_id=target_id,
            transaction_date=transaction_date,
            amount=data.get("amount"),
            description=description,
            reference_no=str(data.get("reference_no", "")).strip() or None,
            user_id=actor["id"],
        )
        # Transfer antar Kas/Bank wajib mempunyai pasangan jurnal GL agar Buku Besar,
        # Neraca Saldo dan Neraca memakai sumber yang sama dengan modul Kas & Bank.
        src=tx.execute("SELECT coa_account_id,account_type FROM cash_accounts WHERE id=?",(source_id,)).fetchone()
        dst=tx.execute("SELECT coa_account_id,account_type FROM cash_accounts WHERE id=?",(target_id,)).fetchone()
        src_coa=int(src["coa_account_id"] or _cash_default_coa(tx,src["account_type"]))
        dst_coa=int(dst["coa_account_id"] or _cash_default_coa(tx,dst["account_type"]))
        amt=Decimal(str(result["source"]["amount"]))
        accounting_service.post_journal(tx,journal_date=transaction_date,description=description,
            source_type="CASH_TRANSFER",source_id=result["transfer_no"],
            reference_no=str(data.get("reference_no", "")).strip() or result["transfer_no"],
            lines=[{"account_id":dst_coa,"debit":amt},{"account_id":src_coa,"credit":amt}],
            user_id=actor["id"])
        audit(
            actor["id"], "CASH_TRANSFER_CREATED", "cash_transfer",
            result["transfer_no"],
            {
                "source_account_id": source_id,
                "target_account_id": target_id,
                "amount": result["source"]["amount"],
            },
            ip, tx,
        )
        return result


def list_cash_transactions(account_id=None, q="", date_from=None, date_to=None, limit=300):
    where = []
    params = []
    if account_id not in (None, ""):
        where.append("t.account_id=?")
        params.append(int(account_id))
    if q:
        term = f"%{q.strip()}%"
        where.append(
            "(t.transaction_no LIKE ? OR t.description LIKE ? "
            "OR COALESCE(t.reference_no,'') LIKE ?)"
        )
        params.extend([term, term, term])
    if date_from:
        where.append("t.transaction_date>=?")
        params.append(str(date_from)[:10])
    if date_to:
        where.append("t.transaction_date<=?")
        params.append(str(date_to)[:10])
    sql = """SELECT t.*,a.code account_code,a.name account_name,u.username
             FROM cash_transactions t
             JOIN cash_accounts a ON a.id=t.account_id
             JOIN users u ON u.id=t.user_id"""
    # Daftar operasional hanya menampilkan transaksi aktif. Baris asli yang sudah
    # dibatalkan dan baris pembalikan internal tetap disimpan untuk audit.
    where.insert(0, "NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='cash' AND v.transaction_key=CAST(t.id AS TEXT))")
    where.insert(1, "NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='cash_reversal' AND v.transaction_key=CAST(t.id AS TEXT))")
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY t.transaction_date DESC,t.id DESC LIMIT ?"
    params.append(max(1, min(int(limit), 1000)))
    conn = connect()
    try:
        return [
            {
                **dict(row),
                "amount": float(row["amount"]),
                "balance_before": float(row["balance_before"]),
                "balance_after": float(row["balance_after"]),
            }
            for row in conn.execute(sql, params).fetchall()
        ]
    finally:
        conn.close()


def cash_summary():
    conn = connect()
    try:
        row = conn.execute(
            """SELECT COUNT(*) account_count,
                      COALESCE(SUM(current_balance),0) total_balance,
                      COALESCE(SUM(CASE WHEN account_type='CASH'
                                        THEN current_balance ELSE 0 END),0) cash_balance,
                      COALESCE(SUM(CASE WHEN account_type='BANK'
                                        THEN current_balance ELSE 0 END),0) bank_balance
               FROM cash_accounts WHERE is_active=1"""
        ).fetchone()
        return {
            "account_count": row["account_count"],
            "total_balance": float(row["total_balance"]),
            "cash_balance": float(row["cash_balance"]),
            "bank_balance": float(row["bank_balance"]),
        }
    finally:
        conn.close()


def company_profile():
    c=connect()
    try:
        r=c.execute("SELECT * FROM company_profile WHERE id=1").fetchone()
        if not r:return {}
        out=dict(r);out.pop("logo_content",None)
        return out
    finally:c.close()

def company_logo_bytes():
    if not IS_POSTGRES:
        p=company_profile().get("logo_path")
        if not p or not Path(p).exists():return None,None
        mime=mimetypes.guess_type(str(p))[0] or "image/png"
        return Path(p).read_bytes(),mime
    c=connect()
    try:
        r=c.execute("SELECT logo_content,logo_mime FROM company_profile WHERE id=1").fetchone()
        if not r or r["logo_content"] is None:return None,None
        return bytes(r["logo_content"]),str(r["logo_mime"] or "image/png")
    finally:c.close()

def update_company_profile(actor,d,ip):
    fields=["company_name","address","city","phone","email","tax_id","website"]
    values={k:str(d.get(k,"")).strip() or None for k in fields}
    if not values["company_name"]: raise ValueError("Nama perusahaan wajib diisi.")
    logo_path=None;logo_binary=None;logo_mime=None
    logo_base64=str(d.get("logo_base64") or "").strip()
    if logo_base64:
        try:
            if "," in logo_base64:logo_base64=logo_base64.split(",",1)[1]
            binary=base64.b64decode(logo_base64,validate=True)
        except Exception as exc:
            raise ValueError("File logo tidak valid.") from exc
        if len(binary)>2_000_000:raise ValueError("Ukuran logo maksimal 2 MB.")
        filename=str(d.get("logo_filename") or "company-logo.png").lower()
        ext=".jpg" if filename.endswith((".jpg",".jpeg")) else ".png"
        logo_mime="image/jpeg" if ext==".jpg" else "image/png"
        if IS_POSTGRES:
            logo_binary=binary
        else:
            logo_dir=DATA_DIR/"company_assets";logo_dir.mkdir(parents=True,exist_ok=True)
            target=logo_dir/("company-logo"+ext)
            target.write_bytes(binary);logo_path=str(target)
    now=utc_now()
    with write_transaction() as tx:
        current=tx.execute("SELECT logo_path,logo_content,logo_mime FROM company_profile WHERE id=1").fetchone()
        final_logo=logo_path or (current["logo_path"] if current else None)
        final_content=logo_binary if logo_binary is not None else (current["logo_content"] if current else None)
        final_mime=logo_mime or (current["logo_mime"] if current else None)
        tx.execute("""UPDATE company_profile SET company_name=?,address=?,city=?,phone=?,
          email=?,tax_id=?,website=?,logo_path=?,logo_content=?,logo_mime=?,updated_at=? WHERE id=1""",
          (values["company_name"],values["address"],values["city"],values["phone"],
           values["email"],values["tax_id"],values["website"],final_logo,final_content,final_mime,now))
        audit(actor["id"],"COMPANY_PROFILE_UPDATED","company","1",
              {**values,"logo_updated":bool(logo_path or logo_binary)},ip,tx)

def list_brands(active_only=False):
    c=connect()
    try:
        sql="SELECT * FROM brands"+(" WHERE is_active=1" if active_only else "")+" ORDER BY code COLLATE NOCASE, name COLLATE NOCASE"
        return [dict(r) for r in c.execute(sql).fetchall()]
    finally:c.close()

def create_brand(actor,d,ip):
    code=str(d.get("code","")).strip().upper();name=str(d.get("name","")).strip()
    if not code or not name:raise ValueError("Kode dan nama merk wajib diisi.")
    now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("INSERT INTO brands(code,name,is_active,created_at,updated_at) VALUES(?,?,1,?,?)",(code,name,now,now))
        audit(actor["id"],"BRAND_CREATED","brand",cur.lastrowid,{"code":code,"name":name},ip,tx);return cur.lastrowid

def list_salespersons(active_only=False):
    c=connect()
    try:
        sql="SELECT * FROM salespersons"+(" WHERE is_active=1" if active_only else "")+" ORDER BY code COLLATE NOCASE, name COLLATE NOCASE"
        rows=c.execute(sql).fetchall();return [{**dict(r),"commission_percent":float(r["commission_percent"])} for r in rows]
    finally:c.close()

def create_salesperson(actor,d,ip):
    code=str(d.get("code","")).strip().upper();name=str(d.get("name","")).strip()
    if not code or not name:raise ValueError("Kode dan nama salesman wajib diisi.")
    commission=_decimal(d.get("commission_percent",0),"Komisi")
    now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("""INSERT INTO salespersons(code,name,phone,email,commission_percent,is_active,created_at,updated_at)
                          VALUES(?,?,?,?,?,1,?,?)""",(code,name,str(d.get("phone","")).strip() or None,str(d.get("email","")).strip() or None,commission,now,now))
        audit(actor["id"],"SALESPERSON_CREATED","salesperson",cur.lastrowid,{"code":code,"name":name},ip,tx);return cur.lastrowid



def _update_simple_master(actor, table, item_id, values, entity, ip):
    if not values: raise ValueError("Tidak ada perubahan.")
    fields=[]; params=[]
    for k,v in values.items(): fields.append(f"{k}=?"); params.append(v)
    fields.append("updated_at=?"); params.append(utc_now()); params.append(item_id)
    with write_transaction() as tx:
        row=tx.execute(f"SELECT id FROM {table} WHERE id=?",(item_id,)).fetchone()
        if not row: raise ValueError("Data tidak ditemukan.")
        tx.execute(f"UPDATE {table} SET {', '.join(fields)} WHERE id=?",params)
        audit(actor["id"],entity.upper()+"_UPDATED",entity,item_id,values,ip,tx)

def _deactivate_simple_master(actor, table, item_id, entity, ip):
    # Tombol Hapus benar-benar menghapus master yang belum pernah dipakai.
    # Bila sudah direferensikan transaksi/master lain, soft-delete agar histori tetap utuh.
    from .database import DB_INTEGRITY_ERRORS
    with write_transaction() as tx:
        row=tx.execute(f"SELECT id FROM {table} WHERE id=?",(item_id,)).fetchone()
        if not row: raise ValueError("Data tidak ditemukan.")
        try:
            tx.execute(f"DELETE FROM {table} WHERE id=?",(item_id,))
            action=entity.upper()+"_DELETED"
        except DB_INTEGRITY_ERRORS:
            tx.execute(f"UPDATE {table} SET is_active=0,updated_at=? WHERE id=?",(utc_now(),item_id))
            action=entity.upper()+"_DEACTIVATED"
        audit(actor["id"],action,entity,item_id,{},ip,tx)

def update_brand(actor,item_id,d,ip):
    vals={}
    if "code" in d: vals["code"]=str(d.get("code","")).strip().upper()
    if "name" in d: vals["name"]=str(d.get("name","")).strip()
    if "is_active" in d: vals["is_active"]=1 if bool(d.get("is_active")) else 0
    if not vals.get("code",True) or not vals.get("name",True): raise ValueError("Kode dan nama merk wajib diisi.")
    _update_simple_master(actor,"brands",item_id,vals,"brand",ip)
def delete_brand(actor,item_id,ip): _deactivate_simple_master(actor,"brands",item_id,"brand",ip)

def update_salesperson(actor,item_id,d,ip):
    vals={}
    for k in ("code","name","phone","email"):
        if k in d: vals[k]=str(d.get(k,"")).strip() or None
    if "code" in vals: vals["code"]=(vals["code"] or "").upper()
    if "commission_percent" in d: vals["commission_percent"]=_decimal(d.get("commission_percent",0),"Komisi")
    if "is_active" in d: vals["is_active"]=1 if bool(d.get("is_active")) else 0
    _update_simple_master(actor,"salespersons",item_id,vals,"salesperson",ip)
def delete_salesperson(actor,item_id,ip): _deactivate_simple_master(actor,"salespersons",item_id,"salesperson",ip)

def update_category(actor,item_id,d,ip):
    vals={}
    if "code" in d: vals["code"]=str(d.get("code","")).strip().upper()
    if "name" in d: vals["name"]=str(d.get("name","")).strip()
    if "is_active" in d: vals["is_active"]=1 if bool(d.get("is_active")) else 0
    _update_simple_master(actor,"item_categories",item_id,vals,"category",ip)
def delete_category(actor,item_id,ip): _deactivate_simple_master(actor,"item_categories",item_id,"category",ip)

def update_unit(actor,item_id,d,ip):
    vals={}
    if "code" in d: vals["code"]=str(d.get("code","")).strip().upper()
    if "name" in d: vals["name"]=str(d.get("name","")).strip()
    if "decimals" in d:
        dec=int(d.get("decimals",0));
        if dec<0 or dec>4: raise ValueError("Desimal satuan harus 0 sampai 4.")
        vals["decimals"]=dec
    if "is_active" in d: vals["is_active"]=1 if bool(d.get("is_active")) else 0
    _update_simple_master(actor,"units",item_id,vals,"unit",ip)
def delete_unit(actor,item_id,ip): _deactivate_simple_master(actor,"units",item_id,"unit",ip)

def delete_partner(actor,item_id,ip): _deactivate_simple_master(actor,"business_partners",item_id,"partner",ip)
def delete_product(actor,item_id,ip):
    item_id=int(item_id)
    with write_transaction() as tx:
        p=tx.execute("SELECT * FROM products WHERE id=?",(item_id,)).fetchone()
        if not p:raise ValueError("Barang tidak ditemukan.")
        refs=[]
        for table,col in (("sales_items","product_id"),("purchase_items","product_id"),("sales_return_items","product_id"),("purchase_return_items","product_id")):
            try:
                if tx.execute(f"SELECT 1 FROM {table} WHERE {col}=? LIMIT 1",(item_id,)).fetchone():refs.append(table)
            except Exception:pass
        try:
            non_open=tx.execute("SELECT 1 FROM inventory_transactions WHERE product_id=? AND COALESCE(reference_type,'') NOT IN ('OPENING','OPENING_EDIT') LIMIT 1",(item_id,)).fetchone()
            if non_open:refs.append("inventory_transactions")
        except Exception:pass
        if refs:
            tx.execute("UPDATE products SET is_active=0,updated_at=? WHERE id=?",(utc_now(),item_id))
            audit(actor["id"],"PRODUCT_DEACTIVATED","product",item_id,{"sku":p["sku"],"reason":"has_transactions"},ip,tx)
            return {"status":"deactivated","message":"Barang memiliki transaksi dan dinonaktifkan agar histori tetap aman."}
        # No business transaction: remove opening-only accounting/inventory data, then master.
        try:
            jids=[x["id"] for x in tx.execute("SELECT id FROM journal_entries WHERE source_type IN ('OPENING_PRODUCT','PRODUCT_OPENING','OPENING_PRODUCT_EDIT','PRODUCT_OPENING_EDIT') AND source_id=?",(item_id,)).fetchall()]
            for jid in jids:tx.execute("DELETE FROM journal_lines WHERE journal_id=?",(jid,))
            tx.execute("DELETE FROM journal_entries WHERE source_type IN ('OPENING_PRODUCT','PRODUCT_OPENING','OPENING_PRODUCT_EDIT','PRODUCT_OPENING_EDIT') AND source_id=?",(item_id,))
        except Exception:pass
        try:tx.execute("DELETE FROM inventory_transactions WHERE product_id=?",(item_id,))
        except Exception:pass
        try:tx.execute("DELETE FROM stock_movements WHERE product_id=?",(item_id,))
        except Exception:pass
        try:tx.execute("DELETE FROM inventory_balances WHERE product_id=?",(item_id,))
        except Exception:pass
        try:tx.execute("DELETE FROM product_price_levels WHERE product_id=?",(item_id,))
        except Exception:pass
        try:tx.execute("DELETE FROM product_units WHERE product_id=?",(item_id,))
        except Exception:pass
        tx.execute("DELETE FROM products WHERE id=?",(item_id,))
        audit(actor["id"],"PRODUCT_DELETED","product",item_id,{"sku":p["sku"]},ip,tx)
        return {"status":"deleted"}


def permissions_catalog():
    return [
      {"group":"Dashboard","code":"dashboard.view","name":"Lihat Dashboard","description":"Melihat ringkasan dan analitik Dashboard."},
      {"group":"Master Data","code":"inventory.view","name":"Lihat Barang & Persediaan","description":"Melihat barang, kategori, satuan, merk, gudang, dan stok."},
      {"group":"Master Data","code":"inventory.manage","name":"Kelola Barang & Persediaan","description":"Tambah/edit barang, gudang, transfer, dan penyesuaian stok."},
      {"group":"Master Data","code":"services.view","name":"Lihat Master Jasa","description":"Melihat kategori jasa dan daftar jasa."},
      {"group":"Master Data","code":"services.manage","name":"Kelola Master Jasa","description":"Tambah, edit, dan hapus kategori serta master jasa."},
      {"group":"Master Data","code":"cost.view","name":"Lihat Harga Modal / Cost Barang","description":"Melihat harga beli, average cost, nilai persediaan, HPP, margin barang, dan cost pada laporan."},
      {"group":"Master Data","code":"inventory.negative_stock","name":"Izinkan Stok Minus","description":"Mengizinkan transaksi penjualan tetap diposting walaupun stok barang tidak mencukupi."},
      {"group":"Proyek & Departemen","code":"projects.view","name":"Lihat Proyek","description":"Melihat master proyek, kartu proyek, progress, termin, budget, dan data proyek."},
      {"group":"Proyek & Departemen","code":"projects.manage","name":"Kelola Proyek","description":"Tambah/edit proyek, progress, termin, budget, dan transaksi proyek."},
      {"group":"Proyek & Departemen","code":"departments.view","name":"Lihat Departemen","description":"Melihat master departemen dan cost center."},
      {"group":"Proyek & Departemen","code":"departments.manage","name":"Kelola Departemen","description":"Tambah, edit, dan menonaktifkan departemen/cost center."},
      {"group":"Master Data","code":"partners.view","name":"Lihat Pelanggan, Pemasok & Salesman","description":"Melihat master pelanggan, pemasok, dan salesman."},
      {"group":"Master Data","code":"partners.manage","name":"Kelola Pelanggan, Pemasok & Salesman","description":"Tambah/edit partner, salesman, dan saldo awal."},
      {"group":"Penjualan","code":"sales.view","name":"Lihat Penjualan","description":"Melihat daftar penjualan dan dokumen."},
      {"group":"Penjualan","code":"sales.manage","name":"Kelola Penjualan","description":"Input, impor, dan cetak transaksi penjualan."},
      {"group":"Penjualan","code":"receivables.view","name":"Lihat Piutang Pelanggan","description":"Melihat piutang dan invoice outstanding."},
      {"group":"Penjualan","code":"receivables.manage","name":"Penerimaan Piutang Pelanggan","description":"Mencatat penerimaan piutang."},
      {"group":"Pembelian","code":"purchases.view","name":"Lihat Pembelian","description":"Melihat daftar pembelian dan dokumen."},
      {"group":"Pembelian","code":"purchases.manage","name":"Kelola Pembelian","description":"Input, impor, dan cetak transaksi pembelian."},
      {"group":"Pembelian","code":"payables.view","name":"Lihat Hutang Pemasok","description":"Melihat hutang dan invoice outstanding."},
      {"group":"Pembelian","code":"payables.manage","name":"Pembayaran Hutang Pemasok","description":"Mencatat pembayaran hutang."},
      {"group":"Kas & Bank","code":"cash.view","name":"Lihat Kas & Bank","description":"Melihat saldo dan mutasi Kas/Bank."},
      {"group":"Kas & Bank","code":"cash.manage","name":"Kelola Kas & Bank","description":"Kas masuk/keluar, transfer, impor Excel, dan rekening koran."},
      {"group":"Akuntansi","code":"accounting.view","name":"Lihat Akuntansi","description":"Melihat COA, jurnal, buku besar, dan neraca saldo."},
      {"group":"Akuntansi","code":"accounting.manage","name":"Kelola Akuntansi","description":"Kelola COA, jurnal manual, dan impor jurnal."},
      {"group":"Laporan","code":"reports.view","name":"Lihat & Ekspor Laporan","description":"Melihat dan mengekspor seluruh laporan."},
      {"group":"Sistem","code":"settings.view","name":"Lihat Pengaturan","description":"Melihat Profil Perusahaan dan Desain Dokumen."},
      {"group":"Sistem","code":"settings.manage","name":"Kelola Pengaturan","description":"Mengubah Profil Perusahaan dan Desain Dokumen."},
      {"group":"Sistem","code":"users.manage","name":"Kelola User & Hak Akses","description":"Mengelola user, role, password, dan permission."},
      {"group":"Sistem","code":"audit.view","name":"Lihat Audit Log","description":"Melihat riwayat aktivitas user."},
    ]

def _valid_permission_codes():
    return {x["code"] for x in permissions_catalog()}

def _normalize_permissions(values):
    if values=="*" or values==["*"]: return ["*"]
    if not isinstance(values,list): raise ValueError("Daftar hak akses tidak valid.")
    valid=_valid_permission_codes()
    result=[]
    for value in values:
        code=str(value or "").strip()
        if code in valid and code not in result:
            result.append(code)
    return result

def create_role(actor,d,ip):
    code=str(d.get("code") or "").strip().upper().replace(" ","_")
    name=str(d.get("name") or "").strip()
    permissions=_normalize_permissions(d.get("permissions") or [])
    if not code or not name: raise ValueError("Kode dan nama role wajib diisi.")
    if code=="ADMIN": raise ValueError("Role Administrator sudah tersedia.")
    if not re.match(r"^[A-Z0-9_]{2,30}$",code):
        raise ValueError("Kode role hanya boleh huruf, angka, dan underscore.")
    with write_transaction() as tx:
        cur=tx.execute(
          "INSERT INTO roles(code,name,permissions_json,created_at) VALUES(?,?,?,?)",
          (code,name,json.dumps(permissions),utc_now()))
        audit(actor["id"],"ROLE_CREATED","role",cur.lastrowid,
          {"code":code,"name":name,"permissions":permissions},ip,tx)
        return cur.lastrowid

def update_role(actor,code,d,ip):
    code=str(code or "").strip().upper()
    name=str(d.get("name") or "").strip()
    if not name: raise ValueError("Nama role wajib diisi.")
    with write_transaction() as tx:
        role=tx.execute("SELECT * FROM roles WHERE code=?",(code,)).fetchone()
        if not role: raise ValueError("Role tidak ditemukan.")
        permissions=["*"] if code=="ADMIN" else _normalize_permissions(d.get("permissions") or [])
        tx.execute("UPDATE roles SET name=?,permissions_json=? WHERE code=?",
          (name,json.dumps(permissions),code))
        audit(actor["id"],"ROLE_UPDATED","role",role["id"],
          {"code":code,"name":name,"permissions":permissions},ip,tx)
    return next(x for x in list_roles() if x["code"]==code)

def update_user(actor,user_id,d,ip):
    user_id=int(user_id)
    full_name=str(d.get("full_name") or "").strip()
    role_code=str(d.get("role_code") or "").strip().upper()
    is_active=1 if bool(d.get("is_active",True)) else 0
    if not full_name or not role_code: raise ValueError("Nama lengkap dan role wajib diisi.")
    if user_id==actor["id"] and not is_active:
        raise ValueError("User yang sedang login tidak dapat dinonaktifkan.")
    with write_transaction() as tx:
        current=tx.execute("""SELECT u.*,r.code role_code FROM users u
          JOIN roles r ON r.id=u.role_id WHERE u.id=?""",(user_id,)).fetchone()
        if not current: raise ValueError("User tidak ditemukan.")
        if not bool(current["is_active"]) and bool(is_active):
            licensing.enforce_user_limit(tx, activating_new_user=True)
        role=tx.execute("SELECT id FROM roles WHERE code=?",(role_code,)).fetchone()
        if not role: raise ValueError("Role tidak ditemukan.")
        if current["role_code"]=="ADMIN" and (role_code!="ADMIN" or not is_active):
            count=tx.execute("""SELECT COUNT(*) jumlah FROM users u JOIN roles r
              ON r.id=u.role_id WHERE r.code='ADMIN' AND u.is_active=1""").fetchone()["jumlah"]
            if count<=1: raise ValueError("Minimal harus ada satu Administrator aktif.")
        tx.execute("""UPDATE users SET full_name=?,role_id=?,is_active=?,updated_at=?
          WHERE id=?""",(full_name,role["id"],is_active,utc_now(),user_id))
        if not is_active:
            tx.execute("DELETE FROM sessions WHERE user_id=?",(user_id,))
        audit(actor["id"],"USER_UPDATED","user",user_id,
          {"username":current["username"],"role_code":role_code,"is_active":bool(is_active)},ip,tx)
    return next(x for x in list_users() if x["id"]==user_id)

def reset_user_password(actor,user_id,new_password,ip):
    user_id=int(user_id)
    new_password=str(new_password or "")
    if len(new_password)<8: raise ValueError("Password baru minimal 8 karakter.")
    with write_transaction() as tx:
        user=tx.execute("SELECT username FROM users WHERE id=?",(user_id,)).fetchone()
        if not user: raise ValueError("User tidak ditemukan.")
        tx.execute("UPDATE users SET password_hash=?,updated_at=? WHERE id=?",
          (hash_password(new_password),utc_now(),user_id))
        tx.execute("DELETE FROM sessions WHERE user_id=?",(user_id,))
        audit(actor["id"],"USER_PASSWORD_RESET","user",user_id,
          {"username":user["username"]},ip,tx)
    return True


def _optional_int(value):
    if value in (None,""):return None
    try:return int(value)
    except (TypeError,ValueError):raise ValueError("ID referensi tidak valid.")

def list_departments(active_only=False):
    c=connect()
    try:
        sql="SELECT * FROM departments"
        if active_only:sql+=" WHERE is_active=1"
        sql+=" ORDER BY code,name"
        return [dict(x) for x in c.execute(sql).fetchall()]
    finally:c.close()

def create_department(actor,d,ip):
    code=str(d.get("code") or "").strip().upper()
    name=str(d.get("name") or "").strip()
    if not code or not name:raise ValueError("Kode dan nama departemen wajib diisi.")
    if not re.match(r"^[A-Z0-9._-]{1,30}$",code):
        raise ValueError("Kode departemen hanya boleh huruf, angka, titik, strip, dan underscore.")
    now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("""INSERT INTO departments(
          code,name,manager_name,notes,is_active,created_at,updated_at
        ) VALUES(?,?,?,?,?,?,?)""",(code,name,
          str(d.get("manager_name") or "").strip() or None,
          str(d.get("notes") or "").strip() or None,
          1 if d.get("is_active",True) else 0,now,now))
        audit(actor["id"],"DEPARTMENT_CREATED","department",cur.lastrowid,
          {"code":code,"name":name},ip,tx)
        return cur.lastrowid

def update_department(actor,department_id,d,ip):
    department_id=int(department_id)
    code=str(d.get("code") or "").strip().upper()
    name=str(d.get("name") or "").strip()
    if not code or not name:raise ValueError("Kode dan nama departemen wajib diisi.")
    with write_transaction() as tx:
        if not tx.execute("SELECT 1 FROM departments WHERE id=?",(department_id,)).fetchone():
            raise ValueError("Departemen tidak ditemukan.")
        tx.execute("""UPDATE departments SET code=?,name=?,manager_name=?,notes=?,
          is_active=?,updated_at=? WHERE id=?""",(code,name,
          str(d.get("manager_name") or "").strip() or None,
          str(d.get("notes") or "").strip() or None,
          1 if d.get("is_active",True) else 0,utc_now(),department_id))
        audit(actor["id"],"DEPARTMENT_UPDATED","department",department_id,
          {"code":code,"name":name},ip,tx)
    return next(x for x in list_departments(False) if x["id"]==department_id)

def list_projects(active_only=False):
    c=connect()
    try:
        sql="""SELECT p.*,bp.code customer_code,bp.name customer_name,
          COALESCE(pb.material_budget,0) material_budget,
          COALESCE(pb.expense_budget,0) expense_budget,
          COALESCE(pb.material_budget,0)+COALESCE(pb.expense_budget,0) total_budget,
          COALESCE((SELECT progress_percent FROM project_progress pp WHERE pp.project_id=p.id ORDER BY progress_date DESC,id DESC LIMIT 1),0) progress_percent,
          COALESCE((SELECT SUM(amount) FROM project_terms pt WHERE pt.project_id=p.id AND pt.status<>'CANCELLED'),0) term_total,
          COALESCE((SELECT SUM(CASE WHEN pt.status='PAID' THEN pt.amount ELSE 0 END) FROM project_terms pt WHERE pt.project_id=p.id),0) term_paid,
          COALESCE((SELECT SUM(retention_amount) FROM project_terms pt WHERE pt.project_id=p.id AND pt.retention_status='PENDING' AND pt.status<>'CANCELLED'),0) retention_pending,
          COALESCE(mi.realization_material,0) realization_material,
          COALESCE(ex.realization_expense,0) realization_expense,
          COALESCE(mi.realization_material,0)+COALESCE(ex.realization_expense,0) realization_total,
          COALESCE(pb.material_budget,0)-COALESCE(mi.realization_material,0)
            remaining_material_budget,
          COALESCE(pb.expense_budget,0)-COALESCE(ex.realization_expense,0)
            remaining_expense_budget,
          COALESCE(pb.material_budget,0)+COALESCE(pb.expense_budget,0)-
            COALESCE(mi.realization_material,0)-COALESCE(ex.realization_expense,0)
            remaining_budget,
          CASE WHEN COALESCE(pb.material_budget,0)>0
            THEN ROUND(COALESCE(mi.realization_material,0)*100.0/
              COALESCE(pb.material_budget,0),2) ELSE 0 END material_usage_percent,
          CASE WHEN COALESCE(pb.expense_budget,0)>0
            THEN ROUND(COALESCE(ex.realization_expense,0)*100.0/
              COALESCE(pb.expense_budget,0),2) ELSE 0 END expense_usage_percent,
          CASE WHEN COALESCE(pb.material_budget,0)+COALESCE(pb.expense_budget,0)>0
            THEN ROUND((COALESCE(mi.realization_material,0)+
              COALESCE(ex.realization_expense,0))*100.0/
              (COALESCE(pb.material_budget,0)+COALESCE(pb.expense_budget,0)),2)
            ELSE 0 END usage_percent
          FROM projects p
          LEFT JOIN business_partners bp ON bp.id=p.customer_id
          LEFT JOIN project_budgets pb ON pb.project_id=p.id
          LEFT JOIN (
            SELECT project_id,COALESCE(SUM(total_cost),0) realization_material
            FROM project_material_issues WHERE status='POSTED' GROUP BY project_id
          ) mi ON mi.project_id=p.id
          LEFT JOIN (
            SELECT jl.project_id,
              COALESCE(SUM(CAST(jl.debit AS REAL)-CAST(jl.credit AS REAL)),0)
                realization_expense
            FROM journal_lines jl
            JOIN journal_entries j ON j.id=jl.journal_id
            JOIN chart_of_accounts coa ON coa.id=jl.account_id
            WHERE j.status='POSTED'
              AND jl.project_id IS NOT NULL
              AND COALESCE(j.source_type,'')<>'PROJECT_MATERIAL_ISSUE'
              AND (coa.account_type='EXPENSE' OR coa.account_subtype='HPP')
            GROUP BY jl.project_id
          ) ex ON ex.project_id=p.id"""
        if active_only:sql+=" WHERE p.is_active=1"
        sql+=" ORDER BY p.code,p.name"
        return [dict(x) for x in c.execute(sql).fetchall()]
    finally:c.close()

def _project_budget_amount(value,label):
    try:
        amount=Decimal(str(value if value not in (None,"") else 0))
    except Exception:
        raise ValueError(f"{label} tidak valid.")
    if amount<0:
        raise ValueError(f"{label} tidak boleh minus.")
    return amount.quantize(Decimal("0.01"))

def list_project_budgets(project_id=None):
    c=connect()
    try:
        sql="""SELECT pb.id,p.id project_id,p.code project_code,p.name project_name,
          bp.name customer_name,p.status project_status,p.is_active project_active,
          pb.material_budget,pb.expense_budget,pb.notes,pb.created_at,pb.updated_at,
          COALESCE(mi.realization_material,0) realization_material,
          COALESCE(ex.realization_expense,0) realization_expense
          FROM project_budgets pb
          JOIN projects p ON p.id=pb.project_id
          LEFT JOIN business_partners bp ON bp.id=p.customer_id
          LEFT JOIN (
            SELECT project_id,COALESCE(SUM(total_cost),0) realization_material
            FROM project_material_issues WHERE status='POSTED' GROUP BY project_id
          ) mi ON mi.project_id=p.id
          LEFT JOIN (
            SELECT jl.project_id,
              COALESCE(SUM(CAST(jl.debit AS REAL)-CAST(jl.credit AS REAL)),0)
                realization_expense
            FROM journal_lines jl
            JOIN journal_entries j ON j.id=jl.journal_id
            JOIN chart_of_accounts coa ON coa.id=jl.account_id
            WHERE j.status='POSTED'
              AND jl.project_id IS NOT NULL
              AND COALESCE(j.source_type,'')<>'PROJECT_MATERIAL_ISSUE'
              AND (coa.account_type='EXPENSE' OR coa.account_subtype='HPP')
            GROUP BY jl.project_id
          ) ex ON ex.project_id=p.id"""
        params=[]
        if project_id not in (None,""):
            sql+=" WHERE p.id=?";params.append(int(project_id))
        sql+=" ORDER BY p.code,p.name"
        rows=[]
        for row in c.execute(sql,params).fetchall():
            item=dict(row)
            material=float(item.get("material_budget") or 0)
            expense=float(item.get("expense_budget") or 0)
            total=material+expense
            realization_material=float(item.get("realization_material") or 0)
            realization_expense=float(item.get("realization_expense") or 0)
            realization_total=realization_material+realization_expense
            item.update({
              "material_budget":material,"expense_budget":expense,
              "total_budget":total,
              "realization_material":realization_material,
              "realization_expense":realization_expense,
              "realization_total":realization_total,
              "remaining_material_budget":material-realization_material,
              "remaining_expense_budget":expense-realization_expense,
              "remaining_budget":total-realization_total,
              "material_usage_percent":round(realization_material*100/material,2) if material>0 else 0.0,
              "expense_usage_percent":round(realization_expense*100/expense,2) if expense>0 else 0.0,
              "usage_percent":round(realization_total*100/total,2) if total>0 else 0.0
            })
            rows.append(item)
        return rows
    finally:c.close()

def project_budget_overview():
    projects=list_projects(False)
    return {
      "items":projects,
      "summary":{
        "project_count":len(projects),
        "budgeted_project_count":sum(1 for x in projects if float(x.get("total_budget") or 0)>0),
        "material_budget":sum(float(x.get("material_budget") or 0) for x in projects),
        "expense_budget":sum(float(x.get("expense_budget") or 0) for x in projects),
        "total_budget":sum(float(x.get("total_budget") or 0) for x in projects),
        "realization_material":sum(float(x.get("realization_material") or 0) for x in projects),
        "realization_expense":sum(float(x.get("realization_expense") or 0) for x in projects),
        "realization_total":sum(float(x.get("realization_total") or 0) for x in projects),
        "remaining_material_budget":sum(float(x.get("remaining_material_budget") or 0) for x in projects),
        "remaining_expense_budget":sum(float(x.get("remaining_expense_budget") or 0) for x in projects),
        "remaining_budget":sum(float(x.get("remaining_budget") or 0) for x in projects),
        "usage_percent":round(
          sum(float(x.get("realization_total") or 0) for x in projects)*100/
          sum(float(x.get("total_budget") or 0) for x in projects),2
        ) if sum(float(x.get("total_budget") or 0) for x in projects)>0 else 0.0
      }
    }

def project_cost_realization_details(project_id,limit=500):
    try:project_id=int(project_id)
    except Exception:raise ValueError("Proyek wajib dipilih.")
    limit=max(1,min(int(limit),2000))
    c=connect()
    try:
        project=c.execute("SELECT id,code,name FROM projects WHERE id=?",(project_id,)).fetchone()
        if not project:raise ValueError("Proyek tidak ditemukan.")
        rows=c.execute("""SELECT j.journal_date,j.journal_no,j.source_type,
          COALESCE(j.reference_no,'-') reference_no,j.description,
          coa.code account_code,coa.name account_name,
          CAST(jl.debit AS REAL) debit,CAST(jl.credit AS REAL) credit,
          CAST(jl.debit AS REAL)-CAST(jl.credit AS REAL) amount,
          d.code department_code,d.name department_name
          FROM journal_lines jl
          JOIN journal_entries j ON j.id=jl.journal_id
          JOIN chart_of_accounts coa ON coa.id=jl.account_id
          LEFT JOIN departments d ON d.id=jl.department_id
          WHERE j.status='POSTED' AND jl.project_id=?
            AND COALESCE(j.source_type,'')<>'PROJECT_MATERIAL_ISSUE'
            AND (coa.account_type='EXPENSE' OR coa.account_subtype='HPP')
          ORDER BY j.journal_date DESC,j.id DESC,jl.id DESC LIMIT ?""",
          (project_id,limit)).fetchall()
        items=[]
        for row in rows:
            item=dict(row)
            item["debit"]=float(item["debit"] or 0)
            item["credit"]=float(item["credit"] or 0)
            item["amount"]=float(item["amount"] or 0)
            items.append(item)
        return {
          "project":{"id":project["id"],"code":project["code"],"name":project["name"]},
          "items":items,
          "summary":{
            "total_debit":sum(x["debit"] for x in items),
            "total_credit":sum(x["credit"] for x in items),
            "net_expense":sum(x["amount"] for x in items)
          }
        }
    finally:c.close()


def create_project_budget(actor,d,ip):
    project_id=_optional_int(d.get("project_id"))
    if not project_id:raise ValueError("Proyek wajib dipilih.")
    material=_project_budget_amount(d.get("material_budget"),"Budget material")
    expense=_project_budget_amount(d.get("expense_budget"),"Budget biaya proyek")
    notes=str(d.get("notes") or "").strip() or None
    now=utc_now()
    with write_transaction() as tx:
        project=tx.execute("SELECT id,name FROM projects WHERE id=?",(project_id,)).fetchone()
        if not project:raise ValueError("Proyek tidak ditemukan.")
        if tx.execute("SELECT id FROM project_budgets WHERE project_id=?",(project_id,)).fetchone():
            raise ValueError("Budget untuk proyek ini sudah tersedia. Gunakan tombol Edit.")
        cur=tx.execute("""INSERT INTO project_budgets(
          project_id,material_budget,expense_budget,notes,created_at,updated_at
        ) VALUES(?,?,?,?,?,?)""",(project_id,str(material),str(expense),notes,now,now))
        audit(actor["id"],"PROJECT_BUDGET_CREATED","project_budget",cur.lastrowid,
          {"project_id":project_id,"material_budget":str(material),
           "expense_budget":str(expense)},ip,tx)
        return cur.lastrowid

def update_project_budget(actor,budget_id,d,ip):
    budget_id=int(budget_id)
    project_id=_optional_int(d.get("project_id"))
    if not project_id:raise ValueError("Proyek wajib dipilih.")
    material=_project_budget_amount(d.get("material_budget"),"Budget material")
    expense=_project_budget_amount(d.get("expense_budget"),"Budget biaya proyek")
    notes=str(d.get("notes") or "").strip() or None
    with write_transaction() as tx:
        current=tx.execute("SELECT * FROM project_budgets WHERE id=?",(budget_id,)).fetchone()
        if not current:raise ValueError("Budget proyek tidak ditemukan.")
        if not tx.execute("SELECT id FROM projects WHERE id=?",(project_id,)).fetchone():
            raise ValueError("Proyek tidak ditemukan.")
        duplicate=tx.execute("""SELECT id FROM project_budgets
          WHERE project_id=? AND id<>?""",(project_id,budget_id)).fetchone()
        if duplicate:raise ValueError("Budget untuk proyek tersebut sudah tersedia.")
        tx.execute("""UPDATE project_budgets SET project_id=?,material_budget=?,
          expense_budget=?,notes=?,updated_at=? WHERE id=?""",
          (project_id,str(material),str(expense),notes,utc_now(),budget_id))
        audit(actor["id"],"PROJECT_BUDGET_UPDATED","project_budget",budget_id,
          {"project_id":project_id,"material_budget":str(material),
           "expense_budget":str(expense)},ip,tx)
    return next(x for x in list_project_budgets() if x["id"]==budget_id)

def delete_project_budget(actor,budget_id,ip):
    budget_id=int(budget_id)
    with write_transaction() as tx:
        current=tx.execute("""SELECT pb.*,p.code project_code,p.name project_name
          FROM project_budgets pb JOIN projects p ON p.id=pb.project_id
          WHERE pb.id=?""",(budget_id,)).fetchone()
        if not current:raise ValueError("Budget proyek tidak ditemukan.")
        tx.execute("DELETE FROM project_budgets WHERE id=?",(budget_id,))
        audit(actor["id"],"PROJECT_BUDGET_DELETED","project_budget",budget_id,
          {"project_id":current["project_id"],"project_code":current["project_code"]},ip,tx)
    return True


def _next_project_material_issue_no(tx,issue_date):
    period=str(issue_date)[:7].replace("-","")
    if len(period)!=6:raise ValueError("Tanggal pengeluaran material tidak valid.")
    key="PROJECT-MATERIAL-"+period
    row=tx.execute("SELECT current_value FROM document_sequences WHERE sequence_key=?",(key,)).fetchone()
    number=int(row["current_value"])+1 if row else 1
    tx.execute("""INSERT INTO document_sequences(sequence_key,current_value,updated_at)
      VALUES(?,?,?) ON CONFLICT(sequence_key) DO UPDATE SET
      current_value=excluded.current_value,updated_at=excluded.updated_at""",
      (key,number,utc_now()))
    return f"PMI-{period}-{number:06d}"

def list_project_material_issues(project_id=None,limit=300):
    c=connect()
    try:
        where=["i.status='POSTED'"];params=[]
        if project_id not in (None,""):
            where.append("i.project_id=?");params.append(int(project_id))
        params.append(max(1,min(int(limit),1000)))
        rows=c.execute(f"""SELECT i.*,p.code project_code,p.name project_name,
          d.code department_code,d.name department_name,
          w.code warehouse_code,w.name warehouse_name,
          coa.code expense_account_code,coa.name expense_account_name,
          u.username
          FROM project_material_issues i
          JOIN projects p ON p.id=i.project_id
          LEFT JOIN departments d ON d.id=i.department_id
          JOIN warehouses w ON w.id=i.warehouse_id
          JOIN chart_of_accounts coa ON coa.id=i.expense_account_id
          JOIN users u ON u.id=i.user_id
          WHERE {' AND '.join(where)}
          ORDER BY i.issue_date DESC,i.id DESC LIMIT ?""",params).fetchall()
        result=[]
        for row in rows:
            item=dict(row);item["total_cost"]=float(item["total_cost"] or 0)
            item["items"]=[dict(x) for x in c.execute("""SELECT mi.*,p.sku,p.name product_name
              FROM project_material_issue_items mi JOIN products p ON p.id=mi.product_id
              WHERE mi.issue_id=? ORDER BY mi.id""",(item["id"],)).fetchall()]
            for detail in item["items"]:
                detail["qty"]=float(detail["qty"]);detail["average_cost"]=float(detail["average_cost"])
                detail["total_cost"]=float(detail["total_cost"])
            result.append(item)
        return result
    finally:c.close()

def create_project_material_issue(actor,d,ip):
    issue_date=str(d.get("issue_date") or date.today().isoformat())[:10]
    project_id=_optional_int(d.get("project_id"))
    department_id=_optional_int(d.get("department_id"))
    warehouse_id=_optional_int(d.get("warehouse_id"))
    expense_account_id=_optional_int(d.get("expense_account_id"))
    notes=str(d.get("notes") or "").strip() or None
    items=d.get("items") or []
    if not project_id:raise ValueError("Proyek wajib dipilih.")
    if not warehouse_id:raise ValueError("Gudang wajib dipilih.")
    if not expense_account_id:raise ValueError("Akun pengeluaran material wajib dipilih.")
    if not isinstance(items,list) or not items:raise ValueError("Minimal satu barang wajib diisi.")
    now=utc_now()
    with write_transaction() as tx:
        project=tx.execute("SELECT id,name FROM projects WHERE id=? AND is_active=1",(project_id,)).fetchone()
        if not project:raise ValueError("Proyek tidak ditemukan atau nonaktif.")
        if department_id and not tx.execute("SELECT 1 FROM departments WHERE id=? AND is_active=1",(department_id,)).fetchone():
            raise ValueError("Departemen tidak ditemukan atau nonaktif.")
        if not tx.execute("SELECT 1 FROM warehouses WHERE id=? AND is_active=1",(warehouse_id,)).fetchone():
            raise ValueError("Gudang tidak ditemukan atau nonaktif.")
        account=tx.execute("""SELECT id,account_type,account_subtype,name
          FROM chart_of_accounts WHERE id=? AND is_active=1""",(expense_account_id,)).fetchone()
        if not account or not (
          account["account_type"]=="ASSET" or
          account["account_subtype"] in ("HPP","OPERATING_EXPENSE") or
          account["account_type"]=="EXPENSE"
        ):
            raise ValueError("Akun pengeluaran hanya boleh tipe HPP, Beban, atau Aset.")
        issue_no=_next_project_material_issue_no(tx,issue_date)
        cur=tx.execute("""INSERT INTO project_material_issues(
          issue_no,issue_date,project_id,department_id,warehouse_id,
          expense_account_id,notes,status,total_cost,user_id,created_at,updated_at
        ) VALUES(?,?,?,?,?,?,?,'POSTED',0,?,?,?)""",
          (issue_no,issue_date,project_id,department_id,warehouse_id,
           expense_account_id,notes,actor["id"],now,now))
        issue_id=cur.lastrowid
        total=Decimal("0.00")
        inventory_groups={}
        normalized=[]
        for index,line in enumerate(items,1):
            product_id=_optional_int(line.get("product_id"))
            if not product_id:raise ValueError(f"Baris {index}: barang wajib dipilih.")
            qty_value=_qty(line.get("qty"))
            if qty_value<=0:raise ValueError(f"Baris {index}: qty harus lebih dari nol.")
            product=tx.execute("""SELECT id,sku,name,product_type,inventory_account_id
              FROM products WHERE id=? AND is_active=1""",(product_id,)).fetchone()
            if not product or product["product_type"]!="STOCK":
                raise ValueError(f"Baris {index}: barang stok tidak valid.")
            if not product["inventory_account_id"]:
                raise ValueError(f"Baris {index}: akun persediaan barang belum diatur.")
            before,avg=inventory_service.get_balance(tx,warehouse_id,product_id)
            movement=inventory_service.post_movement(
              tx,product_id=product_id,warehouse_id=warehouse_id,
              movement_type="PROJECT_OUT",quantity_change=-qty_value,unit_cost=avg,
              reference_type="PROJECT_MATERIAL_ISSUE",reference_no=issue_no,
              reason=f"Pengeluaran material proyek {project['name']}",
              user_id=actor["id"],created_at=now,
              department_id=department_id,project_id=project_id)
            line_total=(qty_value*avg).quantize(Decimal("0.01"))
            total+=line_total
            inventory_groups[int(product["inventory_account_id"])]=(
              inventory_groups.get(int(product["inventory_account_id"]),Decimal("0.00"))+line_total
            )
            normalized.append((product_id,qty_value,avg,line_total,movement["transaction_id"]))
        for product_id,qty_value,avg,line_total,tx_id in normalized:
            tx.execute("""INSERT INTO project_material_issue_items(
              issue_id,product_id,qty,average_cost,total_cost,inventory_transaction_id
            ) VALUES(?,?,?,?,?,?)""",
              (issue_id,product_id,str(qty_value),str(avg),str(line_total),tx_id))
        tx.execute("UPDATE project_material_issues SET total_cost=? WHERE id=?",(str(total),issue_id))
        if total>0:
            lines=[{"account_id":expense_account_id,"debit":total}]
            lines.extend({"account_id":account_id,"credit":amount}
              for account_id,amount in inventory_groups.items() if amount>0)
            accounting_service.post_journal(
              tx,journal_date=issue_date,
              description=f"Pengeluaran material proyek {project['name']}",
              source_type="PROJECT_MATERIAL_ISSUE",source_id=issue_id,
              reference_no=issue_no,lines=lines,user_id=actor["id"],
              department_id=department_id,project_id=project_id)
        audit(actor["id"],"PROJECT_MATERIAL_ISSUED","project_material_issue",issue_id,{
          "issue_no":issue_no,"project_id":project_id,"warehouse_id":warehouse_id,
          "total_cost":str(total),"item_count":len(normalized)
        },ip,tx)
        return {"id":issue_id,"issue_no":issue_no,"total_cost":float(total)}


def create_project(actor,d,ip):
    code=str(d.get("code") or "").strip().upper();name=str(d.get("name") or "").strip()
    if not code or not name:raise ValueError("Kode dan nama proyek wajib diisi.")
    contract_value=Decimal(str(d.get("contract_value") or 0));retention=Decimal(str(d.get("retention_percent") or 0))
    if contract_value<0 or retention<0 or retention>100:raise ValueError("Nilai kontrak/retensi tidak valid.")
    now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("""INSERT INTO projects(code,name,customer_id,start_date,end_date,status,notes,is_active,created_at,updated_at,contract_no,contract_value,location,pic_name,project_type,retention_percent) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(code,name,int(d["customer_id"]) if d.get("customer_id") else None,d.get("start_date") or None,d.get("end_date") or None,d.get("status") or "ACTIVE",str(d.get("notes") or "").strip() or None,1 if d.get("is_active",True) else 0,now,now,str(d.get("contract_no") or "").strip() or None,str(contract_value),str(d.get("location") or "").strip() or None,str(d.get("pic_name") or "").strip() or None,str(d.get("project_type") or "").strip() or None,str(retention)))
        audit(actor["id"],"PROJECT_CREATED","project",cur.lastrowid,{"code":code,"name":name,"contract_value":str(contract_value)},ip,tx)
        return cur.lastrowid

def update_project(actor,project_id,d,ip):
    project_id=int(project_id);code=str(d.get("code") or "").strip().upper();name=str(d.get("name") or "").strip()
    if not code or not name:raise ValueError("Kode dan nama proyek wajib diisi.")
    contract_value=Decimal(str(d.get("contract_value") or 0));retention=Decimal(str(d.get("retention_percent") or 0))
    if contract_value<0 or retention<0 or retention>100:raise ValueError("Nilai kontrak/retensi tidak valid.")
    with write_transaction() as tx:
        if not tx.execute("SELECT 1 FROM projects WHERE id=?",(project_id,)).fetchone():raise ValueError("Proyek tidak ditemukan.")
        tx.execute("""UPDATE projects SET code=?,name=?,customer_id=?,start_date=?,end_date=?,status=?,notes=?,is_active=?,updated_at=?,contract_no=?,contract_value=?,location=?,pic_name=?,project_type=?,retention_percent=? WHERE id=?""",(code,name,int(d["customer_id"]) if d.get("customer_id") else None,d.get("start_date") or None,d.get("end_date") or None,d.get("status") or "ACTIVE",str(d.get("notes") or "").strip() or None,1 if d.get("is_active",True) else 0,utc_now(),str(d.get("contract_no") or "").strip() or None,str(contract_value),str(d.get("location") or "").strip() or None,str(d.get("pic_name") or "").strip() or None,str(d.get("project_type") or "").strip() or None,str(retention),project_id))
        audit(actor["id"],"PROJECT_UPDATED","project",project_id,{"code":code,"name":name},ip,tx)
    return next(x for x in list_projects(False) if x["id"]==project_id)

def _next_assembly_no(tx, assembly_date):
    ym=str(assembly_date)[:7].replace('-','')
    prefix=f"ASM-{ym}-"
    row=tx.execute("SELECT assembly_no FROM assembly_orders WHERE assembly_no LIKE ? ORDER BY id DESC LIMIT 1",(prefix+'%',)).fetchone()
    seq=1
    if row:
        try: seq=int(str(row['assembly_no']).rsplit('-',1)[1])+1
        except Exception: seq=1
    return f"{prefix}{seq:04d}"

def list_assemblies(limit=200):
    c=connect()
    try:
        rows=c.execute("""SELECT a.*,w.name warehouse_name,coa.code wip_code,coa.name wip_name,u.username
          FROM assembly_orders a JOIN warehouses w ON w.id=a.warehouse_id
          JOIN chart_of_accounts coa ON coa.id=a.wip_account_id JOIN users u ON u.id=a.user_id
          ORDER BY a.assembly_date DESC,a.id DESC LIMIT ?""",(max(1,min(int(limit or 200),500)),)).fetchall()
        return [dict(x) for x in rows]
    finally:c.close()

def get_assembly(assembly_id):
    c=connect()
    try:
        h=c.execute("""SELECT a.*,w.name warehouse_name,coa.code wip_code,coa.name wip_name
          FROM assembly_orders a JOIN warehouses w ON w.id=a.warehouse_id
          JOIN chart_of_accounts coa ON coa.id=a.wip_account_id WHERE a.id=?""",(int(assembly_id),)).fetchone()
        if not h:return None
        out=dict(h)
        out['materials']=[dict(x) for x in c.execute("""SELECT m.*,p.sku,p.name product_name,u.code unit_code
          FROM assembly_materials m JOIN products p ON p.id=m.product_id JOIN units u ON u.id=p.unit_id
          WHERE m.assembly_id=? ORDER BY m.id""",(int(assembly_id),)).fetchall()]
        out['costs']=[dict(x) for x in c.execute("""SELECT ac.*,coa.code account_code,coa.name account_name
          FROM assembly_costs ac JOIN chart_of_accounts coa ON coa.id=ac.account_id
          WHERE ac.assembly_id=? ORDER BY ac.id""",(int(assembly_id),)).fetchall()]
        out['outputs']=[dict(x) for x in c.execute("""SELECT o.*,p.sku,p.name product_name,u.code unit_code
          FROM assembly_outputs o JOIN products p ON p.id=o.product_id JOIN units u ON u.id=p.unit_id
          WHERE o.assembly_id=? ORDER BY o.id""",(int(assembly_id),)).fetchall()]
        return out
    finally:c.close()

def create_assembly(actor,d,ip):
    assembly_date=_date_value(d.get('assembly_date') or date.today().isoformat(),'Tanggal assembly')
    warehouse_id=_optional_int(d.get('warehouse_id')); wip_account_id=_optional_int(d.get('wip_account_id'))
    materials=d.get('materials') or []; costs=d.get('costs') or []
    # Gabungkan bahan yang sama agar satu barang hanya diposting satu kali per assembly.
    # Ini juga melindungi transaksi dari baris UI ganda yang tidak sengaja.
    merged_materials={}
    for i,line in enumerate(materials,1):
        pid=_optional_int(line.get('product_id')); qty=_qty(line.get('qty'))
        if not pid or qty<=0:raise ValueError(f'Bahan baris {i}: barang dan qty wajib diisi.')
        merged_materials[pid]=merged_materials.get(pid,Decimal('0'))+qty
    materials=[{'product_id':pid,'qty':str(qty)} for pid,qty in merged_materials.items()]
    if not warehouse_id:raise ValueError('Gudang wajib dipilih.')
    if not wip_account_id:raise ValueError('Akun Persediaan Dalam Proses wajib dipilih.')
    if not materials:raise ValueError('Minimal satu bahan baku wajib diisi.')
    now=utc_now()
    with write_transaction() as tx:
        if not tx.execute("SELECT 1 FROM warehouses WHERE id=? AND is_active=1",(warehouse_id,)).fetchone():raise ValueError('Gudang tidak valid.')
        wip=tx.execute("SELECT * FROM chart_of_accounts WHERE id=? AND is_active=1",(wip_account_id,)).fetchone()
        if not wip or wip['account_type']!='ASSET':raise ValueError('Akun Persediaan Dalam Proses harus bertipe Aset.')
        no=_next_assembly_no(tx,assembly_date)
        cur=tx.execute("""INSERT INTO assembly_orders(assembly_no,assembly_date,warehouse_id,wip_account_id,notes,status,user_id,created_at,updated_at)
          VALUES(?,?,?,?,?,'OPEN',?,?,?)""",(no,assembly_date,warehouse_id,wip_account_id,str(d.get('notes') or '').strip() or None,actor['id'],now,now))
        aid=cur.lastrowid; material_total=Decimal('0.00'); additional_total=Decimal('0.00'); invcredits={}
        for i,line in enumerate(materials,1):
            pid=_optional_int(line.get('product_id')); qty=_qty(line.get('qty'))
            if not pid or qty<=0:raise ValueError(f'Bahan baris {i}: barang dan qty wajib diisi.')
            pr=tx.execute("SELECT id,name,product_type,inventory_account_id FROM products WHERE id=? AND is_active=1",(pid,)).fetchone()
            if not pr or pr['product_type']!='STOCK' or not pr['inventory_account_id']:raise ValueError(f'Bahan baris {i}: barang stok/akun persediaan tidak valid.')
            before,avg=inventory_service.get_balance(tx,warehouse_id,pid)
            mv=inventory_service.post_movement(tx,product_id=pid,warehouse_id=warehouse_id,movement_type='ASSEMBLY_OUT',quantity_change=-qty,unit_cost=avg,reference_type='ASSEMBLY',reference_no=no,reason='Pengambilan bahan baku assembly',user_id=actor['id'],created_at=now)
            total=(qty*avg).quantize(Decimal('0.01')); material_total+=total
            invcredits[int(pr['inventory_account_id'])]=invcredits.get(int(pr['inventory_account_id']),Decimal('0.00'))+total
            tx.execute("INSERT INTO assembly_materials(assembly_id,product_id,qty,unit_cost,total_cost,inventory_transaction_id) VALUES(?,?,?,?,?,?)",(aid,pid,str(qty),str(avg),str(total),mv['transaction_id']))
        costcredits={}
        for i,line in enumerate(costs,1):
            acc=_optional_int(line.get('account_id')); amount=Decimal(str(line.get('amount') or 0)).quantize(Decimal('0.01')); desc=str(line.get('description') or '').strip()
            if amount<=0:continue
            ar=tx.execute("SELECT id,name,account_type FROM chart_of_accounts WHERE id=? AND is_active=1",(acc,)).fetchone() if acc else None
            if not ar:raise ValueError(f'Biaya baris {i}: akun biaya wajib dipilih.')
            additional_total+=amount; costcredits[int(acc)]=costcredits.get(int(acc),Decimal('0.00'))+amount
            tx.execute("INSERT INTO assembly_costs(assembly_id,account_id,description,amount) VALUES(?,?,?,?)",(aid,acc,desc or ar['name'],str(amount)))
        total=material_total+additional_total
        tx.execute("UPDATE assembly_orders SET material_cost=?,additional_cost=?,total_cost=? WHERE id=?",(str(material_total),str(additional_total),str(total),aid))
        lines=[{'account_id':wip_account_id,'debit':total}]
        lines += [{'account_id':k,'credit':v} for k,v in invcredits.items() if v>0]
        lines += [{'account_id':k,'credit':v} for k,v in costcredits.items() if v>0]
        if total>0: accounting_service.post_journal(tx,journal_date=assembly_date,description=f'Assembly tahap 1 {no} - bahan baku dan biaya ke WIP',source_type='ASSEMBLY_ISSUE',source_id=aid,reference_no=no,lines=lines,user_id=actor['id'])
        audit(actor['id'],'ASSEMBLY_CREATED','assembly',aid,{'assembly_no':no,'total_cost':str(total)},ip,tx)
        return {'id':aid,'assembly_no':no,'total_cost':float(total)}

def finish_assembly(actor,assembly_id,d,ip):
    outputs=d.get('outputs') or []
    if not outputs:raise ValueError('Minimal satu barang jadi wajib diisi.')
    finish_date=_date_value(d.get('finish_date') or date.today().isoformat(),'Tanggal finishing')
    now=utc_now(); aid=int(assembly_id)
    with write_transaction() as tx:
        h=tx.execute("SELECT * FROM assembly_orders WHERE id=? AND status='OPEN'",(aid,)).fetchone()
        if not h:raise ValueError('Assembly tidak ditemukan atau sudah selesai.')
        normalized=[]; weights=[]; manual=all(str(x.get('allocation_percent') or '').strip()!='' for x in outputs)
        for i,line in enumerate(outputs,1):
            pid=_optional_int(line.get('product_id')); qty=_qty(line.get('qty'))
            if not pid or qty<=0:raise ValueError(f'Output baris {i}: barang dan qty wajib diisi.')
            pr=tx.execute("SELECT id,name,selling_price,inventory_account_id,product_type FROM products WHERE id=? AND is_active=1",(pid,)).fetchone()
            if not pr or pr['product_type']!='STOCK' or not pr['inventory_account_id']:raise ValueError(f'Output baris {i}: barang jadi/akun persediaan tidak valid.')
            pct=Decimal(str(line.get('allocation_percent') or 0))
            weight=(qty*Decimal(str(pr['selling_price'] or 0))) if not manual else pct
            if not manual and weight<=0:weight=qty
            normalized.append((pid,qty,pr,pct));weights.append(weight)
        total_weight=sum(weights,Decimal('0'))
        if total_weight<=0:raise ValueError('Dasar alokasi biaya tidak valid.')
        if manual and abs(total_weight-Decimal('100'))>Decimal('0.01'):raise ValueError('Total persentase alokasi harus 100%.')
        total=Decimal(str(h['total_cost'] or 0)).quantize(Decimal('0.01')); allocated=Decimal('0'); debitgroups={}
        for idx,(pid,qty,pr,pct) in enumerate(normalized):
            percent=(weights[idx]/total_weight*Decimal('100'))
            cost=(total*weights[idx]/total_weight).quantize(Decimal('0.01')) if idx<len(normalized)-1 else total-allocated
            allocated+=cost; unit=(cost/qty).quantize(Decimal('0.0001'))
            mv=inventory_service.post_movement(tx,product_id=pid,warehouse_id=h['warehouse_id'],movement_type='ASSEMBLY_IN',quantity_change=qty,unit_cost=unit,reference_type='ASSEMBLY_FINISH',reference_no=h['assembly_no'],reason='Finishing barang jadi assembly',user_id=actor['id'],created_at=now)
            debitgroups[int(pr['inventory_account_id'])]=debitgroups.get(int(pr['inventory_account_id']),Decimal('0.00'))+cost
            tx.execute("INSERT INTO assembly_outputs(assembly_id,product_id,qty,allocation_percent,allocated_cost,unit_cost,inventory_transaction_id) VALUES(?,?,?,?,?,?,?)",(aid,pid,str(qty),str(percent.quantize(Decimal('0.0001'))),str(cost),str(unit),mv['transaction_id']))
        lines=[{'account_id':k,'debit':v} for k,v in debitgroups.items() if v>0]+[{'account_id':h['wip_account_id'],'credit':total}]
        accounting_service.post_journal(tx,journal_date=finish_date,description=f"Finishing assembly {h['assembly_no']} - WIP ke barang jadi",source_type='ASSEMBLY_FINISH',source_id=aid,reference_no=h['assembly_no'],lines=lines,user_id=actor['id'])
        tx.execute("UPDATE assembly_orders SET status='FINISHED',finished_at=?,updated_at=? WHERE id=?",(finish_date,now,aid))
        audit(actor['id'],'ASSEMBLY_FINISHED','assembly',aid,{'assembly_no':h['assembly_no'],'total_cost':str(total),'output_count':len(normalized)},ip,tx)
        return {'id':aid,'assembly_no':h['assembly_no'],'total_cost':float(total),'status':'FINISHED'}


def _reverse_assembly_finish(tx, actor, h, now):
    outputs=tx.execute("SELECT * FROM assembly_outputs WHERE assembly_id=? ORDER BY id",(int(h['id']),)).fetchall()
    for o in outputs:
        inventory_service.post_movement(tx,product_id=o['product_id'],warehouse_id=h['warehouse_id'],movement_type='ASSEMBLY_FINISH_VOID',quantity_change=-Decimal(str(o['qty'])),unit_cost=Decimal(str(o['unit_cost'])),reference_type='ASSEMBLY_FINISH_VOID',reference_no=h['assembly_no'],reason='Pembalikan finishing assembly',user_id=actor['id'],created_at=now)
    _void_journals(tx,'ASSEMBLY_FINISH',h['id'])
    tx.execute("DELETE FROM assembly_outputs WHERE assembly_id=?",(int(h['id']),))
    tx.execute("UPDATE assembly_orders SET status='OPEN',finished_at=NULL,updated_at=? WHERE id=?",(now,int(h['id'])))


def _reverse_assembly_issue(tx, actor, h, now):
    materials=tx.execute("SELECT * FROM assembly_materials WHERE assembly_id=? ORDER BY id",(int(h['id']),)).fetchall()
    for m in materials:
        inventory_service.post_movement(tx,product_id=m['product_id'],warehouse_id=h['warehouse_id'],movement_type='ASSEMBLY_ISSUE_VOID',quantity_change=Decimal(str(m['qty'])),unit_cost=Decimal(str(m['unit_cost'])),reference_type='ASSEMBLY_VOID',reference_no=h['assembly_no'],reason='Pembalikan pengambilan bahan assembly',user_id=actor['id'],created_at=now)
    _void_journals(tx,'ASSEMBLY_ISSUE',h['id'])


def update_assembly(actor, assembly_id, d, ip):
    aid=int(assembly_id); now=utc_now()
    with write_transaction() as tx:
        h=tx.execute("SELECT * FROM assembly_orders WHERE id=?",(aid,)).fetchone()
        if not h or h['status']!='OPEN': raise ValueError('Hanya assembly OPEN yang dapat diubah.')
        _reverse_assembly_issue(tx,actor,h,now)
        tx.execute("DELETE FROM assembly_materials WHERE assembly_id=?",(aid,))
        tx.execute("DELETE FROM assembly_costs WHERE assembly_id=?",(aid,))
        assembly_date=_date_value(d.get('assembly_date') or h['assembly_date'],'Tanggal assembly')
        warehouse_id=_optional_int(d.get('warehouse_id')); wip_account_id=_optional_int(d.get('wip_account_id'))
        materials=d.get('materials') or []; costs=d.get('costs') or []
        if not warehouse_id or not wip_account_id or not materials: raise ValueError('Tanggal, gudang, akun WIP, dan bahan wajib diisi.')
        wip=tx.execute("SELECT * FROM chart_of_accounts WHERE id=? AND is_active=1",(wip_account_id,)).fetchone()
        if not wip or wip['account_type']!='ASSET': raise ValueError('Akun WIP harus bertipe Aset.')
        material_total=Decimal('0.00'); additional_total=Decimal('0.00'); invcredits={}; costcredits={}
        for i,line in enumerate(materials,1):
            pid=_optional_int(line.get('product_id')); qty=_qty(line.get('qty'))
            if not pid or qty<=0: raise ValueError(f'Bahan baris {i} tidak valid.')
            pr=tx.execute("SELECT id,name,product_type,inventory_account_id FROM products WHERE id=? AND is_active=1",(pid,)).fetchone()
            if not pr or pr['product_type']!='STOCK' or not pr['inventory_account_id']: raise ValueError(f'Bahan baris {i} bukan barang stok valid.')
            before,avg=inventory_service.get_balance(tx,warehouse_id,pid)
            mv=inventory_service.post_movement(tx,product_id=pid,warehouse_id=warehouse_id,movement_type='ASSEMBLY_OUT',quantity_change=-qty,unit_cost=avg,reference_type='ASSEMBLY',reference_no=h['assembly_no'],reason='Perubahan bahan baku assembly',user_id=actor['id'],created_at=now)
            total=(qty*avg).quantize(Decimal('0.01')); material_total+=total
            invcredits[int(pr['inventory_account_id'])]=invcredits.get(int(pr['inventory_account_id']),Decimal('0.00'))+total
            tx.execute("INSERT INTO assembly_materials(assembly_id,product_id,qty,unit_cost,total_cost,inventory_transaction_id) VALUES(?,?,?,?,?,?)",(aid,pid,str(qty),str(avg),str(total),mv['transaction_id']))
        for i,line in enumerate(costs,1):
            acc=_optional_int(line.get('account_id')); amount=Decimal(str(line.get('amount') or 0)).quantize(Decimal('0.01')); desc=str(line.get('description') or '').strip()
            if amount<=0: continue
            ar=tx.execute("SELECT id,name FROM chart_of_accounts WHERE id=? AND is_active=1",(acc,)).fetchone() if acc else None
            if not ar: raise ValueError(f'Biaya baris {i}: akun tidak valid.')
            additional_total+=amount; costcredits[int(acc)]=costcredits.get(int(acc),Decimal('0.00'))+amount
            tx.execute("INSERT INTO assembly_costs(assembly_id,account_id,description,amount) VALUES(?,?,?,?)",(aid,acc,desc or ar['name'],str(amount)))
        total=material_total+additional_total
        tx.execute("UPDATE assembly_orders SET assembly_date=?,warehouse_id=?,wip_account_id=?,notes=?,material_cost=?,additional_cost=?,total_cost=?,updated_at=? WHERE id=?",(assembly_date,warehouse_id,wip_account_id,str(d.get('notes') or '').strip() or None,str(material_total),str(additional_total),str(total),now,aid))
        lines=[{'account_id':wip_account_id,'debit':total}]+[{'account_id':k,'credit':v} for k,v in invcredits.items() if v>0]+[{'account_id':k,'credit':v} for k,v in costcredits.items() if v>0]
        if total>0: accounting_service.post_journal(tx,journal_date=assembly_date,description=f"Assembly tahap 1 {h['assembly_no']} - revisi bahan dan biaya ke WIP",source_type='ASSEMBLY_ISSUE',source_id=aid,reference_no=h['assembly_no'],lines=lines,user_id=actor['id'])
        audit(actor['id'],'ASSEMBLY_UPDATED','assembly',aid,{'assembly_no':h['assembly_no'],'total_cost':str(total)},ip,tx)
        return {'id':aid,'assembly_no':h['assembly_no'],'total_cost':float(total)}


def delete_assembly(actor, assembly_id, ip):
    aid=int(assembly_id); now=utc_now()
    with write_transaction() as tx:
        h=tx.execute("SELECT * FROM assembly_orders WHERE id=?",(aid,)).fetchone()
        if not h or h['status']=='VOID': raise ValueError('Assembly tidak ditemukan atau sudah dibatalkan.')
        if h['status']=='FINISHED': _reverse_assembly_finish(tx,actor,h,now); h=tx.execute("SELECT * FROM assembly_orders WHERE id=?",(aid,)).fetchone()
        _reverse_assembly_issue(tx,actor,h,now)
        tx.execute("UPDATE assembly_orders SET status='VOID',updated_at=? WHERE id=?",(now,aid))
        audit(actor['id'],'ASSEMBLY_VOIDED','assembly',aid,{'assembly_no':h['assembly_no']},ip,tx)
        return {'id':aid,'status':'VOID'}


def update_assembly_finish(actor, assembly_id, d, ip):
    aid=int(assembly_id); now=utc_now()
    with write_transaction() as tx:
        h=tx.execute("SELECT * FROM assembly_orders WHERE id=?",(aid,)).fetchone()
        if not h or h['status']!='FINISHED': raise ValueError('Finishing assembly tidak ditemukan.')
        _reverse_assembly_finish(tx,actor,h,now)
    return finish_assembly(actor,aid,d,ip)


def delete_assembly_finish(actor, assembly_id, ip):
    aid=int(assembly_id); now=utc_now()
    with write_transaction() as tx:
        h=tx.execute("SELECT * FROM assembly_orders WHERE id=?",(aid,)).fetchone()
        if not h or h['status']!='FINISHED': raise ValueError('Finishing assembly tidak ditemukan.')
        _reverse_assembly_finish(tx,actor,h,now)
        audit(actor['id'],'ASSEMBLY_FINISH_VOIDED','assembly',aid,{'assembly_no':h['assembly_no']},ip,tx)
        return {'id':aid,'status':'OPEN'}


def list_project_cost_accounts():
    c=connect()
    try:
        return [dict(r) for r in c.execute("""SELECT id,code,name,account_type,account_subtype
          FROM chart_of_accounts WHERE is_active=1
          AND (account_type='EXPENSE' OR account_subtype='HPP') ORDER BY code""").fetchall()]
    finally:c.close()

def list_coa(active_only=False):
    c=connect()
    try:
        sql="SELECT * FROM chart_of_accounts"+(" WHERE is_active=1" if active_only else "")+" ORDER BY code"
        out=[]
        for r in c.execute(sql).fetchall():
            x=dict(r);x['display_type']=x.get('account_subtype') or x['account_type'];out.append(x)
        return out
    finally:c.close()

def _sync_cash_account_for_coa(tx, account_id, code, name, subtype, is_active=1, user_id=1):
    """Pastikan setiap COA CASH_BANK mempunyai pasangan rekening operasional Kas/Bank.
    Saldo rekening baru diselaraskan dari saldo GL tanpa membuat jurnal ganda.
    """
    account_id=int(account_id); subtype=str(subtype or '').upper(); active=1 if is_active else 0
    linked=tx.execute("SELECT * FROM cash_accounts WHERE coa_account_id=? ORDER BY id LIMIT 1",(account_id,)).fetchone()
    if subtype!='CASH_BANK':
        if linked:
            tx.execute("UPDATE cash_accounts SET is_active=0,updated_at=? WHERE id=?",(utc_now(),linked['id']))
        return None
    account_type='BANK' if ('BANK' in str(name).upper() or 'BANK' in str(code).upper()) else 'CASH'
    now=utc_now()
    if linked:
        tx.execute("UPDATE cash_accounts SET code=?,name=?,account_type=?,is_active=?,updated_at=? WHERE id=?",
                   (str(code).strip().upper(),str(name).strip(),account_type,active,now,linked['id']))
        return int(linked['id'])
    # Hindari bentrok kode rekening operasional lama yang tidak tertaut ke COA ini.
    cash_code=str(code).strip().upper()
    exists=tx.execute("SELECT id FROM cash_accounts WHERE code=? COLLATE NOCASE",(cash_code,)).fetchone()
    if exists:
        cash_code=f"{cash_code}-GL{account_id}"
    cur=tx.execute("""INSERT INTO cash_accounts(code,name,account_type,bank_name,account_number,opening_balance,current_balance,is_active,created_at,updated_at,coa_account_id)
      VALUES(?,?,?,NULL,NULL,0,0,?,?,?,?)""",(cash_code,str(name).strip(),account_type,active,now,now,account_id))
    cash_id=int(cur.lastrowid)
    # Jika COA sudah memiliki histori GL (mis. diubah dari ASSET menjadi CASH_BANK),
    # cerminkan saldo bersihnya ke subledger kas tanpa posting jurnal baru.
    bal=tx.execute("""SELECT COALESCE(SUM(CASE WHEN je.status='POSTED' THEN CAST(jl.debit AS REAL)-CAST(jl.credit AS REAL) ELSE 0 END),0) v
      FROM journal_lines jl JOIN journal_entries je ON je.id=jl.journal_id WHERE jl.account_id=?""",(account_id,)).fetchone()
    amount=Decimal(str(bal['v'] or 0))
    if amount!=0:
        cash_service.post(tx,account_id=cash_id,transaction_date=now[:10],transaction_type='IN' if amount>0 else 'OUT',amount=abs(amount),
          description='Sinkronisasi saldo awal dari Buku Besar',reference_no=f'COA-SYNC:{account_id}',user_id=int(user_id or 1),allow_negative=True)
    return cash_id


def create_coa(actor,d,ip):
    code=str(d.get('code','')).strip();name=str(d.get('name','')).strip();subtype=str(d.get('account_subtype') or d.get('account_type') or '').upper()
    if not code or not name: raise ValueError('Kode dan nama akun wajib diisi.')
    mapping={'CASH_BANK':('ASSET','DEBIT'),'RECEIVABLE':('ASSET','DEBIT'),'ASSET':('ASSET','DEBIT'),'PAYABLE':('LIABILITY','CREDIT'),'LIABILITY':('LIABILITY','CREDIT'),'EQUITY':('EQUITY','CREDIT'),'REVENUE':('REVENUE','CREDIT'),'HPP':('EXPENSE','DEBIT'),'OPERATING_EXPENSE':('EXPENSE','DEBIT')}
    if subtype not in mapping: raise ValueError('Tipe akun tidak valid.')
    atype,normal=mapping[subtype];now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("INSERT INTO chart_of_accounts(code,name,account_type,account_subtype,normal_balance,parent_id,is_system,is_active,created_at,updated_at) VALUES(?,?,?,?,?,?,0,1,?,?)",(code,name,atype,subtype,normal,d.get('parent_id') or None,now,now))
        _sync_cash_account_for_coa(tx,cur.lastrowid,code,name,subtype,1,actor['id'])
        audit(actor['id'],'COA_CREATED','coa',cur.lastrowid,{'code':code,'name':name,'account_subtype':subtype},ip,tx);return cur.lastrowid

def update_coa(actor, account_id, d, ip):
    account_id=int(account_id); code=str(d.get('code','')).strip(); name=str(d.get('name','')).strip(); subtype=str(d.get('account_subtype') or d.get('account_type') or '').upper()
    if not code or not name: raise ValueError('Kode dan nama akun wajib diisi.')
    mapping={'CASH_BANK':('ASSET','DEBIT'),'RECEIVABLE':('ASSET','DEBIT'),'ASSET':('ASSET','DEBIT'),'PAYABLE':('LIABILITY','CREDIT'),'LIABILITY':('LIABILITY','CREDIT'),'EQUITY':('EQUITY','CREDIT'),'REVENUE':('REVENUE','CREDIT'),'HPP':('EXPENSE','DEBIT'),'OPERATING_EXPENSE':('EXPENSE','DEBIT')}
    if subtype not in mapping: raise ValueError('Tipe akun tidak valid.')
    atype,normal=mapping[subtype]
    with write_transaction() as tx:
        row=tx.execute('SELECT * FROM chart_of_accounts WHERE id=?',(account_id,)).fetchone()
        if not row: raise ValueError('Akun tidak ditemukan.')
        if tx.execute('SELECT 1 FROM chart_of_accounts WHERE code=? AND id<>?',(code,account_id)).fetchone(): raise ValueError('Kode akun sudah digunakan.')
        active=1 if d.get('is_active',True) else 0
        tx.execute('UPDATE chart_of_accounts SET code=?,name=?,account_type=?,account_subtype=?,normal_balance=?,parent_id=?,is_active=?,updated_at=? WHERE id=?',(code,name,atype,subtype,normal,d.get('parent_id') or None,active,utc_now(),account_id))
        _sync_cash_account_for_coa(tx,account_id,code,name,subtype,active,actor['id'])
        audit(actor['id'],'COA_UPDATED','coa',account_id,{'code':code,'name':name,'account_subtype':subtype},ip,tx)
    return account_id

def delete_coa(actor, account_id, ip):
    account_id=int(account_id)
    with write_transaction() as tx:
        row=tx.execute('SELECT * FROM chart_of_accounts WHERE id=?',(account_id,)).fetchone()
        if not row: raise ValueError('Akun tidak ditemukan.')
        used=tx.execute('SELECT 1 FROM journal_lines WHERE account_id=? LIMIT 1',(account_id,)).fetchone()
        refs=[]
        for table,col in [('products','inventory_account_id'),('products','sales_account_id'),('products','cogs_account_id'),('cash_accounts','coa_account_id')]:
            try:
                if tx.execute(f'SELECT 1 FROM {table} WHERE {col}=? LIMIT 1',(account_id,)).fetchone(): refs.append(table)
            except Exception: pass
        if used or refs:
            tx.execute('UPDATE chart_of_accounts SET is_active=0,updated_at=? WHERE id=?',(utc_now(),account_id)); status='deactivated'; action='COA_DEACTIVATED'
        else:
            tx.execute('DELETE FROM chart_of_accounts WHERE id=?',(account_id,)); status='deleted'; action='COA_DELETED'
        audit(actor['id'],action,'coa',account_id,{'code':row['code'],'name':row['name']},ip,tx)
    return {'status':status}

def _sync_manual_journal_cash(tx, journal_id):
    """Mirror baris kas/bank jurnal manual ke cash_transactions tanpa membuat GL kedua."""
    journal_id=int(journal_id); ref=f"MANUAL-JOURNAL:{journal_id}"
    affected={int(r["account_id"]) for r in tx.execute("SELECT DISTINCT account_id FROM cash_transactions WHERE reference_no=?",(ref,)).fetchall()}
    tx.execute("DELETE FROM cash_transactions WHERE reference_no=?",(ref,))
    for aid in affected:
        _rebuild_cash_running_balance(tx,aid)
    j=tx.execute("SELECT journal_date,description,status,user_id FROM journal_entries WHERE id=?",(journal_id,)).fetchone()
    if not j or j["status"]!="POSTED": return
    # Repair database lama: setiap COA CASH_BANK yang dipakai jurnal manual wajib mempunyai
    # rekening operasional agar saldo Penjualan/Pembelian/Kas & Bank ikut bergerak.
    for coa in tx.execute("""SELECT DISTINCT coa.id,coa.code,coa.name,coa.account_subtype,coa.is_active
      FROM journal_lines l JOIN chart_of_accounts coa ON coa.id=l.account_id
      WHERE l.journal_id=? AND coa.account_subtype='CASH_BANK'""",(journal_id,)).fetchall():
        _sync_cash_account_for_coa(tx,coa['id'],coa['code'],coa['name'],coa['account_subtype'],coa['is_active'],j['user_id'])
    rows=tx.execute("""SELECT l.account_id,l.debit,l.credit,ca.id cash_account_id
      FROM journal_lines l JOIN cash_accounts ca ON ca.coa_account_id=l.account_id AND ca.is_active=1
      WHERE l.journal_id=? ORDER BY l.id,ca.id""",(journal_id,)).fetchall()
    # Satu COA Kas/Bank idealnya ditautkan ke satu rekening operasional. Bila ada duplikat,
    # gunakan rekening pertama secara deterministik agar tidak menggandakan saldo.
    seen=set()
    for r in rows:
        coa=int(r["account_id"])
        if coa in seen: continue
        seen.add(coa); debit=Decimal(str(r["debit"] or 0)); credit=Decimal(str(r["credit"] or 0)); delta=debit-credit
        if delta==0: continue
        aid=int(r["cash_account_id"])
        cash_service.post(tx,account_id=aid,transaction_date=j["journal_date"],
            transaction_type="IN" if delta>0 else "OUT",amount=abs(delta),
            description=f"Jurnal manual: {j['description']}",reference_no=ref,user_id=j["user_id"],allow_negative=True)

def create_manual_journal(actor,d,ip):
    journal_date=str(d.get('journal_date',''))[:10];desc=str(d.get('description','')).strip();lines=d.get('lines')
    if len(journal_date)!=10 or not desc: raise ValueError('Tanggal dan keterangan jurnal wajib diisi.')
    with write_transaction() as tx:
        _block_manual_inventory_accounts(tx,lines)
        department_id,project_id=_transaction_dimensions(tx,d)
        # Tampilkan dan simpan nomor otomatis sebagai referensi bila user tidak mengubahnya.
        # journal_no ditentukan di transaksi yang sama agar preview/reference dan nomor bukti konsisten.
        journal_no=accounting_service.next_journal_no(tx,journal_date)
        reference_no=str(d.get('reference_no','')).strip() or journal_no
        result=accounting_service.post_journal(tx,journal_date=journal_date,description=desc,source_type='MANUAL',source_id=None,reference_no=reference_no,lines=lines,user_id=actor['id'],department_id=department_id,project_id=project_id,journal_no=journal_no)
        _sync_manual_journal_cash(tx,result['id'])
        audit(actor['id'],'MANUAL_JOURNAL_CREATED','journal',result['id'],{'journal_no':result['journal_no']},ip,tx);return result


def update_manual_journal(actor,journal_id,d,ip):
    journal_id=int(journal_id)
    journal_date=str(d.get('journal_date',''))[:10]
    desc=str(d.get('description','')).strip()
    lines=d.get('lines') or []
    if len(journal_date)!=10 or not desc: raise ValueError('Tanggal dan keterangan jurnal wajib diisi.')
    if not lines: raise ValueError('Minimal satu baris jurnal wajib diisi.')
    with write_transaction() as tx:
        current=tx.execute("SELECT * FROM journal_entries WHERE id=? AND source_type IN ('MANUAL','MANUAL_EXCEL') AND status='POSTED'",(journal_id,)).fetchone()
        if not current: raise ValueError('Jurnal manual tidak ditemukan atau sudah dibatalkan.')
        # Validate and normalize through accounting service logic, but retain the original number/id.
        normalized=[]; total_debit=Decimal('0'); total_credit=Decimal('0')
        for line in lines:
            aid=int(line.get('account_id'))
            account_row=tx.execute('SELECT account_type,COALESCE(account_subtype,account_type) subtype FROM chart_of_accounts WHERE id=? AND is_active=1',(aid,)).fetchone()
            if not account_row:
                raise ValueError('Akun jurnal tidak valid atau nonaktif.')
            subtype=str(account_row['subtype'] or '').upper()
            partner_id=line.get('partner_id') or None
            if subtype in ('RECEIVABLE','PAYABLE'):
                if not partner_id:
                    raise ValueError('Pelanggan/Pemasok wajib dipilih untuk akun Piutang atau Hutang.')
                allowed='CUSTOMER' if subtype=='RECEIVABLE' else 'SUPPLIER'
                partner=tx.execute("SELECT partner_type FROM business_partners WHERE id=? AND is_active=1",(int(partner_id),)).fetchone()
                if not partner or partner['partner_type'] not in (allowed,'BOTH'):
                    raise ValueError('Partner tidak sesuai dengan tipe akun Piutang/Hutang.')
            if aid in _inventory_control_account_ids(tx):
                raise ValueError('Akun Persediaan yang terhubung ke barang tidak boleh diedit lewat Jurnal Manual. Gunakan modul stok.')
            debit=Decimal(str(_decimal(line.get('debit',0),'Debit'))); credit=Decimal(str(_decimal(line.get('credit',0),'Kredit')))
            if debit<0 or credit<0 or (debit>0 and credit>0) or (debit==0 and credit==0):
                raise ValueError('Setiap baris harus berisi debit atau kredit saja.')
            total_debit+=debit; total_credit+=credit
            normalized.append((aid,partner_id,line.get('department_id') or None,line.get('project_id') or None,str(debit),str(credit),str(line.get('memo') or '').strip() or None))
        if total_debit<=0 or total_debit!=total_credit: raise ValueError('Total debit dan kredit harus sama dan lebih dari nol.')
        now=utc_now()
        tx.execute('DELETE FROM journal_lines WHERE journal_id=?',(journal_id,))
        tx.execute('''UPDATE journal_entries SET journal_date=?,description=?,reference_no=?,department_id=?,project_id=?,total_debit=?,total_credit=?,updated_at=? WHERE id=?''',
            (journal_date,desc,str(d.get('reference_no','')).strip() or None,d.get('department_id') or None,d.get('project_id') or None,str(total_debit),str(total_credit),now,journal_id))
        for aid,partner,department,project,debit,credit,memo in normalized:
            tx.execute('''INSERT INTO journal_lines(journal_id,account_id,partner_id,department_id,project_id,debit,credit,memo) VALUES(?,?,?,?,?,?,?,?)''',
                (journal_id,aid,partner,department,project,debit,credit,memo))
        _sync_manual_journal_cash(tx,journal_id)
        audit(actor['id'],'MANUAL_JOURNAL_UPDATED','journal',journal_id,{'journal_no':current['journal_no']},ip,tx)
        return {'id':journal_id,'journal_no':current['journal_no'],'total_debit':float(total_debit),'total_credit':float(total_credit)}



def _excel_bytes(headers,examples,sheet_name):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
    except ImportError as exc:
        raise ValueError("Fitur Excel membutuhkan openpyxl. Jalankan INSTALL_REPORT_SUPPORT.bat.") from exc
    wb=Workbook();ws=wb.active;ws.title=sheet_name
    ws.append(headers)
    for row in examples:ws.append(row)
    for cell in ws[1]:
        cell.font=Font(bold=True);cell.fill=PatternFill("solid",fgColor="DCEFF1")
        cell.alignment=Alignment(horizontal="center")
    for i in range(1,len(headers)+1):
        values=[str(ws.cell(r,i).value or "") for r in range(1,ws.max_row+1)]
        ws.column_dimensions[chr(64+i)].width=min(36,max(14,max(map(len,values))+2))
    ws.freeze_panes="A2"
    bio=io.BytesIO();wb.save(bio);return bio.getvalue()

def sales_excel_template():
    return _excel_bytes(
      ["Tanggal","No Invoice","No Surat Jalan","Kode Pelanggan","Kode Salesman",
       "Kode Gudang","Metode Pembayaran","Tanggal Jatuh Tempo",
       "Kode Barang","Kode Satuan","Qty","Harga","Diskon Persen","Pajak Persen","Catatan","Kode Kas/Bank"],
      [["2026-07-18","INV-EX-001","SJ-EX-001","CUST001","",
        "GDG01","CREDIT","2026-08-17","BRG001","PCS",2,150000,0,11,"Contoh impor",""]],
      "Penjualan")

def purchase_excel_template():
    return _excel_bytes(
      ["Tanggal","No Invoice Supplier","No Good Receive","Kode Pemasok",
       "Kode Gudang","Metode Pembayaran","Tanggal Jatuh Tempo",
       "Kode Barang","Kode Satuan","Qty","Harga Beli","Diskon Persen","Pajak Persen","Catatan","Kode Kas/Bank"],
      [["2026-07-18","SUP-INV-001","GR-EX-001","SUP001",
        "GDG01","CREDIT","2026-08-17","BRG001","PCS",2,100000,0,11,"Contoh impor",""]],
      "Pembelian")

def cash_excel_template():
    return _excel_bytes(
      ["Tanggal","No Bukti","Jenis","Kode Kas/Bank","Kode Akun Lawan","Nominal","Keterangan","Kode Proyek","Kode Departemen"],
      [["2026-07-18","KM-001","IN","KAS01","4000",1000000,"Penerimaan lain-lain","PRJ001","DPT001"],
       ["2026-07-18","KK-001","OUT","KAS01","5100",250000,"Biaya operasional","PRJ001","DPT001"]],
      "Kas Masuk Keluar")

def _normalize_excel_header(value):
    import re
    text=str(value or '').strip().lower().replace('\n',' ')
    text=re.sub(r'[^a-z0-9]+',' ',text).strip()
    aliases={
      'tanggal transaksi':'tanggal','tgl':'tanggal','tanggal invoice':'tanggal',
      'nomor invoice':'no invoice','invoice no':'no invoice','no faktur':'no invoice',
      'nomor surat jalan':'no surat jalan','surat jalan':'no surat jalan',
      'pelanggan':'kode pelanggan','customer':'kode pelanggan','kode customer':'kode pelanggan',
      'salesman':'kode salesman','sales':'kode salesman','gudang':'kode gudang','warehouse':'kode gudang',
      'metode bayar':'metode pembayaran','cara bayar':'metode pembayaran','jatuh tempo':'tanggal jatuh tempo','due date':'tanggal jatuh tempo',
      'sku':'kode barang','barang':'kode barang','kode produk':'kode barang','quantity':'qty','jumlah':'qty',
      'harga jual':'harga','unit price':'harga','diskon':'diskon persen','discount persen':'diskon persen','discount percent':'diskon persen',
      'ppn':'pajak persen','ppn persen':'pajak persen','tax persen':'pajak persen','tax percent':'pajak persen','keterangan':'catatan','notes':'catatan',
      'nomor invoice supplier':'no invoice supplier','invoice supplier':'no invoice supplier','no faktur supplier':'no invoice supplier',
      'nomor good receive':'no good receive','good receive':'no good receive','no penerimaan barang':'no good receive',
      'pemasok':'kode pemasok','supplier':'kode pemasok','kode supplier':'kode pemasok','harga':'harga beli','unit cost':'harga beli'}
    return aliases.get(text,text)

def _load_excel_rows(excel_base64,expected_headers):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError('Fitur Excel membutuhkan openpyxl. Jalankan INSTALL_REPORT_SUPPORT.bat.') from exc
    try:
        raw=base64.b64decode(excel_base64,validate=True)
        wb=load_workbook(io.BytesIO(raw),data_only=True);ws=wb.active
    except Exception as exc:
        raise ValueError('File Excel tidak valid atau rusak.') from exc
    actual=[_normalize_excel_header(x.value) for x in ws[1]]
    expected=[_normalize_excel_header(x) for x in expected_headers]
    missing=[expected_headers[i] for i,h in enumerate(expected) if h not in actual]
    if missing: raise ValueError('Format kolom Excel tidak sesuai template. Kolom belum ada: '+', '.join(missing))
    indices=[actual.index(h) for h in expected]
    out=[]
    for rawrow in ws.iter_rows(min_row=2,values_only=True):
        if any(v not in (None,'') for v in rawrow): out.append([rawrow[i] if i<len(rawrow) else None for i in indices])
    return out

def _resolve_cash_account_code(conn, code):
    """Resolve Kode Kas/Bank from either operational cash code or linked COA code.
    Template users commonly know COA 1000/1010, while older importer only accepted cash_accounts.code.
    """
    code=str(code or "").strip()
    if not code:return None
    row=conn.execute("SELECT id FROM cash_accounts WHERE code=? COLLATE NOCASE AND is_active=1 ORDER BY id LIMIT 1",(code,)).fetchone()
    if row:return row
    row=conn.execute("""SELECT ca.id FROM chart_of_accounts coa
      JOIN cash_accounts ca ON ca.coa_account_id=coa.id AND ca.is_active=1
      WHERE coa.code=? COLLATE NOCASE AND coa.is_active=1 AND coa.account_subtype='CASH_BANK'
      ORDER BY ca.id LIMIT 1""",(code,)).fetchone()
    return row

def _normalize_import_payment_method(value):
    raw=str(value or 'CREDIT').strip().upper().replace('_',' ').replace('-',' ')
    raw=' '.join(raw.split())
    mapping={'TUNAI':'CASH','CASH':'CASH','KAS':'CASH','TRANSFER':'TRANSFER','BANK TRANSFER':'TRANSFER','TRANSFER BANK':'TRANSFER','KREDIT':'CREDIT','CREDIT':'CREDIT','PIUTANG':'CREDIT','HUTANG':'CREDIT'}
    method=mapping.get(raw,raw)
    if method not in ('CASH','TRANSFER','CREDIT'):
        raise ValueError("Metode pembayaran harus Tunai, Transfer, atau Kredit.")
    return method


def _resolve_import_unit(conn, product_id, code):
    code=str(code or '').strip()
    p=conn.execute("SELECT unit_id FROM products WHERE id=?",(int(product_id),)).fetchone()
    if not p:return None
    if not code:return p['unit_id']
    row=conn.execute("""SELECT u.id FROM units u WHERE u.code=? COLLATE NOCASE AND
      (u.id=? OR EXISTS(SELECT 1 FROM product_units pu WHERE pu.product_id=? AND pu.unit_id=u.id)) LIMIT 1""",(code,p['unit_id'],int(product_id))).fetchone()
    return row['id'] if row else None

def _ensure_no_import_duplicate(conn, table, column, value, label):
    value=str(value or '').strip()
    if value and conn.execute(f"SELECT 1 FROM {table} WHERE {column}=? COLLATE NOCASE LIMIT 1",(value,)).fetchone():
        raise ValueError(f"{label} '{value}' sudah ada di sistem.")

def import_sales_excel(actor,excel_base64,file_name,ip):
    headers=["Tanggal","No Invoice","No Surat Jalan","Kode Pelanggan","Kode Salesman",
      "Kode Gudang","Metode Pembayaran","Tanggal Jatuh Tempo","Kode Barang","Kode Satuan","Qty",
      "Harga","Diskon Persen","Pajak Persen","Catatan","Kode Kas/Bank"]
    rows=_load_excel_rows(excel_base64,headers);groups={};seen_delivery=set()
    c=connect()
    try:
      for no,row in enumerate(rows,2):
        tanggal=str(row[0] or "")[:10];invoice=str(row[1] or "").strip();delivery=str(row[2] or '').strip()
        if not tanggal or not invoice:raise ValueError(f"Baris {no}: Tanggal dan No Invoice wajib.")
        if invoice not in groups:_ensure_no_import_duplicate(c,'sales','invoice_no',invoice,'No Invoice')
        if delivery:
            if delivery.lower() in seen_delivery and groups.get(invoice,{}).get('delivery_no')!=delivery:raise ValueError(f"Baris {no}: No Surat Jalan '{delivery}' duplikat di file impor.")
            if delivery.lower() not in seen_delivery:_ensure_no_import_duplicate(c,'sales','delivery_no',delivery,'No Surat Jalan')
            seen_delivery.add(delivery.lower())
        customer=c.execute("SELECT id FROM business_partners WHERE code=? COLLATE NOCASE AND partner_type IN ('CUSTOMER','BOTH') AND is_active=1",(str(row[3] or "").strip(),)).fetchone()
        if not customer:raise ValueError(f"Baris {no}: Kode pelanggan tidak valid.")
        salesperson=None
        if str(row[4] or "").strip():
            salesperson=c.execute("SELECT id FROM salespersons WHERE code=? COLLATE NOCASE AND is_active=1",(str(row[4]).strip(),)).fetchone()
            if not salesperson:raise ValueError(f"Baris {no}: Kode salesman tidak valid.")
        wh=c.execute("SELECT id FROM warehouses WHERE code=? COLLATE NOCASE AND is_active=1",(str(row[5] or "").strip(),)).fetchone()
        if not wh:raise ValueError(f"Baris {no}: Kode gudang tidak valid.")
        product=c.execute("SELECT id FROM products WHERE sku=? COLLATE NOCASE AND is_active=1",(str(row[8] or "").strip(),)).fetchone()
        if not product:raise ValueError(f"Baris {no}: Kode barang tidak valid.")
        unit_id=_resolve_import_unit(c,product['id'],row[9])
        if not unit_id:raise ValueError(f"Baris {no}: Kode satuan tidak valid untuk barang tersebut.")
        method=_normalize_import_payment_method(row[6]); cash=None;cash_code=str(row[15] or "").strip()
        if cash_code:
            cash=_resolve_cash_account_code(c,cash_code)
            if not cash:raise ValueError(f"Baris {no}: Kode Kas/Bank tidak valid.")
        if method in ("CASH","TRANSFER") and not cash:raise ValueError(f"Baris {no}: Kode Kas/Bank wajib untuk transaksi Tunai/Transfer.")
        group=groups.setdefault(invoice,{"invoice_no":invoice,"delivery_no":delivery,"sale_date":tanggal,"customer_id":customer["id"],"salesperson_id":salesperson["id"] if salesperson else None,"warehouse_id":wh["id"],"payment_method":method,"cash_account_id":cash["id"] if cash else None,"due_date":str(row[7] or "")[:10] or None,"notes":str(row[14] or "").strip(),"tax_percent":row[13] or 0,"paid_amount":0,"items":[]})
        if Decimal(str(group.get("tax_percent") or 0)) != Decimal(str(row[13] or 0)):raise ValueError(f"Baris {no}: Pajak Persen untuk invoice {invoice} harus sama pada semua baris.")
        group["items"].append({"product_id":product["id"],"unit_id":unit_id,"qty":row[10] or 0,"unit_price":row[11] or 0,"discount_mode":"PERCENT","discount_value":row[12] or 0})
    finally:c.close()
    c=connect()
    try: licensing.enforce_transaction_capacity(c,len(groups))
    finally:c.close()
    created=[]
    for data in groups.values():
        if data.get("payment_method") in ("CASH","TRANSFER"):
            base=sum((Decimal(str(i.get("qty") or 0))*Decimal(str(i.get("unit_price") or 0))*(Decimal("1")-Decimal(str(i.get("discount_value") or 0))/Decimal("100")) for i in data["items"]),Decimal("0"));data["paid_amount"]=(base*(Decimal("1")+Decimal(str(data.get("tax_percent") or 0))/Decimal("100"))).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
        created.append(create_sale(actor,data,ip)["invoice_no"])
    return {"transaction_count":len(created),"documents":created}
def import_purchase_excel(actor,excel_base64,file_name,ip):
    headers=["Tanggal","No Invoice Supplier","No Good Receive","Kode Pemasok","Kode Gudang","Metode Pembayaran","Tanggal Jatuh Tempo","Kode Barang","Kode Satuan","Qty","Harga Beli","Diskon Persen","Pajak Persen","Catatan","Kode Kas/Bank"]
    rows=_load_excel_rows(excel_base64,headers);groups={};seen_gr=set();c=connect()
    try:
      for no,row in enumerate(rows,2):
        tanggal=str(row[0] or "")[:10];reference=str(row[1] or "").strip();gr=str(row[2] or '').strip()
        if not tanggal or not reference:raise ValueError(f"Baris {no}: Tanggal dan No Invoice Supplier wajib.")
        if reference not in groups:_ensure_no_import_duplicate(c,'purchases','supplier_invoice_no',reference,'No Invoice Supplier')
        if gr:
            if gr.lower() in seen_gr and groups.get(reference,{}).get('goods_receipt_no')!=gr:raise ValueError(f"Baris {no}: No Good Receive '{gr}' duplikat di file impor.")
            if gr.lower() not in seen_gr:_ensure_no_import_duplicate(c,'purchases','goods_receipt_no',gr,'No Good Receive')
            seen_gr.add(gr.lower())
        supplier=c.execute("SELECT id FROM business_partners WHERE code=? COLLATE NOCASE AND partner_type IN ('SUPPLIER','BOTH') AND is_active=1",(str(row[3] or "").strip(),)).fetchone()
        if not supplier:raise ValueError(f"Baris {no}: Kode pemasok tidak valid.")
        wh=c.execute("SELECT id FROM warehouses WHERE code=? COLLATE NOCASE AND is_active=1",(str(row[4] or "").strip(),)).fetchone()
        if not wh:raise ValueError(f"Baris {no}: Kode gudang tidak valid.")
        product=c.execute("SELECT id FROM products WHERE sku=? COLLATE NOCASE AND is_active=1",(str(row[7] or "").strip(),)).fetchone()
        if not product:raise ValueError(f"Baris {no}: Kode barang tidak valid.")
        unit_id=_resolve_import_unit(c,product['id'],row[8])
        if not unit_id:raise ValueError(f"Baris {no}: Kode satuan tidak valid untuk barang tersebut.")
        method=_normalize_import_payment_method(row[5]);cash=None;cash_code=str(row[14] or "").strip()
        if cash_code:
            cash=_resolve_cash_account_code(c,cash_code)
            if not cash:raise ValueError(f"Baris {no}: Kode Kas/Bank tidak valid.")
        if method in ("CASH","TRANSFER") and not cash:raise ValueError(f"Baris {no}: Kode Kas/Bank wajib untuk transaksi Tunai/Transfer.")
        group=groups.setdefault(reference,{"supplier_invoice_no":reference,"goods_receipt_no":gr,"purchase_date":tanggal,"supplier_id":supplier["id"],"warehouse_id":wh["id"],"payment_method":method,"cash_account_id":cash["id"] if cash else None,"due_date":str(row[6] or "")[:10] or None,"notes":str(row[13] or "").strip(),"tax_percent":row[12] or 0,"paid_amount":0,"items":[]})
        if Decimal(str(group.get("tax_percent") or 0)) != Decimal(str(row[12] or 0)):raise ValueError(f"Baris {no}: Pajak Persen untuk invoice supplier {reference} harus sama pada semua baris.")
        group["items"].append({"product_id":product["id"],"unit_id":unit_id,"qty":row[9] or 0,"unit_cost":row[10] or 0,"discount_mode":"PERCENT","discount_value":row[11] or 0})
    finally:c.close()
    c=connect()
    try: licensing.enforce_transaction_capacity(c,len(groups))
    finally:c.close()
    created=[]
    for data in groups.values():
        if data.get("payment_method") in ("CASH","TRANSFER"):
            base=sum((Decimal(str(i.get("qty") or 0))*Decimal(str(i.get("unit_cost") or 0))*(Decimal("1")-Decimal(str(i.get("discount_value") or 0))/Decimal("100")) for i in data["items"]),Decimal("0"));data["paid_amount"]=(base*(Decimal("1")+Decimal(str(data.get("tax_percent") or 0))/Decimal("100"))).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
        created.append(create_purchase(actor,data,ip)["purchase_no"])
    return {"transaction_count":len(created),"documents":created}
def import_cash_excel(actor,excel_base64,file_name,ip):
    headers=["Tanggal","No Bukti","Jenis","Kode Kas/Bank","Kode Akun Lawan","Nominal","Keterangan"]
    rows=_load_excel_rows(excel_base64,headers);created=[];prepared=[];seen_refs=set()
    c=connect()
    try:
      licensing.enforce_transaction_capacity(c,len(rows))
      for no,row in enumerate(rows,2):
        ref=str(row[1] or '').strip()
        if not ref:raise ValueError(f"Baris {no}: No Bukti wajib diisi.")
        if ref.lower() in seen_refs:raise ValueError(f"Baris {no}: No Bukti '{ref}' duplikat di file impor.")
        if c.execute("SELECT 1 FROM cash_transactions WHERE reference_no=? COLLATE NOCASE LIMIT 1",(ref,)).fetchone():raise ValueError(f"Baris {no}: No Bukti '{ref}' sudah ada di sistem.")
        seen_refs.add(ref.lower())
        cash=_resolve_cash_account_code(c,str(row[3] or '').strip())
        if not cash:raise ValueError(f"Baris {no}: Kode Kas/Bank tidak valid.")
        coa=c.execute("SELECT id,account_subtype FROM chart_of_accounts WHERE code=? COLLATE NOCASE AND is_active=1",(str(row[4] or "").strip(),)).fetchone()
        if not coa or coa["account_subtype"] in ("CASH_BANK","RECEIVABLE","PAYABLE"):
            raise ValueError(f"Baris {no}: Kode akun lawan tidak valid.")
        prepared.append({"transaction_date":str(row[0] or "")[:10],"reference_no":ref,"transaction_type":str(row[2] or "").strip().upper(),"account_id":cash["id"],"counter_account_id":coa["id"],"amount":row[5] or 0,"description":str(row[6] or "").strip()})
    finally:c.close()
    # koneksi baca sudah ditutup sebelum write_transaction agar SQLite tidak saling mengunci.
    for payload in prepared:
        result=create_cash_transaction(actor,payload,ip);created.append(result["transaction_no"])
    return {"transaction_count":len(created),"documents":created}

def manual_journal_excel_template():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
    except ImportError as exc:
        raise ValueError("Fitur Excel membutuhkan openpyxl. Jalankan INSTALL_REPORT_SUPPORT.bat.") from exc
    wb=Workbook();ws=wb.active;ws.title="Jurnal Manual"
    headers=["Tanggal","No Jurnal","Kode Akun","Keterangan","Debit","Kredit","Kode Pelanggan/Pemasok","Kode Proyek","Kode Departemen"]
    ws.append(headers)
    examples=[
      ["2026-07-18","JM-001","1100","Koreksi piutang pelanggan",1000000,0,"CUST001","PRJ001","DPT001"],
      ["2026-07-18","JM-001","4000","Koreksi pendapatan",0,1000000,"","PRJ001","DPT001"],
      ["2026-07-19","JM-002","5100","Biaya listrik",500000,0,"","PRJ001","DPT001"],
      ["2026-07-19","JM-002","1000","Pembayaran tunai",0,500000,"","PRJ001","DPT001"],
    ]
    for row in examples:ws.append(row)
    for cell in ws[1]:
        cell.font=Font(bold=True);cell.fill=PatternFill("solid",fgColor="DCEFF1")
        cell.alignment=Alignment(horizontal="center")
    widths=[14,16,14,36,16,16,28,18,18]
    for i,w in enumerate(widths,1):ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A2"
    bio=io.BytesIO();wb.save(bio)
    return bio.getvalue()

def import_manual_journal_excel(actor,excel_base64,file_name,ip):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Fitur Excel membutuhkan openpyxl. Jalankan INSTALL_REPORT_SUPPORT.bat.") from exc
    try:
        raw=base64.b64decode(excel_base64,validate=True)
        wb=load_workbook(io.BytesIO(raw),data_only=True)
    except Exception as exc:
        raise ValueError("File Excel tidak valid atau rusak.") from exc
    ws=wb.active
    expected=["tanggal","no jurnal","kode akun","keterangan","debit","kredit","kode pelanggan/pemasok","kode proyek","kode departemen"]
    old_project=expected[:-1];legacy=expected[:-2]
    headers=[str(x.value or "").strip().lower() for x in ws[1]]
    has_department=headers[:len(expected)]==expected
    has_project=(not has_department and headers[:len(old_project)]==old_project)
    if not has_department and not has_project and headers[:len(legacy)]!=legacy:
        raise ValueError("Format kolom Excel tidak sesuai template.")
    groups={}
    for row_no,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
        if not any(v not in (None,"") for v in row):continue
        tanggal=str(row[0] or "")[:10]
        no_jurnal=str(row[1] or "").strip()
        kode_akun=str(row[2] or "").strip()
        ket=str(row[3] or "").strip()
        try:debit=Decimal(str(row[4] or 0));kredit=Decimal(str(row[5] or 0))
        except Exception:raise ValueError(f"Baris {row_no}: Debit/Kredit tidak valid.")
        partner_code=str(row[6] or "").strip()
        project_code=str((row[7] if len(row)>7 else "") or "").strip()
        department_code=str((row[8] if len(row)>8 else "") or "").strip()
        if not tanggal or not no_jurnal or not kode_akun or not ket:
            raise ValueError(f"Baris {row_no}: Tanggal, No Jurnal, Kode Akun, dan Keterangan wajib.")
        if debit<0 or kredit<0 or (debit>0)==(kredit>0):
            raise ValueError(f"Baris {row_no}: isi tepat satu nilai Debit atau Kredit.")
        key=(tanggal,no_jurnal)
        groups.setdefault(key,{"description":ket,"lines":[],"project_code":project_code,"department_code":department_code})
        if groups[key].get("project_code")!=project_code:
            raise ValueError(f"Jurnal {no_jurnal}: Kode Proyek harus sama pada semua baris jurnal.")
        if groups[key].get("department_code")!=department_code:
            raise ValueError(f"Jurnal {no_jurnal}: Kode Departemen harus sama pada semua baris jurnal.")
        groups[key]["lines"].append({
          "account_code":kode_akun,"debit":debit,"credit":kredit,
          "memo":ket,"partner_code":partner_code,"row_no":row_no})
    if not groups:raise ValueError("File Excel tidak memiliki baris jurnal.")
    created=[]
    with write_transaction() as tx:
        for (tanggal,no_jurnal),group in groups.items():
            if tx.execute("SELECT 1 FROM journal_entries WHERE (journal_no=? COLLATE NOCASE OR reference_no=? COLLATE NOCASE) LIMIT 1",(no_jurnal,no_jurnal)).fetchone():
                raise ValueError(f"No Jurnal '{no_jurnal}' sudah ada di sistem.")
            lines=[];debit_total=Decimal("0");credit_total=Decimal("0")
            project_id=None;department_id=None
            if group.get("project_code"):
                project=tx.execute("SELECT id FROM projects WHERE code=? COLLATE NOCASE AND is_active=1",(group["project_code"],)).fetchone()
                if not project:raise ValueError(f"Jurnal {no_jurnal}: Kode proyek {group['project_code']} tidak valid.")
                project_id=project["id"]
            if group.get("department_code"):
                department=tx.execute("SELECT id FROM departments WHERE code=? COLLATE NOCASE AND is_active=1",(group["department_code"],)).fetchone()
                if not department:raise ValueError(f"Jurnal {no_jurnal}: Kode departemen {group['department_code']} tidak valid.")
                department_id=department["id"]
            for line in group["lines"]:
                account=tx.execute("""SELECT id,code,name,account_subtype FROM chart_of_accounts
                  WHERE code=? COLLATE NOCASE AND is_active=1""",(line["account_code"],)).fetchone()
                if not account:raise ValueError(f"Baris {line['row_no']}: Kode akun {line['account_code']} tidak ditemukan.")
                if int(account["id"]) in _inventory_control_account_ids(tx):
                    raise ValueError(f"Baris {line['row_no']}: akun {account['code']} adalah akun kontrol Persediaan. Gunakan impor/transaksi stok, bukan jurnal manual.")
                subtype=account["account_subtype"] or ""
                partner_id=None
                if subtype in ("RECEIVABLE","PAYABLE"):
                    if not line["partner_code"]:
                        raise ValueError(f"Baris {line['row_no']}: Kode Pelanggan/Pemasok wajib untuk akun {account['name']}.")
                    required="CUSTOMER" if subtype=="RECEIVABLE" else "SUPPLIER"
                    partner=tx.execute("""SELECT id,partner_type,name FROM business_partners
                      WHERE code=? COLLATE NOCASE AND is_active=1""",(line["partner_code"],)).fetchone()
                    if not partner or partner["partner_type"] not in (required,"BOTH"):
                        label="pelanggan" if required=="CUSTOMER" else "pemasok"
                        raise ValueError(f"Baris {line['row_no']}: kode {label} {line['partner_code']} tidak valid.")
                    partner_id=partner["id"]
                debit_total+=line["debit"];credit_total+=line["credit"]
                lines.append({"account_id":account["id"],"debit":line["debit"],
                  "credit":line["credit"],"memo":line["memo"],"partner_id":partner_id})
            if debit_total!=credit_total:
                raise ValueError(f"Jurnal {no_jurnal}: total debit dan kredit tidak seimbang.")
            result=accounting_service.post_journal(tx,journal_date=tanggal,
              description=group["description"],source_type="MANUAL_EXCEL",
              source_id=None,reference_no=no_jurnal,lines=lines,user_id=actor["id"],project_id=project_id,department_id=department_id)
            _sync_manual_journal_cash(tx,result['id'])
            created.append(result["journal_no"])
        audit(actor["id"],"MANUAL_JOURNAL_EXCEL_IMPORTED","journal_excel",file_name,
              {"journal_count":len(created),"journals":created},ip,tx)
    return {"journal_count":len(created),"journals":created}

def list_journals(q='',date_from=None,date_to=None,limit=300):
    where=[];params=[]
    if q: where.append('(j.journal_no LIKE ? OR j.description LIKE ? OR COALESCE(j.reference_no,\'\') LIKE ?)');like=f'%{q}%';params += [like,like,like]
    if date_from:where.append('j.journal_date>=?');params.append(str(date_from)[:10])
    if date_to:where.append('j.journal_date<=?');params.append(str(date_to)[:10])
    sql="SELECT j.*,u.username FROM journal_entries j JOIN users u ON u.id=j.user_id"
    if where:sql+=' WHERE '+' AND '.join(where)
    sql+=' ORDER BY j.journal_date DESC,j.id DESC LIMIT ?';params.append(max(1,min(int(limit),1000)))
    c=connect()
    try:return [dict(r) for r in c.execute(sql,params).fetchall()]
    finally:c.close()

def get_journal(jid):
    c=connect()
    try:
        h=c.execute('SELECT j.*,u.username FROM journal_entries j JOIN users u ON u.id=j.user_id WHERE j.id=?',(jid,)).fetchone()
        if not h:return None
        out=dict(h);out['lines']=[dict(r) for r in c.execute("SELECT l.*,a.code account_code,a.name account_name FROM journal_lines l JOIN chart_of_accounts a ON a.id=l.account_id WHERE l.journal_id=? ORDER BY l.id",(jid,)).fetchall()];return out
    finally:c.close()


def _date_filter(value, label):
    if value in (None, ''):
        return None
    value=str(value)[:10]
    if len(value)!=10:
        raise ValueError(f'{label} tidak valid.')
    return value

def general_ledger(account_id,date_from=None,date_to=None,partner_id=None):
    try:account_id=int(account_id)
    except Exception:raise ValueError('Akun wajib dipilih.')
    date_from=_date_filter(date_from,'Tanggal mulai');date_to=_date_filter(date_to,'Tanggal akhir');partner_id=int(partner_id) if partner_id not in (None,'') else None
    c=connect()
    try:
        account=c.execute('SELECT * FROM chart_of_accounts WHERE id=? AND is_active=1',(account_id,)).fetchone()
        if not account:raise ValueError('Akun tidak ditemukan atau nonaktif.')
        op_sql="SELECT COALESCE(SUM(l.debit),0) debit,COALESCE(SUM(l.credit),0) credit FROM journal_lines l JOIN journal_entries j ON j.id=l.journal_id WHERE l.account_id=? AND j.status='POSTED'";op_p=[account_id]
        if partner_id:op_sql+=' AND l.partner_id=?';op_p.append(partner_id)
        if date_from:op_sql+=' AND j.journal_date<?';op_p.append(date_from)
        else:op_sql+=' AND 1=0'
        op=c.execute(op_sql,op_p).fetchone();od=float(op['debit']);oc=float(op['credit']);opening=od-oc if account['normal_balance']=='DEBIT' else oc-od
        sql="SELECT j.journal_date,j.journal_no,j.description,j.reference_no,l.debit,l.credit,l.memo,l.partner_id,bp.code partner_code,bp.name partner_name FROM journal_lines l JOIN journal_entries j ON j.id=l.journal_id LEFT JOIN business_partners bp ON bp.id=l.partner_id WHERE l.account_id=? AND j.status='POSTED'";params=[account_id]
        if partner_id:sql+=' AND l.partner_id=?';params.append(partner_id)
        if date_from:sql+=' AND j.journal_date>=?';params.append(date_from)
        if date_to:sql+=' AND j.journal_date<=?';params.append(date_to)
        sql+=' ORDER BY j.journal_date,j.id,l.id';run=opening;items=[];pd=pc=0.0
        for r in c.execute(sql,params).fetchall():
            d=float(r['debit']);cr=float(r['credit']);run+=d-cr if account['normal_balance']=='DEBIT' else cr-d;pd+=d;pc+=cr;x=dict(r);x['debit']=d;x['credit']=cr;x['running_balance']=round(run,2);items.append(x)
        partner=None
        if partner_id:
            q=c.execute('SELECT id,code,name,partner_type FROM business_partners WHERE id=?',(partner_id,)).fetchone();partner=dict(q) if q else None
        return {'account':dict(account),'partner':partner,'date_from':date_from,'date_to':date_to,'opening_balance':round(opening,2),'opening_debit':od,'opening_credit':oc,'period_debit':round(pd,2),'period_credit':round(pc,2),'ending_balance':round(run,2),'items':items}
    finally:c.close()

def trial_balance(date_from=None,date_to=None):
    date_from=_date_filter(date_from,'Tanggal mulai')
    date_to=_date_filter(date_to,'Tanggal akhir')
    if date_from and date_to and date_from>date_to:
        raise ValueError('Tanggal mulai tidak boleh melewati tanggal akhir.')
    c=connect()
    try:
        accounts=c.execute('SELECT * FROM chart_of_accounts WHERE is_active=1 ORDER BY code').fetchall()
        items=[];total_opening_debit=total_opening_credit=total_period_debit=total_period_credit=0.0
        total_ending_debit=total_ending_credit=0.0
        for a in accounts:
            op_sql="""SELECT COALESCE(SUM(l.debit),0) debit,COALESCE(SUM(l.credit),0) credit
                      FROM journal_lines l JOIN journal_entries j ON j.id=l.journal_id
                      WHERE l.account_id=? AND j.status='POSTED'"""
            op_params=[a['id']]
            if date_from: op_sql+=' AND j.journal_date<?';op_params.append(date_from)
            else: op_sql+=' AND 1=0'
            op=c.execute(op_sql,op_params).fetchone();od=float(op['debit']);oc=float(op['credit'])
            period_sql="""SELECT COALESCE(SUM(l.debit),0) debit,COALESCE(SUM(l.credit),0) credit
                           FROM journal_lines l JOIN journal_entries j ON j.id=l.journal_id
                           WHERE l.account_id=? AND j.status='POSTED'"""
            period_params=[a['id']]
            if date_from: period_sql+=' AND j.journal_date>=?';period_params.append(date_from)
            if date_to: period_sql+=' AND j.journal_date<=?';period_params.append(date_to)
            pr=c.execute(period_sql,period_params).fetchone();pd=float(pr['debit']);pc=float(pr['credit'])
            net=(od+pd)-(oc+pc)
            ed=max(net,0.0);ec=max(-net,0.0)
            total_opening_debit+=od;total_opening_credit+=oc;total_period_debit+=pd;total_period_credit+=pc
            total_ending_debit+=ed;total_ending_credit+=ec
            items.append({
                'id':a['id'],'code':a['code'],'name':a['name'],'account_type':a['account_type'],
                'normal_balance':a['normal_balance'],'opening_debit':round(od,2),'opening_credit':round(oc,2),
                'period_debit':round(pd,2),'period_credit':round(pc,2),
                'ending_debit':round(ed,2),'ending_credit':round(ec,2),
                'balance':round(net if a['normal_balance']=='DEBIT' else -net,2)
            })
        return {
            'date_from':date_from,'date_to':date_to,'items':items,
            'total_opening_debit':round(total_opening_debit,2),'total_opening_credit':round(total_opening_credit,2),
            'total_period_debit':round(total_period_debit,2),'total_period_credit':round(total_period_credit,2),
            'total_ending_debit':round(total_ending_debit,2),'total_ending_credit':round(total_ending_credit,2),
            'balanced':round(total_ending_debit-total_ending_credit,2)==0
        }
    finally:c.close()


def _next_settlement_no(tx,prefix,payment_date):
    period=str(payment_date)[:7].replace('-','');key=f'{prefix}-{period}'
    row=tx.execute('SELECT current_value FROM document_sequences WHERE sequence_key=?',(key,)).fetchone();n=int(row['current_value'])+1 if row else 1
    tx.execute("INSERT INTO document_sequences(sequence_key,current_value,updated_at) VALUES(?,?,?) ON CONFLICT(sequence_key) DO UPDATE SET current_value=excluded.current_value,updated_at=excluded.updated_at",(key,n,utc_now()))
    return f'{prefix}-{period}-{n:06d}'

def _opening_settlement_paid(c, kind, partner_id, as_of=None):
    table='opening_receivable_payments' if kind=='customer' else 'opening_payable_payments'
    col='customer_id' if kind=='customer' else 'supplier_id'
    where=[f'{col}=?'];params=[int(partner_id)]
    if as_of: where.append('payment_date<=?');params.append(str(as_of)[:10])
    void_kind='receivable' if kind=='customer' else 'payable'
    where.append(f"NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind=? AND v.transaction_key=('OPENING-'||CAST({table}.id AS TEXT)))")
    params.append(void_kind)
    settled=float(c.execute(f"SELECT COALESCE(SUM(CAST(amount AS REAL)),0) v FROM {table} WHERE {' AND '.join(where)}",params).fetchone()['v'] or 0)
    typ='CUSTOMER' if kind=='customer' else 'SUPPLIER'; alloc_where=["allocation_type=?","invoice_id=?","status='POSTED'"];alloc_params=[typ,-int(partner_id)]
    if as_of:alloc_where.append('allocation_date<=?');alloc_params.append(str(as_of)[:10])
    allocated=float(c.execute("SELECT COALESCE(SUM(CAST(amount AS REAL)),0) v FROM downpayment_allocations WHERE "+' AND '.join(alloc_where),alloc_params).fetchone()['v'] or 0)
    return settled+allocated

def open_receivables(customer_id=None,as_of=None):
    where=["s.status='POSTED'"];params=[]
    if customer_id not in (None,''):where.append('s.customer_id=?');params.append(int(customer_id))
    if as_of:where.append('s.sale_date<=?');params.append(str(as_of)[:10])
    c=connect()
    try:
        rows=c.execute("""SELECT s.id,s.invoice_no,s.sale_date,s.due_date,s.total_amount,bp.id customer_id,bp.code customer_code,bp.name customer_name,
          s.paid_amount stored_paid,
          COALESCE((SELECT SUM(rp.amount) FROM receivable_payments rp WHERE rp.sale_id=s.id AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='receivable' AND v.transaction_key=CAST(rp.id AS TEXT))),0) cash_paid,
          COALESCE((SELECT SUM(a.amount) FROM downpayment_allocations a WHERE a.allocation_type='CUSTOMER' AND a.invoice_id=s.id AND a.status='POSTED'),0) dp_allocated, COALESCE((SELECT SUM(sr.total_sales) FROM sales_returns sr WHERE sr.sale_id=s.id AND sr.status='POSTED'),0) returned_amount,
          CAST(julianday(COALESCE(?,date('now')))-julianday(COALESCE(s.due_date,s.sale_date)) AS INTEGER) days_overdue
          FROM sales s JOIN business_partners bp ON bp.id=s.customer_id WHERE """+' AND '.join(where)+" ORDER BY COALESCE(s.due_date,s.sale_date),s.id",([str(as_of)[:10] if as_of else None]+params)).fetchall()
        out=[]
        for row in rows:
            x=dict(row);net_total=max(0.0,float(x['total_amount'] or 0)-float(x.get('returned_amount') or 0));paid=(float(x.get('stored_paid') or 0) if not as_of else max(float(x.get('stored_paid') or 0) if float(x.get('stored_paid') or 0)>=net_total else 0.0,float(x["cash_paid"] or 0)+float(x["dp_allocated"] or 0)));due=max(0.0,net_total-paid);x['original_total_amount']=float(x['total_amount'] or 0);x['total_amount']=net_total
            if due>0:x["paid_amount"]=paid;x["balance_due"]=due;x["days_overdue"]=max(0,int(x["days_overdue"] or 0));x['is_opening_balance']=False;out.append(x)
        pwhere=["is_active=1","partner_type IN ('CUSTOMER','BOTH')","COALESCE(opening_balance,0)>0"];pargs=[]
        if customer_id not in (None,''):pwhere.append('id=?');pargs.append(int(customer_id))
        if as_of:pwhere.append('COALESCE(opening_balance_date,?)<=?');pargs.extend([str(as_of)[:10],str(as_of)[:10]])
        for r in c.execute("SELECT * FROM business_partners WHERE "+' AND '.join(pwhere)+" ORDER BY code COLLATE NOCASE",pargs).fetchall():
            paid=_opening_settlement_paid(c,'customer',r['id'],as_of); total=float(r['opening_balance'] or 0); due=max(0.0,total-paid)
            if due<=0: continue
            od=str(r['opening_balance_date'] or (as_of or date.today().isoformat()))[:10]
            due_date=(date.fromisoformat(od)+timedelta(days=int(r['payment_term_days'] or 0))).isoformat()
            days=max(0,(date.fromisoformat(str(as_of)[:10] if as_of else date.today().isoformat())-date.fromisoformat(due_date)).days)
            out.append({'id':f'OPENING:{r["id"]}','invoice_no':f'SALDO AWAL-{r["code"]}','sale_date':od,'due_date':due_date,'total_amount':total,'customer_id':r['id'],'customer_code':r['code'],'customer_name':r['name'],'cash_paid':paid,'dp_allocated':0,'paid_amount':paid,'balance_due':due,'days_overdue':days,'is_opening_balance':True})
        return sorted(out,key=lambda x:(x.get('due_date') or x.get('sale_date') or '',str(x.get('invoice_no') or '')))
    finally:c.close()

def open_payables(supplier_id=None,as_of=None):
    where=["p.status='POSTED'"];params=[]
    if supplier_id not in (None,''):where.append('p.supplier_id=?');params.append(int(supplier_id))
    if as_of:where.append('p.purchase_date<=?');params.append(str(as_of)[:10])
    c=connect()
    try:
        rows=c.execute("""SELECT p.id,p.purchase_no,p.purchase_date,p.due_date,p.total_amount,bp.id supplier_id,bp.code supplier_code,bp.name supplier_name,
          p.paid_amount stored_paid,
          COALESCE((SELECT SUM(pp.amount) FROM payable_payments pp WHERE pp.purchase_id=p.id AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='payable' AND v.transaction_key=CAST(pp.id AS TEXT))),0) cash_paid,
          COALESCE((SELECT SUM(a.amount) FROM downpayment_allocations a WHERE a.allocation_type='SUPPLIER' AND a.invoice_id=p.id AND a.status='POSTED'),0) dp_allocated, COALESCE((SELECT SUM(pr.total_payable) FROM purchase_returns pr WHERE pr.purchase_id=p.id AND pr.status='POSTED'),0) returned_amount,
          CAST(julianday(COALESCE(?,date('now')))-julianday(COALESCE(p.due_date,p.purchase_date)) AS INTEGER) days_overdue
          FROM purchases p JOIN business_partners bp ON bp.id=p.supplier_id WHERE """+' AND '.join(where)+" ORDER BY COALESCE(p.due_date,p.purchase_date),p.id",([str(as_of)[:10] if as_of else None]+params)).fetchall()
        out=[]
        for row in rows:
            x=dict(row);net_total=max(0.0,float(x['total_amount'] or 0)-float(x.get('returned_amount') or 0));paid=(float(x.get('stored_paid') or 0) if not as_of else max(float(x.get('stored_paid') or 0) if float(x.get('stored_paid') or 0)>=net_total else 0.0,float(x["cash_paid"] or 0)+float(x["dp_allocated"] or 0)));due=max(0.0,net_total-paid);x['original_total_amount']=float(x['total_amount'] or 0);x['total_amount']=net_total
            if due>0:x["paid_amount"]=paid;x["balance_due"]=due;x["days_overdue"]=max(0,int(x["days_overdue"] or 0));x['is_opening_balance']=False;out.append(x)
        pwhere=["is_active=1","partner_type IN ('SUPPLIER','BOTH')","COALESCE(opening_balance,0)>0"];pargs=[]
        if supplier_id not in (None,''):pwhere.append('id=?');pargs.append(int(supplier_id))
        if as_of:pwhere.append('COALESCE(opening_balance_date,?)<=?');pargs.extend([str(as_of)[:10],str(as_of)[:10]])
        for r in c.execute("SELECT * FROM business_partners WHERE "+' AND '.join(pwhere)+" ORDER BY code COLLATE NOCASE",pargs).fetchall():
            paid=_opening_settlement_paid(c,'supplier',r['id'],as_of); total=float(r['opening_balance'] or 0); due=max(0.0,total-paid)
            if due<=0: continue
            od=str(r['opening_balance_date'] or (as_of or date.today().isoformat()))[:10]
            due_date=(date.fromisoformat(od)+timedelta(days=int(r['payment_term_days'] or 0))).isoformat()
            days=max(0,(date.fromisoformat(str(as_of)[:10] if as_of else date.today().isoformat())-date.fromisoformat(due_date)).days)
            out.append({'id':f'OPENING:{r["id"]}','purchase_no':f'SALDO AWAL-{r["code"]}','purchase_date':od,'due_date':due_date,'total_amount':total,'supplier_id':r['id'],'supplier_code':r['code'],'supplier_name':r['name'],'cash_paid':paid,'dp_allocated':0,'paid_amount':paid,'balance_due':due,'days_overdue':days,'is_opening_balance':True})
        return sorted(out,key=lambda x:(x.get('due_date') or x.get('purchase_date') or '',str(x.get('purchase_no') or '')))
    finally:c.close()

def _aging(items):
    buckets={'CURRENT':0.0,'1_30':0.0,'31_60':0.0,'61_90':0.0,'OVER_90':0.0,'TOTAL':0.0}
    for x in items:
        v=float(x['balance_due']);d=int(x['days_overdue']);b='CURRENT' if d<=0 else '1_30' if d<=30 else '31_60' if d<=60 else '61_90' if d<=90 else 'OVER_90';buckets[b]+=v;buckets['TOTAL']+=v
    return buckets

def receivables_aging(as_of=None):return {'as_of':as_of or date.today().isoformat(),'buckets':_aging(open_receivables(as_of=as_of)),'items':open_receivables(as_of=as_of)}
def payables_aging(as_of=None):return {'as_of':as_of or date.today().isoformat(),'buckets':_aging(open_payables(as_of=as_of)),'items':open_payables(as_of=as_of)}

def _effective_invoice_balance(tx,kind,invoice_id):
    customer=kind=="customer";inv="sales" if customer else "purchases";pk="sale_id" if customer else "purchase_id";pay="receivable_payments" if customer else "payable_payments";voidkind="receivable" if customer else "payable";atype="CUSTOMER" if customer else "SUPPLIER"
    row=tx.execute(f"SELECT total_amount,paid_amount,balance_due,payment_method FROM {inv} WHERE id=? AND status='POSTED'",(int(invoice_id),)).fetchone()
    if not row:return Decimal("0"),Decimal("0")
    cash=_money(tx.execute(f"SELECT COALESCE(SUM(p.amount),0) v FROM {pay} p WHERE p.{pk}=? AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind=? AND v.transaction_key=CAST(p.id AS TEXT))",(int(invoice_id),voidkind)).fetchone()["v"])
    dp=_money(tx.execute("SELECT COALESCE(SUM(amount),0) v FROM downpayment_allocations WHERE allocation_type=? AND invoice_id=? AND status='POSTED'",(atype,int(invoice_id))).fetchone()["v"])
    # paid_amount adalah sumber kumulatif resmi invoice (termasuk pembayaran saat invoice dibuat).
    # Importer/DP lama hanya menghitung tabel settlement sehingga invoice TUNAI terlihat outstanding.
    stored=_money(row['paid_amount']); paid=max(stored,cash+dp)
    if customer:
        returned=_money(tx.execute("SELECT COALESCE(SUM(total_sales),0) v FROM sales_returns WHERE sale_id=? AND status='POSTED'",(int(invoice_id),)).fetchone()['v'])
    else:
        returned=_money(tx.execute("SELECT COALESCE(SUM(total_payable),0) v FROM purchase_returns WHERE purchase_id=? AND status='POSTED'",(int(invoice_id),)).fetchone()['v'])
    total=max(Decimal("0"),_money(row['total_amount'])-returned)
    return paid,max(Decimal("0"),total-paid)

def receive_receivable(actor,d,ip):
    payment_date=str(d.get('payment_date') or date.today().isoformat())[:10]
    sale_key=str(d.get('sale_id') or '').strip()
    try:cash_account_id=int(d.get('cash_account_id'))
    except:raise ValueError('Invoice dan akun kas/bank wajib dipilih.')
    opening_partner_id=int(sale_key.split(':',1)[1]) if sale_key.startswith('OPENING:') else None
    try:sale_id=None if opening_partner_id else int(sale_key)
    except:raise ValueError('Invoice piutang tidak valid.')
    amount=_money(d.get('amount'));notes=str(d.get('notes','')).strip() or None
    if amount<=0:raise ValueError('Jumlah penerimaan harus lebih dari nol.')
    now=utc_now()
    with write_transaction() as tx:
        if opening_partner_id:
            partner=tx.execute("SELECT * FROM business_partners WHERE id=? AND is_active=1 AND partner_type IN ('CUSTOMER','BOTH')",(opening_partner_id,)).fetchone()
            if not partner: raise ValueError('Saldo awal pelanggan tidak ditemukan.')
            paid=_money(tx.execute("SELECT COALESCE(SUM(amount),0) v FROM opening_receivable_payments op WHERE customer_id=? AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='receivable' AND v.transaction_key=('OPENING-'||CAST(op.id AS TEXT)))",(opening_partner_id,)).fetchone()['v'])+_money(tx.execute("SELECT COALESCE(SUM(amount),0) v FROM downpayment_allocations WHERE allocation_type='CUSTOMER' AND invoice_id=? AND status='POSTED'",(-opening_partner_id,)).fetchone()['v'])
            due=max(Decimal('0'),_money(partner['opening_balance'])-paid)
            if due<=0: raise ValueError('Saldo awal piutang sudah lunas.')
            if amount>due: raise ValueError('Jumlah penerimaan melebihi saldo awal piutang.')
            no=_next_settlement_no(tx,'RP',payment_date)
            cur=tx.execute('INSERT INTO opening_receivable_payments(payment_no,payment_date,customer_id,cash_account_id,amount,notes,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)',(no,payment_date,opening_partner_id,cash_account_id,str(amount),notes,actor['id'],now))
            cash_service.post(tx,account_id=cash_account_id,transaction_date=payment_date,transaction_type='IN',amount=amount,description=f'Penerimaan saldo awal piutang {partner["code"]}',reference_no=no,user_id=actor['id'])
            accounting_service.post_receivable_payment(tx,payment_id=cur.lastrowid,payment_no=no,payment_date=payment_date,amount=amount,cash_account_id=cash_account_id,customer_id=opening_partner_id,user_id=actor['id'],receivable_account_id=partner['receivable_account_id'])
            audit(actor['id'],'OPENING_RECEIVABLE_RECEIVED','partner',opening_partner_id,{'payment_no':no,'amount':float(amount),'remaining':float(due-amount)},ip,tx)
            return {'id':cur.lastrowid,'payment_no':no,'remaining_balance':float(due-amount)}
        sale=tx.execute("SELECT * FROM sales WHERE id=? AND status='POSTED'",(sale_id,)).fetchone()
        if not sale:raise ValueError('Invoice tidak ditemukan.')
        effective_paid,effective_due=_effective_invoice_balance(tx,'customer',sale_id)
        if effective_due<=0:raise ValueError('Invoice sudah lunas.')
        if amount>effective_due:raise ValueError('Jumlah penerimaan melebihi saldo piutang setelah alokasi DP.')
        no=_next_settlement_no(tx,'RP',payment_date);new_paid=effective_paid+amount;new_due=effective_due-amount
        cur=tx.execute('INSERT INTO receivable_payments(payment_no,payment_date,sale_id,customer_id,cash_account_id,amount,notes,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(no,payment_date,sale_id,sale['customer_id'],cash_account_id,str(amount),notes,actor['id'],now))
        tx.execute('UPDATE sales SET paid_amount=?,balance_due=?,updated_at=? WHERE id=?',(str(new_paid),str(new_due),now,sale_id))
        cash_service.post(tx,account_id=cash_account_id,transaction_date=payment_date,transaction_type='IN',amount=amount,description=f'Penerimaan piutang {sale["invoice_no"]}',reference_no=no,user_id=actor['id'])
        accounting_service.post_receivable_payment(tx,payment_id=cur.lastrowid,payment_no=no,payment_date=payment_date,amount=amount,cash_account_id=cash_account_id,customer_id=sale['customer_id'],user_id=actor['id'],receivable_account_id=(tx.execute('SELECT receivable_account_id FROM business_partners WHERE id=?',(sale['customer_id'],)).fetchone() or {'receivable_account_id':None})['receivable_account_id'])
        audit(actor['id'],'RECEIVABLE_RECEIVED','sale',sale_id,{'payment_no':no,'amount':float(amount),'remaining':float(new_due)},ip,tx)
        return {'id':cur.lastrowid,'payment_no':no,'remaining_balance':float(new_due)}

def pay_payable(actor,d,ip):
    payment_date=str(d.get('payment_date') or date.today().isoformat())[:10]
    purchase_key=str(d.get('purchase_id') or '').strip()
    try:cash_account_id=int(d.get('cash_account_id'))
    except:raise ValueError('Pembelian dan akun kas/bank wajib dipilih.')
    opening_partner_id=int(purchase_key.split(':',1)[1]) if purchase_key.startswith('OPENING:') else None
    try:purchase_id=None if opening_partner_id else int(purchase_key)
    except:raise ValueError('Pembelian hutang tidak valid.')
    amount=_money(d.get('amount'));notes=str(d.get('notes','')).strip() or None
    if amount<=0:raise ValueError('Jumlah pembayaran harus lebih dari nol.')
    now=utc_now()
    with write_transaction() as tx:
        if opening_partner_id:
            partner=tx.execute("SELECT * FROM business_partners WHERE id=? AND is_active=1 AND partner_type IN ('SUPPLIER','BOTH')",(opening_partner_id,)).fetchone()
            if not partner: raise ValueError('Saldo awal pemasok tidak ditemukan.')
            paid=_money(tx.execute("SELECT COALESCE(SUM(amount),0) v FROM opening_payable_payments op WHERE supplier_id=? AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='payable' AND v.transaction_key=('OPENING-'||CAST(op.id AS TEXT)))",(opening_partner_id,)).fetchone()['v'])+_money(tx.execute("SELECT COALESCE(SUM(amount),0) v FROM downpayment_allocations WHERE allocation_type='SUPPLIER' AND invoice_id=? AND status='POSTED'",(-opening_partner_id,)).fetchone()['v'])
            due=max(Decimal('0'),_money(partner['opening_balance'])-paid)
            if due<=0: raise ValueError('Saldo awal hutang sudah lunas.')
            if amount>due: raise ValueError('Jumlah pembayaran melebihi saldo awal hutang.')
            no=_next_settlement_no(tx,'PH',payment_date)
            cur=tx.execute('INSERT INTO opening_payable_payments(payment_no,payment_date,supplier_id,cash_account_id,amount,notes,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)',(no,payment_date,opening_partner_id,cash_account_id,str(amount),notes,actor['id'],now))
            cash_service.post(tx,account_id=cash_account_id,transaction_date=payment_date,transaction_type='OUT',amount=amount,description=f'Pembayaran saldo awal hutang {partner["code"]}',reference_no=no,user_id=actor['id'],allow_negative=True)
            accounting_service.post_payable_payment(tx,payment_id=cur.lastrowid,payment_no=no,payment_date=payment_date,amount=amount,cash_account_id=cash_account_id,supplier_id=opening_partner_id,user_id=actor['id'],payable_account_id=partner['payable_account_id'])
            audit(actor['id'],'OPENING_PAYABLE_PAID','partner',opening_partner_id,{'payment_no':no,'amount':float(amount),'remaining':float(due-amount)},ip,tx)
            return {'id':cur.lastrowid,'payment_no':no,'remaining_balance':float(due-amount)}
        p=tx.execute("SELECT * FROM purchases WHERE id=? AND status='POSTED'",(purchase_id,)).fetchone()
        if not p:raise ValueError('Pembelian tidak ditemukan.')
        effective_paid,effective_due=_effective_invoice_balance(tx,'supplier',purchase_id)
        if effective_due<=0:raise ValueError('Pembelian sudah lunas.')
        if amount>effective_due:raise ValueError('Jumlah pembayaran melebihi saldo hutang setelah alokasi DP.')
        no=_next_settlement_no(tx,'PH',payment_date);new_paid=effective_paid+amount;new_due=effective_due-amount
        cur=tx.execute('INSERT INTO payable_payments(payment_no,payment_date,purchase_id,supplier_id,cash_account_id,amount,notes,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(no,payment_date,purchase_id,p['supplier_id'],cash_account_id,str(amount),notes,actor['id'],now))
        tx.execute('UPDATE purchases SET paid_amount=?,balance_due=?,updated_at=? WHERE id=?',(str(new_paid),str(new_due),now,purchase_id))
        cash_service.post(tx,account_id=cash_account_id,transaction_date=payment_date,transaction_type='OUT',amount=amount,description=f'Pembayaran hutang {p["purchase_no"]}',reference_no=no,user_id=actor['id'],allow_negative=True)
        accounting_service.post_payable_payment(tx,payment_id=cur.lastrowid,payment_no=no,payment_date=payment_date,amount=amount,cash_account_id=cash_account_id,supplier_id=p['supplier_id'],user_id=actor['id'],payable_account_id=(tx.execute('SELECT payable_account_id FROM business_partners WHERE id=?',(p['supplier_id'],)).fetchone() or {'payable_account_id':None})['payable_account_id'])
        audit(actor['id'],'PAYABLE_PAID','purchase',purchase_id,{'payment_no':no,'amount':float(amount),'remaining':float(new_due)},ip,tx)
        return {'id':cur.lastrowid,'payment_no':no,'remaining_balance':float(new_due)}

def list_customers(q="",active=True):
    return [x for x in list_partners(q,"CUSTOMER",active) if x["partner_type"] in ("CUSTOMER","BOTH")]
def list_suppliers(q="",active=True):
    return [x for x in list_partners(q,"SUPPLIER",active) if x["partner_type"] in ("SUPPLIER","BOTH")]

def dashboard_analytics(as_of=None):
    asof=date.fromisoformat(as_of) if as_of else date.today();start=asof.replace(day=1).isoformat();end=asof.isoformat();week_end=(asof+timedelta(days=7)).isoformat()
    c=connect()
    try:
        # Dashboard P&L wajib memakai sumber angka yang sama dengan Laporan Laba Rugi.
        # Sebelumnya pendapatan berasal dari sales_items dan beban memakai account_type,
        # sehingga nilai dapat berbeda dari Neraca Saldo/Laporan Keuangan saat subtype COA
        # disesuaikan pengguna atau ada jurnal manual/kas.
        financial=income_statement(start,end)
        revenue=float(financial.get("revenue") or 0)
        cogs=float(financial.get("hpp") or 0)
        expense=float(financial.get("operating_expense") or 0)
        # Penyusutan bulanan diposting pada akhir bulan. Saat dashboard dibuka sebelum
        # tanggal akhir bulan, jurnal penyusutan periode berjalan tetap harus masuk
        # kartu Beban bulan ini meskipun tanggal jurnalnya berada sesudah as-of.
        period=asof.strftime("%Y-%m")
        future_dep=c.execute("""SELECT COALESCE(SUM(CAST(d.amount AS REAL)),0) v
          FROM fixed_asset_depreciations d JOIN journal_entries j ON j.id=d.journal_id
          WHERE d.period=? AND d.depreciation_date>? AND j.status='POSTED'""",(period,end)).fetchone()
        expense += float(future_dep['v'] or 0)
        top_items=[dict(x) for x in c.execute("SELECT si.product_name name,SUM(si.qty) qty,SUM(si.line_total) value FROM sales_items si JOIN sales s ON s.id=si.sale_id WHERE s.status='POSTED' AND s.sale_date BETWEEN ? AND ? GROUP BY si.product_id,si.product_name ORDER BY qty DESC,value DESC LIMIT 5",(start,end))]
        top_margin_items=[dict(x) for x in c.execute("SELECT si.product_name name,SUM(si.line_total-(si.qty*si.purchase_price_snapshot)) margin_value,CASE WHEN SUM(si.line_total)>0 THEN (SUM(si.line_total-(si.qty*si.purchase_price_snapshot))*100.0/SUM(si.line_total)) ELSE 0 END margin_percent,SUM(si.qty) qty FROM sales_items si JOIN sales s ON s.id=si.sale_id WHERE s.status='POSTED' AND s.sale_date BETWEEN ? AND ? GROUP BY si.product_id,si.product_name HAVING SUM(si.line_total)>0 ORDER BY margin_value DESC,margin_percent DESC LIMIT 5",(start,end))]
        top_sales=[dict(x) for x in c.execute("SELECT COALESCE(sp.name,'Tanpa Salesman') name,SUM(s.total_amount) value,COUNT(*) transactions FROM sales s LEFT JOIN salespersons sp ON sp.id=s.salesperson_id WHERE s.status='POSTED' AND s.sale_date BETWEEN ? AND ? GROUP BY sp.id,sp.name ORDER BY value DESC LIMIT 5",(start,end))]
        receivables_due_week=[dict(x) for x in c.execute("SELECT s.invoice_no number,bp.name partner,s.due_date,s.balance_due amount,CAST(julianday(s.due_date)-julianday(?) AS INTEGER) days_left FROM sales s JOIN business_partners bp ON bp.id=s.customer_id WHERE s.status='POSTED' AND s.balance_due>0 AND s.due_date IS NOT NULL AND s.due_date BETWEEN ? AND ? ORDER BY s.due_date,s.balance_due DESC",(end,end,week_end))]
        payables_due_week=[dict(x) for x in c.execute("SELECT p.purchase_no number,bp.name partner,p.due_date,p.balance_due amount,CAST(julianday(p.due_date)-julianday(?) AS INTEGER) days_left FROM purchases p JOIN business_partners bp ON bp.id=p.supplier_id WHERE p.status='POSTED' AND p.balance_due>0 AND p.due_date IS NOT NULL AND p.due_date BETWEEN ? AND ? ORDER BY p.due_date,p.balance_due DESC",(end,end,week_end))]
        overdue_receivables=[dict(x) for x in c.execute("SELECT s.invoice_no number,bp.name partner,s.due_date,s.balance_due amount FROM sales s JOIN business_partners bp ON bp.id=s.customer_id WHERE s.status='POSTED' AND s.balance_due>0 AND s.due_date IS NOT NULL AND s.due_date<? ORDER BY s.due_date",(end,))]
        overdue_payables=[dict(x) for x in c.execute("SELECT p.purchase_no number,bp.name partner,p.due_date,p.balance_due amount FROM purchases p JOIN business_partners bp ON bp.id=p.supplier_id WHERE p.status='POSTED' AND p.balance_due>0 AND p.due_date IS NOT NULL AND p.due_date<? ORDER BY p.due_date",(end,))]
        # Saldo awal partner juga mempunyai jatuh tempo: tanggal saldo awal + termin partner.
        # Masukkan ke panel dashboard agar saldo awal lama tidak luput dari notifikasi AR/AP.
        opening_ar=[dict(x) for x in c.execute("""SELECT 'SALDO AWAL-'||bp.code number,bp.name partner,
          date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day') due_date,
          MAX(0,bp.opening_balance-COALESCE((SELECT SUM(amount) FROM opening_receivable_payments op WHERE op.customer_id=bp.id AND op.payment_date<=? AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='receivable' AND v.transaction_key=('OPENING-'||CAST(op.id AS TEXT)))),0)) amount
          FROM business_partners bp WHERE bp.is_active=1 AND bp.partner_type IN ('CUSTOMER','BOTH') AND bp.opening_balance>0
          AND date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day')<?
          ORDER BY due_date""",(end,end,end,end)).fetchall()]
        opening_ap=[dict(x) for x in c.execute("""SELECT 'SALDO AWAL-'||bp.code number,bp.name partner,
          date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day') due_date,
          MAX(0,bp.opening_balance-COALESCE((SELECT SUM(amount) FROM opening_payable_payments op WHERE op.supplier_id=bp.id AND op.payment_date<=? AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='payable' AND v.transaction_key=('OPENING-'||CAST(op.id AS TEXT)))),0)) amount
          FROM business_partners bp WHERE bp.is_active=1 AND bp.partner_type IN ('SUPPLIER','BOTH') AND bp.opening_balance>0
          AND date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day')<?
          ORDER BY due_date""",(end,end,end,end)).fetchall()]
        overdue_receivables=sorted(overdue_receivables+opening_ar,key=lambda x:(x.get('due_date') or ''))
        overdue_payables=sorted(overdue_payables+opening_ap,key=lambda x:(x.get('due_date') or ''))
        opening_ar_week=[dict(x) for x in c.execute("""SELECT 'SALDO AWAL-'||bp.code number,bp.name partner,
          date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day') due_date,MAX(0,bp.opening_balance-COALESCE((SELECT SUM(amount) FROM opening_receivable_payments op WHERE op.customer_id=bp.id AND op.payment_date<=? AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='receivable' AND v.transaction_key=('OPENING-'||CAST(op.id AS TEXT)))),0)) amount,
          CAST(julianday(date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day'))-julianday(?) AS INTEGER) days_left
          FROM business_partners bp WHERE bp.is_active=1 AND bp.partner_type IN ('CUSTOMER','BOTH') AND bp.opening_balance>0
          AND date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day') BETWEEN ? AND ?""",(end,end,end,end,end,end,week_end)).fetchall()]
        opening_ap_week=[dict(x) for x in c.execute("""SELECT 'SALDO AWAL-'||bp.code number,bp.name partner,
          date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day') due_date,MAX(0,bp.opening_balance-COALESCE((SELECT SUM(amount) FROM opening_payable_payments op WHERE op.supplier_id=bp.id AND op.payment_date<=? AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='payable' AND v.transaction_key=('OPENING-'||CAST(op.id AS TEXT)))),0)) amount,
          CAST(julianday(date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day'))-julianday(?) AS INTEGER) days_left
          FROM business_partners bp WHERE bp.is_active=1 AND bp.partner_type IN ('SUPPLIER','BOTH') AND bp.opening_balance>0
          AND date(COALESCE(bp.opening_balance_date,?), '+'||COALESCE(bp.payment_term_days,0)||' day') BETWEEN ? AND ?""",(end,end,end,end,end,end,week_end)).fetchall()]
        receivables_due_week=sorted(receivables_due_week+opening_ar_week,key=lambda x:(x.get('due_date') or ''))
        payables_due_week=sorted(payables_due_week+opening_ap_week,key=lambda x:(x.get('due_date') or ''))
        low_stock=[dict(x) for x in c.execute("SELECT p.sku,p.name,COALESCE(SUM(b.quantity),0) stock,p.minimum_stock,COALESCE((SELECT SUM(si.qty) FROM sales_items si JOIN sales s ON s.id=si.sale_id WHERE si.product_id=p.id AND s.status='POSTED' AND s.sale_date BETWEEN ? AND ?),0) sold_month FROM products p LEFT JOIN inventory_balances b ON b.product_id=p.id WHERE p.is_active=1 AND p.product_type='STOCK' GROUP BY p.id HAVING COALESCE(SUM(b.quantity),0)<=p.minimum_stock ORDER BY stock LIMIT 10",(start,end))]
        elapsed=max((asof-asof.replace(day=1)).days+1,1)
        for item in low_stock:
            daily=float(item.get('sold_month') or 0)/elapsed
            item['estimated_days_left']=round(float(item.get('stock') or 0)/daily,1) if daily>0 else None
        cutoff=(asof-timedelta(days=90)).isoformat()
        dead_stock=[dict(x) for x in c.execute("SELECT p.sku,p.name,COALESCE(SUM(b.quantity),0) stock,MAX(CASE WHEN t.quantity_change<0 THEN substr(t.created_at,1,10) END) last_out FROM products p LEFT JOIN inventory_balances b ON b.product_id=p.id LEFT JOIN inventory_transactions t ON t.product_id=p.id WHERE p.is_active=1 AND p.product_type='STOCK' GROUP BY p.id HAVING COALESCE(SUM(b.quantity),0)>0 AND (MAX(CASE WHEN t.quantity_change<0 THEN substr(t.created_at,1,10) END) IS NULL OR MAX(CASE WHEN t.quantity_change<0 THEN substr(t.created_at,1,10) END)<?) ORDER BY stock DESC LIMIT 10",(cutoff,))]
        for item in dead_stock:
            item['inactive_days']=(asof-date.fromisoformat(item['last_out'])).days if item.get('last_out') else None
        top_expenses=[dict(x) for x in c.execute("SELECT a.code,a.name,SUM(jl.debit-jl.credit) amount FROM journal_lines jl JOIN journal_entries je ON je.id=jl.journal_id JOIN chart_of_accounts a ON a.id=jl.account_id WHERE je.status='POSTED' AND a.account_type='EXPENSE' AND a.code<>'5000' AND je.journal_date BETWEEN ? AND ? GROUP BY a.id,a.code,a.name HAVING SUM(jl.debit-jl.credit)>0 ORDER BY amount DESC LIMIT 5",(start,end))]
        return {"period":{"from":start,"to":end},"profit_loss":{"revenue":revenue,"cogs":cogs,"expense":expense,"profit":revenue-cogs-expense},"top_items":top_items,"top_margin_items":top_margin_items,"top_sales":top_sales,"receivables_due_week":receivables_due_week,"payables_due_week":payables_due_week,"overdue_receivables":overdue_receivables,"overdue_payables":overdue_payables,"low_stock":low_stock,"dead_stock":dead_stock,"top_expenses":top_expenses}
    finally:c.close()

def _document_default_options(dtype):
    return {
      "show_sku": True,
      "show_unit": True,
      "show_warehouse": dtype in ("DELIVERY_ORDER","GOODS_RECEIPT"),
      "show_discount": dtype in ("SALES_INVOICE","PURCHASE_INVOICE"),
      "show_tax": dtype in ("SALES_INVOICE","PURCHASE_INVOICE"),
      "show_payment_terms": dtype in ("SALES_INVOICE","PURCHASE_INVOICE"),
      "show_payment_account_note": False,
      "show_partner_contact": True,
      "show_amount_words": True,
      "payment_account_note": "",
      # Nilai berikut adalah bobot lebar relatif. Saat dicetak akan dinormalisasi
      # menjadi persen sesuai kolom yang memang sedang ditampilkan.
      "column_widths": {
        "no": 5, "sku": 13, "description": 37, "qty": 9,
        "unit": 10, "price": 13, "amount": 13
      },
      "special_note": "",
      "left_signer": "Disiapkan oleh",
      "right_signer": "Diterima oleh",
    }

def document_templates():
    defaults={
      "SALES_INVOICE":("INVOICE PENJUALAN",1),
      "DELIVERY_ORDER":("SURAT JALAN",0),
      "GOODS_RECEIPT":("GOOD RECEIVE",0),
      "PURCHASE_INVOICE":("FAKTUR PEMBELIAN",1),
    }
    with write_transaction() as tx:
        for dtype,(title,show_prices) in defaults.items():
            if not tx.execute("SELECT 1 FROM document_templates WHERE document_type=?",(dtype,)).fetchone():
                tx.execute("""INSERT INTO document_templates(
                  document_type,title,header_text,footer_text,show_logo,
                  show_prices,paper_size,options_json,updated_at
                ) VALUES(?,?,NULL,NULL,1,?,'A5',?,?)""",
                (dtype,title,show_prices,json.dumps(_document_default_options(dtype)),utc_now()))
        rows=tx.execute("""SELECT * FROM document_templates
          ORDER BY CASE document_type
          WHEN 'SALES_INVOICE' THEN 1 WHEN 'DELIVERY_ORDER' THEN 2
          WHEN 'GOODS_RECEIPT' THEN 3 WHEN 'PURCHASE_INVOICE' THEN 4 ELSE 9 END"""
        ).fetchall()
        items=[]
        for row in rows:
            item=dict(row)
            options=_document_default_options(item["document_type"])
            try:
                stored=json.loads(item.get("options_json") or "{}")
                if isinstance(stored,dict):options.update(stored)
            except Exception:
                pass
            item["options"]=options
            items.append(item)
        return items

def update_document_template(actor,dtype,d,ip):
    dtype=str(dtype or "").upper().strip()
    if dtype not in {"SALES_INVOICE","DELIVERY_ORDER","GOODS_RECEIPT","PURCHASE_INVOICE"}:
        raise ValueError("Jenis dokumen tidak valid.")
    title=str(d.get("title") or "").strip()
    if not title:raise ValueError("Judul dokumen wajib diisi.")
    paper=str(d.get("paper_size") or "A5").upper().strip()
    if paper not in ("A5","A4","LETTER","STRUK58","STRUK80"):raise ValueError("Ukuran kertas tidak valid.")
    no_price=dtype in ("DELIVERY_ORDER","GOODS_RECEIPT")
    default_widths=_document_default_options(dtype)["column_widths"]
    incoming_widths=d.get("column_widths") if isinstance(d.get("column_widths"),dict) else {}
    column_widths={}
    for key,default in default_widths.items():
        try:value=float(incoming_widths.get(key,default))
        except (TypeError,ValueError):value=float(default)
        column_widths[key]=max(2.0,min(80.0,value))
    options={
      "show_sku":bool(d.get("show_sku",True)),
      "show_unit":bool(d.get("show_unit",True)),
      "show_warehouse":bool(d.get("show_warehouse",True)),
      "show_discount":False if no_price else bool(d.get("show_discount",True)),
      "show_tax":False if no_price else bool(d.get("show_tax",True)),
      "show_payment_terms":False if no_price else bool(d.get("show_payment_terms",True)),
      "show_payment_account_note":False if no_price else bool(d.get("show_payment_account_note",False)),
      "show_partner_contact":bool(d.get("show_partner_contact",True)),
      "show_amount_words":False if no_price else bool(d.get("show_amount_words",True)),
      "payment_account_note":str(d.get("payment_account_note") or "").strip(),
      "column_widths":column_widths,
      "special_note":str(d.get("special_note") or "").strip(),
      "left_signer":str(d.get("left_signer") or "Disiapkan oleh").strip(),
      "right_signer":str(d.get("right_signer") or "Diterima oleh").strip(),
    }
    values=(
      title,str(d.get("header_text") or "").strip() or None,
      str(d.get("footer_text") or "").strip() or None,
      1 if d.get("show_logo",True) else 0,
      0 if no_price else (1 if d.get("show_prices",True) else 0),
      paper,json.dumps(options),utc_now(),dtype
    )
    with write_transaction() as tx:
        if tx.execute("SELECT 1 FROM document_templates WHERE document_type=?",(dtype,)).fetchone():
            tx.execute("""UPDATE document_templates SET title=?,header_text=?,
              footer_text=?,show_logo=?,show_prices=?,paper_size=?,
              options_json=?,updated_at=? WHERE document_type=?""",values)
        else:
            tx.execute("""INSERT INTO document_templates(
              title,header_text,footer_text,show_logo,show_prices,paper_size,
              options_json,updated_at,document_type
            ) VALUES(?,?,?,?,?,?,?,?,?)""",values)
        audit(actor["id"],"DOCUMENT_TEMPLATE_UPDATED","document_template",dtype,
          {"title":title,"paper_size":paper,"options":options},ip,tx)
    return next(x for x in document_templates() if x["document_type"]==dtype)


FLEX_INVOICE_BLOCKS=[
 ("company_header","Header Perusahaan"),("customer_info","Informasi Customer"),
 ("project_info","Informasi Proyek"),("contract_info","Informasi Kontrak / SPK"),
 ("term_summary","Ringkasan Termin"),("invoice_items","Rincian Item Invoice"),
 ("project_materials","Rincian Material Proyek"),("project_progress","Progress Pekerjaan"),
 ("retention","Retensi"),("tax_totals","Pajak & Total Tagihan"),
 ("custom_text","Teks Custom"),("transaction_notes","Catatan Transaksi"),
 ("signatures","Tanda Tangan"),("footer","Footer")
]
def _flex_default_blocks(kind="STANDARD"):
    kind=str(kind).upper()
    if kind=="TERM":
        keys=["company_header","customer_info","project_info","contract_info","term_summary","retention","tax_totals","custom_text","transaction_notes","signatures","footer"]
    elif kind=="TERM_MATERIAL":
        keys=["company_header","customer_info","project_info","contract_info","project_materials","term_summary","retention","tax_totals","custom_text","transaction_notes","signatures","footer"]
    else:
        keys=["company_header","customer_info","project_info","invoice_items","tax_totals","transaction_notes","signatures","footer"]
    labels=dict(FLEX_INVOICE_BLOCKS);return [{"key":k,"label":labels[k],"enabled":True} for k in keys]
def _flex_default_options(kind="STANDARD"):
    return {"paper_size":"A4","orientation":"portrait","show_logo":True,"show_sku":True,"show_unit":True,
      "show_prices":True,"show_material_values":True,
      "custom_text":"Pembayaran mohon dilakukan sesuai jatuh tempo yang tercantum pada invoice.",
      "left_signer":"Dibuat oleh","right_signer":"Disetujui / Diterima oleh",
      "title":"INVOICE" if str(kind).upper()=="STANDARD" else "INVOICE TERMIN"}
def _seed_flexible_invoice_templates():
    defaults=[("INV-STANDARD","Invoice Penjualan Standar","STANDARD",1),("INV-TERM","Invoice Termin Proyek","TERM",0),("INV-TERM-MAT","Invoice Termin + Rincian Material","TERM_MATERIAL",0)]
    with write_transaction() as tx:
        if tx.execute("SELECT COUNT(*) n FROM flexible_invoice_templates").fetchone()["n"]:return
        now=utc_now()
        for code,name,kind,is_default in defaults:
            tx.execute("""INSERT INTO flexible_invoice_templates(code,name,template_kind,blocks_json,options_json,is_default,is_active,created_at,updated_at)
              VALUES(?,?,?,?,?,?,1,?,?)""",(code,name,kind,json.dumps(_flex_default_blocks(kind)),json.dumps(_flex_default_options(kind)),is_default,now,now))
def flexible_invoice_templates(active_only=False):
    c=connect()
    try:
        sql="SELECT * FROM flexible_invoice_templates"+(" WHERE is_active=1" if active_only else "")+" ORDER BY is_default DESC,name,id"
        rows=[dict(x) for x in c.execute(sql).fetchall()]
    finally:c.close()
    if not rows:_seed_flexible_invoice_templates();return flexible_invoice_templates(active_only)
    for x in rows:
        try:x["blocks"]=json.loads(x.get("blocks_json") or "[]")
        except:x["blocks"]=[]
        try:x["options"]=json.loads(x.get("options_json") or "{}")
        except:x["options"]={}
        x["is_default"]=bool(x["is_default"]);x["is_active"]=bool(x["is_active"])
    return rows
def save_flexible_invoice_template(actor,d,ip,template_id=None):
    code=str(d.get("code") or "").strip().upper();name=str(d.get("name") or "").strip();kind=str(d.get("template_kind") or "STANDARD").upper()
    if not code or not name:raise ValueError("Kode dan nama template wajib diisi.")
    if kind not in {"STANDARD","TERM","TERM_MATERIAL","CUSTOM"}:raise ValueError("Jenis template tidak valid.")
    valid=dict(FLEX_INVOICE_BLOCKS);clean=[];seen=set()
    for b in (d.get("blocks") if isinstance(d.get("blocks"),list) else _flex_default_blocks(kind)):
        if not isinstance(b,dict):continue
        key=str(b.get("key") or "")
        if key in valid and key not in seen:seen.add(key);clean.append({"key":key,"label":valid[key],"enabled":bool(b.get("enabled",True))})
    if not clean:clean=_flex_default_blocks(kind)
    options=_flex_default_options(kind)
    if isinstance(d.get("options"),dict):options.update(d["options"])
    options["paper_size"]=str(options.get("paper_size") or "A4").upper()
    if options["paper_size"] not in {"A4","A5","LETTER"}:options["paper_size"]="A4"
    options["orientation"]="landscape" if str(options.get("orientation")).lower()=="landscape" else "portrait"
    for k in ("title","custom_text","left_signer","right_signer"):options[k]=str(options.get(k) or "")
    is_default=1 if d.get("is_default") else 0;is_active=1 if d.get("is_active",True) else 0;now=utc_now()
    with write_transaction() as tx:
        if is_default:tx.execute("UPDATE flexible_invoice_templates SET is_default=0")
        if template_id:
            if not tx.execute("SELECT 1 FROM flexible_invoice_templates WHERE id=?",(int(template_id),)).fetchone():raise ValueError("Template tidak ditemukan.")
            tx.execute("""UPDATE flexible_invoice_templates SET code=?,name=?,template_kind=?,blocks_json=?,options_json=?,is_default=?,is_active=?,updated_at=? WHERE id=?""",
              (code,name,kind,json.dumps(clean),json.dumps(options),is_default,is_active,now,int(template_id)));tid=int(template_id);action="FLEX_INVOICE_TEMPLATE_UPDATED"
        else:
            cur=tx.execute("""INSERT INTO flexible_invoice_templates(code,name,template_kind,blocks_json,options_json,is_default,is_active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)""",
              (code,name,kind,json.dumps(clean),json.dumps(options),is_default,is_active,now,now));tid=cur.lastrowid;action="FLEX_INVOICE_TEMPLATE_CREATED"
        audit(actor["id"],action,"flexible_invoice_template",tid,{"code":code,"name":name},ip,tx)
    return next(x for x in flexible_invoice_templates(False) if int(x["id"])==tid)
def delete_flexible_invoice_template(actor,template_id,ip):
    with write_transaction() as tx:
        row=tx.execute("SELECT * FROM flexible_invoice_templates WHERE id=?",(int(template_id),)).fetchone()
        if not row:raise ValueError("Template tidak ditemukan.")
        tx.execute("DELETE FROM flexible_invoice_templates WHERE id=?",(int(template_id),))
        if row["is_default"]:
            x=tx.execute("SELECT id FROM flexible_invoice_templates WHERE is_active=1 ORDER BY id LIMIT 1").fetchone()
            if x:tx.execute("UPDATE flexible_invoice_templates SET is_default=1 WHERE id=?",(x["id"],))
        audit(actor["id"],"FLEX_INVOICE_TEMPLATE_DELETED","flexible_invoice_template",template_id,{"code":row["code"]},ip,tx)
def _invoice_project_context(data):
    pid=data.get("project_id")
    if not pid:return {"project":None,"term":None,"progress":None,"materials":[]}
    c=connect()
    try:
        p=c.execute("SELECT * FROM projects WHERE id=?",(pid,)).fetchone()
        term=c.execute("SELECT * FROM project_terms WHERE project_id=? AND sale_id=? ORDER BY id DESC LIMIT 1",(pid,data["id"])).fetchone()
        prog=c.execute("SELECT * FROM project_progress WHERE project_id=? ORDER BY progress_date DESC,id DESC LIMIT 1",(pid,)).fetchone()
    finally:c.close()
    mats=[]
    for issue in list_project_material_issues(pid,1000):
        for item in issue.get("items") or []:
            mats.append({"date":issue.get("issue_date"),"issue_no":issue.get("issue_no"),"sku":item.get("sku"),"name":item.get("product_name"),"qty":float(item.get("qty") or 0),"unit_cost":float(item.get("average_cost") or 0),"total_cost":float(item.get("total_cost") or 0)})
    return {"project":dict(p) if p else None,"term":dict(term) if term else None,"progress":dict(prog) if prog else None,"materials":mats}
def _flex_money(v):return "Rp"+format(float(v or 0),",.0f")
def _flex_replace(text,data,ctx):
    p=ctx.get("project") or {};t=ctx.get("term") or {};g=ctx.get("progress") or {}
    vals={"invoice_no":data.get("invoice_no") or "","invoice_date":data.get("sale_date") or "","customer_name":data.get("customer_name") or "",
      "project_name":p.get("name") or "","project_code":p.get("code") or "","contract_no":p.get("contract_no") or "",
      "contract_value":_flex_money(p.get("contract_value")),"term_no":t.get("term_no") or "","term_percent":t.get("percentage") or 0,
      "term_value":_flex_money(t.get("amount")),"retention_value":_flex_money(t.get("retention_amount")),"total_invoice":_flex_money(data.get("total_amount")),
      "due_date":data.get("due_date") or "","progress_percent":g.get("progress_percent") or 0}
    out=str(text or "")
    for k,v in vals.items():out=out.replace("{{"+k+"}}",str(v))
    return out
def render_flexible_invoice_html(sale_id,template_id=None):
    data=get_sale(sale_id)
    if not data:raise ValueError("Penjualan tidak ditemukan.")
    templates=flexible_invoice_templates(True)
    tpl=next((x for x in templates if template_id not in (None,"") and int(x["id"])==int(template_id)),None)
    if not tpl:tpl=next((x for x in templates if x["is_default"]),templates[0] if templates else None)
    if not tpl:raise ValueError("Belum ada template invoice aktif.")
    opt=_flex_default_options(tpl["template_kind"]);opt.update(tpl.get("options") or {});ctx=_invoice_project_context(data)
    p=ctx.get("project") or {};term=ctx.get("term") or {};progress=ctx.get("progress") or {}
    c=connect()
    try:profile=dict(c.execute("SELECT * FROM company_profile WHERE id=1").fetchone())
    finally:c.close()
    esc=lambda x:html.escape(str(x or ""));logo=_logo_data(profile.get("logo_path")) if opt.get("show_logo",True) else "";logo_html=f"<img class='logo' src='{logo}'>" if logo else ""
    enabled=[b["key"] for b in tpl.get("blocks") or [] if b.get("enabled")];blocks=[]
    def add(key,body):
        if key in enabled:blocks.append(f"<section class='block block-{key}'>{body}</section>")
    add("company_header",f"<div class='head'>{logo_html}<div><h2>{esc(profile.get('company_name') or PRODUCT_NAME)}</h2><p>{esc(profile.get('address'))}</p><p>{esc(profile.get('phone'))}</p></div><div class='title'><h1>{esc(opt.get('title'))}</h1><b>{esc(data.get('invoice_no'))}</b><small>{esc(tpl.get('name'))}</small></div></div>")
    add("customer_info",f"<div class='grid'><div><span>Pelanggan</span><b>{esc(data.get('customer_name') or 'Umum')}</b></div><div><span>Tanggal</span><b>{esc(data.get('sale_date'))}</b></div><div><span>Jatuh Tempo</span><b>{esc(data.get('due_date') or '-')}</b></div><div><span>Metode</span><b>{esc(data.get('payment_method') or '-')}</b></div></div>")
    if p:
        add("project_info",f"<h3>INFORMASI PROYEK</h3><div class='grid'><div><span>Kode</span><b>{esc(p.get('code'))}</b></div><div><span>Proyek</span><b>{esc(p.get('name'))}</b></div><div><span>Lokasi</span><b>{esc(p.get('location') or '-')}</b></div><div><span>PIC</span><b>{esc(p.get('pic_name') or '-')}</b></div></div>")
        add("contract_info",f"<h3>KONTRAK / SPK</h3><div class='grid'><div><span>No. Kontrak</span><b>{esc(p.get('contract_no') or '-')}</b></div><div><span>Nilai Kontrak</span><b>{_flex_money(p.get('contract_value'))}</b></div><div><span>Retensi Default</span><b>{float(p.get('retention_percent') or 0):g}%</b></div><div><span>Jenis Proyek</span><b>{esc(p.get('project_type') or '-')}</b></div></div>")
    if term:
        add("term_summary",f"<h3>RINGKASAN TERMIN</h3><table><tr><th>Termin</th><th>Uraian</th><th>Persentase</th><th>Nilai</th></tr><tr><td>{esc(term.get('term_no'))}</td><td>{esc(term.get('description') or '-')}</td><td class='num'>{float(term.get('percentage') or 0):g}%</td><td class='num'>{_flex_money(term.get('amount'))}</td></tr></table>")
        add("retention",f"<div class='summaryline'><span>Retensi {float(term.get('retention_percent') or 0):g}%</span><b>{_flex_money(term.get('retention_amount'))}</b></div>")
    if "invoice_items" in enabled:
        heads=["No"]+(["SKU"] if opt.get("show_sku",True) else [])+["Deskripsi","Qty"]+(["Sat."] if opt.get("show_unit",True) else [])+(["Harga","Jumlah"] if opt.get("show_prices",True) else [])
        rows=[]
        for i,x in enumerate(data.get("items") or [],1):
            vals=[i]+([esc(x.get("sku") or "-")] if opt.get("show_sku",True) else [])+[esc(x.get("product_name") or "-"),f"{float(x.get('qty') or 0):g}"]+([esc(x.get("unit_code") or "")] if opt.get("show_unit",True) else [])+([_flex_money(x.get("unit_price")),_flex_money(x.get("line_total"))] if opt.get("show_prices",True) else [])
            rows.append("<tr>"+"".join(f"<td>{v}</td>" for v in vals)+"</tr>")
        blocks.append("<section class='block'><h3>RINCIAN INVOICE</h3><table><thead><tr>"+"".join(f"<th>{x}</th>" for x in heads)+"</tr></thead><tbody>"+"".join(rows)+"</tbody></table></section>")
    if "project_materials" in enabled:
        heads=["No","Tanggal","Dok. Keluar","SKU","Material","Qty"]+(["Cost","Nilai"] if opt.get("show_material_values",True) else []);rows=[]
        for i,x in enumerate(data.get("invoice_materials") or [],1):
            vals=[i,esc(x.get("date")),esc(x.get("issue_no")),esc(x.get("sku")),esc(x.get("name")),f"{float(x.get('qty') or 0):g}"]+([_flex_money(x.get("unit_cost")),_flex_money(x.get("total_cost"))] if opt.get("show_material_values",True) else [])
            rows.append("<tr>"+"".join(f"<td>{v}</td>" for v in vals)+"</tr>")
        add("project_materials","<h3>RINCIAN MATERIAL PROYEK</h3><table><thead><tr>"+"".join(f"<th>{x}</th>" for x in heads)+"</tr></thead><tbody>"+("".join(rows) if rows else f"<tr><td colspan='{len(heads)}'>Belum ada material proyek.</td></tr>")+"</tbody></table>")
    if progress:add("project_progress",f"<h3>PROGRESS PEKERJAAN</h3><div class='progress'><div style='width:{min(100,max(0,float(progress.get('progress_percent') or 0)))}%'></div></div><p><b>{float(progress.get('progress_percent') or 0):g}%</b> per {esc(progress.get('progress_date'))} · {esc(progress.get('notes') or '')}</p>")
    add("tax_totals",f"<div class='totals'><div><span>Subtotal</span><b>{_flex_money(data.get('subtotal'))}</b></div><div><span>Diskon</span><b>{_flex_money(data.get('discount_amount'))}</b></div><div><span>Pajak</span><b>{_flex_money(data.get('tax_amount'))}</b></div><div class='grand'><span>TOTAL TAGIHAN</span><b>{_flex_money(data.get('total_amount'))}</b></div></div>")
    custom=_flex_replace(opt.get("custom_text"),data,ctx)
    if custom:add("custom_text",f"<div class='note'>{esc(custom).replace(chr(10),'<br>')}</div>")
    if data.get("notes"):add("transaction_notes",f"<div class='note'><b>Catatan:</b><br>{esc(data.get('notes')).replace(chr(10),'<br>')}</div>")
    add("signatures",f"<div class='sign'><div>{esc(opt.get('left_signer'))}<span></span></div><div>{esc(opt.get('right_signer'))}<span></span></div></div>")
    add("footer",f"<footer>{esc(profile.get('company_name') or PRODUCT_NAME)} · {esc(profile.get('phone') or '')}</footer>")
    page={"A5":"A5","LETTER":"Letter"}.get(str(opt.get("paper_size") or "A4").upper(),"A4");orient="landscape" if opt.get("orientation")=="landscape" else "portrait"
    return f"""<!doctype html><html><head><meta charset='utf-8'><title>{esc(opt.get('title'))} - {esc(data.get('invoice_no'))}</title><style>
@page{{size:{page} {orient};margin:10mm}}*{{box-sizing:border-box}}body{{font:10.5px Arial;color:#203038;margin:0;background:#eef2f3}}.toolbar{{max-width:210mm;margin:8px auto;text-align:right}}button{{background:#087f8c;color:white;border:0;border-radius:6px;padding:8px 14px}}.sheet{{max-width:210mm;margin:0 auto;background:white;padding:10mm;box-shadow:0 2px 12px #8998a055}}.block{{margin-bottom:10px;break-inside:avoid}}.head{{display:flex;gap:10px;border-bottom:2px solid #087f8c;padding-bottom:8px;align-items:center}}.logo{{max-width:30mm;max-height:18mm}}.head>div:nth-child(2){{flex:1}}h1,h2,h3,p{{margin:2px 0}}h3{{font-size:10px;color:#087f8c;margin:8px 0 4px}}.title{{text-align:right}}.title small{{display:block;color:#789}}.grid{{display:grid;grid-template-columns:1fr 1fr;gap:5px 14px;background:#f4f8f9;padding:7px}}.grid div{{display:flex;justify-content:space-between;gap:8px}}.grid span{{color:#607680}}table{{width:100%;border-collapse:collapse;table-layout:fixed}}th{{background:#087f8c;color:#fff;padding:5px}}td{{border-bottom:1px solid #d4dfe2;padding:5px;vertical-align:top;overflow-wrap:anywhere}}.num{{text-align:right}}.summaryline{{display:flex;justify-content:flex-end;gap:30px;background:#f2f7f8;padding:7px}}.totals{{width:75mm;margin-left:auto}}.totals div{{display:flex;justify-content:space-between;padding:4px 6px}}.totals .grand{{background:#e1f2f4;border-top:2px solid #087f8c;font-size:12px}}.note{{border-left:3px solid #087f8c;background:#f4f8f9;padding:7px}}.progress{{height:9px;background:#e2e9ec;border-radius:9px;overflow:hidden}}.progress div{{height:100%;background:#087f8c}}.sign{{display:grid;grid-template-columns:1fr 1fr;gap:30mm;text-align:center;margin-top:18px}}.sign span{{display:block;height:22mm;border-bottom:1px solid #536970}}footer{{text-align:center;color:#73868d;border-top:1px solid #ccd7da;padding-top:6px}}@media print{{body{{background:white}}.toolbar{{display:none}}.sheet{{box-shadow:none;margin:0;max-width:none;padding:0}}}}</style></head><body><div class='toolbar'><button onclick='window.print()'>Cetak / Simpan PDF</button></div><div class='sheet'>{''.join(blocks)}</div></body></html>"""
def _logo_data(path):
    if IS_POSTGRES:
        content,mime=company_logo_bytes()
        if not content:return ""
        return "data:"+(mime or "image/png")+";base64,"+base64.b64encode(content).decode()
    if not path:return ""
    p=Path(path)
    if not p.exists():return ""
    return "data:"+(mimetypes.guess_type(str(p))[0] or "image/png")+";base64,"+base64.b64encode(p.read_bytes()).decode()
def _terbilang_id(value):
    try:n=int(round(float(value or 0)))
    except Exception:return "Nol Rupiah"
    words=["","Satu","Dua","Tiga","Empat","Lima","Enam","Tujuh","Delapan","Sembilan","Sepuluh","Sebelas"]
    def say(x):
        if x<12:return words[x]
        if x<20:return say(x-10)+" Belas"
        if x<100:return say(x//10)+" Puluh"+(" "+say(x%10) if x%10 else "")
        if x<200:return "Seratus"+(" "+say(x-100) if x>100 else "")
        if x<1000:return say(x//100)+" Ratus"+(" "+say(x%100) if x%100 else "")
        if x<2000:return "Seribu"+(" "+say(x-1000) if x>1000 else "")
        if x<1000000:return say(x//1000)+" Ribu"+(" "+say(x%1000) if x%1000 else "")
        if x<1000000000:return say(x//1000000)+" Juta"+(" "+say(x%1000000) if x%1000000 else "")
        if x<1000000000000:return say(x//1000000000)+" Miliar"+(" "+say(x%1000000000) if x%1000000000 else "")
        return say(x//1000000000000)+" Triliun"+(" "+say(x%1000000000000) if x%1000000000000 else "")
    return (say(abs(n)) if n else "Nol")+" Rupiah"

def render_document_html(kind,eid):
    cfg={
      "sales-invoice":("SALES_INVOICE",True,"SALE"),
      "delivery-order":("DELIVERY_ORDER",False,"SALE"),
      "purchase-invoice":("PURCHASE_INVOICE",True,"PURCHASE"),
      "goods-receipt":("GOODS_RECEIPT",False,"PURCHASE")}
    if kind not in cfg:raise ValueError("Jenis cetakan tidak valid.")
    dtype,price_document,entity=cfg[kind]
    c=connect()
    try:
        profile=dict(c.execute("SELECT * FROM company_profile WHERE id=1").fetchone())
        row=c.execute("SELECT * FROM document_templates WHERE document_type=?",(dtype,)).fetchone()
    finally:c.close()
    if not row:
        document_templates()
        c=connect()
        try:row=c.execute("SELECT * FROM document_templates WHERE document_type=?",(dtype,)).fetchone()
        finally:c.close()
    tpl=dict(row)
    options=_document_default_options(dtype)
    try:
        stored=json.loads(tpl.get("options_json") or "{}")
        if isinstance(stored,dict):options.update(stored)
    except Exception:
        pass

    data=get_sale(eid) if entity=="SALE" else get_purchase(eid)
    if not data:raise ValueError("Transaksi tidak ditemukan.")
    show_prices=price_document and bool(tpl.get("show_prices"))
    show_sku=bool(options.get("show_sku",True))
    show_unit=bool(options.get("show_unit",True))
    show_warehouse=bool(options.get("show_warehouse",True))
    show_discount=show_prices and bool(options.get("show_discount",True))
    show_tax=show_prices and bool(options.get("show_tax",True))
    show_payment_terms=price_document and bool(options.get("show_payment_terms",True))
    show_payment_account_note=price_document and bool(options.get("show_payment_account_note",False))
    show_partner_contact=bool(options.get("show_partner_contact",True))
    show_amount_words=price_document and bool(options.get("show_amount_words",True))

    if kind=="delivery-order":number=data.get("delivery_no") or data.get("invoice_no")
    elif kind=="goods-receipt":number=data.get("goods_receipt_no") or data.get("purchase_no")
    elif kind=="purchase-invoice":number=data.get("supplier_invoice_no") or data.get("purchase_no")
    else:number=data.get("invoice_no")

    txdate=data.get("sale_date") or data.get("purchase_date")
    partner=data.get("customer_name") or data.get("supplier_name") or "-"
    partner_address=data.get("customer_address") or data.get("supplier_address") or ""
    partner_city=data.get("customer_city") or data.get("supplier_city") or ""
    partner_phone=data.get("customer_phone") or data.get("supplier_phone") or ""
    partner_email=data.get("customer_email") or data.get("supplier_email") or ""
    warehouse=data.get("warehouse_name") or "-"
    logo=_logo_data(profile.get("logo_path")) if tpl.get("show_logo") else ""
    logo_html=f"<img class='logo' src='{logo}'>" if logo else ""
    special_note=html.escape(str(options.get("special_note") or "")).replace("\n","<br>")
    payment_account_note=html.escape(str(options.get("payment_account_note") or "")).replace("\n","<br>")
    left_signer=html.escape(str(options.get("left_signer") or "Disiapkan oleh"))
    right_signer=html.escape(str(options.get("right_signer") or "Diterima oleh"))

    # Lebar kolom mengikuti rancangan dokumen. Bobot dinormalisasi agar selalu
    # pas 100% walaupun beberapa kolom (SKU/Satuan/Harga) disembunyikan.
    default_widths=_document_default_options(dtype)["column_widths"]
    raw_widths=options.get("column_widths") if isinstance(options.get("column_widths"),dict) else {}
    def _width(key):
        try:return max(2.0,min(80.0,float(raw_widths.get(key,default_widths[key]))))
        except (TypeError,ValueError):return float(default_widths[key])
    active_columns=["no"]
    if show_sku:active_columns.append("sku")
    active_columns.extend(["description","qty"])
    if show_unit:active_columns.append("unit")
    if show_prices:active_columns.extend(["price","amount"])
    total_width=sum(_width(k) for k in active_columns) or 1.0
    colgroup="<colgroup>"+"".join(
        f"<col style='width:{(_width(k)/total_width)*100:.3f}%'>" for k in active_columns
    )+"</colgroup>"

    heads=["<th>No</th>"]
    if show_sku:heads.append("<th>SKU</th>")
    heads.append("<th>Nama Barang / Jasa</th>")
    heads.append("<th>Qty</th>")
    if show_unit:heads.append("<th>Satuan</th>")
    if show_prices:heads.extend(["<th>Harga</th>","<th>Jumlah</th>"])

    rows=[]
    for index,item in enumerate(data.get("items") or [],1):
        cells=[f"<td class='ctr'>{index}</td>"]
        if show_sku:cells.append(f"<td>{html.escape(str(item.get('sku') or '-'))}</td>")
        cells.append(f"<td class='item-desc'>{html.escape(str(item.get('product_name') or '-'))}</td>")
        cells.append(f"<td class='num'>{float(item.get('qty') or 0):g}</td>")
        if show_unit:cells.append(f"<td>{html.escape(str(item.get('unit_code') or item.get('unit_name') or 'PCS'))}</td>")
        if show_prices:
            price=float(item.get("unit_price",item.get("unit_cost",0)) or 0)
            cells.append(f"<td class='num'>Rp{price:,.0f}</td>")
            cells.append(f"<td class='num'>Rp{float(item.get('line_total') or 0):,.0f}</td>")
        rows.append("<tr>"+"".join(cells)+"</tr>")

    total_html=""
    if show_prices:
        totals=[f"<div><span>Subtotal</span><b>Rp{float(data.get('subtotal',0)):,.0f}</b></div>"]
        if show_discount:totals.append(f"<div><span>Diskon</span><b>Rp{float(data.get('discount_amount',0)):,.0f}</b></div>")
        if show_tax:totals.append(f"<div><span>Pajak</span><b>Rp{float(data.get('tax_amount',0)):,.0f}</b></div>")
        totals.append(f"<div class='grand'><span>TOTAL</span><b>Rp{float(data.get('total_amount',0)):,.0f}</b></div>")
        total_html="<div class='totals'>"+"".join(totals)+"</div>"
        if show_amount_words:
            total_html+=f"<div class='amount-words'><b>Terbilang:</b> {html.escape(_terbilang_id(data.get('total_amount',0)))}</div>"

    paper=str(tpl.get("paper_size") or "A5").upper()
    page_size="58mm auto" if paper=="STRUK58" else ("80mm auto" if paper=="STRUK80" else ("A4" if paper=="A4" else ("Letter" if paper=="LETTER" else "A5")))
    width="58mm" if paper=="STRUK58" else ("80mm" if paper=="STRUK80" else ("210mm" if paper=="A4" else ("216mm" if paper=="LETTER" else "148mm")))
    receipt=paper in ("STRUK58","STRUK80")
    company=html.escape(profile.get("company_name") or "StokLedger Pro")
    address=html.escape(profile.get("address") or "")
    phone=html.escape(profile.get("phone") or "")
    taxid=html.escape(profile.get("tax_id") or "")
    title=html.escape(tpl.get("title") or dtype)
    header=html.escape(tpl.get("header_text") or "").replace("\n","<br>")
    footer=html.escape(tpl.get("footer_text") or "").replace("\n","<br>")
    note=html.escape(data.get("notes") or "").replace("\n","<br>")
    warehouse_row=f"<tr><td>Gudang</td><td>{html.escape(str(warehouse))}</td><td>Status</td><td>{html.escape(str(data.get('payment_method') or data.get('payment_type') or '-'))}</td></tr>" if show_warehouse else ""
    partner_contact=[]
    if partner_address: partner_contact.append(html.escape(str(partner_address)))
    if partner_city: partner_contact.append(html.escape(str(partner_city)))
    contact_bits=[x for x in (partner_phone,partner_email) if x]
    if contact_bits: partner_contact.append(html.escape(" · ".join(map(str,contact_bits))))
    partner_contact_row=(f"<tr><td>Alamat / Kontak</td><td colspan='3'>{'<br>'.join(partner_contact)}</td></tr>" if show_partner_contact and partner_contact else "")
    payment_method=str(data.get('payment_method') or data.get('payment_type') or '-').upper()
    payment_label={'CASH':'Tunai','TRANSFER':'Transfer','CREDIT':'Kredit'}.get(payment_method,payment_method.title() if payment_method else '-')
    due_raw=str(data.get('due_date') or '').strip()[:10]
    if due_raw:
        # due_date yang tersimpan pada transaksi adalah sumber utama. Jika pengguna
        # mengisi tanggal jatuh tempo manual, tanggal itulah yang dicetak; jika kosong,
        # backend transaksi sudah mengisinya dari termin pembayaran.
        try:
            due_label=date.fromisoformat(due_raw).strftime('%d/%m/%Y')
        except Exception:
            due_label=due_raw
        term_text=f"{payment_label} · Jatuh tempo {due_label}"
    else:
        term_text=payment_label
    term_block=f"<div class='payment-info'><b>Termin Pembayaran:</b> {html.escape(term_text)}</div>" if show_payment_terms else ""
    account_label=str(data.get('cash_account_name') or '').strip()
    account_line=(f"<br><b>Akun transaksi:</b> {html.escape(account_label)}" if account_label else '')
    payment_note_block=(f"<div class='payment-info'><b>Catatan Rekening Pembayaran:</b><br>{payment_account_note}{account_line}</div>" if show_payment_account_note and (payment_account_note or account_line) else "")
    header_block=f"<div class='intro'>{header}</div>" if header else ""
    special_block=f"<div class='special-note'><b>Keterangan Khusus:</b><br>{special_note}</div>" if special_note else ""
    footer_block=f"<footer>{footer}</footer>" if footer else ""

    return f"""<!doctype html><html><head><meta charset='utf-8'>
<title>{title} - {html.escape(str(number or ''))}</title><style>
@page{{size:{page_size} portrait;margin:8mm}}*{{box-sizing:border-box}}
body{{font:{'8px' if receipt else '10.5px'} Arial,Helvetica,sans-serif;color:#1d2b32;margin:0;background:#eef2f3}}
.sheet{{width:{width};min-height:{'auto' if receipt else '200mm'};margin:8px auto;background:#fff;padding:{'3mm' if receipt else '8mm'};box-shadow:0 2px 12px #8998a055}}
.toolbar{{width:{width};margin:8px auto;text-align:right}}button{{border:0;background:#087f8c;color:#fff;padding:8px 14px;border-radius:6px}}
header{{display:flex;gap:10px;align-items:center;border-bottom:2px solid #087f8c;padding-bottom:8px}}
.logo{{max-width:{'20mm' if receipt else '34mm'};max-height:{'12mm' if receipt else '20mm'}}}.company{{flex:1}}.company h2{{font-size:15px;margin:0 0 3px;color:#075f68}}
.company p{{margin:1px 0;color:#4b626b}}.doc-title{{text-align:right}}.doc-title h1{{font-size:16px;margin:0}}
.meta{{width:100%;border-collapse:collapse;margin:10px 0}}.meta td{{padding:3px 5px}}.meta td:nth-child(odd){{color:#52666e}}
.intro,.special-note,.payment-info{{background:#f1f8f8;border-left:3px solid #087f8c;padding:7px 9px;margin:8px 0}}
.items{{width:100%;border-collapse:collapse;margin-top:8px;table-layout:fixed}}.items th{{background:#087f8c;color:#fff;padding:5px 4px;overflow-wrap:anywhere}}
.items td{{border-bottom:1px solid #cbd8dc;padding:5px 4px;vertical-align:top;max-width:0;overflow-wrap:anywhere;word-break:break-word;white-space:normal}}.items .item-desc{{overflow-wrap:anywhere;word-break:break-word;white-space:normal;line-height:1.35}}.ctr{{text-align:center}}.num{{text-align:right;white-space:nowrap}}
.totals{{width:66mm;margin-left:auto;margin-top:8px}}.totals div{{display:flex;justify-content:space-between;padding:3px 5px}}
.totals .grand{{border-top:2px solid #087f8c;background:#eaf6f7;font-size:12px;padding:6px}}
.amount-words{{margin:6px 0 8px auto;max-width:90mm;padding:6px 8px;background:#f6fbfb;border-left:3px solid #087f8c;font-style:italic}}
.notes{{min-height:15mm;margin-top:8px;border:1px solid #d4dee1;border-radius:5px;padding:6px}}
.signatures{{display:{'none' if receipt else 'grid'};grid-template-columns:1fr 1fr;gap:25mm;margin-top:15px;text-align:center}}
.signatures div{{padding-top:20mm;border-bottom:1px solid #536970}}
footer{{margin-top:10px;padding-top:6px;border-top:1px solid #b9c9cd;text-align:center;color:#65777e}}
@media print{{body{{background:#fff}}.toolbar{{display:none}}.sheet{{margin:0;box-shadow:none;width:auto;min-height:auto;padding:0}}}}
</style></head><body><div class='toolbar'><button onclick='window.print()'>Cetak Dokumen</button></div>
<div class='sheet'><header>{logo_html}<div class='company'><h2>{company}</h2><p>{address}</p>
<p>{phone}{' · NPWP '+taxid if taxid else ''}</p></div><div class='doc-title'><h1>{title}</h1>
<b>{html.escape(str(number or '-'))}</b></div></header>{header_block}
<table class='meta'><tr><td>Tanggal</td><td>{html.escape(str(txdate or '-'))}</td>
<td>Partner</td><td>{html.escape(str(partner))}</td></tr>{partner_contact_row}{warehouse_row}</table>
<table class='items'>{colgroup}<thead><tr>{''.join(heads)}</tr></thead><tbody>{''.join(rows)}</tbody></table>
{total_html}{term_block}{payment_note_block}<div class='notes'><b>Catatan Transaksi:</b><br>{note or '-'}</div>{special_block}
<div class='signatures'><div>{left_signer}</div><div>{right_signer}</div></div>{footer_block}</div></body></html>"""

def preview_bank_csv(actor,csv_text,file_name="rekening.csv",bank_name=None):
    reader=csv.DictReader(io.StringIO(csv_text));fields={str(x).strip().lower() for x in (reader.fieldnames or [])}
    if not {"tanggal","keterangan","debit","kredit"}.issubset(fields):raise ValueError("CSV wajib: tanggal,keterangan,debit,kredit. saldo opsional.")
    items=[]
    def num(v):
        t=str(v or "0").strip()
        if not t:return 0.0
        return float(t.replace(".","").replace(",","."))
    for no,row in enumerate(reader,2):
        r={str(k).strip().lower():v for k,v in row.items()};debit=num(r.get("debit"));credit=num(r.get("kredit"))
        if (debit>0)==(credit>0):raise ValueError(f"Baris {no}: isi salah satu debit atau kredit.")
        items.append({"transaction_date":str(r.get("tanggal",""))[:10],"description":str(r.get("keterangan","")).strip(),"debit":debit,"credit":credit,"balance":num(r.get("saldo")) if str(r.get("saldo","")).strip() else None})
    with write_transaction() as tx:
        cur=tx.execute("INSERT INTO bank_import_batches(file_name,bank_name,row_count,posted_count,status,user_id,created_at) VALUES(?,?,?,0,'PREVIEW',?,?)",(file_name,bank_name,len(items),actor["id"],utc_now()));bid=cur.lastrowid
        for x in items:tx.execute("INSERT INTO bank_import_rows(batch_id,transaction_date,description,debit,credit,balance,status) VALUES(?,?,?,?,?,?,'PREVIEW')",(bid,x["transaction_date"],x["description"],x["debit"],x["credit"],x["balance"]))
    return {"batch_id":bid,"row_count":len(items),"items":items}
def bank_import_rows(batch_id):
    c=connect()
    try:return [dict(x) for x in c.execute("SELECT * FROM bank_import_rows WHERE batch_id=? ORDER BY id",(batch_id,))]
    finally:c.close()
def post_bank_import(actor,bid,cash_id,mappings,ip):
    bid=int(bid);cash_id=int(cash_id)
    if not isinstance(mappings,list) or not mappings:raise ValueError("Belum ada baris dipilih.")
    mapped={}
    for m in mappings:
        rid=int(m.get("row_id"));mapped[rid]={"coa":int(m.get("counter_coa_id")),"partner":int(m["partner_id"]) if m.get("partner_id") not in (None,"") else None}
    with write_transaction() as tx:
        licensing.enforce_transaction_capacity(tx,len(mapped))
        bank=tx.execute("SELECT * FROM cash_accounts WHERE id=? AND is_active=1",(cash_id,)).fetchone()
        if not bank:raise ValueError("Akun Kas/Bank tidak valid.")
        marks=",".join("?" for _ in mapped)
        rows=tx.execute(f"SELECT * FROM bank_import_rows WHERE batch_id=? AND status='PREVIEW' AND id IN ({marks}) ORDER BY id",[bid,*mapped]).fetchall()
        if len(rows)!=len(mapped):raise ValueError("Baris preview tidak lengkap atau sudah diposting.")
        bank_coa=bank["coa_account_id"] or accounting_service.account_id(tx,"1010")
        posted=ins=outs=0
        for row in rows:
            m=mapped[row["id"]]
            coa=tx.execute("SELECT id,name,account_subtype FROM chart_of_accounts WHERE id=? AND is_active=1",(m["coa"],)).fetchone()
            if not coa or coa["account_subtype"]=="CASH_BANK":raise ValueError(f"Baris {row['id']}: akun lawan tidak valid.")
            partner=m["partner"]
            if coa["account_subtype"] in ("RECEIVABLE","PAYABLE"):
                need="CUSTOMER" if coa["account_subtype"]=="RECEIVABLE" else "SUPPLIER"
                p=tx.execute("SELECT partner_type FROM business_partners WHERE id=? AND is_active=1",(partner,)).fetchone() if partner else None
                if not p or p["partner_type"] not in (need,"BOTH"):raise ValueError(f"Baris {row['id']}: partner wajib dipilih.")
            else:partner=None
            try:
                credit_value=cash_service.money(row["credit"],f"Bank Masuk baris {row['id']}")
                debit_value=cash_service.money(row["debit"],f"Bank Keluar baris {row['id']}")
            except ValueError as exc:
                raise ValueError(str(exc)) from exc
            if (credit_value > 0) == (debit_value > 0):
                raise ValueError(f"Baris {row['id']}: isi tepat salah satu Bank Masuk atau Bank Keluar.")
            incoming=credit_value > 0
            value=credit_value if incoming else debit_value
            ref=f"BANK-{bid}-{row['id']}"
            cash=cash_service.post(tx,account_id=cash_id,transaction_date=row["transaction_date"],transaction_type="IN" if incoming else "OUT",amount=value,description=row["description"],reference_no=ref,user_id=actor["id"],allow_negative=True)
            other={"account_id":coa["id"],"partner_id":partner}
            lines=[{"account_id":bank_coa,"debit":value},{**other,"credit":value}] if incoming else [{**other,"debit":value},{"account_id":bank_coa,"credit":value}]
            accounting_service.post_journal(tx,journal_date=row["transaction_date"],description=row["description"],source_type="BANK_IMPORT",source_id=row["id"],reference_no=ref,lines=lines,user_id=actor["id"])
            tx.execute("UPDATE bank_import_rows SET cash_account_id=?,counter_coa_id=?,posted_transaction_id=?,status='POSTED' WHERE id=?",(cash_id,coa["id"],cash["id"],row["id"]))
            posted+=1;ins+=1 if incoming else 0;outs+=0 if incoming else 1
        left=tx.execute("SELECT COUNT(*) n FROM bank_import_rows WHERE batch_id=? AND status='PREVIEW'",(bid,)).fetchone()["n"]
        total=tx.execute("SELECT COUNT(*) n FROM bank_import_rows WHERE batch_id=? AND status='POSTED'",(bid,)).fetchone()["n"]
        tx.execute("UPDATE bank_import_batches SET posted_count=?,status=? WHERE id=?",(total,"POSTED" if left==0 else "PARTIAL",bid))
        return {"posted_count":posted,"incoming_count":ins,"outgoing_count":outs,"remaining_count":left}

def owner_mobile_dashboard():
    d=dashboard_analytics();d["cash"]=cash_summary();d["sales"]=sales_summary(d["period"]["from"],d["period"]["to"]);d["purchases"]=purchases_summary(d["period"]["from"],d["period"]["to"]);return d



def _suggest_bank_counter_account(tx, description, direction):
    text=(description or "").upper()
    rules=[
      (("BIAYA ADM","ADMINISTRASI","ADM "),("5100",0.99,"Biaya administrasi bank")),
      (("GAJI","SALARY","SPV ","DIVISI ","BONUS","REWARD"),("5100",0.86,"Biaya operasional / penggajian")),
      (("PLN","TELKOM","INDIH","AETRA","BPJS","LISTRIK","AIR"),("5100",0.90,"Biaya utilitas")),
      (("BENSIN","SPBU","TOL","PARKIR","TRANSPORT","REM BUS","REMBUS"),("5100",0.85,"Biaya transportasi")),
      (("TOKOPEDIA","TRAVELOKA","DANA","QRIS","KARTU DEBIT","FLAZZ","TARIKAN ATM"),("5100",0.72,"Biaya operasional")),
      (("SEWA","RENT"),("5100",0.88,"Biaya sewa")),
      (("VPS","HOSTING","ZOOM","DROPBOX","VPN","ACCURATE ONL"),("5100",0.88,"Biaya software / layanan")),
      (("INV/","PELUNASAN","PEMBAYARAN","TERMIN","ACCURATE","TRAINING","SUPPORT","DONGLE","TOOLS"),("1100",0.78,"Kemungkinan penerimaan piutang pelanggan")),
      (("SETORAN TUNAI","BI-FAST CR","KR OTOMATIS","TRSF E-BANKING CR"),("4000",0.60,"Kemungkinan pendapatan atau penerimaan")),
    ]
    for keywords,(code,score,reason) in rules:
        if any(k in text for k in keywords):
            row=tx.execute("SELECT id,code,name,account_subtype FROM chart_of_accounts WHERE code=? AND is_active=1",(code,)).fetchone()
            if row:return {"counter_coa_id":row["id"],"account_code":row["code"],"account_name":row["name"],"account_subtype":row["account_subtype"],"score":score,"reason":reason}
    fallback="4000" if direction=="CREDIT" else "5100"
    row=tx.execute("SELECT id,code,name,account_subtype FROM chart_of_accounts WHERE code=? AND is_active=1",(fallback,)).fetchone()
    if not row:return None
    return {"counter_coa_id":row["id"],"account_code":row["code"],"account_name":row["name"],"account_subtype":row["account_subtype"],"score":0.45,"reason":"Saran umum berdasarkan arah mutasi bank"}

def preview_bank_pdf(actor,pdf_base64,file_name="rekening_koran.pdf"):
    try:data=base64.b64decode(pdf_base64,validate=True)
    except Exception as exc:raise ValueError("Data PDF tidak valid.") from exc
    parsed=bank_pdf_parser.parse_statement_pdf(data);saved=[]
    with write_transaction() as tx:
        b=tx.execute("INSERT INTO bank_import_batches(file_name,bank_name,row_count,posted_count,status,user_id,created_at) VALUES(?,?,?,0,'PREVIEW',?,?)",(file_name,parsed["bank_name"],len(parsed["items"]),actor["id"],utc_now()))
        for x in parsed["items"]:
            r=tx.execute("INSERT INTO bank_import_rows(batch_id,transaction_date,description,debit,credit,balance,status) VALUES(?,?,?,?,?,?,'PREVIEW')",(b.lastrowid,x["transaction_date"],x["description"],x["debit"],x["credit"],x["balance"]))
            suggestion=_suggest_bank_counter_account(tx,x["description"],x["direction"])
            saved.append({**x,"id":r.lastrowid,"suggestion":suggestion})
    return {**parsed,"batch_id":b.lastrowid,"items":saved}


def income_statement(date_from=None,date_to=None):
    date_from=_date_filter(date_from,'Tanggal mulai')
    date_to=_date_filter(date_to,'Tanggal akhir')
    c=connect()
    try:
        where=["j.status='POSTED'"];params=[]
        if date_from:where.append("j.journal_date>=?");params.append(date_from)
        if date_to:where.append("j.journal_date<=?");params.append(date_to)
        rows=c.execute("""SELECT a.id,a.code,a.name,a.account_type,a.account_subtype,
          COALESCE(SUM(CASE WHEN """+" AND ".join(where)+""" THEN l.debit ELSE 0 END),0) debit,
          COALESCE(SUM(CASE WHEN """+" AND ".join(where)+""" THEN l.credit ELSE 0 END),0) credit
          FROM chart_of_accounts a
          LEFT JOIN journal_lines l ON l.account_id=a.id
          LEFT JOIN journal_entries j ON j.id=l.journal_id
          WHERE (a.account_subtype IN ('REVENUE','HPP','OPERATING_EXPENSE') OR a.account_type IN ('REVENUE','EXPENSE'))
          GROUP BY a.id ORDER BY a.code""",params+params).fetchall()
        revenue_items=[];hpp_items=[];expense_items=[]
        for r in rows:
            subtype=r["account_subtype"] or ""; atype=r["account_type"] or ""
            bucket=("REVENUE" if subtype=="REVENUE" or atype=="REVENUE" else ("HPP" if subtype=="HPP" else "OPERATING_EXPENSE"))
            amount=(float(r["credit"])-float(r["debit"])) if bucket=="REVENUE" else (float(r["debit"])-float(r["credit"]))
            item={"code":r["code"],"name":r["name"],"amount":round(amount,2)}
            if bucket=="REVENUE":revenue_items.append(item)
            elif bucket=="HPP":hpp_items.append(item)
            else:expense_items.append(item)
        revenue=sum(x["amount"] for x in revenue_items)
        hpp=sum(x["amount"] for x in hpp_items)
        gross=revenue-hpp
        expense=sum(x["amount"] for x in expense_items)
        return {"date_from":date_from,"date_to":date_to,
          "revenue_items":revenue_items,"hpp_items":hpp_items,"expense_items":expense_items,
          "revenue":round(revenue,2),"hpp":round(hpp,2),"gross_profit":round(gross,2),
          "operating_expense":round(expense,2),"net_profit":round(gross-expense,2)}
    finally:c.close()


# --- Fixed Assets v1.9.4 ---
def _fa_date(value,label):
    return _date_value(value,label)

def create_fixed_asset(actor,d,ip):
    code=str(d.get("asset_code","")).strip().upper(); name=str(d.get("asset_name","")).strip()
    if not code: raise ValueError("Kode aktiva tetap wajib diisi.")
    if not name: raise ValueError("Nama aktiva tetap wajib diisi.")
    acq=_fa_date(d.get("acquisition_date"),"Tanggal perolehan")
    cost=_decimal(d.get("acquisition_cost",0),"Harga perolehan")
    residual=_decimal(d.get("residual_value",0),"Nilai residu")
    if cost<=0: raise ValueError("Harga perolehan harus lebih dari nol.")
    if residual<0 or residual>=cost: raise ValueError("Nilai residu harus lebih kecil dari harga perolehan.")
    try: life=int(d.get("useful_life_months",0))
    except: raise ValueError("Masa manfaat harus berupa jumlah bulan.")
    if life<0: raise ValueError("Masa manfaat tidak boleh negatif. Gunakan 0 untuk tanah/aset yang tidak disusutkan.")
    ids=[]
    for key,label in [("asset_account_id","Akun aktiva tetap"),("accumulated_depreciation_account_id","Akun akumulasi penyusutan"),("depreciation_expense_account_id","Akun beban penyusutan"),("contra_account_id","Akun lawan")]:
        try: aid=int(d.get(key))
        except: raise ValueError(label+" wajib dipilih.")
        ids.append(aid)
    now=utc_now()
    with write_transaction() as tx:
        rows=[]
        for aid in ids:
            row=tx.execute("SELECT id,account_subtype FROM chart_of_accounts WHERE id=? AND is_active=1",(aid,)).fetchone()
            if not row: raise ValueError("Salah satu akun tidak ditemukan atau nonaktif.")
            if row["account_subtype"] in ("RECEIVABLE","PAYABLE"): raise ValueError("Akun Piutang/Hutang Usaha tidak dapat dipakai pada aktiva tetap karena membutuhkan partner.")
            rows.append(row)
        cur=tx.execute("""INSERT INTO fixed_assets(asset_code,asset_name,acquisition_date,acquisition_cost,residual_value,useful_life_months,asset_account_id,accumulated_depreciation_account_id,depreciation_expense_account_id,contra_account_id,notes,status,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,'ACTIVE',?,?)""",(code,name,acq,cost,residual,life,*ids,str(d.get("notes","")).strip() or None,now,now))
        asset_id=cur.lastrowid
        jid=accounting_service.post_journal(tx,journal_date=acq,description=f"Perolehan aktiva tetap {code} - {name}",source_type="FIXED_ASSET_ACQUISITION",source_id=asset_id,reference_no=code,lines=[{"account_id":ids[0],"debit":cost},{"account_id":ids[3],"credit":cost}],user_id=actor["id"])
        audit(actor["id"],"FIXED_ASSET_CREATED","fixed_asset",asset_id,{"code":code,"cost":float(cost),"journal_id":jid},ip,tx)
        return asset_id

def list_fixed_assets():
    c=connect()
    try:
        rows=c.execute("""SELECT fa.*,aa.code||' - '||aa.name asset_account_name,ad.code||' - '||ad.name accumulated_account_name,de.code||' - '||de.name expense_account_name,ca.code||' - '||ca.name contra_account_name,COALESCE((SELECT SUM(amount) FROM fixed_asset_depreciations d WHERE d.asset_id=fa.id),0) accumulated_depreciation FROM fixed_assets fa JOIN chart_of_accounts aa ON aa.id=fa.asset_account_id JOIN chart_of_accounts ad ON ad.id=fa.accumulated_depreciation_account_id JOIN chart_of_accounts de ON de.id=fa.depreciation_expense_account_id JOIN chart_of_accounts ca ON ca.id=fa.contra_account_id ORDER BY fa.acquisition_date DESC,fa.id DESC""").fetchall()
        out=[]
        for r in rows:
            x=dict(r); x['acquisition_cost']=float(x['acquisition_cost']); x['residual_value']=float(x['residual_value']); x['accumulated_depreciation']=float(x['accumulated_depreciation']); x['book_value']=round(x['acquisition_cost']-x['accumulated_depreciation'],2); x['monthly_depreciation']=0.0 if int(x['useful_life_months'] or 0)==0 else round((x['acquisition_cost']-x['residual_value'])/x['useful_life_months'],2); out.append(x)
        return out
    finally:c.close()

def get_fixed_asset(asset_id):
    c=connect()
    try:
        row=c.execute("SELECT * FROM fixed_assets WHERE id=?",(int(asset_id),)).fetchone()
        return dict(row) if row else None
    finally:c.close()

def list_fixed_asset_depreciations(asset_id=None):
    c=connect()
    try:
        sql="""SELECT d.*,fa.asset_code,fa.asset_name,j.journal_no FROM fixed_asset_depreciations d JOIN fixed_assets fa ON fa.id=d.asset_id LEFT JOIN journal_entries j ON j.id=d.journal_id"""
        args=[]
        if asset_id: sql+=" WHERE d.asset_id=?";args.append(int(asset_id))
        sql+=" ORDER BY d.period DESC,d.id DESC"
        return [dict(x) for x in c.execute(sql,args).fetchall()]
    finally:c.close()

def update_fixed_asset(actor,asset_id,d,ip):
    asset_id=int(asset_id); old=get_fixed_asset(asset_id)
    if not old: raise ValueError("Aktiva tetap tidak ditemukan.")
    code=str(d.get("asset_code",old['asset_code'])).strip().upper();name=str(d.get("asset_name",old['asset_name'])).strip()
    acq=str(d.get("acquisition_date",old['acquisition_date']));cost=_money(d.get("acquisition_cost",old['acquisition_cost']));residual=_money(d.get("residual_value",old['residual_value']));life=int(d.get("useful_life_months",old['useful_life_months']))
    if not code or not name or cost<=0 or residual<0 or residual>=cost or life<0: raise ValueError("Data aktiva tetap tidak valid. Masa manfaat 0 diperbolehkan untuk aset yang tidak disusutkan.")
    ids=[int(d.get(k,old[k])) for k in ('asset_account_id','accumulated_depreciation_account_id','depreciation_expense_account_id','contra_account_id')]
    with write_transaction() as tx:
        if tx.execute("SELECT 1 FROM fixed_assets WHERE asset_code=? COLLATE NOCASE AND id<>?",(code,asset_id)).fetchone():
            raise ValueError('Kode aktiva sudah digunakan oleh aktiva lain.')
        dep=tx.execute("SELECT COALESCE(SUM(amount),0) v FROM fixed_asset_depreciations WHERE asset_id=?",(asset_id,)).fetchone()['v']
        if Decimal(str(dep))>0 and (cost!=_money(old['acquisition_cost']) or acq!=old['acquisition_date'] or ids[0]!=old['asset_account_id'] or ids[3]!=old['contra_account_id']):
            raise ValueError("Harga, tanggal, dan akun perolehan tidak dapat diubah setelah penyusutan diproses. Hapus penyusutan terlebih dahulu.")
        tx.execute("""UPDATE fixed_assets SET asset_code=?,asset_name=?,acquisition_date=?,acquisition_cost=?,residual_value=?,useful_life_months=?,asset_account_id=?,accumulated_depreciation_account_id=?,depreciation_expense_account_id=?,contra_account_id=?,notes=?,updated_at=? WHERE id=?""",(code,name,acq,float(cost),float(residual),life,*ids,str(d.get('notes',old.get('notes') or '')).strip() or None,utc_now(),asset_id))
        if Decimal(str(dep))==0:
            j=tx.execute("SELECT * FROM journal_entries WHERE source_type='FIXED_ASSET' AND CAST(source_id AS TEXT)=? ORDER BY id LIMIT 1",(str(asset_id),)).fetchone()
            if j:
                tx.execute("DELETE FROM journal_lines WHERE journal_id=?",(j['id'],));tx.execute("INSERT INTO journal_lines(journal_id,account_id,debit,credit) VALUES(?,?,?,0)",(j['id'],ids[0],str(cost)));tx.execute("INSERT INTO journal_lines(journal_id,account_id,debit,credit) VALUES(?,?,0,?)",(j['id'],ids[3],str(cost)));tx.execute("UPDATE journal_entries SET journal_date=?,description=?,reference_no=?,status='POSTED',total_debit=?,total_credit=?,updated_at=? WHERE id=?",(acq,f'Perolehan aktiva {code} - {name}',code,str(cost),str(cost),utc_now(),j['id']))
        audit(actor['id'],'FIXED_ASSET_UPDATED','fixed_asset',asset_id,{'code':code},ip,tx)
    return get_fixed_asset(asset_id)

def delete_fixed_asset_depreciation(actor,dep_id,ip):
    with write_transaction() as tx:
        d=tx.execute("SELECT * FROM fixed_asset_depreciations WHERE id=?",(int(dep_id),)).fetchone()
        if not d: raise ValueError("Penyusutan tidak ditemukan.")
        if d['journal_id']: tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE id=?",(utc_now(),d['journal_id']))
        tx.execute("DELETE FROM fixed_asset_depreciations WHERE id=?",(int(dep_id),))
        audit(actor['id'],'FIXED_ASSET_DEPRECIATION_DELETED','fixed_asset_depreciation',dep_id,{'period':d['period']},ip,tx)
    return {'id':int(dep_id)}

def delete_fixed_asset(actor,asset_id,ip):
    with write_transaction() as tx:
        a=tx.execute("SELECT * FROM fixed_assets WHERE id=?",(int(asset_id),)).fetchone()
        if not a: raise ValueError("Aktiva tetap tidak ditemukan.")
        if tx.execute("SELECT 1 FROM fixed_asset_depreciations WHERE asset_id=?",(int(asset_id),)).fetchone(): raise ValueError("Hapus seluruh penyusutan aktiva ini terlebih dahulu.")
        tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE source_type='FIXED_ASSET_ACQUISITION' AND source_id=?",(utc_now(),int(asset_id)))
        tx.execute("DELETE FROM fixed_assets WHERE id=?",(int(asset_id),))
        audit(actor['id'],'FIXED_ASSET_DELETED','fixed_asset',asset_id,{'code':a['asset_code']},ip,tx)
    return {'id':int(asset_id)}

def run_fixed_asset_depreciation(actor,period,ip):
    period=str(period or "").strip()
    if not re.match(r'^\d{4}-\d{2}$',period): raise ValueError("Periode harus berformat YYYY-MM.")
    from calendar import monthrange
    target_y,target_m=map(int,period.split('-')); target_index=target_y*12+target_m-1
    created=[]
    with write_transaction() as tx:
        assets=tx.execute("SELECT * FROM fixed_assets WHERE status='ACTIVE' AND acquisition_date<=? ORDER BY id",(f"{target_y:04d}-{target_m:02d}-{monthrange(target_y,target_m)[1]:02d}",)).fetchall()
        for a in assets:
            acq=date.fromisoformat(a['acquisition_date']); start_index=acq.year*12+acq.month-1
            life=int(a['useful_life_months'])
            if life<=0:
                # Tanah / aset non-depreciable: tetap tercatat sebagai aktiva, tanpa jurnal penyusutan.
                continue
            maxdep=Decimal(str(a['acquisition_cost']))-Decimal(str(a['residual_value']))
            monthly=(maxdep/Decimal(life)).quantize(Decimal('0.01'))
            last_index=min(target_index,start_index+life-1)
            for idx in range(start_index,last_index+1):
                yy=idx//12;mm=idx%12+1;per=f"{yy:04d}-{mm:02d}"
                if tx.execute("SELECT 1 FROM fixed_asset_depreciations WHERE asset_id=? AND period=?",(a['id'],per)).fetchone():continue
                done=Decimal(str(tx.execute("SELECT COALESCE(SUM(amount),0) v FROM fixed_asset_depreciations WHERE asset_id=?",(a['id'],)).fetchone()['v']))
                amount=min(monthly,maxdep-done)
                if amount<=0:break
                dep_date=f"{yy:04d}-{mm:02d}-{monthrange(yy,mm)[1]:02d}"
                jid=accounting_service.post_journal(tx,journal_date=dep_date,description=f"Penyusutan {per} {a['asset_code']} - {a['asset_name']}",source_type="FIXED_ASSET_DEPRECIATION",source_id=None,reference_no=f"DEP-{per}-{a['asset_code']}",lines=[{"account_id":a['depreciation_expense_account_id'],"debit":amount},{"account_id":a['accumulated_depreciation_account_id'],"credit":amount}],user_id=actor['id'])
                cur=tx.execute("INSERT INTO fixed_asset_depreciations(asset_id,period,depreciation_date,amount,journal_id,created_at) VALUES(?,?,?,?,?,?)",(a['id'],per,dep_date,float(amount),jid['id'],utc_now()))
                created.append({'asset_id':a['id'],'asset_code':a['asset_code'],'period':per,'amount':float(amount)})
        audit(actor['id'],'FIXED_ASSET_DEPRECIATION_RUN','fixed_asset',None,{'period':period,'count':len(created),'total':sum(x['amount'] for x in created)},ip,tx)
    return {'period':period,'count':len(created),'total':round(sum(x['amount'] for x in created),2),'items':created}


def _sync_invoice_balance_after_return(tx,kind,invoice_id):
    if not invoice_id:return
    customer=kind=='sales'; table='sales' if customer else 'purchases'; pk='sale_id' if customer else 'purchase_id'
    ret='sales_returns' if customer else 'purchase_returns'; total_col='total_sales' if customer else 'total_payable'
    row=tx.execute(f"SELECT total_amount,paid_amount FROM {table} WHERE id=? AND status='POSTED'",(int(invoice_id),)).fetchone()
    if not row:return
    returned=_money(tx.execute(f"SELECT COALESCE(SUM({total_col}),0) v FROM {ret} WHERE {pk}=? AND status='POSTED'",(int(invoice_id),)).fetchone()['v'])
    net=max(Decimal('0'),_money(row['total_amount'])-returned); paid=_money(row['paid_amount'])
    tx.execute(f"UPDATE {table} SET balance_due=?,updated_at=? WHERE id=?",(str(max(Decimal('0'),net-paid)),utc_now(),int(invoice_id)))

def create_purchase_return(actor,d,ip):
    with write_transaction() as tx:
        result=returns_service.create_purchase_return(tx,actor,d,ip,audit)
        _sync_invoice_balance_after_return(tx,'purchase',d.get('purchase_id'))
        return result
def create_sales_return(actor,d,ip):
    with write_transaction() as tx:
        result=returns_service.create_sales_return(tx,actor,d,ip,audit)
        _sync_invoice_balance_after_return(tx,'sales',d.get('sale_id'))
        return result
def list_purchase_returns():
    c=connect()
    try:return [dict(x) for x in c.execute("SELECT r.*,bp.name supplier_name,w.name warehouse_name FROM purchase_returns r JOIN business_partners bp ON bp.id=r.supplier_id JOIN warehouses w ON w.id=r.warehouse_id ORDER BY r.return_date DESC,r.id DESC LIMIT 300")]
    finally:c.close()
def list_sales_returns():
    c=connect()
    try:return [dict(x) for x in c.execute("SELECT r.*,bp.name customer_name,w.name warehouse_name FROM sales_returns r JOIN business_partners bp ON bp.id=r.customer_id JOIN warehouses w ON w.id=r.warehouse_id ORDER BY r.return_date DESC,r.id DESC LIMIT 300")]
    finally:c.close()



def get_purchase_return(return_id):
    c=connect()
    try:
        row=c.execute("SELECT r.*,bp.name supplier_name,w.name warehouse_name FROM purchase_returns r JOIN business_partners bp ON bp.id=r.supplier_id JOIN warehouses w ON w.id=r.warehouse_id WHERE r.id=?",(int(return_id),)).fetchone()
        if not row:return None
        d=dict(row);d['items']=[dict(x) for x in c.execute("SELECT * FROM purchase_return_items WHERE return_id=? ORDER BY id",(int(return_id),)).fetchall()];return d
    finally:c.close()

def get_sales_return(return_id):
    c=connect()
    try:
        row=c.execute("SELECT r.*,bp.name customer_name,w.name warehouse_name FROM sales_returns r JOIN business_partners bp ON bp.id=r.customer_id JOIN warehouses w ON w.id=r.warehouse_id WHERE r.id=?",(int(return_id),)).fetchone()
        if not row:return None
        d=dict(row);d['items']=[dict(x) for x in c.execute("SELECT * FROM sales_return_items WHERE return_id=? ORDER BY id",(int(return_id),)).fetchall()];return d
    finally:c.close()

def delete_purchase_return(actor,return_id,ip):
    with write_transaction() as tx:
        row=tx.execute("SELECT * FROM purchase_returns WHERE id=?",(int(return_id),)).fetchone()
        if not row:raise ValueError('Retur pembelian tidak ditemukan.')
        if row['status']=='VOID':raise ValueError('Retur pembelian sudah dibatalkan.')
        items=tx.execute("SELECT * FROM purchase_return_items WHERE return_id=?",(int(return_id),)).fetchall();now=utc_now()
        for i in items:
            inventory_service.post_movement(tx,product_id=i['product_id'],warehouse_id=row['warehouse_id'],movement_type='PURCHASE_RETURN_VOID',quantity_change=Decimal(str(i['qty'])),unit_cost=Decimal(str(i['average_cost'])),reference_type='PURCHASE_RETURN_VOID',reference_no='VOID-'+row['return_no'],reason='Pembalikan '+row['return_no'],user_id=actor['id'],created_at=now)
        tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE source_type='PURCHASE_RETURN' AND CAST(source_id AS TEXT)=?",(now,str(return_id)))
        tx.execute("UPDATE purchase_returns SET status='VOID' WHERE id=?",(int(return_id),));_sync_invoice_balance_after_return(tx,'purchase',row['purchase_id']);audit(actor['id'],'PURCHASE_RETURN_VOIDED','purchase_return',return_id,{'return_no':row['return_no']},ip,tx)
    return {'id':int(return_id),'status':'VOID'}

def delete_sales_return(actor,return_id,ip):
    with write_transaction() as tx:
        row=tx.execute("SELECT * FROM sales_returns WHERE id=?",(int(return_id),)).fetchone()
        if not row:raise ValueError('Retur penjualan tidak ditemukan.')
        if row['status']=='VOID':raise ValueError('Retur penjualan sudah dibatalkan.')
        items=tx.execute("SELECT * FROM sales_return_items WHERE return_id=?",(int(return_id),)).fetchall();now=utc_now()
        for i in items:
            if i['product_type']=='STOCK':
                before,_=inventory_service.get_balance(tx,row['warehouse_id'],i['product_id'])
                q=Decimal(str(i['qty']))
                if before<q:raise ValueError(f"Stok {i['product_name']} tidak cukup untuk membatalkan retur penjualan.")
                inventory_service.post_movement(tx,product_id=i['product_id'],warehouse_id=row['warehouse_id'],movement_type='SALES_RETURN_VOID',quantity_change=-q,unit_cost=Decimal(str(i['average_cost'])),reference_type='SALES_RETURN_VOID',reference_no='VOID-'+row['return_no'],reason='Pembalikan '+row['return_no'],user_id=actor['id'],created_at=now)
        tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE source_type='SALES_RETURN' AND CAST(source_id AS TEXT)=?",(now,str(return_id)))
        tx.execute("UPDATE sales_returns SET status='VOID' WHERE id=?",(int(return_id),));_sync_invoice_balance_after_return(tx,'sales',row['sale_id']);audit(actor['id'],'SALES_RETURN_VOIDED','sales_return',return_id,{'return_no':row['return_no']},ip,tx)
    return {'id':int(return_id),'status':'VOID'}

def update_purchase_return(actor,return_id,data,ip):
    delete_purchase_return(actor,return_id,ip);return create_purchase_return(actor,data,ip)

def update_sales_return(actor,return_id,data,ip):
    delete_sales_return(actor,return_id,ip);return create_sales_return(actor,data,ip)

def close_prior_fiscal_years():
    current=date.today().year
    with write_transaction() as tx:
      years=[int(x['y']) for x in tx.execute("SELECT DISTINCT CAST(substr(journal_date,1,4) AS INTEGER) y FROM journal_entries WHERE status='POSTED' AND CAST(substr(journal_date,1,4) AS INTEGER)<?",(current,)).fetchall()]
      for y in years:
       if tx.execute("SELECT 1 FROM fiscal_year_closings WHERE fiscal_year=?",(y,)).fetchone():continue
       rows=tx.execute("SELECT a.id,a.code,a.account_type,COALESCE(SUM(l.debit),0) d,COALESCE(SUM(l.credit),0) c FROM journal_entries j JOIN journal_lines l ON l.journal_id=j.id JOIN chart_of_accounts a ON a.id=l.account_id WHERE j.status='POSTED' AND substr(j.journal_date,1,4)=? AND a.account_type IN ('REVENUE','EXPENSE') GROUP BY a.id",(str(y),)).fetchall()
       lines=[]; profit=Decimal('0')
       for r in rows:
        bal=Decimal(str(r['c']))-Decimal(str(r['d'])) if r['account_type']=='REVENUE' else Decimal(str(r['d']))-Decimal(str(r['c']))
        if bal==0:continue
        if r['account_type']=='REVENUE': lines.append({'account_id':r['id'],'debit':bal});profit+=bal
        else: lines.append({'account_id':r['id'],'credit':bal});profit-=bal
       if not lines:continue
       if profit>=0:lines.append({'account_code':'3100','credit':profit})
       else:lines.append({'account_code':'3100','debit':-profit})
       j=accounting_service.post_journal(tx,journal_date=f'{y}-12-31',description=f'Penutupan laba rugi tahun {y}',source_type='FISCAL_YEAR_CLOSE',source_id=y,reference_no=f'CLOSE-{y}',lines=lines,user_id=1)
       tx.execute("INSERT INTO fiscal_year_closings(fiscal_year,closing_date,net_profit,journal_id,created_at) VALUES(?,?,?,?,?)",(y,f'{y}-12-31',str(profit),j['id'],utc_now()))


# --- Transaction maintenance lists / edit / delete (v1.10.6) ---
def _voided_keys(conn, kind):
    return {str(r["transaction_key"]) for r in conn.execute(
        "SELECT transaction_key FROM transaction_voids WHERE transaction_kind=?",(kind,)).fetchall()}

def list_transaction_maintenance(kind, limit=500):
    limit=max(1,min(int(limit),1000)); c=connect()
    try:
        voided=_voided_keys(c,kind)
        if kind=='cash':
            rows=c.execute("""SELECT t.id key,t.transaction_no number,t.transaction_date date,
              a.name account_name,t.transaction_type type,t.amount,t.description,t.reference_no,
              t.account_id,t.department_id,t.project_id,
              (SELECT jl.account_id FROM journal_entries je JOIN journal_lines jl ON jl.journal_id=je.id
                WHERE je.source_type='CASH_TRANSACTION' AND CAST(je.source_id AS TEXT)=CAST(t.id AS TEXT)
                  AND je.status='POSTED'
                  AND jl.account_id<>COALESCE(a.coa_account_id,-1)
                ORDER BY je.id DESC,jl.id LIMIT 1) counter_account_id,
              (SELECT ca.code FROM journal_entries je JOIN journal_lines jl ON jl.journal_id=je.id
                 JOIN chart_of_accounts ca ON ca.id=jl.account_id
                WHERE je.source_type='CASH_TRANSACTION' AND CAST(je.source_id AS TEXT)=CAST(t.id AS TEXT)
                  AND je.status='POSTED' AND jl.account_id<>COALESCE(a.coa_account_id,-1)
                ORDER BY je.id DESC,jl.id LIMIT 1) counter_account_code,
              (SELECT ca.name FROM journal_entries je JOIN journal_lines jl ON jl.journal_id=je.id
                 JOIN chart_of_accounts ca ON ca.id=jl.account_id
                WHERE je.source_type='CASH_TRANSACTION' AND CAST(je.source_id AS TEXT)=CAST(t.id AS TEXT)
                  AND je.status='POSTED' AND jl.account_id<>COALESCE(a.coa_account_id,-1)
                ORDER BY je.id DESC,jl.id LIMIT 1) counter_account_name
              FROM cash_transactions t JOIN cash_accounts a ON a.id=t.account_id
              WHERE t.transaction_type IN ('IN','OUT')
                -- Daftar Kas Masuk/Keluar harus murni transaksi yang dibuat dari
                -- menu Kas Masuk/Keluar. Modul lain (AR/AP, penjualan, pembelian,
                -- DP, jurnal manual, saldo awal, import bank) juga menulis ke
                -- cash_transactions untuk menjaga saldo, tetapi bukan bagian
                -- dari daftar maintenance Kas Masuk/Keluar.
                AND EXISTS(SELECT 1 FROM journal_entries src
                  WHERE src.source_type='CASH_TRANSACTION'
                    AND CAST(src.source_id AS TEXT)=CAST(t.id AS TEXT)
                    AND src.status='POSTED')
                AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='cash' AND v.transaction_key=CAST(t.id AS TEXT))
                AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='cash_reversal' AND v.transaction_key=CAST(t.id AS TEXT))
              ORDER BY t.transaction_date DESC,t.id DESC LIMIT ?""",(limit,)).fetchall()
        elif kind=='transfer':
            rows=c.execute("""SELECT MIN(t.id) key,t.transfer_group number,MIN(t.transaction_date) date,
              MAX(CASE WHEN t.transaction_type='TRANSFER_OUT' THEN a.name END) source_account,
              MAX(CASE WHEN t.transaction_type='TRANSFER_IN' THEN a.name END) target_account,
              MAX(t.amount) amount,MAX(t.description) description,MAX(t.reference_no) reference_no,
              MAX(CASE WHEN t.transaction_type='TRANSFER_OUT' THEN t.account_id END) source_account_id,
              MAX(CASE WHEN t.transaction_type='TRANSFER_IN' THEN t.account_id END) target_account_id
              FROM cash_transactions t JOIN cash_accounts a ON a.id=t.account_id
              WHERE t.transfer_group IS NOT NULL GROUP BY t.transfer_group
              ORDER BY MIN(t.transaction_date) DESC,MIN(t.id) DESC LIMIT ?""",(limit,)).fetchall()
        elif kind=='journal':
            rows=c.execute("""SELECT j.id key,j.journal_no number,j.journal_date date,j.description,j.reference_no,
              j.total_debit amount,j.status FROM journal_entries j
              WHERE j.source_type IN ('MANUAL','MANUAL_EXCEL') ORDER BY j.journal_date DESC,j.id DESC LIMIT ?""",(limit,)).fetchall()
        elif kind=='adjustment':
            rows=c.execute("""SELECT t.id key,COALESCE(t.reference_no,'ADJ-'||t.id) number,substr(t.created_at,1,10) date,
              p.sku,p.name product_name,w.name warehouse_name,t.quantity_change amount,t.reason description,
              t.reference_no,t.product_id,t.warehouse_id,t.department_id,t.project_id,
              (SELECT jl.account_id FROM journal_entries je JOIN journal_lines jl ON jl.journal_id=je.id
                WHERE je.source_type='STOCK_ADJUSTMENT' AND CAST(je.source_id AS TEXT)=CAST(t.id AS TEXT)
                  AND je.status='POSTED' AND jl.account_id<>p.inventory_account_id ORDER BY jl.id LIMIT 1) adjustment_account_id
              FROM inventory_transactions t JOIN products p ON p.id=t.product_id JOIN warehouses w ON w.id=t.warehouse_id
              WHERE t.reference_type='ADJUSTMENT' ORDER BY t.id DESC LIMIT ?""",(limit,)).fetchall()
        elif kind=='warehouse_transfer':
            rows=c.execute("""SELECT o.id key,COALESCE(o.reference_no,'TRF-'||o.id) number,substr(o.created_at,1,10) date,
              p.sku,p.name product_name,sw.name source_warehouse,tw.name target_warehouse,
              ABS(o.quantity_change) amount,o.reason description,o.reference_no,o.product_id,
              o.warehouse_id source_warehouse_id,i.warehouse_id target_warehouse_id,i.id target_transaction_id
              FROM inventory_transactions o
              JOIN inventory_transactions i ON i.product_id=o.product_id AND i.reference_type='TRANSFER'
                AND i.movement_type='TRANSFER_IN' AND i.created_at=o.created_at
                AND COALESCE(i.reference_no,'')=COALESCE(o.reference_no,'')
              JOIN products p ON p.id=o.product_id JOIN warehouses sw ON sw.id=o.warehouse_id
              JOIN warehouses tw ON tw.id=i.warehouse_id
              WHERE o.reference_type='TRANSFER' AND o.movement_type='TRANSFER_OUT'
              ORDER BY o.id DESC LIMIT ?""",(limit,)).fetchall()
        elif kind=='receivable':
            rows=c.execute("""SELECT * FROM (
              SELECT CAST(r.id AS TEXT) key,r.payment_no number,r.payment_date date,b.name partner_name,
                s.invoice_no invoice_no,c.name account_name,r.amount,r.notes description,
                CAST(r.sale_id AS TEXT) sale_id,r.cash_account_id,r.customer_id,0 is_opening
              FROM receivable_payments r JOIN sales s ON s.id=r.sale_id
              JOIN business_partners b ON b.id=r.customer_id JOIN cash_accounts c ON c.id=r.cash_account_id
              WHERE NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='receivable' AND v.transaction_key=CAST(r.id AS TEXT))
              UNION ALL
              SELECT 'OPENING-'||CAST(r.id AS TEXT) key,r.payment_no number,r.payment_date date,b.name partner_name,
                'SALDO AWAL-'||b.code invoice_no,c.name account_name,r.amount,r.notes description,
                'OPENING:'||CAST(r.customer_id AS TEXT) sale_id,r.cash_account_id,r.customer_id,1 is_opening
              FROM opening_receivable_payments r JOIN business_partners b ON b.id=r.customer_id JOIN cash_accounts c ON c.id=r.cash_account_id
              WHERE NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='receivable' AND v.transaction_key=('OPENING-'||CAST(r.id AS TEXT)))
            ) z ORDER BY date DESC,number DESC LIMIT ?""",(limit,)).fetchall()
        elif kind=='payable':
            rows=c.execute("""SELECT * FROM (
              SELECT CAST(r.id AS TEXT) key,r.payment_no number,r.payment_date date,b.name partner_name,
                p.purchase_no invoice_no,c.name account_name,r.amount,r.notes description,
                CAST(r.purchase_id AS TEXT) purchase_id,r.cash_account_id,r.supplier_id,0 is_opening
              FROM payable_payments r JOIN purchases p ON p.id=r.purchase_id
              JOIN business_partners b ON b.id=r.supplier_id JOIN cash_accounts c ON c.id=r.cash_account_id
              WHERE NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='payable' AND v.transaction_key=CAST(r.id AS TEXT))
              UNION ALL
              SELECT 'OPENING-'||CAST(r.id AS TEXT) key,r.payment_no number,r.payment_date date,b.name partner_name,
                'SALDO AWAL-'||b.code invoice_no,c.name account_name,r.amount,r.notes description,
                'OPENING:'||CAST(r.supplier_id AS TEXT) purchase_id,r.cash_account_id,r.supplier_id,1 is_opening
              FROM opening_payable_payments r JOIN business_partners b ON b.id=r.supplier_id JOIN cash_accounts c ON c.id=r.cash_account_id
              WHERE NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='payable' AND v.transaction_key=('OPENING-'||CAST(r.id AS TEXT)))
            ) z ORDER BY date DESC,number DESC LIMIT ?""",(limit,)).fetchall()
        else: raise ValueError('Jenis daftar transaksi tidak valid.')
        return [{**dict(r),'is_void':str(r['key'] if kind!='transfer' else r['number']) in voided} for r in rows]
    finally:c.close()

def get_transaction_maintenance(kind,key):
    items=list_transaction_maintenance(kind,1000)
    item=next((x for x in items if str(x['key'])==str(key) or (kind=='transfer' and str(x['number'])==str(key))),None)
    if not item:return None
    if kind=='journal': item=get_journal(int(key))
    return item

def _mark_void(tx,kind,key,actor,ip,reason,replacement_key=None):
    tx.execute("INSERT INTO transaction_voids(transaction_kind,transaction_key,reason,replacement_key,user_id,created_at) VALUES(?,?,?,?,?,?)",
      (kind,str(key),reason or 'Dihapus pengguna',str(replacement_key) if replacement_key else None,actor['id'],utc_now()))
    audit(actor['id'],'TRANSACTION_VOIDED',kind,str(key),{'reason':reason,'replacement_key':replacement_key},ip,tx)

def _void_journals(tx,source_type,source_id):
    journal_ids=[row['id'] for row in tx.execute("SELECT id FROM journal_entries WHERE source_type=? AND CAST(source_id AS TEXT)=? AND status='POSTED'",(source_type,str(source_id))).fetchall()]
    for journal_id in journal_ids:
        licensing.release_transaction_usage(tx,source_table='journal_entries',source_id=journal_id)
    tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE source_type=? AND CAST(source_id AS TEXT)=? AND status='POSTED'",
      (utc_now(),source_type,str(source_id)))

def _cash_decimal(value, label="Nilai", allow_negative=True):
    """Parser khusus nilai database lama/import; tidak bergantung format tampilan."""
    from decimal import Decimal, InvalidOperation
    if isinstance(value, bytes):
        value=value.decode("utf-8", errors="ignore")
    try:
        result=Decimal(cash_service._normalize_decimal_text(value))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise ValueError(f"{label} tidak dapat dibaca: {value!r}.") from exc
    if not result.is_finite() or (not allow_negative and result < 0):
        raise ValueError(f"{label} tidak valid: {value!r}.")
    return result.quantize(MONEY, rounding=ROUND_HALF_UP)


def _cash_reversal_amount(row):
    """Ambil nominal kas dari amount atau selisih saldo transaksi lama."""
    errors=[]
    for value,label in ((row['amount'],'amount'),):
        try:
            amount=abs(_cash_decimal(value, f'Nominal {label}'))
            if amount>0:return amount
        except Exception as exc:
            errors.append(str(exc))
    try:
        before=_cash_decimal(row['balance_before'],'Saldo sebelum')
        after=_cash_decimal(row['balance_after'],'Saldo sesudah')
        amount=abs(after-before)
        if amount>0:return amount
    except Exception as exc:
        errors.append(str(exc))
    raise ValueError('Nominal transaksi kas tidak dapat dibaca. '+(' | '.join(errors) if errors else ''))


def _reverse_cash_row(tx,row,actor,now):
    """Pembalikan kas langsung dari nilai database, termasuk transaksi Smart Import."""
    account=tx.execute("SELECT * FROM cash_accounts WHERE id=?",(row['account_id'],)).fetchone()
    if not account:raise ValueError('Akun kas/bank transaksi tidak ditemukan.')
    amount=_cash_reversal_amount(row)
    current=_cash_decimal(account['current_balance'],'Saldo kas/bank aktif')
    opposite='OUT' if row['transaction_type']=='IN' else 'IN'
    new_balance=current-amount if opposite=='OUT' else current+amount
    reversal_no=cash_service.next_number(tx,now[:10],'VR')
    tx.execute("UPDATE cash_accounts SET current_balance=?,updated_at=? WHERE id=?",
      (str(new_balance),now,row['account_id']))
    cur=tx.execute("""INSERT INTO cash_transactions(
      transaction_no,transaction_date,account_id,transaction_type,amount,
      balance_before,balance_after,description,reference_no,transfer_group,
      department_id,project_id,user_id,created_at)
      VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
      (reversal_no,now[:10],row['account_id'],opposite,str(amount),str(current),str(new_balance),
       'Pembalikan '+str(row['transaction_no']),str(row['transaction_no']),None,
       row['department_id'],row['project_id'],actor['id'],now))
    return {'id':cur.lastrowid,'transaction_no':reversal_no,'amount':amount}


def delete_transaction_maintenance(actor,kind,key,ip,reason=''):
    now=utc_now()
    with write_transaction() as tx:
        exists=tx.execute("SELECT 1 FROM transaction_voids WHERE transaction_kind=? AND transaction_key=?",(kind,str(key))).fetchone()
        if exists:raise ValueError('Transaksi sudah dihapus/dibatalkan.')
        if kind=='cash':
            try:cash_id=int(key)
            except (TypeError,ValueError):raise ValueError('ID transaksi kas tidak valid.')
            r=tx.execute("SELECT * FROM cash_transactions WHERE id=? AND transaction_type IN ('IN','OUT')",(cash_id,)).fetchone()
            if not r:raise ValueError('Transaksi kas tidak ditemukan.')
            reversal=_reverse_cash_row(tx,r,actor,now)

            # Smart Import menyimpan referensi di bank_import_rows dan jurnal BANK_IMPORT.
            import_rows=tx.execute("SELECT id FROM bank_import_rows WHERE posted_transaction_id=?",(cash_id,)).fetchall()
            if import_rows:
                for imported in import_rows:
                    tx.execute("UPDATE bank_import_rows SET status='VOID' WHERE id=?",(imported['id'],))
                    _void_journals(tx,'BANK_IMPORT',imported['id'])
            else:
                _void_journals(tx,'CASH_TRANSACTION',cash_id)

            licensing.release_transaction_usage(tx,source_table='cash_transactions',source_id=cash_id);_mark_void(tx,kind,key,actor,ip,reason,reversal['id'])
            tx.execute("INSERT INTO transaction_voids(transaction_kind,transaction_key,reason,replacement_key,user_id,created_at) VALUES(?,?,?,?,?,?)",
              ('cash_reversal',str(reversal['id']),'Pembalikan internal transaksi '+str(r['transaction_no']),str(key),actor['id'],now))
        elif kind=='transfer':
            rows=tx.execute("SELECT * FROM cash_transactions WHERE transfer_group=? ORDER BY id",(str(key),)).fetchall()
            if len(rows)!=2:raise ValueError('Transfer tidak ditemukan.')
            out=next(r for r in rows if r['transaction_type']=='TRANSFER_OUT'); inn=next(r for r in rows if r['transaction_type']=='TRANSFER_IN')
            cash_service.transfer(tx,source_account_id=inn['account_id'],target_account_id=out['account_id'],transaction_date=now[:10],
              amount=out['amount'],description='Pembalikan '+str(key),reference_no=str(key),user_id=actor['id'])
            _void_journals(tx,'CASH_TRANSFER',str(key));licensing.release_transaction_usage(tx,event_key='CASH-TRANSFER-'+str(key));_mark_void(tx,kind,key,actor,ip,reason)
        elif kind=='journal':
            r=tx.execute("SELECT * FROM journal_entries WHERE id=? AND source_type IN ('MANUAL','MANUAL_EXCEL')",(int(key),)).fetchone()
            if not r:raise ValueError('Jurnal manual tidak ditemukan.')
            licensing.release_transaction_usage(tx,source_table='journal_entries',source_id=int(key));tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE id=?",(now,int(key)));_sync_manual_journal_cash(tx,int(key));_mark_void(tx,kind,key,actor,ip,reason)
        elif kind=='adjustment':
            r=tx.execute("SELECT * FROM inventory_transactions WHERE id=? AND reference_type='ADJUSTMENT'",(int(key),)).fetchone()
            if not r:raise ValueError('Penyesuaian stok tidak ditemukan.')
            rev=inventory_service.adjust(tx,product_id=r['product_id'],warehouse_id=r['warehouse_id'],quantity_change=-Decimal(str(r['quantity_change'])),
              reason='Pembalikan adjustment '+str(key),reference_no='VOID-'+str(key),user_id=actor['id'],created_at=now,
              department_id=r['department_id'],project_id=r['project_id'])
            _void_journals(tx,'STOCK_ADJUSTMENT',r['id']);_mark_void(tx,kind,key,actor,ip,reason,rev['transaction_id'])
        elif kind=='warehouse_transfer':
            out=tx.execute("SELECT * FROM inventory_transactions WHERE id=? AND reference_type='TRANSFER' AND movement_type='TRANSFER_OUT'",(int(key),)).fetchone()
            if not out:raise ValueError('Transfer gudang tidak ditemukan.')
            inn=tx.execute("SELECT * FROM inventory_transactions WHERE product_id=? AND reference_type='TRANSFER' AND movement_type='TRANSFER_IN' AND created_at=? AND COALESCE(reference_no,'')=COALESCE(?, '') ORDER BY id LIMIT 1",(out['product_id'],out['created_at'],out['reference_no'])).fetchone()
            if not inn:raise ValueError('Pasangan transfer gudang tidak ditemukan.')
            rev=inventory_service.transfer(tx,product_id=out['product_id'],source_warehouse_id=inn['warehouse_id'],target_warehouse_id=out['warehouse_id'],quantity=abs(Decimal(str(out['quantity_change']))),reference_no='VOID-'+str(out['reference_no'] or key),reason='Pembalikan transfer gudang '+str(key),user_id=actor['id'],created_at=now)
            _mark_void(tx,kind,key,actor,ip,reason,rev['source']['transaction_id'])
        elif kind=='receivable':
            if str(key).startswith('OPENING-'):
                pid=int(str(key).split('-',1)[1]);r=tx.execute("SELECT * FROM opening_receivable_payments WHERE id=?",(pid,)).fetchone()
                if not r:raise ValueError('Penerimaan saldo awal piutang tidak ditemukan.')
                amt=_money(r['amount']);cash_service.post(tx,account_id=r['cash_account_id'],transaction_date=now[:10],transaction_type='OUT',amount=amt,description='Pembalikan '+r['payment_no'],reference_no=r['payment_no'],user_id=actor['id'],allow_negative=True)
                jids=[x['id'] for x in tx.execute("SELECT id FROM journal_entries WHERE source_type='RECEIVABLE_PAYMENT' AND CAST(source_id AS TEXT)=? AND reference_no=? AND status='POSTED'",(str(pid),r['payment_no'])).fetchall()]
                for jid in jids: licensing.release_transaction_usage(tx,source_table='journal_entries',source_id=jid)
                if jids: tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE id IN (%s)" % ','.join('?'*len(jids)),[now,*jids])
                _mark_void(tx,kind,key,actor,ip,reason)
            else:
                r=tx.execute("SELECT * FROM receivable_payments WHERE id=?",(int(key),)).fetchone()
                if not r:raise ValueError('Penerimaan piutang tidak ditemukan.')
                sale=tx.execute("SELECT * FROM sales WHERE id=?",(r['sale_id'],)).fetchone(); amt=_money(r['amount'])
                tx.execute("UPDATE sales SET paid_amount=?,balance_due=?,updated_at=? WHERE id=?",(str(_money(sale['paid_amount'])-amt),str(_money(sale['balance_due'])+amt),now,r['sale_id']))
                cash_service.post(tx,account_id=r['cash_account_id'],transaction_date=now[:10],transaction_type='OUT',amount=amt,description='Pembalikan '+r['payment_no'],reference_no=r['payment_no'],user_id=actor['id'],allow_negative=True)
                _void_journals(tx,'RECEIVABLE_PAYMENT',r['id']);_mark_void(tx,kind,key,actor,ip,reason)
        elif kind=='payable':
            if str(key).startswith('OPENING-'):
                pid=int(str(key).split('-',1)[1]);r=tx.execute("SELECT * FROM opening_payable_payments WHERE id=?",(pid,)).fetchone()
                if not r:raise ValueError('Pembayaran saldo awal hutang tidak ditemukan.')
                amt=_money(r['amount']);cash_service.post(tx,account_id=r['cash_account_id'],transaction_date=now[:10],transaction_type='IN',amount=amt,description='Pembalikan '+r['payment_no'],reference_no=r['payment_no'],user_id=actor['id'])
                jids=[x['id'] for x in tx.execute("SELECT id FROM journal_entries WHERE source_type='PAYABLE_PAYMENT' AND CAST(source_id AS TEXT)=? AND reference_no=? AND status='POSTED'",(str(pid),r['payment_no'])).fetchall()]
                for jid in jids: licensing.release_transaction_usage(tx,source_table='journal_entries',source_id=jid)
                if jids: tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE id IN (%s)" % ','.join('?'*len(jids)),[now,*jids])
                _mark_void(tx,kind,key,actor,ip,reason)
            else:
                r=tx.execute("SELECT * FROM payable_payments WHERE id=?",(int(key),)).fetchone()
                if not r:raise ValueError('Pembayaran hutang tidak ditemukan.')
                pur=tx.execute("SELECT * FROM purchases WHERE id=?",(r['purchase_id'],)).fetchone(); amt=_money(r['amount'])
                tx.execute("UPDATE purchases SET paid_amount=?,balance_due=?,updated_at=? WHERE id=?",(str(_money(pur['paid_amount'])-amt),str(_money(pur['balance_due'])+amt),now,r['purchase_id']))
                cash_service.post(tx,account_id=r['cash_account_id'],transaction_date=now[:10],transaction_type='IN',amount=amt,description='Pembalikan '+r['payment_no'],reference_no=r['payment_no'],user_id=actor['id'])
                _void_journals(tx,'PAYABLE_PAYMENT',r['id']);_mark_void(tx,kind,key,actor,ip,reason)
        else:raise ValueError('Jenis transaksi tidak valid.')
    return {'status':'voided'}

def _rebuild_cash_running_balance(tx,account_id):
    """Recalculate operational cash balance and per-row running balances from cash_transactions."""
    rows=tx.execute("""SELECT ct.id,ct.transaction_type,ct.amount FROM cash_transactions ct
      WHERE ct.account_id=?
        AND NOT EXISTS(SELECT 1 FROM transaction_voids v
          WHERE v.transaction_key=CAST(ct.id AS TEXT)
            AND v.transaction_kind IN ('cash','cash_reversal'))
      ORDER BY ct.transaction_date,ct.id""",(int(account_id),)).fetchall()
    bal=Decimal("0.00")
    for row in rows:
        amount=_money(row["amount"])
        before=bal
        if row["transaction_type"] in ("IN","TRANSFER_IN"):bal+=amount
        elif row["transaction_type"] in ("OUT","TRANSFER_OUT"):bal-=amount
        tx.execute("UPDATE cash_transactions SET balance_before=?,balance_after=? WHERE id=?",
          (str(before),str(bal),row["id"]))
    tx.execute("UPDATE cash_accounts SET current_balance=?,updated_at=? WHERE id=?",
      (str(bal),utc_now(),int(account_id)))
    return bal


def _update_cash_transaction_in_place(actor,cash_id,data,ip):
    """Edit Kas Masuk/Keluar secara atomik pada ID yang sama.

    Tidak memakai VOID + replacement agar laporan Kas/Bank, GL, Neraca Saldo,
    Laba Rugi dan dashboard selalu membaca satu transaksi aktif yang sama.
    """
    transaction_date=str(data.get('transaction_date') or '')[:10]
    if len(transaction_date)!=10:raise ValueError('Tanggal wajib diisi.')
    try:account_id=int(data.get('account_id'))
    except Exception:raise ValueError('Akun kas/bank wajib dipilih.')
    transaction_type=str(data.get('transaction_type') or 'IN').upper()
    if transaction_type not in ('IN','OUT'):raise ValueError('Jenis transaksi kas tidak valid.')
    try:counter_account_id=int(data.get('counter_account_id'))
    except Exception:raise ValueError('Akun lawan transaksi wajib dipilih.')
    description=str(data.get('description') or '').strip()
    if not description:raise ValueError('Keterangan wajib diisi.')
    amount=_cash_decimal(data.get('amount'),'Jumlah',allow_negative=False)
    if amount<=0:raise ValueError('Jumlah harus lebih besar dari 0.')
    reference_no=str(data.get('reference_no') or '').strip() or None
    with write_transaction() as tx:
        row=tx.execute("SELECT * FROM cash_transactions WHERE id=? AND transaction_type IN ('IN','OUT')",(int(cash_id),)).fetchone()
        if not row:raise ValueError('Transaksi kas tidak ditemukan.')
        if tx.execute("SELECT 1 FROM transaction_voids WHERE transaction_kind='cash' AND transaction_key=?",(str(cash_id),)).fetchone():
            raise ValueError('Transaksi sudah dihapus/dibatalkan.')
        counter=tx.execute("SELECT id,account_subtype FROM chart_of_accounts WHERE id=? AND is_active=1",(counter_account_id,)).fetchone()
        if not counter or counter['account_subtype'] in ('CASH_BANK','RECEIVABLE','PAYABLE'):
            raise ValueError('Akun lawan kas/bank tidak valid.')
        account=tx.execute("SELECT id,coa_account_id FROM cash_accounts WHERE id=? AND is_active=1",(account_id,)).fetchone()
        if not account:raise ValueError('Akun kas/bank tidak valid.')
        department_id,project_id=_transaction_dimensions(tx,data)
        old_account_id=int(row['account_id'])
        tx.execute("""UPDATE cash_transactions SET transaction_date=?,account_id=?,transaction_type=?,amount=?,
          description=?,reference_no=?,department_id=?,project_id=? WHERE id=?""",
          (transaction_date,account_id,transaction_type,str(amount),description,reference_no,department_id,project_id,int(cash_id)))
        cash_coa=account['coa_account_id'] or accounting_service.account_id(tx,'1000')
        lines=([{'account_id':cash_coa,'debit':amount},{'account_id':counter_account_id,'credit':amount}]
               if transaction_type=='IN' else
               [{'account_id':counter_account_id,'debit':amount},{'account_id':cash_coa,'credit':amount}])
        # Edit jurnal sumber yang sama secara in-place. Jangan membuat jurnal VOID + jurnal baru,
        # karena beberapa laporan/rekonsiliasi historis masih menautkan source journal pertama.
        j=tx.execute("SELECT * FROM journal_entries WHERE source_type='CASH_TRANSACTION' AND CAST(source_id AS TEXT)=? ORDER BY CASE status WHEN 'POSTED' THEN 0 ELSE 1 END,id LIMIT 1",(str(cash_id),)).fetchone()
        if j:
            jid=int(j['id']);tx.execute("DELETE FROM journal_lines WHERE journal_id=?",(jid,))
            for ln in lines:
                tx.execute("INSERT INTO journal_lines(journal_id,account_id,debit,credit) VALUES(?,?,?,?)",(jid,int(ln['account_id']),str(ln.get('debit',0)),str(ln.get('credit',0))))
            tx.execute("""UPDATE journal_entries SET journal_date=?,description=?,reference_no=?,department_id=?,project_id=?,status='POSTED',total_debit=?,total_credit=?,updated_at=? WHERE id=?""",
              (transaction_date,description,reference_no or row['transaction_no'],department_id,project_id,str(amount),str(amount),utc_now(),jid))
        else:
            accounting_service.post_journal(tx,journal_date=transaction_date,description=description,
              source_type='CASH_TRANSACTION',source_id=int(cash_id),reference_no=reference_no or row['transaction_no'],
              lines=lines,user_id=actor['id'],department_id=department_id,project_id=project_id)
        _rebuild_cash_running_balance(tx,old_account_id)
        if account_id!=old_account_id:_rebuild_cash_running_balance(tx,account_id)
        audit(actor['id'],'CASH_TRANSACTION_UPDATED','cash_transaction',int(cash_id),
          {'transaction_no':row['transaction_no'],'old_amount':float(_money(row['amount'])),'amount':float(amount),
           'old_account_id':old_account_id,'account_id':account_id,'transaction_type':transaction_type},ip,tx)
        return {'id':int(cash_id),'transaction_no':row['transaction_no'],'amount':float(amount)}


def _update_settlement_in_place(actor,kind,key,data,ip):
    receivable=(kind=='receivable'); target_field='sale_id' if receivable else 'purchase_id'
    payment_date=str(data.get('payment_date') or date.today().isoformat())[:10]
    try: cash_account_id=int(data.get('cash_account_id'))
    except Exception: raise ValueError('Akun kas/bank wajib dipilih.')
    amount=_money(data.get('amount')); notes=str(data.get('notes') or '').strip() or None
    if amount<=0: raise ValueError('Jumlah pembayaran harus lebih dari nol.')
    opening=str(key).startswith('OPENING-'); pid=int(str(key).split('-',1)[1]) if opening else int(key); now=utc_now()
    with write_transaction() as tx:
        if opening:
            table='opening_receivable_payments' if receivable else 'opening_payable_payments'; partner_col='customer_id' if receivable else 'supplier_id'
            r=tx.execute(f'SELECT * FROM {table} WHERE id=?',(pid,)).fetchone()
            if not r: raise ValueError('Pembayaran saldo awal tidak ditemukan.')
            original='OPENING:'+str(r[partner_col])
            if str(data.get(target_field) or original)!=original: raise ValueError('Partner saldo awal tidak dapat diganti saat edit.')
            partner=tx.execute('SELECT * FROM business_partners WHERE id=?',(r[partner_col],)).fetchone(); total=_money(partner['opening_balance'])
            sql=f"SELECT COALESCE(SUM(amount),0) v FROM {table} x WHERE x.{partner_col}=? AND x.id<>? AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind=? AND v.transaction_key=('OPENING-'||CAST(x.id AS TEXT)))"
            paid_other=_money(tx.execute(sql,(r[partner_col],pid,kind)).fetchone()['v'])
            if amount>max(Decimal('0'),total-paid_other): raise ValueError('Jumlah edit melebihi sisa saldo awal.')
            tx.execute(f'UPDATE {table} SET payment_date=?,cash_account_id=?,amount=?,notes=? WHERE id=?',(payment_date,cash_account_id,str(amount),notes,pid))
            partner_id=r[partner_col]; payment_no=r['payment_no']
        else:
            table='receivable_payments' if receivable else 'payable_payments'; fk='sale_id' if receivable else 'purchase_id'; partner_col='customer_id' if receivable else 'supplier_id'
            r=tx.execute(f'SELECT * FROM {table} WHERE id=?',(pid,)).fetchone()
            if not r: raise ValueError('Pembayaran tidak ditemukan.')
            if str(data.get(target_field) or r[fk])!=str(r[fk]): raise ValueError('Invoice tidak dapat diganti saat edit.')
            inv_table='sales' if receivable else 'purchases'; inv=tx.execute(f"SELECT * FROM {inv_table} WHERE id=? AND status='POSTED'",(r[fk],)).fetchone()
            if not inv: raise ValueError('Invoice transaksi tidak ditemukan.')
            old=_money(r['amount']); maximum=_money(inv['balance_due'])+old
            if amount>maximum: raise ValueError('Jumlah edit melebihi sisa tagihan.')
            tx.execute(f'UPDATE {table} SET payment_date=?,cash_account_id=?,amount=?,notes=? WHERE id=?',(payment_date,cash_account_id,str(amount),notes,pid))
            tx.execute(f'UPDATE {inv_table} SET paid_amount=?,balance_due=?,updated_at=? WHERE id=?',(str(_money(inv['paid_amount'])-old+amount),str(max(Decimal('0'),_money(inv['balance_due'])+old-amount)),now,r[fk]))
            partner_id=r[partner_col]; payment_no=r['payment_no']
        direction='IN' if receivable else 'OUT'; cashrow=tx.execute('SELECT * FROM cash_transactions WHERE reference_no=? AND transaction_type=? ORDER BY id LIMIT 1',(payment_no,direction)).fetchone(); old_cash=int(cashrow['account_id']) if cashrow else cash_account_id
        desc=('Penerimaan piutang ' if receivable else 'Pembayaran hutang ')+payment_no
        if cashrow: tx.execute('UPDATE cash_transactions SET transaction_date=?,account_id=?,amount=?,description=? WHERE id=?',(payment_date,cash_account_id,str(amount),desc,cashrow['id']))
        else: cash_service.post(tx,account_id=cash_account_id,transaction_date=payment_date,transaction_type=direction,amount=amount,description=desc,reference_no=payment_no,user_id=actor['id'],allow_negative=True)
        _rebuild_cash_running_balance(tx,old_cash)
        if cash_account_id!=old_cash:_rebuild_cash_running_balance(tx,cash_account_id)
        source='RECEIVABLE_PAYMENT' if receivable else 'PAYABLE_PAYMENT'
        # Hapus tombstone stale hanya jika record pembayaran yang sedang diedit nyata/aktif.
        tomb_key=('OPENING-'+str(pid)) if opening else str(pid)
        tx.execute("DELETE FROM transaction_voids WHERE transaction_kind=? AND transaction_key=?",(kind,tomb_key))
        # Bentuk ulang jurnal lalu fold ke journal source pertama agar edit tidak menciptakan rantai VOID.
        oldj=tx.execute("SELECT * FROM journal_entries WHERE source_type=? AND CAST(source_id AS TEXT)=? AND (?='' OR reference_no=?) ORDER BY CASE status WHEN 'POSTED' THEN 0 ELSE 1 END,id LIMIT 1",(source,str(pid),payment_no,payment_no)).fetchone()
        if oldj:
            ojid=int(oldj['id']);tx.execute("DELETE FROM journal_lines WHERE journal_id=?",(ojid,))
            cash=tx.execute("SELECT coa_account_id,account_type FROM cash_accounts WHERE id=?",(cash_account_id,)).fetchone();cash_coa=int(cash['coa_account_id'] or _cash_default_coa(tx,cash['account_type']))
            arap=accounting_service.account_id(tx,'1100' if receivable else '2000')
            lines=([{'account_id':cash_coa,'debit':amount},{'account_id':arap,'credit':amount,'partner_id':partner_id}] if receivable else [{'account_id':arap,'debit':amount,'partner_id':partner_id},{'account_id':cash_coa,'credit':amount}])
            for ln in lines:tx.execute("INSERT INTO journal_lines(journal_id,account_id,partner_id,debit,credit) VALUES(?,?,?,?,?)",(ojid,int(ln['account_id']),ln.get('partner_id'),str(ln.get('debit',0)),str(ln.get('credit',0))))
            tx.execute("UPDATE journal_entries SET journal_date=?,description=?,reference_no=?,status='POSTED',total_debit=?,total_credit=?,updated_at=? WHERE id=?",(payment_date,desc,payment_no,str(amount),str(amount),now,ojid))
        elif receivable: accounting_service.post_receivable_payment(tx,payment_id=pid,payment_no=payment_no,payment_date=payment_date,amount=amount,cash_account_id=cash_account_id,customer_id=partner_id,user_id=actor['id'],receivable_account_id=(tx.execute('SELECT receivable_account_id FROM business_partners WHERE id=?',(partner_id,)).fetchone() or {'receivable_account_id':None})['receivable_account_id'])
        else: accounting_service.post_payable_payment(tx,payment_id=pid,payment_no=payment_no,payment_date=payment_date,amount=amount,cash_account_id=cash_account_id,supplier_id=partner_id,user_id=actor['id'],payable_account_id=(tx.execute('SELECT payable_account_id FROM business_partners WHERE id=?',(partner_id,)).fetchone() or {'payable_account_id':None})['payable_account_id'])
        audit(actor['id'],'RECEIVABLE_PAYMENT_UPDATED' if receivable else 'PAYABLE_PAYMENT_UPDATED','payment',pid,{'payment_no':payment_no,'amount':float(amount)},ip,tx)
        return {'id':pid,'payment_no':payment_no,'amount':float(amount)}

def update_transaction_maintenance(actor,kind,key,data,ip):
    if kind=='journal':
        return update_manual_journal(actor,int(key),data,ip)
    if kind=='cash':
        return _update_cash_transaction_in_place(actor,int(key),data,ip)
    if kind in ('receivable','payable'):
        return _update_settlement_in_place(actor,kind,key,data,ip)

    delete_transaction_maintenance(actor,kind,key,ip,'Diedit dan diganti transaksi baru')
    if kind=='transfer': return create_cash_transfer(actor,data,ip)
    if kind=='adjustment': return adjust_stock(actor,data,ip)
    if kind=='receivable': return receive_receivable(actor,data,ip)
    if kind=='payable': return pay_payable(actor,data,ip)
    if kind=='warehouse_transfer': return transfer_stock(actor,data,ip)
    raise ValueError('Jenis transaksi tidak valid.')

def _replay_inventory_state_tx(tx):
    """Rebuild quantity/value state without synthetic edit/reversal rows.

    SALE movements are re-costed from the moving-average immediately before the
    original SALE row. This is essential after an opening-cost edit: the same
    SALE row stays in history, but its cost/HPP follows the corrected valuation.
    """
    now=utc_now();pairs=0;movements=0
    tx.execute("DELETE FROM inventory_balances")
    for pair in tx.execute("SELECT DISTINCT product_id,warehouse_id FROM inventory_transactions ORDER BY product_id,warehouse_id").fetchall():
        q=Decimal('0.0000');value=Decimal('0.00');avg=Decimal('0.000000')
        moves=tx.execute("SELECT * FROM inventory_transactions WHERE product_id=? AND warehouse_id=? ORDER BY id",
          (pair['product_id'],pair['warehouse_id'])).fetchall()
        for m in moves:
            change=Decimal(str(m['quantity_change'] or 0)).quantize(Decimal('0.0001'))
            stored_cost=Decimal(str(m['unit_cost'] or 0)).quantize(Decimal('0.000001'))
            # Penjualan selalu keluar dengan moving-average yang berlaku tepat
            # sebelum baris SALE tersebut. Edit saldo awal otomatis merambat ke HPP.
            cost=(avg if change<0 and str(m['reference_type'] or '')=='SALE' else stored_cost).quantize(Decimal('0.000001'))
            before=q;vb=value;after=(before+change).quantize(Decimal('0.0001'))
            vc=(change*cost).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)
            va=(vb+vc).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)
            ab=avg
            if after>0:avg=(va/after).quantize(Decimal('0.000001'),rounding=ROUND_HALF_UP)
            elif after==0:avg=Decimal('0.000000')
            tx.execute("""UPDATE inventory_transactions SET quantity_before=?,quantity_after=?,unit_cost=?,
              average_cost_before=?,average_cost_after=?,value_change=?,value_before=?,value_after=? WHERE id=?""",
              (str(before),str(after),str(cost),str(ab),str(avg),str(vc),str(vb),str(va),m['id']))
            q=after;value=va;movements+=1
        tx.execute("""INSERT INTO inventory_balances(warehouse_id,product_id,quantity,average_cost,book_value,updated_at)
          VALUES(?,?,?,?,?,?)""",(pair['warehouse_id'],pair['product_id'],str(q),str(avg),str(value),now));pairs+=1
    tx.execute("UPDATE products SET stock_qty=COALESCE((SELECT SUM(quantity) FROM inventory_balances b WHERE b.product_id=products.id),0),updated_at=?",(now,))
    return {'pairs':pairs,'movements':movements}


def _replay_cash_state_tx(tx):
    """Recalculate operational Cash/Bank balances after an in-place Sales edit."""
    now=utc_now()
    for a in tx.execute("SELECT id,opening_balance FROM cash_accounts ORDER BY id").fetchall():
        bal=Decimal(str(a['opening_balance'] or 0)).quantize(Decimal('0.01'))
        rows=tx.execute("SELECT id,transaction_type,amount FROM cash_transactions WHERE account_id=? ORDER BY transaction_date,id",(a['id'],)).fetchall()
        for r in rows:
            amt=Decimal(str(r['amount'] or 0)).quantize(Decimal('0.01'))
            before=bal
            bal=(bal + amt if r['transaction_type'] in ('IN','TRANSFER_IN') else bal-amt).quantize(Decimal('0.01'))
            tx.execute("UPDATE cash_transactions SET balance_before=?,balance_after=? WHERE id=?",(str(before),str(bal),r['id']))
        tx.execute("UPDATE cash_accounts SET current_balance=?,updated_at=? WHERE id=?",(str(bal),now,a['id']))


def _replace_opening_product_journal_tx(tx, product_id, sku, opening_date, inventory_account_id, amount, user_id):
    """Keep one opening journal and update that same journal in place."""
    amount=Decimal(str(amount or 0)).quantize(Decimal('0.01'))
    rows=tx.execute("""SELECT id FROM journal_entries WHERE source_id=? AND source_type IN
      ('OPENING_PRODUCT','PRODUCT_OPENING','OPENING_PRODUCT_EDIT','PRODUCT_OPENING_EDIT') ORDER BY id""",(str(product_id),)).fetchall()
    keep=int(rows[0]['id']) if rows else None
    for r in rows[1:]:
        tx.execute("DELETE FROM journal_entries WHERE id=?",(r['id'],))
    now=utc_now()
    if amount<=0:
        if keep:
            tx.execute("DELETE FROM journal_lines WHERE journal_id=?",(keep,))
            tx.execute("UPDATE journal_entries SET journal_date=?,description=?,source_type='OPENING_PRODUCT',source_id=?,reference_no=?,status='VOID',total_debit=0,total_credit=0,updated_at=? WHERE id=?",
              (str(opening_date)[:10],f"Saldo awal {sku}",str(product_id),sku,now,keep))
        return keep
    eq_aid=accounting_service.account_id(tx,'3200')
    if keep:
        tx.execute("DELETE FROM journal_lines WHERE journal_id=?",(keep,))
        tx.execute("UPDATE journal_entries SET journal_date=?,description=?,source_type='OPENING_PRODUCT',source_id=?,reference_no=?,status='POSTED',total_debit=?,total_credit=?,updated_at=? WHERE id=?",
          (str(opening_date)[:10],f"Saldo awal {sku}",str(product_id),sku,str(amount),str(amount),now,keep))
        tx.execute("INSERT INTO journal_lines(journal_id,account_id,debit,credit) VALUES(?,?,?,0)",(keep,int(inventory_account_id),str(amount)))
        tx.execute("INSERT INTO journal_lines(journal_id,account_id,debit,credit) VALUES(?,?,0,?)",(keep,eq_aid,str(amount)))
        return keep
    j=accounting_service.post_opening_balance(tx,entity_type='PRODUCT',entity_id=product_id,reference_no=sku,amount=amount,user_id=user_id,opening_date=opening_date,inventory_account_id=inventory_account_id)
    return j['id'] if j else None


def _sync_sale_cogs_journals_tx(tx, invoice_no=None):
    """Synchronize SALE item cost snapshots and HPP/Persediaan journal lines to valuation."""
    params=[];where="s.status='POSTED'"
    if invoice_no is not None:where+=' AND s.invoice_no=?';params.append(str(invoice_no))
    sales=tx.execute(f"SELECT s.id,s.invoice_no FROM sales s WHERE {where}",params).fetchall()
    changed=0
    for sale in sales:
        j=tx.execute("SELECT id FROM journal_entries WHERE source_type='SALE' AND source_id=? AND status='POSTED' ORDER BY id LIMIT 1",(str(sale['id']),)).fetchone()
        if not j:continue
        moves=tx.execute("SELECT product_id,quantity_change,unit_cost FROM inventory_transactions WHERE reference_type='SALE' AND reference_no=? ORDER BY id",(sale['invoice_no'],)).fetchall()
        by_product={}
        for m in moves:
            qty=abs(Decimal(str(m['quantity_change'] or 0)));cost=Decimal(str(m['unit_cost'] or 0));
            q,v=by_product.get(int(m['product_id']),(Decimal('0'),Decimal('0')));by_product[int(m['product_id'])]=(q+qty,v+(qty*cost))
        groups={}; cogs_ids=set(); inv_ids=set()
        for item in tx.execute("SELECT si.product_id,si.qty,p.cogs_account_id,p.inventory_account_id FROM sales_items si JOIN products p ON p.id=si.product_id WHERE si.sale_id=? AND si.product_type='STOCK'",(sale['id'],)).fetchall():
            pid=int(item['product_id']);q,v=by_product.get(pid,(Decimal('0'),Decimal('0')))
            if q>0:tx.execute("UPDATE sales_items SET purchase_price_snapshot=? WHERE sale_id=? AND product_id=?",(str((v/q).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)),sale['id'],pid))
            ca=int(item['cogs_account_id'] or accounting_service.account_id(tx,'5000'));ia=int(item['inventory_account_id'] or accounting_service.account_id(tx,'1200'))
            cogs_ids.add(ca);inv_ids.add(ia);groups[(ca,ia)]=groups.get((ca,ia),Decimal('0'))+v
        jid=int(j['id'])
        if cogs_ids:
            marks=','.join('?'*len(cogs_ids));tx.execute(f"DELETE FROM journal_lines WHERE journal_id=? AND account_id IN ({marks}) AND CAST(debit AS NUMERIC)>0",[jid,*cogs_ids])
        if inv_ids:
            marks=','.join('?'*len(inv_ids));tx.execute(f"DELETE FROM journal_lines WHERE journal_id=? AND account_id IN ({marks}) AND CAST(credit AS NUMERIC)>0",[jid,*inv_ids])
        for (ca,ia),amt in groups.items():
            amt=amt.quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)
            if amt<=0:continue
            tx.execute("INSERT INTO journal_lines(journal_id,account_id,debit,credit) VALUES(?,?,?,0)",(jid,ca,str(amt)))
            tx.execute("INSERT INTO journal_lines(journal_id,account_id,debit,credit) VALUES(?,?,0,?)",(jid,ia,str(amt)))
        sums=tx.execute("SELECT COALESCE(SUM(CAST(debit AS NUMERIC)),0)d,COALESCE(SUM(CAST(credit AS NUMERIC)),0)c FROM journal_lines WHERE journal_id=?",(jid,)).fetchone()
        tx.execute("UPDATE journal_entries SET total_debit=?,total_credit=?,updated_at=? WHERE id=?",(str(Decimal(str(sums['d'])).quantize(Decimal('0.01'))),str(Decimal(str(sums['c'])).quantize(Decimal('0.01'))),utc_now(),jid));changed+=1
    return changed


def delete_sale(actor,sale_id,ip,_mode="DELETE"):
    """Hapus jejak inventory Sales, bukan membuat movement reversal.

    Jika invoice Sales dihapus, movement OUT SALE milik invoice tersebut ikut
    dihapus dari inventory history. Jejak legacy SALE_VOID/SALE_DELETE/
    SALE_EDIT_REVERSAL untuk invoice yang sama juga dibersihkan. Sesudah itu
    valuation direplay dari movement yang memang masih sah. Karena itu laporan
    Rincian Valuasi tidak menampilkan baris pembatalan sintetis.
    """
    from .services import cash_service
    mode=str(_mode or "DELETE").upper()
    with write_transaction() as tx:
        s=tx.execute("SELECT * FROM sales WHERE id=?",(sale_id,)).fetchone()
        if not s:raise ValueError('Penjualan tidak ditemukan.')
        if s['status']!='POSTED':raise ValueError('Penjualan sudah dibatalkan.')
        if tx.execute("SELECT 1 FROM sales_returns WHERE sale_id=? AND status='POSTED'",(sale_id,)).fetchone():raise ValueError('Batalkan retur penjualan terlebih dahulu.')
        if mode!='EDIT' and tx.execute("""SELECT 1 FROM receivable_payments r WHERE r.sale_id=?
            AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='receivable' AND v.transaction_key=CAST(r.id AS TEXT))""",(sale_id,)).fetchone():raise ValueError('Batalkan penerimaan piutang terlebih dahulu.')

        invoice=str(s['invoice_no'])
        # Sales deletion/edit must leave no synthetic inventory reversal. Remove
        # the source SALE movement itself plus any reversal produced by older builds.
        deleted_moves=tx.execute("""SELECT COUNT(*) n FROM inventory_transactions
          WHERE reference_no=? AND reference_type IN ('SALE','SALE_VOID','SALE_DELETE','SALE_EDIT_REVERSAL')""",(invoice,)).fetchone()['n']
        tx.execute("""DELETE FROM inventory_transactions
          WHERE reference_no=? AND reference_type IN ('SALE','SALE_VOID','SALE_DELETE','SALE_EDIT_REVERSAL')""",(invoice,))
        replay=_replay_inventory_state_tx(tx)

        if _money(s['paid_amount'] or 0)>0 and s['cash_account_id']:
            cash_service.post(tx,account_id=s['cash_account_id'],transaction_date=s['sale_date'],
              transaction_type='OUT',amount=_money(s['paid_amount'] or 0),
              description=('Koreksi edit ' if mode=='EDIT' else 'Hapus ') + invoice,
              reference_no=invoice,user_id=actor['id'],allow_negative=True)
        _void_journals(tx,'SALE',sale_id)
        tx.execute("UPDATE sales SET status='VOID',updated_at=? WHERE id=?",(utc_now(),sale_id))
        if 'sales_order_id' in s.keys() and s['sales_order_id']:
            tx.execute("UPDATE sales_orders SET status='OPEN',updated_at=? WHERE id=? AND status='COMPLETED'",(utc_now(),s['sales_order_id']))
        action='SALE_EDIT_RETIRED' if mode=='EDIT' else 'SALE_DELETED'
        audit(actor['id'],action,'sale',sale_id,{'invoice_no':invoice,'inventory_rows_deleted':int(deleted_moves or 0),**replay},ip,tx)
        return {'id':sale_id,'effective_date':s['sale_date'],'inventory_rows_deleted':int(deleted_moves or 0)}

def delete_purchase(actor,purchase_id,ip):
    from .services import inventory_service,cash_service
    with write_transaction() as tx:
        p=tx.execute("SELECT * FROM purchases WHERE id=?",(purchase_id,)).fetchone()
        if not p:raise ValueError('Pembelian tidak ditemukan.')
        if p['status']!='POSTED':raise ValueError('Pembelian sudah dibatalkan.')
        if tx.execute("SELECT 1 FROM purchase_returns WHERE purchase_id=? AND status='POSTED'",(purchase_id,)).fetchone():raise ValueError('Batalkan retur pembelian terlebih dahulu.')
        if tx.execute("""SELECT 1 FROM payable_payments r WHERE r.purchase_id=?
            AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='payable' AND v.transaction_key=CAST(r.id AS TEXT))""",(purchase_id,)).fetchone():raise ValueError('Batalkan pembayaran hutang terlebih dahulu.')
        for i in tx.execute("SELECT * FROM purchase_items WHERE purchase_id=?",(purchase_id,)).fetchall():
            prod=tx.execute("SELECT product_type FROM products WHERE id=?",(i['product_id'],)).fetchone()
            if prod and prod['product_type']=='STOCK': inventory_service.stock_out(tx,product_id=i['product_id'],warehouse_id=p['warehouse_id'],quantity=i['qty'],reference_type='PURCHASE_VOID',reference_no=p['purchase_no'],reason='Pembatalan pembelian',user_id=actor['id'])
        if _money(p['paid_amount'] or 0)>0 and p['cash_account_id']: cash_service.post(tx,account_id=p['cash_account_id'],transaction_date=p['purchase_date'],transaction_type='IN',amount=_money(p['paid_amount'] or 0),description='Pembatalan '+p['purchase_no'],reference_no=p['purchase_no'],user_id=actor['id'])
        _void_journals(tx,'PURCHASE',purchase_id);tx.execute("UPDATE purchases SET status='VOID',updated_at=? WHERE id=?",(utc_now(),purchase_id));
        if 'purchase_order_id' in p.keys() and p['purchase_order_id']: tx.execute("UPDATE purchase_orders SET status='OPEN',updated_at=? WHERE id=? AND status='COMPLETED'",(utc_now(),p['purchase_order_id']))
        audit(actor['id'],'PURCHASE_VOIDED','purchase',purchase_id,{'purchase_no':p['purchase_no']},ip,tx);return {'id':purchase_id}


# ---------------- TINGKATAN HARGA ----------------
def list_price_levels(active=None):
    where=[];params=[]
    if active is not None: where.append("is_active=?");params.append(1 if active else 0)
    clause=(" WHERE "+" AND ".join(where)) if where else ""
    c=connect()
    try:return [dict(r) for r in c.execute("SELECT * FROM price_levels"+clause+" ORDER BY code COLLATE NOCASE, name COLLATE NOCASE",params)]
    finally:c.close()

def create_price_level(actor,d,ip):
    code=str(d.get("code","")).strip().upper();name=str(d.get("name","")).strip()
    if not code or not name:raise ValueError("Kode dan nama tingkatan harga wajib diisi.")
    now=utc_now()
    with write_transaction() as tx:
        cur=tx.execute("INSERT INTO price_levels(code,name,description,is_active,created_at,updated_at) VALUES(?,?,?,?,?,?)",(code,name,str(d.get("description","")).strip() or None,1,now,now))
        audit(actor["id"],"PRICE_LEVEL_CREATED","price_level",cur.lastrowid,{"code":code,"name":name},ip,tx);return cur.lastrowid

def update_price_level(actor,item_id,d,ip):
    code=str(d.get("code","")).strip().upper();name=str(d.get("name","")).strip()
    if not code or not name:raise ValueError("Kode dan nama tingkatan harga wajib diisi.")
    with write_transaction() as tx:
        if not tx.execute("SELECT 1 FROM price_levels WHERE id=?",(item_id,)).fetchone():raise ValueError("Tingkatan harga tidak ditemukan.")
        tx.execute("UPDATE price_levels SET code=?,name=?,description=?,is_active=?,updated_at=? WHERE id=?",(code,name,str(d.get("description","")).strip() or None,1 if d.get("is_active",True) else 0,utc_now(),item_id))
        audit(actor["id"],"PRICE_LEVEL_UPDATED","price_level",item_id,{"code":code,"name":name},ip,tx)

def delete_price_level(actor,item_id,ip):
    with write_transaction() as tx:
        tx.execute("UPDATE price_levels SET is_active=0,updated_at=? WHERE id=?",(utc_now(),item_id))
        audit(actor["id"],"PRICE_LEVEL_DISABLED","price_level",item_id,{},ip,tx)

def _save_product_price_levels(tx,product_id,values):
    if values is None:return
    if not isinstance(values,(list,dict)):raise ValueError("Data harga bertingkat tidak valid.")
    rows=values.items() if isinstance(values,dict) else [(x.get("price_level_id"),x.get("selling_price")) for x in values]
    for level_id,price in rows:
        try: level_id=int(level_id)
        except: continue
        if not tx.execute("SELECT 1 FROM price_levels WHERE id=? AND is_active=1",(level_id,)).fetchone():continue
        amount=_decimal(price or 0,"Harga jual bertingkat")
        tx.execute("INSERT INTO product_price_levels(product_id,price_level_id,selling_price,updated_at) VALUES(?,?,?,?) ON CONFLICT(product_id,price_level_id) DO UPDATE SET selling_price=excluded.selling_price,updated_at=excluded.updated_at",(product_id,level_id,amount,utc_now()))

def product_price_map(product_id):
    c=connect()
    try:return {str(r['price_level_id']):float(r['selling_price']) for r in c.execute("SELECT price_level_id,selling_price FROM product_price_levels WHERE product_id=?",(product_id,))}
    finally:c.close()

def update_sale(actor,sale_id,d,ip):
    """Atomic in-place Sales edit, including invoices that already have installments.

    The original invoice/journal/movement ids stay active.  A temporary replacement
    is built inside the SAME SQLite transaction and folded back into the original
    records.  Any validation error rolls the whole edit back, so merely attempting
    an edit can never make the original journal disappear.
    """
    with write_transaction() as tx:
        old=tx.execute("SELECT * FROM sales WHERE id=?",(sale_id,)).fetchone()
        if not old or old['status']!='POSTED':raise ValueError("Penjualan tidak ditemukan atau sudah dibatalkan.")
        if tx.execute("SELECT 1 FROM sales_returns WHERE sale_id=? AND status='POSTED'",(sale_id,)).fetchone():
            raise ValueError('Batalkan retur penjualan terlebih dahulu sebelum mengubah invoice.')
        invoice=str(old['invoice_no']); delivery=str(old['delivery_no'] or '')
        old_created=old['created_at']
        old_moves=[dict(x) for x in tx.execute("SELECT * FROM inventory_transactions WHERE reference_type='SALE' AND reference_no=? ORDER BY id",(invoice,)).fetchall()]
        old_j=tx.execute("SELECT * FROM journal_entries WHERE source_type='SALE' AND source_id=? ORDER BY CASE status WHEN 'POSTED' THEN 0 ELSE 1 END,id LIMIT 1",(str(sale_id),)).fetchone()
        old_j=dict(old_j) if old_j else None
        old_cash=tx.execute("SELECT * FROM cash_transactions WHERE reference_no=? AND description=? ORDER BY id LIMIT 1",(invoice,f'Penerimaan penjualan {invoice}')).fetchone()
        old_cash=dict(old_cash) if old_cash else None
        upfront=_money(old_cash['amount'] if old_cash else 0)
        receipts=_money(tx.execute("""SELECT COALESCE(SUM(r.amount),0) v FROM receivable_payments r WHERE r.sale_id=?
          AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='receivable' AND v.transaction_key=CAST(r.id AS TEXT))""",(sale_id,)).fetchone()['v'])
        dp=_money(tx.execute("SELECT COALESCE(SUM(amount),0) v FROM downpayment_allocations WHERE allocation_type='CUSTOMER' AND invoice_id=? AND status='POSTED'",(sale_id,)).fetchone()['v'])

        data=dict(d); data['invoice_no']=invoice; data['delivery_no']=delivery
        if 'sales_order_id' in old.keys() and old['sales_order_id'] and not data.get('sales_order_id'):data['sales_order_id']=old['sales_order_id']
        method=str(data.get('payment_method',data.get('payment_type',old['payment_method'] or 'CREDIT'))).upper()
        data['_skip_credit_limit']=True
        data['paid_amount']=str(Decimal('999999999999999.00') if method in ('CASH','TRANSFER') else upfront)

        # Temporarily retire only inside this transaction.  On any error SQLite rolls everything back.
        if old_j: licensing.release_transaction_usage(tx,source_table='journal_entries',source_id=int(old_j['id']))
        tx.execute("UPDATE sales SET status='VOID',invoice_no=?,delivery_no=? WHERE id=?",(f'TMP-EDIT-{sale_id}-{invoice}',f'TMP-EDIT-{sale_id}-{delivery or "SJ"}',sale_id))
        tx.execute("DELETE FROM inventory_transactions WHERE reference_no=? AND reference_type IN ('SALE','SALE_VOID','SALE_DELETE','SALE_EDIT_REVERSAL')",(invoice,))
        if old_cash:tx.execute("DELETE FROM cash_transactions WHERE id=?",(old_cash['id'],))

        new=_create_sale_tx(tx,actor,data,ip); new_id=int(new['id'])
        ns=tx.execute("SELECT * FROM sales WHERE id=?",(new_id,)).fetchone()
        total=_money(ns['total_amount'])
        settled=(upfront+receipts+dp).quantize(Decimal('0.01'))
        if settled>total:
            raise ValueError(f"Total invoice hasil edit ({float(total):,.2f}) tidak boleh lebih kecil dari jumlah yang sudah dibayar/alokasikan ({float(settled):,.2f}).")
        final_due=(total-settled).quantize(Decimal('0.01'))

        # Free unique keys on the temporary row before restoring the original row.
        tx.execute("UPDATE sales SET invoice_no=?,delivery_no=? WHERE id=?",(f'TMP-MERGE-{new_id}-{invoice}',f'TMP-MERGE-{new_id}-{delivery or "SJ"}',new_id))
        cols=['invoice_no','delivery_no','sale_date','customer_id','salesperson_id','warehouse_id','department_id','project_id','payment_type','payment_method','cash_account_id','due_date','tax_percent','tax_amount','subtotal','discount_amount','total_amount','notes','updated_at']
        vals=[invoice,delivery,ns['sale_date'],ns['customer_id'],ns['salesperson_id'],ns['warehouse_id'],ns['department_id'],ns['project_id'],ns['payment_type'],ns['payment_method'],ns['cash_account_id'],ns['due_date'],ns['tax_percent'],ns['tax_amount'],ns['subtotal'],ns['discount_amount'],ns['total_amount'],ns['notes'],utc_now()]
        tx.execute('UPDATE sales SET '+','.join(c+'=?' for c in cols)+",status='POSTED',paid_amount=?,balance_due=? WHERE id=?",(*vals,str(settled),str(final_due),sale_id))
        tx.execute("DELETE FROM sales_items WHERE sale_id=?",(sale_id,));tx.execute("UPDATE sales_items SET sale_id=? WHERE sale_id=?",(sale_id,new_id))

        new_moves=[dict(x) for x in tx.execute("SELECT * FROM inventory_transactions WHERE reference_type='SALE' AND reference_no=? ORDER BY id",(invoice,)).fetchall()]
        for idx,nm in enumerate(new_moves):
            if idx<len(old_moves):
                oid=int(old_moves[idx]['id']);tx.execute("DELETE FROM inventory_transactions WHERE id=?",(nm['id'],))
                tx.execute("""INSERT INTO inventory_transactions(id,product_id,warehouse_id,movement_type,quantity_change,quantity_before,quantity_after,unit_cost,average_cost_before,average_cost_after,reference_type,reference_no,reason,department_id,project_id,user_id,created_at)
                  VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(oid,nm['product_id'],nm['warehouse_id'],nm['movement_type'],nm['quantity_change'],0,0,nm['unit_cost'],0,0,'SALE',invoice,'Penjualan',nm['department_id'],nm['project_id'],nm['user_id'],old_moves[idx].get('created_at') or old_created))
            else:tx.execute("UPDATE inventory_transactions SET created_at=? WHERE id=?",(old_created,nm['id']))
        _replay_inventory_state_tx(tx)

        nj=tx.execute("SELECT * FROM journal_entries WHERE source_type='SALE' AND source_id=? AND status='POSTED'",(str(new_id),)).fetchone()
        if old_j and nj:
            ojid=int(old_j['id']);njid=int(nj['id'])
            licensing.release_transaction_usage(tx,source_table='journal_entries',source_id=njid)
            licensing.record_transaction_usage(tx,event_key=f'JOURNAL-{ojid}',event_type='SALE',reference_no=invoice,source_table='journal_entries',source_id=ojid,units=1)
            tx.execute("DELETE FROM journal_lines WHERE journal_id=?",(ojid,));tx.execute("UPDATE journal_lines SET journal_id=? WHERE journal_id=?",(ojid,njid));tx.execute("DELETE FROM journal_entries WHERE id=?",(njid,))
            tx.execute("""UPDATE journal_entries SET journal_no=?,journal_date=?,description=?,source_type='SALE',source_id=?,reference_no=?,department_id=?,project_id=?,status='POSTED',total_debit=?,total_credit=?,updated_at=? WHERE id=?""",
              (old_j['journal_no'],nj['journal_date'],nj['description'],str(sale_id),invoice,nj['department_id'],nj['project_id'],nj['total_debit'],nj['total_credit'],utc_now(),ojid))
        elif nj:tx.execute("UPDATE journal_entries SET source_id=? WHERE id=?",(str(sale_id),nj['id']))
        _sync_sale_cogs_journals_tx(tx,invoice)

        new_cash=tx.execute("SELECT * FROM cash_transactions WHERE reference_no=? AND description=? ORDER BY id DESC LIMIT 1",(invoice,f'Penerimaan penjualan {invoice}')).fetchone()
        if old_cash and new_cash:
            if int(new_cash['id'])!=int(old_cash['id']):tx.execute("DELETE FROM cash_transactions WHERE id=?",(new_cash['id'],))
            tx.execute("""INSERT OR REPLACE INTO cash_transactions(id,transaction_no,transaction_date,account_id,transaction_type,amount,balance_before,balance_after,description,reference_no,transfer_group,department_id,project_id,user_id,created_at)
              VALUES(?,?,?,?,?,?,0,0,?,?,?,?,?,?,?)""",(old_cash['id'],old_cash['transaction_no'],ns['sale_date'],ns['cash_account_id'],'IN',str(upfront),f'Penerimaan penjualan {invoice}',invoice,None,ns['department_id'],ns['project_id'],ns['user_id'],old_cash['created_at']))
        elif old_cash and not new_cash:
            tx.execute("""INSERT OR REPLACE INTO cash_transactions(id,transaction_no,transaction_date,account_id,transaction_type,amount,balance_before,balance_after,description,reference_no,transfer_group,department_id,project_id,user_id,created_at)
              VALUES(?,?,?,?,?,?,0,0,?,?,?,?,?,?,?)""",(old_cash['id'],old_cash['transaction_no'],old['sale_date'],old_cash['account_id'],'IN',str(upfront),f'Penerimaan penjualan {invoice}',invoice,None,old['department_id'],old['project_id'],old['user_id'],old_cash['created_at']))
        _replay_cash_state_tx(tx)
        tx.execute("DELETE FROM sales WHERE id=?",(new_id,))
        audit(actor['id'],'SALE_UPDATED_IN_PLACE','sale',sale_id,{'invoice_no':invoice,'settled':float(settled),'balance_due':float(final_due)},ip,tx)
    result=get_sale(sale_id) or {'id':sale_id,'invoice_no':invoice};result['updated_from_id']=sale_id;return result


def update_purchase(actor,purchase_id,d,ip):
    """Atomic in-place Purchase edit, preserving installments and original journal."""
    with write_transaction() as tx:
        old=tx.execute("SELECT * FROM purchases WHERE id=?",(purchase_id,)).fetchone()
        if not old or old['status']!='POSTED':raise ValueError("Pembelian tidak ditemukan atau sudah dibatalkan.")
        if tx.execute("SELECT 1 FROM purchase_returns WHERE purchase_id=? AND status='POSTED'",(purchase_id,)).fetchone():
            raise ValueError('Batalkan retur pembelian terlebih dahulu sebelum mengubah pembelian.')
        pno=str(old['purchase_no']); gr=str(old['goods_receipt_no'] or '')
        old_created=old['created_at']
        old_moves=[dict(x) for x in tx.execute("SELECT * FROM inventory_transactions WHERE reference_type='PURCHASE' AND reference_no=? ORDER BY id",(pno,)).fetchall()]
        old_j=tx.execute("SELECT * FROM journal_entries WHERE source_type='PURCHASE' AND source_id=? ORDER BY CASE status WHEN 'POSTED' THEN 0 ELSE 1 END,id LIMIT 1",(str(purchase_id),)).fetchone();old_j=dict(old_j) if old_j else None
        old_cash=tx.execute("SELECT * FROM cash_transactions WHERE reference_no=? AND description=? ORDER BY id LIMIT 1",(pno,f'Pembayaran pembelian {pno}')).fetchone();old_cash=dict(old_cash) if old_cash else None
        upfront=_money(old_cash['amount'] if old_cash else 0)
        payments=_money(tx.execute("""SELECT COALESCE(SUM(r.amount),0) v FROM payable_payments r WHERE r.purchase_id=?
          AND NOT EXISTS(SELECT 1 FROM transaction_voids v WHERE v.transaction_kind='payable' AND v.transaction_key=CAST(r.id AS TEXT))""",(purchase_id,)).fetchone()['v'])
        dp=_money(tx.execute("SELECT COALESCE(SUM(amount),0) v FROM downpayment_allocations WHERE allocation_type='SUPPLIER' AND invoice_id=? AND status='POSTED'",(purchase_id,)).fetchone()['v'])
        data=dict(d);data['purchase_no']=pno;data['goods_receipt_no']=gr
        if 'purchase_order_id' in old.keys() and old['purchase_order_id'] and not data.get('purchase_order_id'):data['purchase_order_id']=old['purchase_order_id']
        method=str(data.get('payment_method',data.get('payment_type',old['payment_method'] or 'CREDIT'))).upper()
        data['paid_amount']=str(Decimal('999999999999999.00') if method in ('CASH','TRANSFER') else upfront)

        if old_j:licensing.release_transaction_usage(tx,source_table='journal_entries',source_id=int(old_j['id']))
        tx.execute("UPDATE purchases SET status='VOID',purchase_no=?,goods_receipt_no=? WHERE id=?",(f'TMP-EDIT-{purchase_id}-{pno}',f'TMP-EDIT-{purchase_id}-{gr or "GR"}',purchase_id))
        tx.execute("DELETE FROM inventory_transactions WHERE reference_type IN ('PURCHASE','PURCHASE_VOID') AND reference_no=?",(pno,))
        if old_cash:tx.execute("DELETE FROM cash_transactions WHERE id=?",(old_cash['id'],))

        new=purchase_service.create_purchase(tx,actor=actor,data=data,client_ip=ip,audit_callback=audit);new_id=int(new['id'])
        np=tx.execute("SELECT * FROM purchases WHERE id=?",(new_id,)).fetchone();total=_money(np['total_amount'])
        settled=(upfront+payments+dp).quantize(Decimal('0.01'))
        if settled>total:raise ValueError(f"Total pembelian hasil edit ({float(total):,.2f}) tidak boleh lebih kecil dari jumlah yang sudah dibayar/alokasikan ({float(settled):,.2f}).")
        final_due=(total-settled).quantize(Decimal('0.01'))

        tx.execute("UPDATE purchases SET purchase_no=?,goods_receipt_no=? WHERE id=?",(f'TMP-MERGE-{new_id}-{pno}',f'TMP-MERGE-{new_id}-{gr or "GR"}',new_id))
        cols=['purchase_no','supplier_invoice_no','goods_receipt_no','purchase_date','supplier_id','warehouse_id','department_id','project_id','payment_type','payment_method','cash_account_id','due_date','tax_percent','tax_amount','subtotal','discount_amount','total_amount','notes','updated_at']
        vals=[pno,np['supplier_invoice_no'],gr,np['purchase_date'],np['supplier_id'],np['warehouse_id'],np['department_id'],np['project_id'],np['payment_type'],np['payment_method'],np['cash_account_id'],np['due_date'],np['tax_percent'],np['tax_amount'],np['subtotal'],np['discount_amount'],np['total_amount'],np['notes'],utc_now()]
        tx.execute('UPDATE purchases SET '+','.join(c+'=?' for c in cols)+",status='POSTED',paid_amount=?,balance_due=? WHERE id=?",(*vals,str(settled),str(final_due),purchase_id))
        tx.execute("DELETE FROM purchase_items WHERE purchase_id=?",(purchase_id,));tx.execute("UPDATE purchase_items SET purchase_id=? WHERE purchase_id=?",(purchase_id,new_id))

        new_moves=[dict(x) for x in tx.execute("SELECT * FROM inventory_transactions WHERE reference_type='PURCHASE' AND reference_no=? ORDER BY id",(pno,)).fetchall()]
        for idx,nm in enumerate(new_moves):
            if idx<len(old_moves):
                oid=int(old_moves[idx]['id']);tx.execute("DELETE FROM inventory_transactions WHERE id=?",(nm['id'],))
                tx.execute("""INSERT INTO inventory_transactions(id,product_id,warehouse_id,movement_type,quantity_change,quantity_before,quantity_after,unit_cost,average_cost_before,average_cost_after,reference_type,reference_no,reason,department_id,project_id,user_id,created_at)
                  VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(oid,nm['product_id'],nm['warehouse_id'],nm['movement_type'],nm['quantity_change'],0,0,nm['unit_cost'],0,0,'PURCHASE',pno,'Pembelian',nm['department_id'],nm['project_id'],nm['user_id'],old_moves[idx].get('created_at') or old_created))
            else:tx.execute("UPDATE inventory_transactions SET created_at=? WHERE id=?",(old_created,nm['id']))
        _replay_inventory_state_tx(tx);_sync_sale_cogs_journals_tx(tx)

        nj=tx.execute("SELECT * FROM journal_entries WHERE source_type='PURCHASE' AND source_id=? AND status='POSTED'",(str(new_id),)).fetchone()
        if old_j and nj:
            ojid=int(old_j['id']);njid=int(nj['id']);licensing.release_transaction_usage(tx,source_table='journal_entries',source_id=njid)
            licensing.record_transaction_usage(tx,event_key=f'JOURNAL-{ojid}',event_type='PURCHASE',reference_no=pno,source_table='journal_entries',source_id=ojid,units=1)
            tx.execute("DELETE FROM journal_lines WHERE journal_id=?",(ojid,));tx.execute("UPDATE journal_lines SET journal_id=? WHERE journal_id=?",(ojid,njid));tx.execute("DELETE FROM journal_entries WHERE id=?",(njid,))
            tx.execute("""UPDATE journal_entries SET journal_no=?,journal_date=?,description=?,source_type='PURCHASE',source_id=?,reference_no=?,department_id=?,project_id=?,status='POSTED',total_debit=?,total_credit=?,updated_at=? WHERE id=?""",
              (old_j['journal_no'],nj['journal_date'],nj['description'],str(purchase_id),pno,nj['department_id'],nj['project_id'],nj['total_debit'],nj['total_credit'],utc_now(),ojid))
        elif nj:tx.execute("UPDATE journal_entries SET source_id=? WHERE id=?",(str(purchase_id),nj['id']))

        new_cash=tx.execute("SELECT * FROM cash_transactions WHERE reference_no=? AND description=? ORDER BY id DESC LIMIT 1",(pno,f'Pembayaran pembelian {pno}')).fetchone()
        if old_cash and new_cash:
            if int(new_cash['id'])!=int(old_cash['id']):tx.execute("DELETE FROM cash_transactions WHERE id=?",(new_cash['id'],))
            tx.execute("""INSERT OR REPLACE INTO cash_transactions(id,transaction_no,transaction_date,account_id,transaction_type,amount,balance_before,balance_after,description,reference_no,transfer_group,department_id,project_id,user_id,created_at)
              VALUES(?,?,?,?,?,?,0,0,?,?,?,?,?,?,?)""",(old_cash['id'],old_cash['transaction_no'],np['purchase_date'],np['cash_account_id'],'OUT',str(upfront),f'Pembayaran pembelian {pno}',pno,None,np['department_id'],np['project_id'],np['user_id'],old_cash['created_at']))
        elif old_cash and not new_cash:
            tx.execute("""INSERT OR REPLACE INTO cash_transactions(id,transaction_no,transaction_date,account_id,transaction_type,amount,balance_before,balance_after,description,reference_no,transfer_group,department_id,project_id,user_id,created_at)
              VALUES(?,?,?,?,?,?,0,0,?,?,?,?,?,?,?)""",(old_cash['id'],old_cash['transaction_no'],old['purchase_date'],old_cash['account_id'],'OUT',str(upfront),f'Pembayaran pembelian {pno}',pno,None,old['department_id'],old['project_id'],old['user_id'],old_cash['created_at']))
        _replay_cash_state_tx(tx)
        tx.execute("DELETE FROM purchases WHERE id=?",(new_id,))
        audit(actor['id'],'PURCHASE_UPDATED_IN_PLACE','purchase',purchase_id,{'purchase_no':pno,'settled':float(settled),'balance_due':float(final_due)},ip,tx)
    result=get_purchase(purchase_id) or {'id':purchase_id,'purchase_no':pno};result['updated_from_id']=purchase_id;return result

def repair_legacy_opening_edit_rows():
    """Collapse legacy OPENING_EDIT rows/journals into the original opening source."""
    with write_transaction() as tx:
        tx.execute("CREATE TABLE IF NOT EXISTS system_repairs(repair_key TEXT PRIMARY KEY,applied_at TEXT NOT NULL,details TEXT)")
        key='2026-08-18-opening-edit-in-place-r1'
        if tx.execute("SELECT 1 FROM system_repairs WHERE repair_key=?",(key,)).fetchone():return {'status':'already_applied'}
        changed=0
        products=tx.execute("SELECT * FROM products WHERE product_type='STOCK' AND (COALESCE(opening_stock_qty,0)<>0 OR EXISTS(SELECT 1 FROM inventory_transactions i WHERE i.product_id=products.id AND i.reference_type='OPENING_EDIT'))").fetchall()
        for p in products:
            target=Decimal(str(p['opening_stock_qty'] or 0));cost=Decimal(str(p['purchase_price'] or 0));wh=int(p['opening_warehouse_id'] or inventory_service.get_default_warehouse_id(tx));od=str(p['opening_balance_date'] or utc_now()[:10])[:10]
            rows=tx.execute("SELECT * FROM inventory_transactions WHERE product_id=? AND reference_type IN ('OPENING','OPENING_EDIT') ORDER BY id",(p['id'],)).fetchall();keep=rows[0] if rows else None
            for r in rows[1:]:tx.execute("DELETE FROM inventory_transactions WHERE id=?",(r['id'],));changed+=1
            if target>0:
                if keep:tx.execute("UPDATE inventory_transactions SET warehouse_id=?,movement_type='IN',quantity_change=?,unit_cost=?,reference_type='OPENING',reference_no='OPENING',reason='Stok awal' WHERE id=?",(wh,str(target),str(cost),keep['id']))
                else:tx.execute("INSERT INTO inventory_transactions(product_id,warehouse_id,movement_type,quantity_change,quantity_before,quantity_after,unit_cost,average_cost_before,average_cost_after,reference_type,reference_no,reason,user_id,created_at) VALUES(?,?,'IN',?,0,0,?,0,0,'OPENING','OPENING','Stok awal',1,?)",(p['id'],wh,str(target),str(cost),od+'T00:00:00Z'))
            elif keep:tx.execute("DELETE FROM inventory_transactions WHERE id=?",(keep['id'],));changed+=1
            _replace_opening_product_journal_tx(tx,p['id'],p['sku'],od,int(p['inventory_account_id'] or accounting_service.account_id(tx,'1200')),(target*cost).quantize(Decimal('0.01')),1)
        replay=_replay_inventory_state_tx(tx);synced=_sync_sale_cogs_journals_tx(tx)
        tx.execute("INSERT INTO system_repairs(repair_key,applied_at,details) VALUES(?,?,?)",(key,utc_now(),json.dumps({'opening_rows':changed,'sales_synced':synced,**replay})))
        return {'status':'applied','opening_rows':changed,'sales_synced':synced,**replay}


CORE_RECONCILIATION_KEY="2026-08-12-core-ledger-r1"
def reconcile_core_linkages():
    """Legacy repair excluding Inventory/HPP. Inventory is transaction-driven and repaired separately."""
    with write_transaction() as tx:
        tx.execute("CREATE TABLE IF NOT EXISTS system_repairs(repair_key TEXT PRIMARY KEY,applied_at TEXT NOT NULL,details TEXT)")
        tx.execute("UPDATE products SET stock_qty=COALESCE((SELECT SUM(quantity) FROM inventory_balances b WHERE b.product_id=products.id),0)")
        key="2026-08-13-core-noninventory-r1"
        if tx.execute("SELECT 1 FROM system_repairs WHERE repair_key=?",(key,)).fetchone():return {"status":"already_applied"}
        now=utc_now();ur=tx.execute("SELECT id FROM users WHERE is_active=1 ORDER BY CASE WHEN username='admin' THEN 0 ELSE 1 END,id LIMIT 1").fetchone();uid=ur["id"] if ur else 1
        stats={"depreciation":0,"opening_partner":0}
        for x in tx.execute("SELECT * FROM business_partners WHERE opening_balance>0 AND partner_type IN ('CUSTOMER','SUPPLIER')").fetchall():
            stype="OPENING_"+x["partner_type"]
            if not tx.execute("SELECT 1 FROM journal_entries WHERE source_type=? AND CAST(source_id AS TEXT)=? AND status='POSTED'",(stype,str(x["id"]))).fetchone():
                accounting_service.post_opening_balance(tx,entity_type=x["partner_type"],entity_id=x["id"],reference_no=x["code"],amount=x["opening_balance"],user_id=uid,opening_date=x["opening_balance_date"]);stats["opening_partner"]+=1
        for d in tx.execute("""SELECT d.*,fa.asset_code,fa.depreciation_expense_account_id,fa.accumulated_depreciation_account_id
          FROM fixed_asset_depreciations d JOIN fixed_assets fa ON fa.id=d.asset_id LEFT JOIN journal_entries j ON j.id=d.journal_id
          WHERE d.journal_id IS NULL OR j.id IS NULL OR j.status<>'POSTED'""").fetchall():
            amt=Decimal(str(d["amount"] or 0))
            if amt<=0:continue
            j=accounting_service.post_journal(tx,journal_date=d["depreciation_date"],description=f"Penyusutan {d['period']} {d['asset_code']}",
              source_type="FIXED_ASSET_DEPRECIATION",source_id=d["id"],reference_no=f"DEP-{d['period']}-{d['asset_code']}",
              lines=[{"account_id":d["depreciation_expense_account_id"],"debit":amt},{"account_id":d["accumulated_depreciation_account_id"],"credit":amt}],user_id=uid)
            tx.execute("UPDATE fixed_asset_depreciations SET journal_id=? WHERE id=?",(j["id"],d["id"]));stats["depreciation"]+=1
        tx.execute("INSERT INTO system_repairs(repair_key,applied_at,details) VALUES(?,?,?)",(key,now,json.dumps(stats)))
        return {"status":"applied",**stats}

def repair_legacy_sale_void_valuation():
    """Remove legacy Sales reversal movements and retire deleted Sales movements.

    Rules:
    - Deleted invoice with no active replacement: remove SALE OUT and all legacy
      SALE_VOID/SALE_DELETE/SALE_EDIT_REVERSAL rows.
    - Edited invoice with an active replacement using the same invoice number:
      remove the old SALE batch that precedes the first reversal and remove all
      reversal rows, while retaining the replacement SALE batch.
    Then replay inventory valuation from the cleaned movement history.
    """
    with write_transaction() as tx:
        tx.execute("CREATE TABLE IF NOT EXISTS system_repairs(repair_key TEXT PRIMARY KEY,applied_at TEXT NOT NULL,details TEXT)")
        key="2026-08-18-sale-delete-remove-movements-r2"
        if tx.execute("SELECT 1 FROM system_repairs WHERE repair_key=?",(key,)).fetchone():
            return {'status':'already_applied'}
        refs=[r['reference_no'] for r in tx.execute("""SELECT DISTINCT reference_no FROM inventory_transactions
          WHERE reference_type IN ('SALE_VOID','SALE_DELETE','SALE_EDIT_REVERSAL') AND COALESCE(reference_no,'')<>''""").fetchall()]
        deleted_source=0;deleted_reversal=0;refs_fixed=0
        for ref in refs:
            reversal=tx.execute("""SELECT MIN(id) first_id FROM inventory_transactions
              WHERE reference_no=? AND reference_type IN ('SALE_VOID','SALE_DELETE','SALE_EDIT_REVERSAL')""",(ref,)).fetchone()
            first_rev=int(reversal['first_id']) if reversal and reversal['first_id'] is not None else None
            active=tx.execute("SELECT id FROM sales WHERE invoice_no=? AND status='POSTED' ORDER BY id DESC LIMIT 1",(ref,)).fetchone()
            if active and first_rev is not None:
                cur=tx.execute("DELETE FROM inventory_transactions WHERE reference_no=? AND reference_type='SALE' AND id<?",(ref,first_rev))
                deleted_source+=max(cur.rowcount,0)
            elif not active:
                cur=tx.execute("DELETE FROM inventory_transactions WHERE reference_no=? AND reference_type='SALE'",(ref,))
                deleted_source+=max(cur.rowcount,0)
            cur=tx.execute("""DELETE FROM inventory_transactions WHERE reference_no=?
              AND reference_type IN ('SALE_VOID','SALE_DELETE','SALE_EDIT_REVERSAL')""",(ref,))
            deleted_reversal+=max(cur.rowcount,0);refs_fixed+=1

        # Also clean deleted Sales created by a build that did not leave a reversal row.
        void_sales=tx.execute("SELECT id,invoice_no FROM sales WHERE status='VOID'").fetchall()
        for sale in void_sales:
            inv=str(sale['invoice_no'] or '')
            # update_sale renames the retired row as VOID-<id>-<original invoice>.
            prefix=f"VOID-{sale['id']}-"
            ref=inv[len(prefix):] if inv.startswith(prefix) else inv
            if not ref:continue
            active=tx.execute("SELECT 1 FROM sales WHERE invoice_no=? AND status='POSTED' LIMIT 1",(ref,)).fetchone()
            if not active:
                cur=tx.execute("DELETE FROM inventory_transactions WHERE reference_no=? AND reference_type='SALE'",(ref,))
                deleted_source+=max(cur.rowcount,0)
                cur=tx.execute("""DELETE FROM inventory_transactions WHERE reference_no=?
                  AND reference_type IN ('SALE_VOID','SALE_DELETE','SALE_EDIT_REVERSAL')""",(ref,))
                deleted_reversal+=max(cur.rowcount,0)

        replay=_replay_inventory_state_tx(tx)
        now=utc_now()
        details={'references_cleaned':refs_fixed,'sale_rows_deleted':deleted_source,'reversal_rows_deleted':deleted_reversal,**replay}
        tx.execute("INSERT INTO system_repairs(repair_key,applied_at,details) VALUES(?,?,?)",(key,now,json.dumps(details)))
        return {'status':'applied',**details}

def rebuild_inventory_valuation_state():
    """Replay movement history and rebuild exact monetary inventory book value."""
    with write_transaction() as tx:
        tx.execute("CREATE TABLE IF NOT EXISTS system_repairs(repair_key TEXT PRIMARY KEY,applied_at TEXT NOT NULL,details TEXT)")
        key="2026-08-13-inventory-book-value-r4"
        if tx.execute("SELECT 1 FROM system_repairs WHERE repair_key=?",(key,)).fetchone():return {"status":"already_applied"}
        now=utc_now();pairs=0;movements_count=0
        for pair in tx.execute("SELECT DISTINCT product_id,warehouse_id FROM inventory_transactions ORDER BY product_id,warehouse_id").fetchall():
            q=Decimal("0.0000");value=Decimal("0.00");avg=Decimal("0.000000")
            moves=tx.execute("SELECT * FROM inventory_transactions WHERE product_id=? AND warehouse_id=? ORDER BY id",(pair["product_id"],pair["warehouse_id"])).fetchall()
            for m in moves:
                change=Decimal(str(m["quantity_change"] or 0)).quantize(Decimal("0.0001"))
                cost=Decimal(str(m["unit_cost"] or 0)).quantize(Decimal("0.000001"))
                before=q;vb=value;after=(before+change).quantize(Decimal("0.0001"))
                vc=(change*cost).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
                va=(vb+vc).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
                ab=avg
                if after>0:avg=(va/after).quantize(Decimal("0.000001"),rounding=ROUND_HALF_UP)
                elif after==0:avg=Decimal("0.000000")
                tx.execute("""UPDATE inventory_transactions SET quantity_before=?,quantity_after=?,
                  average_cost_before=?,average_cost_after=?,value_change=?,value_before=?,value_after=? WHERE id=?""",
                  (str(before),str(after),str(ab),str(avg),str(vc),str(vb),str(va),m["id"]))
                q=after;value=va;movements_count+=1
            tx.execute("""INSERT INTO inventory_balances(warehouse_id,product_id,quantity,average_cost,book_value,updated_at)
              VALUES(?,?,?,?,?,?) ON CONFLICT(warehouse_id,product_id) DO UPDATE SET
              quantity=excluded.quantity,average_cost=excluded.average_cost,book_value=excluded.book_value,updated_at=excluded.updated_at""",
              (pair["warehouse_id"],pair["product_id"],str(q),str(avg),str(value),now));pairs+=1
        tx.execute("UPDATE products SET stock_qty=COALESCE((SELECT SUM(quantity) FROM inventory_balances b WHERE b.product_id=products.id),0),updated_at=?",(now,))
        tx.execute("INSERT INTO system_repairs(repair_key,applied_at,details) VALUES(?,?,?)",(key,now,json.dumps({"pairs":pairs,"movements":movements_count})))
        return {"status":"applied","pairs":pairs,"movements":movements_count}


PROJECT_DOCUMENT_CATEGORIES = [
    "KONTRAK & SPK","KEUANGAN","PROGRESS","TEKNIS / DRAWING",
    "LEGAL","PENGADAAN","SERAH TERIMA","FOTO PROGRESS","LAIN-LAIN"
]
PROJECT_DOCUMENT_ALLOWED_EXT = {
    ".pdf",".xlsx",".xls",".docx",".doc",".jpg",".jpeg",".png",".webp",
    ".txt",".csv",".zip",".rar",".7z",".dwg",".dxf"
}
PROJECT_DOCUMENT_MAX_BYTES = 20 * 1024 * 1024

def _project_document_root():
    from .config import DB_PATH
    root=DB_PATH.parent/"project_documents"
    root.mkdir(parents=True,exist_ok=True)
    return root

def _safe_document_filename(name):
    name=Path(str(name or "dokumen")).name
    stem=re.sub(r"[^A-Za-z0-9._ -]+","_",Path(name).stem).strip(" ._") or "dokumen"
    ext=Path(name).suffix.lower()
    return stem[:100]+ext

def list_project_documents(project_id=None,category=None,q=None,include_history=False):
    c=connect()
    try:
        sql="""SELECT d.*,p.code project_code,p.name project_name,u.username uploaded_by_username
          FROM project_documents d JOIN projects p ON p.id=d.project_id
          LEFT JOIN users u ON u.id=d.uploaded_by WHERE 1=1"""
        params=[]
        if project_id not in (None,""):
            sql+=" AND d.project_id=?";params.append(int(project_id))
        if category:
            sql+=" AND d.category=?";params.append(str(category))
        if not include_history:
            sql+=" AND d.is_current=1"
        if q:
            term="%"+str(q).strip()+"%"
            sql+=" AND (d.title LIKE ? OR COALESCE(d.document_no,'') LIKE ? OR d.original_filename LIKE ? OR COALESCE(d.description,'') LIKE ?)"
            params.extend([term,term,term,term])
        sql+=" ORDER BY COALESCE(d.document_date,d.created_at) DESC,d.id DESC"
        return [dict(x) for x in c.execute(sql,params).fetchall()]
    finally:c.close()

def get_project_document(document_id):
    c=connect()
    try:
        x=c.execute("""SELECT d.*,p.code project_code,p.name project_name,u.username uploaded_by_username
          FROM project_documents d JOIN projects p ON p.id=d.project_id
          LEFT JOIN users u ON u.id=d.uploaded_by WHERE d.id=?""",(int(document_id),)).fetchone()
        return dict(x) if x else None
    finally:c.close()

def save_project_document(actor,d,ip,document_id=None):
    project_id=int(d.get("project_id") or 0)
    if not project_id:raise ValueError("Proyek wajib dipilih.")
    title=str(d.get("title") or "").strip()
    if not title:raise ValueError("Judul dokumen wajib diisi.")
    category=str(d.get("category") or "LAIN-LAIN").strip().upper()
    if category not in PROJECT_DOCUMENT_CATEGORIES:category="LAIN-LAIN"
    original=str(d.get("file_name") or "").strip()
    file_data=str(d.get("file_base64") or "")
    new_file=bool(file_data)
    current=get_project_document(document_id) if document_id else None
    if document_id and not current:raise ValueError("Dokumen proyek tidak ditemukan.")
    if not new_file and not current:raise ValueError("Pilih file dokumen.")

    c=connect()
    try: project=c.execute("SELECT id,code,name FROM projects WHERE id=?",(project_id,)).fetchone()
    finally:c.close()
    if not project:raise ValueError("Proyek tidak ditemukan.")

    now=utc_now()
    version_no=int(current.get("version_no") or 1) if current else 1
    replaces_id=None
    relative_path=current.get("relative_path") if current else None
    stored_filename=current.get("stored_filename") if current else None
    mime_type=current.get("mime_type") if current else None
    file_size=int(current.get("file_size") or 0) if current else 0
    original_filename=current.get("original_filename") if current else original
    raw_content=None

    if new_file:
        if "," in file_data:file_data=file_data.split(",",1)[1]
        try:raw=base64.b64decode(file_data,validate=True)
        except Exception:raise ValueError("File dokumen tidak valid.")
        if not raw:raise ValueError("File dokumen kosong.")
        if len(raw)>PROJECT_DOCUMENT_MAX_BYTES:raise ValueError("Ukuran file dokumen maksimal 20 MB.")
        safe=_safe_document_filename(original or (current.get("original_filename") if current else "dokumen"))
        ext=Path(safe).suffix.lower()
        if ext not in PROJECT_DOCUMENT_ALLOWED_EXT:raise ValueError("Jenis file tidak diizinkan.")
        import uuid as _uuid
        project_folder=re.sub(r"[^A-Za-z0-9._-]+","_",str(project["code"]))[:60] or str(project_id)
        stored_filename=f"{_uuid.uuid4().hex}_{safe}"
        relative_path=str(Path(project_folder)/stored_filename)
        mime_type=mimetypes.guess_type(safe)[0] or "application/octet-stream"
        file_size=len(raw);original_filename=safe
        raw_content=raw
        if not IS_POSTGRES:
            folder=_project_document_root()/project_folder
            folder.mkdir(parents=True,exist_ok=True)
            (folder/stored_filename).write_bytes(raw)
        if current:
            version_no=int(current.get("version_no") or 1)+1
            replaces_id=int(current["id"])

    with write_transaction() as tx:
        if new_file and current:
            tx.execute("UPDATE project_documents SET is_current=0,updated_at=? WHERE id=?",(now,int(current["id"])))
            cur=tx.execute("""INSERT INTO project_documents(
              project_id,category,title,document_no,document_date,valid_until,description,
              original_filename,stored_filename,relative_path,mime_type,file_size,file_content,version_no,replaces_id,is_current,
              uploaded_by,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
              (project_id,category,title,str(d.get("document_no") or "").strip() or None,
               d.get("document_date") or None,d.get("valid_until") or None,str(d.get("description") or "").strip() or None,
               original_filename,stored_filename,relative_path,mime_type,file_size,raw_content,version_no,replaces_id,1,
               actor["id"],now,now))
            new_id=cur.lastrowid;action="PROJECT_DOCUMENT_VERSION_CREATED"
        elif current:
            tx.execute("""UPDATE project_documents SET project_id=?,category=?,title=?,document_no=?,
              document_date=?,valid_until=?,description=?,updated_at=? WHERE id=?""",
              (project_id,category,title,str(d.get("document_no") or "").strip() or None,
               d.get("document_date") or None,d.get("valid_until") or None,str(d.get("description") or "").strip() or None,
               now,int(document_id)))
            new_id=int(document_id);action="PROJECT_DOCUMENT_UPDATED"
        else:
            cur=tx.execute("""INSERT INTO project_documents(
              project_id,category,title,document_no,document_date,valid_until,description,
              original_filename,stored_filename,relative_path,mime_type,file_size,file_content,version_no,replaces_id,is_current,
              uploaded_by,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
              (project_id,category,title,str(d.get("document_no") or "").strip() or None,
               d.get("document_date") or None,d.get("valid_until") or None,str(d.get("description") or "").strip() or None,
               original_filename,stored_filename,relative_path,mime_type,file_size,raw_content,1,None,1,actor["id"],now,now))
            new_id=cur.lastrowid;action="PROJECT_DOCUMENT_CREATED"
        audit(actor["id"],action,"project_document",new_id,{"project_id":project_id,"title":title,"version":version_no},ip,tx)
    return new_id

def delete_project_document(actor,document_id,ip):
    item=get_project_document(document_id)
    if not item:raise ValueError("Dokumen proyek tidak ditemukan.")
    with write_transaction() as tx:
        tx.execute("DELETE FROM project_documents WHERE id=?",(int(document_id),))
        if item.get("is_current") and item.get("replaces_id"):
            tx.execute("UPDATE project_documents SET is_current=1,updated_at=? WHERE id=?",(utc_now(),int(item["replaces_id"])))
        audit(actor["id"],"PROJECT_DOCUMENT_DELETED","project_document",document_id,{"project_id":item["project_id"]},ip,tx)
    if not IS_POSTGRES:
        try:
            path=_project_document_root()/str(item["relative_path"])
            if path.exists():path.unlink()
        except OSError:pass
    return True

def project_document_file(document_id):
    item=get_project_document(document_id)
    if not item:raise ValueError("Dokumen proyek tidak ditemukan.")
    if IS_POSTGRES:
        c=connect()
        try:r=c.execute("SELECT file_content FROM project_documents WHERE id=?",(int(document_id),)).fetchone()
        finally:c.close()
        content=(r["file_content"] if r else None)
        if content is None:raise ValueError("File dokumen belum tersimpan di PostgreSQL.")
        return item,bytes(content)
    root=_project_document_root().resolve()
    path=(root/str(item["relative_path"])).resolve()
    if root not in path.parents:raise ValueError("Lokasi dokumen tidak valid.")
    if not path.exists():raise ValueError("File dokumen tidak ditemukan di penyimpanan.")
    return item,path.read_bytes()



def _inventory_control_account_ids(tx):
    return {int(x["inventory_account_id"]) for x in tx.execute("SELECT DISTINCT inventory_account_id FROM products WHERE product_type='STOCK' AND inventory_account_id IS NOT NULL").fetchall()}

def _block_manual_inventory_accounts(tx,lines):
    controls=_inventory_control_account_ids(tx)
    used=[]
    for line in (lines or []):
        aid=line.get("account_id")
        if aid not in (None,"") and int(aid) in controls:used.append(int(aid))
    if used:
        raise ValueError("Akun Persediaan yang terhubung ke barang tidak boleh diposting lewat Jurnal Manual. Gunakan Saldo Awal Barang, Pembelian, Retur, atau Penyesuaian Stok agar nilai barang dan Neraca selalu sama.")

def _inventory_source_journal(tx,reference_type,reference_no):
    """Resolve inventory movement to its exact business journal.
    Never use a broad reference_no match because reference can be blank/reused.
    """
    rt=str(reference_type or "").upper();ref=reference_no
    source_type=None;source_id=None
    if rt=="PURCHASE":
        row=tx.execute("SELECT id FROM purchases WHERE purchase_no=? AND status='POSTED' ORDER BY id DESC LIMIT 1",(ref,)).fetchone()
        if row:source_type="PURCHASE";source_id=row["id"]
    elif rt=="SALE":
        row=tx.execute("SELECT id FROM sales WHERE invoice_no=? AND status='POSTED' ORDER BY id DESC LIMIT 1",(ref,)).fetchone()
        if row:source_type="SALE";source_id=row["id"]
    elif rt=="PURCHASE_RETURN":
        row=tx.execute("SELECT id FROM purchase_returns WHERE return_no=? AND status='POSTED' ORDER BY id DESC LIMIT 1",(ref,)).fetchone()
        if row:source_type="PURCHASE_RETURN";source_id=row["id"]
    elif rt=="SALES_RETURN":
        row=tx.execute("SELECT id FROM sales_returns WHERE return_no=? AND status='POSTED' ORDER BY id DESC LIMIT 1",(ref,)).fetchone()
        if row:source_type="SALES_RETURN";source_id=row["id"]
    elif rt=="PROJECT_MATERIAL_ISSUE":
        row=tx.execute("SELECT id FROM project_material_issues WHERE issue_no=? AND status='POSTED' ORDER BY id DESC LIMIT 1",(ref,)).fetchone()
        if row:source_type="PROJECT_MATERIAL_ISSUE";source_id=row["id"]
    if source_type is None:return None
    return tx.execute("""SELECT id,source_type,source_id,reference_no FROM journal_entries
      WHERE status='POSTED' AND source_type=? AND CAST(source_id AS TEXT)=?
      ORDER BY id LIMIT 1""",(source_type,str(source_id))).fetchone()

def sync_inventory_gl_by_source():
    """Canonical inventory GL cleanup for databases that passed through older reconciliation builds.

    Rules:
    - legacy blanket/system inventory reconciliations are VOID;
    - duplicate INVENTORY_SOURCE_* journals from the old reference matcher are VOID;
    - old injected 'Sinkron nilai persediaan' pairs are removed;
    - manual journals may not control item inventory;
    - ADJUSTMENT is matched by inventory transaction id, even if reference_no is blank;
    - all other movements are matched to the exact source record id, not a broad reference_no.
    """
    with write_transaction() as tx:
        tx.execute("CREATE TABLE IF NOT EXISTS system_repairs(repair_key TEXT PRIMARY KEY,applied_at TEXT NOT NULL,details TEXT)")
        key="2026-08-13-inventory-canonical-source-r2"
        now=utc_now()
        if tx.execute("SELECT 1 FROM system_repairs WHERE repair_key=?",(key,)).fetchone():
            # Old duplicate repair journals must remain disabled even after restore/import.
            tx.execute("""UPDATE journal_entries SET status='VOID',updated_at=?
              WHERE status='POSTED' AND (source_type IN ('SYSTEM_INVENTORY_RECON','SYSTEM_HPP_RECON')
              OR source_type LIKE 'INVENTORY_SOURCE_%')""",(now,))
            return {"status":"already_applied"}

        stats={"voided_duplicate_repair_journals":0,"removed_old_sync_lines":0,
               "manual_reclassified":0,"source_corrected":0,"source_missing_rebuilt":0}
        controls=_inventory_control_account_ids(tx)
        variance=accounting_service.account_id(tx,"5100")
        hpp_default=accounting_service.account_id(tx,"5000")

        # 1) Disable every reconciliation/duplicate source journal produced by older builds.
        old=tx.execute("""SELECT id FROM journal_entries WHERE status='POSTED' AND
          (source_type IN ('SYSTEM_INVENTORY_RECON','SYSTEM_HPP_RECON') OR source_type LIKE 'INVENTORY_SOURCE_%')""").fetchall()
        for j in old:
            tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE id=?",(now,j["id"]))
        stats["voided_duplicate_repair_journals"]=len(old)

        # 2) Remove correction pairs injected into ORIGINAL journals by v1.22.25/v1.22.28.
        # Both inventory and offset lines carry these exact memo prefixes, so deleting both keeps the journal balanced.
        sync_lines=tx.execute("""SELECT jl.id FROM journal_lines jl JOIN journal_entries je ON je.id=jl.journal_id
          WHERE je.status='POSTED' AND
          (COALESCE(jl.memo,'') LIKE 'Sinkron nilai persediaan %'
           OR COALESCE(jl.memo,'') LIKE 'Lawan sinkron nilai persediaan %')""").fetchall()
        for line in sync_lines:tx.execute("DELETE FROM journal_lines WHERE id=?",(line["id"],))
        stats["removed_old_sync_lines"]=len(sync_lines)

        # Refresh totals after removing old correction pairs.
        tx.execute("""UPDATE journal_entries SET
          total_debit=COALESCE((SELECT SUM(debit) FROM journal_lines WHERE journal_id=journal_entries.id),0),
          total_credit=COALESCE((SELECT SUM(credit) FROM journal_lines WHERE journal_id=journal_entries.id),0),
          updated_at=? WHERE status='POSTED'""",(now,))

        # 3) A manual journal has no SKU/warehouse. It must not alter an item-controlled inventory account.
        if controls:
            qs=",".join("?" for _ in controls)
            bad=tx.execute(f"""SELECT jl.id FROM journal_lines jl JOIN journal_entries je ON je.id=jl.journal_id
              WHERE je.status='POSTED' AND je.source_type IN ('MANUAL','MANUAL_EXCEL')
              AND jl.account_id IN ({qs})""",list(controls)).fetchall()
            for line in bad:
                tx.execute("""UPDATE journal_lines SET account_id=?,
                  memo=TRIM(COALESCE(memo,'')||' [Legacy inventory manual -> Penyesuaian Stok]')
                  WHERE id=?""",(variance,line["id"]))
            stats["manual_reclassified"]=len(bad)

        # 4) Canonicalize non-opening movement sources.
        groups=tx.execute("""SELECT it.reference_type,it.reference_no,p.inventory_account_id,
          COALESCE(SUM(COALESCE(it.value_change,it.quantity_change*it.unit_cost)),0) expected,
          MIN(it.id) first_tx_id,MIN(substr(it.created_at,1,10)) trx_date
          FROM inventory_transactions it JOIN products p ON p.id=it.product_id
          WHERE p.inventory_account_id IS NOT NULL AND COALESCE(it.reference_type,'') NOT IN
            ('TRANSFER','WAREHOUSE_REASSIGN','OPENING','OPENING_EDIT',
             'SALE_VOID','SALE_DELETE','SALE_EDIT_REVERSAL','PURCHASE_VOID','PURCHASE_RETURN_VOID','SALES_RETURN_VOID')
          GROUP BY it.reference_type,it.reference_no,p.inventory_account_id
          ORDER BY MIN(it.id)""").fetchall()

        for g in groups:
            rt=str(g["reference_type"] or "").upper()
            aid=int(g["inventory_account_id"])
            expected=Decimal(str(g["expected"] or 0)).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)

            if rt=="ADJUSTMENT":
                # Stock adjustment journal source_id is the exact inventory transaction id.
                # This fixes the old bug when reference_no was NULL.
                tids=tx.execute("""SELECT it.id,COALESCE(it.value_change,it.quantity_change*it.unit_cost) value_change
                  FROM inventory_transactions it JOIN products p ON p.id=it.product_id
                  WHERE it.reference_type='ADJUSTMENT'
                    AND ((it.reference_no=? ) OR (it.reference_no IS NULL AND ? IS NULL))
                    AND p.inventory_account_id=? ORDER BY it.id""",(g["reference_no"],g["reference_no"],aid)).fetchall()
                targets=[]
                for movement in tids:
                    j=tx.execute("""SELECT id FROM journal_entries WHERE status='POSTED'
                      AND source_type='STOCK_ADJUSTMENT' AND CAST(source_id AS TEXT)=? ORDER BY id LIMIT 1""",
                      (str(movement["id"]),)).fetchone()
                    targets.append((j,Decimal(str(movement["value_change"] or 0)).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)))
            else:
                j=_inventory_source_journal(tx,rt,g["reference_no"])
                targets=[(j,expected)]

            for j,exp in targets:
                if not j:
                    # Do NOT blindly create a journal when source cannot be resolved.
                    # Missing source will be reported and left for transaction-integrity repair.
                    stats["source_missing_rebuilt"]+=1
                    continue
                actual=Decimal(str(tx.execute("""SELECT COALESCE(SUM(debit-credit),0) v
                  FROM journal_lines WHERE journal_id=? AND account_id=?""",(j["id"],aid)).fetchone()["v"] or 0)).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
                delta=(exp-actual).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
                if abs(delta)<=Decimal("0.01"):continue

                # Source-specific balancing: inventory delta is corrected inside the ORIGINAL source journal.
                # SALE/returns use HPP; stock adjustments/purchases use Penyesuaian Stok.
                offset=hpp_default if rt in ("SALE","SALES_RETURN","PURCHASE_RETURN") else variance
                if delta>0:
                    tx.execute("""INSERT INTO journal_lines(journal_id,account_id,debit,credit,memo)
                      VALUES(?,?,?,?,?)""",(j["id"],aid,str(delta),"0",f"Canonical inventory {rt}"))
                    tx.execute("""INSERT INTO journal_lines(journal_id,account_id,debit,credit,memo)
                      VALUES(?,?,?,?,?)""",(j["id"],offset,"0",str(delta),f"Canonical inventory offset {rt}"))
                else:
                    amt=abs(delta)
                    tx.execute("""INSERT INTO journal_lines(journal_id,account_id,debit,credit,memo)
                      VALUES(?,?,?,?,?)""",(j["id"],aid,"0",str(amt),f"Canonical inventory {rt}"))
                    tx.execute("""INSERT INTO journal_lines(journal_id,account_id,debit,credit,memo)
                      VALUES(?,?,?,?,?)""",(j["id"],offset,str(amt),"0",f"Canonical inventory offset {rt}"))
                tx.execute("""UPDATE journal_entries SET
                  total_debit=COALESCE((SELECT SUM(debit) FROM journal_lines WHERE journal_id=?),0),
                  total_credit=COALESCE((SELECT SUM(credit) FROM journal_lines WHERE journal_id=?),0),updated_at=?
                  WHERE id=?""",(j["id"],j["id"],now,j["id"]))
                stats["source_corrected"]+=1

        # Final diagnostic: inventory ledger and GL per control account.
        diagnostics=[]
        for aid in sorted(controls):
            stock=Decimal(str(tx.execute("""SELECT COALESCE(SUM(COALESCE(ib.book_value,ib.quantity*ib.average_cost)),0) v
              FROM inventory_balances ib JOIN products p ON p.id=ib.product_id
              WHERE p.inventory_account_id=?""",(aid,)).fetchone()["v"] or 0)).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
            gl=Decimal(str(tx.execute("""SELECT COALESCE(SUM(jl.debit-jl.credit),0) v FROM journal_lines jl
              JOIN journal_entries je ON je.id=jl.journal_id WHERE je.status='POSTED' AND jl.account_id=?""",
              (aid,)).fetchone()["v"] or 0)).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
            diagnostics.append({"account_id":aid,"inventory_value":float(stock),"gl_value":float(gl),"difference":float(stock-gl)})
        stats["accounts"]=diagnostics
        tx.execute("INSERT INTO system_repairs(repair_key,applied_at,details) VALUES(?,?,?)",
          (key,now,json.dumps(stats)))
        return {"status":"applied",**stats}


def enforce_inventory_transaction_integrity():
    """Make inventory GL follow stock movements; never plug inventory differences to retained earnings."""
    with write_transaction() as tx:
        tx.execute("CREATE TABLE IF NOT EXISTS system_repairs(repair_key TEXT PRIMARY KEY,applied_at TEXT NOT NULL,details TEXT)")
        key="2026-08-13-inventory-transaction-driven-r2"
        # Always remove/void blanket reconciliation entries from prior builds.
        old=tx.execute("""SELECT id FROM journal_entries WHERE source_type IN ('SYSTEM_INVENTORY_RECON','SYSTEM_HPP_RECON') AND status='POSTED'""").fetchall()
        for j in old:tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE id=?",(utc_now(),j["id"]))
        tx.execute("UPDATE products SET stock_qty=COALESCE((SELECT SUM(quantity) FROM inventory_balances b WHERE b.product_id=products.id),0)")
        if tx.execute("SELECT 1 FROM system_repairs WHERE repair_key=?",(key,)).fetchone():
            return {"status":"already_applied","voided_legacy_reconciliation":len(old)}
        now=utc_now();ur=tx.execute("SELECT id FROM users WHERE is_active=1 ORDER BY CASE WHEN username='admin' THEN 0 ELSE 1 END,id LIMIT 1").fetchone();uid=ur["id"] if ur else 1
        stats={"voided_legacy_reconciliation":len(old),"opening_normalized":0,"sale_hpp_normalized":0,"unresolved_inventory_accounts":0}

        # Consolidate every product's OPENING/OPENING_EDIT movements into one source-linked opening journal.
        products=tx.execute("""SELECT p.id,p.sku,p.inventory_account_id,p.opening_balance_date
          FROM products p WHERE p.product_type='STOCK'""").fetchall()
        retained=accounting_service.account_id(tx,"3200")
        for p in products:
            expected=Decimal(str(tx.execute("""SELECT COALESCE(SUM(quantity_change*unit_cost),0) v FROM inventory_transactions
              WHERE product_id=? AND reference_type IN ('OPENING','OPENING_EDIT')""",(p["id"],)).fetchone()["v"] or 0)).quantize(Decimal("0.01"))
            jrows=tx.execute("""SELECT id FROM journal_entries WHERE source_id=? AND source_type IN
              ('OPENING_PRODUCT','PRODUCT_OPENING','OPENING_PRODUCT_EDIT','PRODUCT_OPENING_EDIT') AND status='POSTED'""",(p["id"],)).fetchall()
            existing=Decimal("0")
            for j in jrows:
                existing+=Decimal(str(tx.execute("""SELECT COALESCE(SUM(jl.debit-jl.credit),0) v FROM journal_lines jl
                  WHERE jl.journal_id=? AND jl.account_id=?""",(j["id"],p["inventory_account_id"])).fetchone()["v"] or 0))
            if abs(existing-expected)>Decimal("0.01") or len(jrows)>1:
                for j in jrows:tx.execute("UPDATE journal_entries SET status='VOID',updated_at=? WHERE id=?",(now,j["id"]))
                if expected!=0:
                    lines=[{"account_id":p["inventory_account_id"],"debit":expected},{"account_id":retained,"credit":expected}] if expected>0 else [{"account_id":retained,"debit":abs(expected)},{"account_id":p["inventory_account_id"],"credit":abs(expected)}]
                    accounting_service.post_journal(tx,journal_date=p["opening_balance_date"] or now[:10],description=f"Saldo awal {p['sku']}",
                      source_type="OPENING_PRODUCT",source_id=p["id"],reference_no=p["sku"],lines=lines,user_id=uid)
                stats["opening_normalized"]+=1

        # Rewrite only the inventory/HPP pair inside each SALE journal from actual SALE movements.
        for s in tx.execute("SELECT id,invoice_no FROM sales WHERE status='POSTED'").fetchall():
            groups=tx.execute("""SELECT p.cogs_account_id,p.inventory_account_id,
              COALESCE(SUM(ABS(it.quantity_change)*it.unit_cost),0) amount
              FROM inventory_transactions it JOIN products p ON p.id=it.product_id
              WHERE it.reference_type='SALE' AND it.reference_no=?
              GROUP BY p.cogs_account_id,p.inventory_account_id""",(s["invoice_no"],)).fetchall()
            if not groups:continue
            j=tx.execute("SELECT id FROM journal_entries WHERE source_type='SALE' AND source_id=? AND status='POSTED' ORDER BY id LIMIT 1",(s["id"],)).fetchone()
            if not j:continue
            account_ids=set()
            for g in groups:
                account_ids.add(int(g["cogs_account_id"] or accounting_service.account_id(tx,"5000")))
                account_ids.add(int(g["inventory_account_id"] or accounting_service.account_id(tx,"1200")))
            current=Decimal("0")
            for aid in account_ids:
                current+=abs(Decimal(str(tx.execute("SELECT COALESCE(SUM(debit-credit),0) v FROM journal_lines WHERE journal_id=? AND account_id=?",(j["id"],aid)).fetchone()["v"] or 0)))
            expected_pairs=sum(Decimal(str(g["amount"] or 0))*2 for g in groups)
            if abs(current-expected_pairs)>Decimal("0.01"):
                qs=",".join("?" for _ in account_ids)
                tx.execute(f"DELETE FROM journal_lines WHERE journal_id=? AND account_id IN ({qs})",[j["id"],*account_ids])
                for g in groups:
                    amt=Decimal(str(g["amount"] or 0)).quantize(Decimal("0.01"))
                    if amt<=0:continue
                    ca=int(g["cogs_account_id"] or accounting_service.account_id(tx,"5000"));ia=int(g["inventory_account_id"] or accounting_service.account_id(tx,"1200"))
                    tx.execute("INSERT INTO journal_lines(journal_id,account_id,debit,credit,memo,partner_id,partner_type,department_id,project_id) VALUES(?,?,?,?,?,?,?,?,?)",
                      (j["id"],ca,str(amt),"0","HPP",None,None,None,None))
                    tx.execute("INSERT INTO journal_lines(journal_id,account_id,debit,credit,memo,partner_id,partner_type,department_id,project_id) VALUES(?,?,?,?,?,?,?,?,?)",
                      (j["id"],ia,"0",str(amt),"Persediaan keluar",None,None,None,None))
                stats["sale_hpp_normalized"]+=1

        # Report residual differences only; do not auto-create a balancing/reconciliation journal.
        for a in tx.execute("SELECT DISTINCT inventory_account_id id FROM products WHERE product_type='STOCK' AND inventory_account_id IS NOT NULL").fetchall():
            aid=int(a["id"])
            stock=Decimal(str(tx.execute("""SELECT COALESCE(SUM(COALESCE(ib.book_value,ib.quantity*ib.average_cost)),0) v FROM inventory_balances ib
              JOIN products p ON p.id=ib.product_id WHERE p.inventory_account_id=?""",(aid,)).fetchone()["v"] or 0)).quantize(Decimal("0.01"))
            gl=Decimal(str(tx.execute("""SELECT COALESCE(SUM(jl.debit-jl.credit),0) v FROM journal_lines jl JOIN journal_entries je ON je.id=jl.journal_id
              WHERE je.status='POSTED' AND jl.account_id=?""",(aid,)).fetchone()["v"] or 0)).quantize(Decimal("0.01"))
            if abs(stock-gl)>Decimal("0.01"):stats["unresolved_inventory_accounts"]+=1
        tx.execute("INSERT INTO system_repairs(repair_key,applied_at,details) VALUES(?,?,?)",(key,now,json.dumps(stats)))
        return {"status":"applied",**stats}
