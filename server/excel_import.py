import base64
from io import BytesIO
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import re

from .database import connect, write_transaction, utc_now
from . import repository as repo
from .product import IS_ULTIMA

SPECS = {
    "accounts": {"title":"Daftar Akun (COA)","headers":["Kode","Nama","Tipe Akun","Subtipe","Saldo Normal","Kode Induk","Aktif"],"sample":["6100","Biaya Administrasi","EXPENSE","OPERATING_EXPENSE","DEBIT","","YA"]},
    "categories": {"title":"Kategori Barang / Jasa","optional_headers":["Jenis Kategori"],"headers":["Kode","Nama","Jenis Kategori"],"sample":["ELEK","Elektronik","BARANG"]},
    "units": {"title":"Satuan","headers":["Kode","Nama","Desimal"],"sample":["PCS","Pieces",0]},
    "brands": {"title":"Merk","headers":["Kode","Nama"],"sample":["GEN","Generic"]},
    "warehouses": {"title":"Gudang","headers":["Kode","Nama","Alamat","Gudang Utama"],"sample":["GDG2","Gudang Cabang","Jakarta","TIDAK"]},
    "customers": {"title":"Pelanggan","optional_headers":["Kode Akun Piutang"],"headers":["Kode","Nama","NPWP","Telepon","Email","Alamat","Kota","Batas Kredit","Termin Hari","Saldo Awal","Tanggal Saldo Awal","Grup","Kontak","Catatan","Kode Akun Piutang"],"sample":["CUST001","PT Contoh Pelanggan","","021000000","finance@example.com","Jl. Contoh","Jakarta",0,30,0,date.today().isoformat(),"Retail","Budi","","1100"]},
    "suppliers": {"title":"Pemasok","optional_headers":["Kode Akun Hutang"],"headers":["Kode","Nama","NPWP","Telepon","Email","Alamat","Kota","Batas Kredit","Termin Hari","Saldo Awal","Tanggal Saldo Awal","Grup","Kontak","Catatan","Kode Akun Hutang"],"sample":["SUP001","PT Contoh Pemasok","","021000001","ap@example.com","Jl. Pemasok","Jakarta",0,30,0,date.today().isoformat(),"Material","Sari","","2000"]},
    "products": {"title":"Barang dan Saldo Awal","optional_headers":["Barcode","Satuan Alternatif 1","Rasio 1","Harga Beli Alt 1","Harga Jual Alt 1","Satuan Alternatif 2","Rasio 2","Harga Beli Alt 2","Harga Jual Alt 2","Satuan Alternatif 3","Rasio 3","Harga Beli Alt 3","Harga Jual Alt 3"],"headers":["SKU","Nama","Barcode","Kode Kategori","Kode Merk","Kode Satuan","Harga Beli","Harga Jual","Stok Minimum","Stok Awal","Kode Gudang","Tanggal Saldo Awal","Kode Akun Persediaan","Kode Akun Penjualan","Kode Akun HPP","Catatan","Satuan Alternatif 1","Rasio 1","Harga Beli Alt 1","Harga Jual Alt 1","Satuan Alternatif 2","Rasio 2","Harga Beli Alt 2","Harga Jual Alt 2","Satuan Alternatif 3","Rasio 3","Harga Beli Alt 3","Harga Jual Alt 3"],"sample":["BRG001","Barang Contoh","","","","KG",100000,120000,5,0,"UTAMA",date.today().isoformat(),"1200","4000","5000","","GRAM",0.001,100,120,"","","","","","","",""]},
    "services": {"title":"Jasa","optional_headers":["Kode Satuan"],"headers":["Kode Jasa","Nama Jasa","Kode Kategori Jasa","Kode Satuan","Harga Beli","Harga Jual","PPN Persen","Kode Akun Pembelian","Kode Akun Penjualan","Catatan"],"sample":["JSA001","Jasa Instalasi","","JAM",0,250000,11,"5100","4000",""]},
    "cash_accounts": {"title":"Kas dan Bank + Saldo Awal","headers":["Kode","Nama","Jenis","Nama Bank","Nomor Rekening","Saldo Awal","Tanggal Saldo Awal","Kode COA Kas/Bank"],"sample":["BANK2","Bank Operasional","BANK","BCA","1234567890",1000000,date.today().isoformat(),"1010"]},
    "salespersons": {"title":"Salesman","headers":["Kode","Nama","Telepon","Email","Komisi Persen","Aktif"],"sample":["SLS001","Salesman Contoh","08123456789","sales@example.com",2.5,"YA"]},
    "fixed_assets": {"title":"Aktiva Tetap","headers":["Kode Aktiva","Nama Aktiva","Tanggal Perolehan","Harga Perolehan","Nilai Residu","Masa Manfaat Bulan","Kode Akun Aktiva","Kode Akun Akumulasi","Kode Akun Beban Penyusutan","Kode Akun Lawan","Catatan"],"sample":["FA001","Komputer Kantor",date.today().isoformat(),12000000,0,48,"1400","1410","5200","1000",""]},
    "departments": {"title":"Departemen (Ultima)","headers":["Kode","Nama","Manajer","Catatan","Aktif"],"sample":["OPS","Operasional","Budi","","YA"],"ultima":True},
    "projects": {"title":"Project (Ultima)","headers":["Kode","Nama","Kode Pelanggan","Tanggal Mulai","Tanggal Selesai","Status","Catatan","Aktif"],"sample":["PRJ001","Proyek Contoh","CUST001",date.today().isoformat(),"","ACTIVE","","YA"],"ultima":True},
}

ALIASES={
 'kode barang':'sku','kode produk':'sku','nama barang':'nama','nama produk':'nama','unit':'kode satuan','satuan':'kode satuan',
 'kategori':'kode kategori','merk':'kode merk','gudang':'kode gudang','harga modal':'harga beli','harga beli default':'harga beli',
 'harga jual default':'harga jual','minimum stok':'stok minimum','qty awal':'stok awal','jumlah awal':'stok awal',
 'tanggal awal':'tanggal saldo awal','akun persediaan':'kode akun persediaan','akun penjualan':'kode akun penjualan','akun hpp':'kode akun hpp',
 'kode customer':'kode','nama customer':'nama','kode pelanggan':'kode','nama pelanggan':'nama','kode supplier':'kode','nama supplier':'nama',
 'akun piutang':'kode akun piutang','coa piutang':'kode akun piutang','akun hutang':'kode akun hutang','coa hutang':'kode akun hutang','jenis':'jenis kategori','tipe kategori':'jenis kategori','kategori jenis':'jenis kategori','jenis data':'jenis kategori','tipe data':'jenis kategori','termin':'termin hari','jatuh tempo hari':'termin hari','saldo':'saldo awal','kode jasa/layanan':'kode jasa','nama layanan':'nama jasa'
}

def available_types(): return [{"key":k,"title":v["title"]} for k,v in SPECS.items() if IS_ULTIMA or not v.get("ultima")]

def _wb():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        return Workbook, load_workbook, Font, PatternFill, Alignment
    except ImportError as exc: raise ValueError("Impor Excel membutuhkan openpyxl.") from exc

def template(kind):
    if kind not in SPECS or (SPECS[kind].get("ultima") and not IS_ULTIMA): raise ValueError("Jenis template tidak tersedia.")
    Workbook,_,Font,PatternFill,Alignment=_wb();spec=SPECS[kind];wb=Workbook();ws=wb.active;ws.title='Template'
    ws.append(spec['headers']);ws.append(spec['sample'])
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='087F8C');c.alignment=Alignment(horizontal='center')
    from openpyxl.utils import get_column_letter
    for i,h in enumerate(spec['headers'],1): ws.column_dimensions[get_column_letter(i)].width=max(14,min(32,len(h)+5))
    note=wb.create_sheet('Petunjuk');note.append([spec['title']]);note.append(['Kolom boleh berurutan berbeda, tetapi nama header wajib dikenali. Tanggal gunakan YYYY-MM-DD. Data dengan kode yang sama akan diperbarui.'])
    bio=BytesIO();wb.save(bio);return bio.getvalue(),f'Template_{kind}.xlsx'

def _text(v): return str(v if v is not None else '').strip()
def _norm(v):
    s=re.sub(r'[^a-z0-9]+',' ',_text(v).lower()).strip();return ALIASES.get(s,s)
def _yes(v,default=True):
    s=_text(v).upper();return default if not s else s in ('YA','Y','YES','1','TRUE','AKTIF','ACTIVE')
def _num(v,default=0):
    if v in (None,''): return default
    if isinstance(v,(int,float,Decimal)): return float(v)
    s=_text(v).replace('Rp','').replace('rp','').replace(' ','')
    neg=s.startswith('(') and s.endswith(')');s=s.strip('()')
    if ',' in s and '.' in s:
        if s.rfind(',')>s.rfind('.'): s=s.replace('.','').replace(',','.')
        else: s=s.replace(',','')
    elif ',' in s:
        p=s.split(',');s=(p[0].replace('.','')+'.'+p[1]) if len(p)==2 and len(p[1])<=4 else ''.join(p)
    elif s.count('.')>1 or (s.count('.')==1 and len(s.rsplit('.',1)[1])==3): s=s.replace('.','')
    try: n=float(s)
    except Exception: raise ValueError(f"Nilai angka tidak valid: {v}")
    return -n if neg else n

def _date(v):
    if not v:return date.today().isoformat()
    if isinstance(v,datetime):return v.date().isoformat()
    if isinstance(v,date):return v.isoformat()
    s=_text(v)
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%d-%m-%Y','%m/%d/%Y'):
        try:return datetime.strptime(s[:10],fmt).date().isoformat()
        except:pass
    raise ValueError(f'Tanggal tidak valid: {s}')

def _lookup(table,code_col,code,label,optional=False):
    code=_text(code).upper()
    if not code and optional:return None
    c=connect()
    try:r=c.execute(f'SELECT id FROM {table} WHERE UPPER({code_col})=UPPER(?)',(code,)).fetchone()
    finally:c.close()
    if not r:raise ValueError(f"{label} '{code}' tidak ditemukan.")
    return int(r['id'])
def _clean_code(value):
    s=_text(value)
    if re.fullmatch(r'[-+]?\d+\.0+',s): s=s.split('.',1)[0]
    return s

def _account(code,label):return _lookup('chart_of_accounts','code',_clean_code(code),label)
def _partner_account(code,subtype,label):
    code=_clean_code(code)
    if not code:return None
    c=connect()
    try:r=c.execute("SELECT id,account_subtype FROM chart_of_accounts WHERE UPPER(code)=UPPER(?) AND is_active=1",(code,)).fetchone()
    finally:c.close()
    if not r:raise ValueError(f"{label} '{code}' tidak ditemukan.")
    if str(r['account_subtype'] or '').upper()!=subtype:raise ValueError(f"{label} '{code}' harus bertipe {'Piutang' if subtype=='RECEIVABLE' else 'Hutang'}.")
    return int(r['id'])

def _service_category(code):
    code=_text(code).upper()
    if not code:return None
    c=connect()
    try:
        row=c.execute('SELECT id FROM service_categories WHERE UPPER(code)=UPPER(?)',(code,)).fetchone()
        if row:return int(row['id'])
        item=c.execute('SELECT code,name,is_active FROM item_categories WHERE UPPER(code)=UPPER(?)',(code,)).fetchone()
    finally:c.close()
    if not item:raise ValueError(f"Kategori jasa '{code}' tidak ditemukan.")
    now=utc_now()
    with write_transaction() as tx:
        row=tx.execute('SELECT id FROM service_categories WHERE UPPER(code)=UPPER(?)',(code,)).fetchone()
        if row:return int(row['id'])
        cur=tx.execute('INSERT INTO service_categories(code,name,notes,is_active,created_at,updated_at) VALUES(?,?,NULL,?,?,?)',(item['code'],item['name'],int(item['is_active'] if item['is_active'] is not None else 1),now,now))
        return int(cur.lastrowid)
def _existing(table,col,value):
    c=connect()
    try:return c.execute(f'SELECT id FROM {table} WHERE UPPER({col})=UPPER(?)',(_text(value),)).fetchone()
    finally:c.close()

def _create_account(actor,row,ip):
    code,name,atype,subtype,normal,parent_code,active=[_text(x) for x in row[:7]];atype=atype.upper();subtype=subtype.upper();normal=normal.upper()
    shorthand={
      'CASH_BANK':('ASSET','CASH_BANK'),'RECEIVABLE':('ASSET','RECEIVABLE'),'PAYABLE':('LIABILITY','PAYABLE'),
      'HPP':('EXPENSE','HPP'),'OPERATING_EXPENSE':('EXPENSE','OPERATING_EXPENSE'),'REVENUE':('REVENUE','REVENUE')}
    if atype in shorthand:
        mapped=shorthand[atype];atype,subtype=mapped
    if atype not in ('ASSET','LIABILITY','EQUITY','REVENUE','EXPENSE'):raise ValueError('Tipe akun harus ASSET/LIABILITY/EQUITY/REVENUE/EXPENSE atau CASH_BANK/RECEIVABLE/PAYABLE/HPP/OPERATING_EXPENSE.')
    if normal not in ('DEBIT','CREDIT'):normal='DEBIT' if atype in ('ASSET','EXPENSE') else 'CREDIT'
    parent=_lookup('chart_of_accounts','code',_clean_code(parent_code),'Akun induk',True);now=utc_now()
    with write_transaction() as tx:
        old=tx.execute('SELECT id FROM chart_of_accounts WHERE UPPER(code)=UPPER(?)',(_clean_code(code),)).fetchone()
        if old:
            account_id=int(old['id'])
            tx.execute('UPDATE chart_of_accounts SET name=?,account_type=?,account_subtype=?,normal_balance=?,parent_id=?,is_active=?,updated_at=? WHERE id=?',(name,atype,subtype or None,normal,parent,1 if _yes(active) else 0,now,account_id))
        else:
            cur=tx.execute('INSERT INTO chart_of_accounts(code,name,account_type,account_subtype,normal_balance,parent_id,is_system,is_active,created_at,updated_at) VALUES(?,?,?,?,?,?,0,?,?,?)',(_clean_code(code).upper(),name,atype,subtype or None,normal,parent,1 if _yes(active) else 0,now,now));account_id=cur.lastrowid
        # COA CASH_BANK wajib langsung tersedia sebagai rekening operasional.
        if subtype=='CASH_BANK': repo._sync_cash_account_for_coa(tx,account_id,_clean_code(code).upper(),name,subtype,1 if _yes(active) else 0,actor['id'])
        return account_id

def _process(kind,actor,row,ip):
    if kind=='accounts':return _create_account(actor,row,ip)
    if kind=='categories':
        code=_text(row[0]).upper(); name=_text(row[1]); category_kind=_text(row[2] if len(row)>2 else '').upper() or 'BOTH'
        if category_kind in ('ITEM','STOCK','PRODUCT'): category_kind='BARANG'
        if category_kind in ('SERVICE','LAYANAN'): category_kind='JASA'
        if category_kind not in ('BARANG','JASA','BOTH','KEDUANYA'): raise ValueError("Jenis Kategori harus BARANG, JASA, atau BOTH.")
        result=None
        if category_kind in ('BARANG','BOTH','KEDUANYA'):
            old=_existing('item_categories','code',code); result=repo.update_category(actor,old['id'],{'code':code,'name':name,'is_active':True},ip) if old else repo.create_category(actor,{'code':code,'name':name},ip)
        if category_kind in ('JASA','BOTH','KEDUANYA'):
            old=_existing('service_categories','code',code); result=repo.update_service_category(actor,old['id'],{'code':code,'name':name,'is_active':True},ip) if old else repo.create_service_category(actor,{'code':code,'name':name,'is_active':True},ip)
        return result
    if kind=='units':
        old=_existing('units','code',row[0]);p={'code':row[0],'name':row[1],'decimals':int(_num(row[2],0))};return repo.update_unit(actor,old['id'],p,ip) if old else repo.create_unit(actor,p,ip)
    if kind=='brands':
        old=_existing('brands','code',row[0]);p={'code':row[0],'name':row[1]};return repo.update_brand(actor,old['id'],p,ip) if old else repo.create_brand(actor,p,ip)
    if kind=='warehouses':
        old=_existing('warehouses','code',row[0]);p={'code':row[0],'name':row[1],'address':row[2],'is_default':_yes(row[3],False)};return repo.update_warehouse(actor,old['id'],p,ip) if old else repo.create_warehouse(actor,p,ip)
    if kind in ('customers','suppliers'):
        payload={'partner_type':'CUSTOMER' if kind=='customers' else 'SUPPLIER','code':row[0],'name':row[1],'tax_id':row[2],'phone':row[3],'email':row[4],'address':row[5],'city':row[6],'credit_limit':_num(row[7]),'payment_term_days':int(_num(row[8])),'opening_balance':_num(row[9]),'opening_balance_date':_date(row[10]),'partner_group':row[11],'contact_person':row[12],'notes':row[13],'receivable_account_id':_partner_account(row[14] if len(row)>14 else None,'RECEIVABLE','Akun Piutang') if kind=='customers' else None,'payable_account_id':_partner_account(row[14] if len(row)>14 else None,'PAYABLE','Akun Hutang') if kind=='suppliers' else None,'is_active':True}
        old=_existing('business_partners','code',row[0]);return repo.update_partner(actor,old['id'],payload,ip) if old else repo.create_partner(actor,payload,ip)
    if kind=='products':
        base_unit=_lookup('units','code',row[5],'Satuan')
        units=[{'unit_id':base_unit,'conversion_ratio':1,'purchase_price':_num(row[6]),'selling_price':_num(row[7]),'is_base':True}]
        for start in (16,20,24):
            if len(row)>start and _text(row[start]):
                alt_id=_lookup('units','code',row[start],'Satuan alternatif')
                ratio=_num(row[start+1],1)
                if ratio<=0: raise ValueError('Rasio satuan alternatif harus lebih dari nol.')
                units.append({'unit_id':alt_id,'conversion_ratio':ratio,'purchase_price':_num(row[start+2]),'selling_price':_num(row[start+3]),'is_base':False})
        payload={'sku':row[0],'name':row[1],'barcode':_text(row[2]) or None,'category_id':_lookup('item_categories','code',row[3],'Kategori',True),'brand_id':_lookup('brands','code',row[4],'Merk',True),'unit_id':base_unit,'product_type':'STOCK','purchase_price':_num(row[6]),'selling_price':_num(row[7]),'minimum_stock':_num(row[8]),'initial_stock':_num(row[9]),'warehouse_id':_lookup('warehouses','code',row[10],'Gudang',True),'opening_balance_date':_date(row[11]),'inventory_account_id':_account(row[12] or '1200','Akun persediaan'),'sales_account_id':_account(row[13] or '4000','Akun penjualan'),'cogs_account_id':_account(row[14] or '5000','Akun HPP'),'notes':row[15],'units':units,'is_active':True}
        old=_existing('products','sku',row[0])
        if old:return repo.update_product(actor,old['id'],payload,ip)
        return repo.create_product(actor,payload,ip)
    if kind=='services':
        p={'service_code':row[0],'service_name':row[1],'category_id':_service_category(row[2]),'unit_id':_lookup('units','code',row[3],'Satuan',True),'purchase_price':_num(row[4]),'selling_price':_num(row[5]),'tax_percent':_num(row[6]),'purchase_account_id':_account(row[7] or '5100','Akun pembelian jasa'),'sales_account_id':_account(row[8] or '4000','Akun penjualan jasa'),'notes':row[9],'is_active':True}
        old=_existing('services','service_code',row[0]);return repo.update_service(actor,old['id'],p,ip) if old else repo.create_service(actor,p,ip)
    if kind=='cash_accounts':
        old=_existing('cash_accounts','code',row[0]);p={'code':row[0],'name':row[1],'account_type':_text(row[2]).upper() or 'CASH','bank_name':row[3],'account_number':row[4],'opening_balance':_num(row[5]),'opening_balance_date':_date(row[6]) if row[6] else date.today().isoformat(),'coa_account_id':_account(row[7] or ('1010' if (_text(row[2]).upper() or 'CASH')=='BANK' else '1000'),'COA Kas/Bank')};return repo.update_cash_account(actor,old['id'],p,ip) if old else repo.create_cash_account(actor,p,ip)
    if kind=='salespersons':
        p={'code':row[0],'name':row[1],'phone':row[2],'email':row[3],'commission_percent':_num(row[4]),'is_active':_yes(row[5])}
        old=_existing('salespersons','code',row[0]);return repo.update_salesperson(actor,old['id'],p,ip) if old else repo.create_salesperson(actor,p,ip)
    if kind=='fixed_assets':
        p={'asset_code':row[0],'asset_name':row[1],'acquisition_date':_date(row[2]),'acquisition_cost':_num(row[3]),'residual_value':_num(row[4]),'useful_life_months':int(_num(row[5])),'asset_account_id':_account(row[6],'Akun aktiva'),'accumulated_depreciation_account_id':_account(row[7],'Akun akumulasi'),'depreciation_expense_account_id':_account(row[8],'Akun beban penyusutan'),'contra_account_id':_account(row[9],'Akun lawan'),'notes':row[10]}
        old=_existing('fixed_assets','asset_code',row[0]);return repo.update_fixed_asset(actor,old['id'],p,ip) if old else repo.create_fixed_asset(actor,p,ip)
    if kind=='departments':
        old=_existing('departments','code',row[0]);p={'code':row[0],'name':row[1],'manager_name':row[2],'notes':row[3],'is_active':_yes(row[4])};return repo.update_department(actor,old['id'],p,ip) if old else repo.create_department(actor,p,ip)
    if kind=='projects':
        old=_existing('projects','code',row[0]);p={'code':row[0],'name':row[1],'customer_id':_lookup('business_partners','code',row[2],'Pelanggan',True),'start_date':_date(row[3]) if row[3] else None,'end_date':_date(row[4]) if row[4] else None,'status':_text(row[5]).upper() or 'ACTIVE','notes':row[6],'is_active':_yes(row[7])};return repo.update_project(actor,old['id'],p,ip) if old else repo.create_project(actor,p,ip)
    raise ValueError('Jenis impor belum didukung.')

def _header_mapping(ws,headers,optional_headers=None):
    expected=[_norm(x) for x in headers]; optional={_norm(x) for x in (optional_headers or [])}
    required=[x for x in expected if x not in optional]
    for rowno in range(1,min(ws.max_row,10)+1):
        values=[_norm(c.value) for c in ws[rowno]]
        index={v:i for i,v in enumerate(values) if v}
        if all(x in index for x in required):return rowno,[index.get(x) for x in expected]
    found=[_text(c.value) for c in ws[1]]
    raise ValueError('Header template tidak dikenali. Gunakan Download Template. Kolom ditemukan: '+', '.join(found))

def import_excel(actor,kind,excel_base64,file_name,ip):
    if kind not in SPECS or (SPECS[kind].get('ultima') and not IS_ULTIMA):raise ValueError('Jenis impor tidak tersedia.')
    if not excel_base64:raise ValueError('Isi file Excel kosong.')
    _,load_workbook,Font,PatternFill,_=_wb()
    try:
        raw=base64.b64decode(excel_base64.split(',',1)[-1],validate=True)
        if len(raw)>25*1024*1024:raise ValueError('Ukuran file maksimal 25 MB.')
        wb=load_workbook(BytesIO(raw),data_only=True,read_only=True);ws=wb.active
    except ValueError:raise
    except Exception as exc:raise ValueError('File Excel tidak valid atau rusak. Gunakan format .xlsx dari template aplikasi.') from exc
    headers=SPECS[kind]['headers'];header_row,mapping=_header_mapping(ws,headers,SPECS[kind].get('optional_headers'))
    ok=0;errors=[];total=0
    for idx,values in enumerate(ws.iter_rows(min_row=header_row+1,values_only=True),header_row+1):
        row=[values[i] if i is not None and i<len(values) else None for i in mapping]
        if not any(v not in (None,'') for v in row):continue
        total+=1
        try:_process(kind,actor,row,ip);ok+=1
        except Exception as exc:errors.append((idx,row,str(exc)))
    if total==0:raise ValueError('Tidak ada baris data yang dapat diproses.')
    err_b64=None
    if errors:
        from openpyxl import Workbook
        ew=Workbook();es=ew.active;es.title='Data Gagal';es.append(['Baris']+headers+['Error'])
        for c in es[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='B42318')
        for idx,row,message in errors:es.append([idx]+row+[message])
        bio=BytesIO();ew.save(bio);err_b64=base64.b64encode(bio.getvalue()).decode()
    repo.audit(actor['id'],'MASTER_EXCEL_IMPORTED','master_import',kind,{'file':file_name,'total':total,'success':ok,'failed':len(errors)},ip)
    return {'type':kind,'total_rows':total,'success_count':ok,'failed_count':len(errors),'error_file_base64':err_b64,'error_filename':f'Error_Import_{kind}.xlsx'}
